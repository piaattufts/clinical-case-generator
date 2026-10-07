# ruff: noqa: E501
"""Round 2 version 4 keeps sources immutable and separates the two reviews."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from app.services.round2_v4 import (
    ALEX_DOCX,
    ALEX_EXPECTED,
    ALL_IDS,
    SET1_IDS,
    SET2,
    SET2_IDS,
    V3,
    leak_hits,
    old_design_hits,
    relationship_counts,
    sha256_file,
    write_package,
)
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_v4_package_gates_and_preserves_sources(tmp_path: Path) -> None:
    watched = [ALEX_DOCX, *V3.glob("VAL-*_resident.json"), *V3.glob("VAL-*_evaluator.json")]
    watched.extend(SET2.glob("VAL-*_resident.json"))
    watched.extend(SET2.glob("VAL-*_evaluator.json"))
    before = {path: _sha(path) for path in watched}
    result = write_package(tmp_path)
    assert {path: _sha(path) for path in watched} == before
    assert result["mismatches"] == []
    assert result["held"] == []
    assert result["set1"] == list(SET1_IDS)
    assert result["set2"] == list(SET2_IDS)

    residents = {}
    for case_id in ALL_IDS:
        resident = json.loads((tmp_path / "cases" / f"{case_id}_resident.json").read_text(encoding="utf-8"))
        evaluator = json.loads((tmp_path / "cases" / f"{case_id}_evaluator.json").read_text(encoding="utf-8"))
        residents[case_id] = resident
        assert "reference_discharge_plan" not in resident
        assert "v4_provenance" not in resident
        assert "reference_discharge_plan" in evaluator
        assert evaluator["v4_provenance"]["synthetic_facts"] or case_id in {"VAL-817", "VAL-818", "VAL-820"}
        assert leak_hits(resident) == []
        assert old_design_hits(resident) == []
        for medication in evaluator["reference_discharge_plan"]["medications"]:
            assert medication["evidence_class"] in {
                "SUFFICIENT_VISIBLE_EVIDENCE",
                "WEAK_VISIBLE_EVIDENCE",
                "CLINICALLY_AMBIGUOUS",
            }
    assert "38.4" not in json.dumps(residents["VAL-809"])
    assert "viridans" not in json.dumps(residents["VAL-809"]).lower()
    assert "mycophenolate" not in json.dumps(residents["VAL-813"]).lower()
    assert "intravenous" not in json.dumps(residents["VAL-805"]).lower() or "no intravenous" in json.dumps(residents["VAL-805"]).lower()

    alex = json.loads((tmp_path / "alex_feedback" / "VAL-801_alex_feedback.json").read_text(encoding="utf-8"))
    assert alex["c1_overall"]["selected"] == ["Fail"]
    assert alex["c2"]["selected"] == []
    assert alex["recommendation"]["selected"] == []
    assert "baseline" in alex["c1_comment"].lower()
    for case_id in ALEX_EXPECTED:
        assert (tmp_path / "alex_feedback" / f"{case_id}_alex_feedback.json").is_file()
    assert sha256_file(ALEX_DOCX) == before[ALEX_DOCX]

    set1 = Document(str(tmp_path / "codebooks" / "CliniProof_Round2_v4_Set1_Clinician_Validation.docx"))
    text = "\n".join(paragraph.text for paragraph in set1.paragraphs)
    assert text.index("CURRENT RESIDENT-FACING CASE") < text.index("CLINICIAN VALIDATION REFERENCE")
    assert text.index("CLINICIAN VALIDATION REFERENCE") < text.index("HISTORICAL ROUND 1 FEEDBACK — KO")
    assert text.index("HISTORICAL C1 FEEDBACK — ALEX / TEAM KAPOw") < text.index("CLINICAL REVISION RECORD")
    assert "Clinical-sufficiency rubric applied; no Alex case-specific feedback." in text
    assert "VAL-804" not in text
    set2 = Document(str(tmp_path / "codebooks" / "CliniProof_Round2_v4_Set2_Clinician_Validation.docx"))
    set2_text = "\n".join(paragraph.text for paragraph in set2.paragraphs)
    assert "ALEX / TEAM KAPOw" not in set2_text
    assert "VAL-801" not in set2_text
    counts = relationship_counts()
    assert counts["CONCORDANT"] >= 1
    assert counts["COMPLEMENTARY"] >= 1
    assert counts["INDEPENDENT"] >= 18
