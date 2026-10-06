"""KO review packages stay separated and do not alter frozen VAL files."""

from __future__ import annotations

import json
from pathlib import Path

from app.services.ko_review_sets import (
    FRESH_UNTOUCHED_IDS,
    FROZEN_RESIDENT,
    REPAIR_IDS,
    REVISED_IDS,
    SOURCE_DIR,
    write_ko_review_sets,
)
from docx import Document

FORBIDDEN_PACKET_PHRASES = ("error-bearing", "clean control", "clean controls")


def _tree_bytes(path: Path) -> dict[str, bytes]:
    files = [item for item in sorted(path.rglob("*")) if item.is_file()]
    return {item.relative_to(path).as_posix(): item.read_bytes() for item in files}


def test_ko_sets_cover_each_seed_case_once_and_hide_the_reference(tmp_path: Path) -> None:
    frozen_before = FROZEN_RESIDENT.read_bytes()
    frozen_trees = {
        name: _tree_bytes(Path("data/case_sets") / name) for name in ("seed_guided", "balanced")
    }
    source_before = _tree_bytes(SOURCE_DIR)
    revised = tmp_path / "revised"
    fresh = tmp_path / "fresh"
    held = tmp_path / "held"
    summary = write_ko_review_sets(revised, fresh, held)
    assert FROZEN_RESIDENT.read_bytes() == frozen_before
    for name, before in frozen_trees.items():
        assert _tree_bytes(Path("data/case_sets") / name) == before
    assert _tree_bytes(SOURCE_DIR) == source_before
    assert summary["revised"] == list(REVISED_IDS)
    assert summary["held"] == []
    assert set(summary["fresh"]) == set(FRESH_UNTOUCHED_IDS) | set(REPAIR_IDS)
    assert set(REVISED_IDS).isdisjoint(summary["fresh"])
    assert set(summary["revised"]) | set(summary["fresh"]) | set(summary["held"]) == {
        f"VAL-{number}" for number in range(801, 825)
    }
    assert summary["repair_status"] == {
        case_id: "READY_FOR_CLINICIAN_REVIEW" for case_id in REPAIR_IDS
    }
    assert not list(held.glob("*.docx"))
    held_audit = (held / "KO_HELD_CASES_AUDIT.md").read_text(encoding="utf-8")
    assert "No case remains held." in held_audit

    for case_id in REVISED_IDS:
        resident = json.loads((revised / f"{case_id}_resident.json").read_text(encoding="utf-8"))
        evaluator = json.loads((revised / f"{case_id}_evaluator.json").read_text(encoding="utf-8"))
        _assert_split(resident, evaluator)
        original = json.loads(
            (SOURCE_DIR / f"{case_id}_resident.json").read_text(encoding="utf-8")
        )
        assert resident["ClinicalCase"]["presentation"]["hpi"] != original["ClinicalCase"][
            "presentation"
        ]["hpi"]
        classes = [item["evidence_class"] for item in evaluator["evidence_trace"]]
        assert classes
        assert set(classes) <= {"SUFFICIENT_EVIDENCE"}

    for case_id in FRESH_UNTOUCHED_IDS:
        resident = json.loads((fresh / f"{case_id}_resident.json").read_text(encoding="utf-8"))
        original = json.loads(
            (SOURCE_DIR / f"{case_id}_resident.json").read_text(encoding="utf-8")
        )
        assert resident == original
        evaluator = json.loads((fresh / f"{case_id}_evaluator.json").read_text(encoding="utf-8"))
        _assert_split(resident, evaluator)

    for case_id in REPAIR_IDS:
        resident = json.loads((fresh / f"{case_id}_resident.json").read_text(encoding="utf-8"))
        evaluator = json.loads((fresh / f"{case_id}_evaluator.json").read_text(encoding="utf-8"))
        original = json.loads(
            (SOURCE_DIR / f"{case_id}_resident.json").read_text(encoding="utf-8")
        )
        assert resident != original
        _assert_split(resident, evaluator)
        classes = [item["evidence_class"] for item in evaluator["evidence_trace"]]
        assert set(classes) <= {"SUFFICIENT_EVIDENCE", "WEAK_EVIDENCE"}
        assert "HIDDEN_ANSWER_DEPENDENCY" not in classes
        assert "CLINICALLY_INCONSISTENT" not in classes
        blob = json.dumps(resident).casefold()
        assert "resume when holding" not in blob
        assert "intended to restart" not in blob
        assert "intended regimen" not in blob
        assert "in this profile" not in blob
        assert "expected state" not in blob

    statin = json.loads((revised / "VAL-802_resident.json").read_text(encoding="utf-8"))
    assert any("atorvastatin" in row["drug"] for row in statin["CaseMedication"])
    supply = json.loads((revised / "VAL-803_evaluator.json").read_text(encoding="utf-8"))
    lisinopril = next(
        item
        for item in supply["reference_discharge_plan"]["medications"]
        if item["medication"].startswith("lisinopril")
    )
    assert lisinopril["duration"] == "30 days"
    ibuprofen = next(
        item
        for item in json.loads((revised / "VAL-801_evaluator.json").read_text(encoding="utf-8"))[
            "reference_discharge_plan"
        ]["medications"]
        if item["medication"].startswith("ibuprofen")
    )
    assert ibuprofen["action"] == "continue"

    lisinopril_806 = json.loads((fresh / "VAL-806_evaluator.json").read_text(encoding="utf-8"))
    held_lisinopril = next(
        item
        for item in lisinopril_806["reference_discharge_plan"]["medications"]
        if item["medication"].startswith("lisinopril")
    )
    assert held_lisinopril["action"] == "restart"
    assert next(
        item["evidence_class"]
        for item in lisinopril_806["evidence_trace"]
        if item["medication"].startswith("lisinopril")
    ) == "WEAK_EVIDENCE"
    chart_806 = json.loads((fresh / "VAL-806_resident.json").read_text(encoding="utf-8"))
    assert "restart lisinopril" not in json.dumps(chart_806).casefold()
    potassium = json.loads((fresh / "VAL-807_resident.json").read_text(encoding="utf-8"))
    assert "potassium repletion" not in json.dumps(potassium).casefold()
    diuretic = json.loads((fresh / "VAL-808_resident.json").read_text(encoding="utf-8"))
    inpatient_loop = next(
        row
        for row in diuretic["CaseMedication"]
        if row["drug"].startswith("furosemide") and row["context"] == "inpatient"
    )
    home_loop = next(
        row
        for row in diuretic["CaseMedication"]
        if row["drug"].startswith("furosemide") and row["context"] == "home"
    )
    assert inpatient_loop["route"] == home_loop["route"] == "oral"
    assert inpatient_loop["frequency"] == home_loop["frequency"] == "once daily"
    discharge_weight = next(
        row for row in diuretic["CaseWeight"] if row["timepoint"] == "discharge"
    )
    assert discharge_weight["weight_kg"] == "92.000"
    assert discharge_weight["dry_weight_kg"] == "92.000"
    cmv = json.loads((fresh / "VAL-814_resident.json").read_text(encoding="utf-8"))
    assert not any(
        "valganciclovir" in row["drug"] and row["context"] == "home"
        for row in cmv["CaseMedication"]
    )
    tacrolimus = json.loads((fresh / "VAL-815_resident.json").read_text(encoding="utf-8"))
    doses = {
        row["context"]: row["dose"]
        for row in tacrolimus["CaseMedication"]
        if "tacrolimus" in row["drug"]
    }
    assert doses["home"] == doses["inpatient"] == "1 MG"
    assert "no dose change" in json.dumps(tacrolimus).casefold()
    fracture = json.loads((fresh / "VAL-817_resident.json").read_text(encoding="utf-8"))
    assert "enoxaparin" not in json.dumps(fracture).casefold()
    prophylaxis = json.loads((fresh / "VAL-820_resident.json").read_text(encoding="utf-8"))
    enoxaparin = next(row for row in prophylaxis["CaseMedication"] if "enoxaparin" in row["drug"])
    assert enoxaparin["indication"] == "Inpatient venous-thromboembolism prophylaxis"
    apixaban = json.loads((fresh / "VAL-821_resident.json").read_text(encoding="utf-8"))
    assert "resume apixaban" not in json.dumps(apixaban).casefold()

    revised_text = "\n".join(
        paragraph.text
        for paragraph in Document(str(revised / "KO_REVISED_CASES_REVIEW.docx")).paragraphs
    ).casefold()
    fresh_text = "\n".join(
        paragraph.text
        for paragraph in Document(str(fresh / "KO_CLEAN_CASES_REVIEW.docx")).paragraphs
    ).casefold()
    for text in (revised_text, fresh_text):
        for phrase in FORBIDDEN_PACKET_PHRASES:
            assert phrase not in text
        assert "clinician validation reference" in text
        assert "not shown to residents" in text
        assert "c1 — clinical plausibility" in text
        assert "c4 — alternative acceptable answers" in text
    assert "which additional heart-failure therapies" in revised_text
    assert "temporarily held" in revised_text
    assert "clean cases ready for clinician review" in fresh_text
    assert "does not mean clinically validated" in fresh_text
    for case_id in summary["held"]:
        assert case_id.casefold() not in fresh_text


def _assert_split(resident: dict[str, object], evaluator: dict[str, object]) -> None:
    assert "reference_discharge_plan" not in resident
    assert "reference_discharge_plan" in evaluator
    blob = json.dumps(resident).casefold()
    assert "exactly as listed" not in blob
    assert "error_category" not in blob
    assert "clean_expected_state" not in blob
    medications = resident["CaseMedication"]
    assert isinstance(medications, list)
    assert all(row["context"] != "discharge" for row in medications)
