"""Resident Word exports match the current active case sources and hide answer keys."""

from __future__ import annotations

import ast
import hashlib
import json
import re
from datetime import date
from pathlib import Path

from app.services.word_export import export_all, load_active_batches, load_resident_cases
from app.services.word_export_qa import audit_document
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
EXPORT_DATE = date(2026, 9, 29)
WORD_DIR = ROOT / "exports" / "word"
ARCHIVED = (
    "CLINIPROOF_TAXONOMY_V1",
    "CLINIPROOF_BALANCED_V2",
    "CLINIPROOF_BALANCED_V3",
    "CLINIPROOF_SEEDCASES_V1",
    "CLINIPROOF_SEEDCASES_V2",
    "VAL-201",
    "VAL-301",
    "VAL-401",
    "VAL-501",
    "VAL-601",
)
RUBRIC_PHRASES = (
    "A rating of 1 means implausible.",
    "A rating of 2 means questionable and requires revision.",
    "A rating of 3 means plausible with minor concern.",
    "A rating of 4 means fully plausible.",
    "Easy, Moderate, Hard, or Inappropriate / outlier",
    "Actual difficulty should ultimately be determined from resident performance.",
    "Could this hospitalization occur as charted?",
    "Is the predetermined discrepancy actually present, and does it match the answer key?",
    "without the chart announcing the answer",
    "could be an alternative answer",
    "Accept. The case is suitable for use without clinically meaningful revision.",
    "Revise. The case requires one or more changes before it should be used.",
    "Exclude. The case should not be used",
    "planned",
    "There is no resident-review application in this repository.",
)


def _hashes() -> dict[str, str]:
    paths = [
        ROOT / "data" / "case_sets" / "balanced" / "resident_validation_cases.json",
        ROOT / "data" / "case_sets" / "seed_guided" / "resident_validation_cases.json",
        ROOT / "data" / "case_sets" / "balanced" / "investigator_answer_key.json",
        ROOT / "data" / "case_sets" / "seed_guided" / "investigator_answer_key.json",
    ]
    return {
        str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest() for path in paths
    }


def test_active_batches_are_current_and_archives_are_excluded(tmp_path: Path) -> None:
    before = _hashes()
    balanced, seed = load_active_batches(ROOT)
    assert balanced.batch_code == "CLINIPROOF_BALANCED_V4"
    assert balanced.first_id == "VAL-701" and balanced.last_id == "VAL-724"
    assert seed.batch_code == "CLINIPROOF_SEEDCASES_V3"
    assert seed.first_id == "VAL-801" and seed.last_id == "VAL-824"
    written = export_all(ROOT, tmp_path, export_date=EXPORT_DATE)
    assert _hashes() == before
    names = {path.name for path in written.values()}
    assert "CliniProof_Balanced_Case_Set.docx" in names
    assert "CliniProof_Seed_Guided_Case_Set.docx" in names
    assert "CliniProof_Resident_Validation_Codebook.docx" in names
    for path in written.values():
        text = "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
        for token in ARCHIVED:
            assert token not in text


def test_case_documents_match_source_without_answer_keys(tmp_path: Path) -> None:
    balanced, seed = load_active_batches(ROOT)
    export_all(ROOT, tmp_path, export_date=EXPORT_DATE)
    for batch, filename in (
        (balanced, "CliniProof_Balanced_Case_Set.docx"),
        (seed, "CliniProof_Seed_Guided_Case_Set.docx"),
    ):
        cases = load_resident_cases(batch)
        audit = audit_document(tmp_path / filename, batch, cases)
        assert audit.exported_ids == [str(case["case_id_code"]) for case in cases]
        assert audit.missing == []
        assert audit.duplicates == []
        assert audit.answer_key_hits == []
        assert audit.passed, _failure_sample(audit)
        rendered = _document_text(tmp_path / filename)
        assert "ClinicalCase" not in rendered
        assert '"CaseMedication"' not in rendered
        for case in cases:
            for meds in case["CaseMedication"]:
                assert meds["dose"] in rendered
                assert meds["route"] in rendered
                assert meds["frequency"] in rendered
            for lab in case["CaseLab"]:
                assert str(lab["test_name"]) in rendered
                assert json.dumps(lab["value"]) in rendered
                assert str(lab["unit"]) in rendered


def test_committed_word_files_pass_the_same_audit() -> None:
    balanced, seed = load_active_batches(ROOT)
    for batch, filename in (
        (balanced, "CliniProof_Balanced_Case_Set.docx"),
        (seed, "CliniProof_Seed_Guided_Case_Set.docx"),
    ):
        path = WORD_DIR / filename
        coded = WORD_DIR / batch.coded_filename
        assert path.is_file()
        assert coded.read_bytes() == path.read_bytes()
        audit = audit_document(path, batch, load_resident_cases(batch))
        assert audit.passed, _failure_sample(audit)
        Document(str(path))


def test_codebook_follows_the_rubric_and_hides_case_targets() -> None:
    path = WORD_DIR / "CliniProof_Resident_Validation_Codebook.docx"
    document = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for phrase in RUBRIC_PHRASES:
        assert phrase in text
    ids = set(re.findall(r"VAL-(\d{3})", text))
    assert ids <= {"701", "724", "801", "824"}
    for token in ("f1_", "f2_", "error_category", "error_family", "correct_action"):
        assert token not in text
    assert "not implemented" in text.casefold() or "planned" in text.casefold()


def test_exporter_source_does_not_call_a_language_model() -> None:
    for relative in (
        "app/services/word_export.py",
        "app/services/word_facts.py",
        "app/services/word_export_qa.py",
    ):
        source = (ROOT / relative).read_text(encoding="utf-8")
        tree = ast.parse(source)
        imported: set[str] = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported.update(alias.name.split(".")[0] for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported.add(node.module.split(".")[0])
        assert "openai" not in imported
        assert "httpx" not in imported
        lowered = source.casefold()
        assert "chat.completions" not in lowered
        assert "responses.create" not in lowered


def _document_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _failure_sample(audit: object) -> str:
    lines: list[str] = []
    cases = getattr(audit, "cases", [])
    for case in cases:
        if getattr(case, "ok", True):
            continue
        lines.append(case.case_id)
        lines.extend(case.medication_mismatches[:5])
        lines.extend(case.numeric_mismatches[:8])
        lines.extend(case.unsupported[:12])
        if len(lines) > 40:
            break
    return "\n".join(lines)
