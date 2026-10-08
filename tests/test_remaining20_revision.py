"""The remaining twenty clean cases follow the frozen framework."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from docx import Document

REPO = Path(__file__).resolve().parents[1]
CLEAN = REPO / "data" / "case_sets" / "seed_guided" / "CLEAN_BASE"
REVISED = REPO / "data" / "case_sets" / "seed_guided" / "REVISED" / "remaining_20"
OVERLAP = ("VAL-801", "VAL-805", "VAL-809", "VAL-813")
REMAINING = (
    "VAL-802",
    "VAL-803",
    "VAL-804",
    "VAL-806",
    "VAL-807",
    "VAL-808",
    "VAL-810",
    "VAL-811",
    "VAL-812",
    "VAL-814",
    "VAL-815",
    "VAL-816",
    "VAL-817",
    "VAL-818",
    "VAL-819",
    "VAL-820",
    "VAL-821",
    "VAL-822",
    "VAL-823",
    "VAL-824",
)
DOMAINS = (
    "presentation_diagnosis_coherence",
    "causal_context",
    "diagnostic_workup",
    "treatment_trajectory",
    "laboratory_vital_trend",
    "hospital_course_completeness",
    "medication_decision_support",
    "discharge_stability_chronology",
    "internal_consistency",
    "unsupported_hidden_reference_action",
)
NO_COMMENT = "Framework-guided investigator revision; no case-specific clinician feedback."
LEAKS = ("reference_discharge_plan", "revision_provenance", "synthetic_facts", "answer_key")


def _load(path: Path) -> dict[str, Any]:
    payload: Any = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise AssertionError(path)
    return payload


def test_framework_is_frozen_and_covers_only_reviewed_defects() -> None:
    text = (REPO / "docs" / "revision" / "revision_framework.md").read_text(encoding="utf-8")
    assert "This framework is frozen." in text
    assert "Answer-revealing wording is not a domain" in text
    for case_id in OVERLAP:
        assert case_id in text
    for domain in (
        "Presentation and diagnosis coherence",
        "Causal context",
        "Diagnostic work-up",
        "Treatment trajectory",
        "Laboratory and vital trend",
        "Hospital-course completeness",
        "Medication-decision support",
        "Discharge stability and chronology",
        "Internal consistency",
        "Unsupported hidden-reference action",
    ):
        assert domain in text


def test_clean_base_is_unchanged_and_overlap_cases_are_not_in_this_folder() -> None:
    integrity = (REPO / "docs" / "source_integrity.md").read_text(encoding="utf-8")
    rows = re.findall(
        r"`(data/case_sets/seed_guided/CLEAN_BASE/[^`]+\.json)` \| `([0-9a-f]{64})`",
        integrity,
    )
    assert len(rows) == 48
    for relative, digest in rows:
        actual = hashlib.sha256((REPO / relative).read_bytes()).hexdigest()
        assert actual == digest, relative
    for case_id in OVERLAP:
        assert not (REVISED / f"{case_id}_resident.json").exists()


def test_every_remaining_case_hides_the_reference_and_records_provenance() -> None:
    for case_id in REMAINING:
        resident = _load(REVISED / f"{case_id}_resident.json")
        evaluator = _load(REVISED / f"{case_id}_evaluator.json")
        blob = json.dumps(resident)
        for leak in LEAKS:
            assert leak not in blob
        assert all(row.get("context") != "discharge" for row in resident["CaseMedication"])
        stripped = dict(evaluator)
        reference = stripped.pop("reference_discharge_plan")
        provenance = stripped.pop("revision_provenance")
        assert stripped == resident
        assert provenance["clinically_validated"] is False
        assert provenance["status"] == "FRAMEWORK-REVISED AND READY FOR CLINICIAN REVIEW"
        assert provenance["synthetic_facts"]
        chart = blob.casefold()
        for item in reference["medications"]:
            token = str(item["medication"]).split()[0].casefold()
            assert token in chart, f"{case_id} {token}"
        if case_id in {"VAL-802", "VAL-803"}:
            assert provenance["reviewer_1_comment"]
            assert provenance["statement"] != NO_COMMENT
        else:
            assert provenance["reviewer_1_comment"] is None
            assert provenance["statement"] == NO_COMMENT


def test_changed_references_match_the_visible_decisions() -> None:
    expected = {
        "VAL-802": ("atorvastatin", "continue"),
        "VAL-803": ("hydrochlorothiazide", "stop"),
        "VAL-806": ("lisinopril", "hold"),
        "VAL-808": ("furosemide", "change"),
        "VAL-814": ("mycophenolate", "hold"),
        "VAL-814-dose": ("valganciclovir", "start"),
        "VAL-821": ("apixaban", "hold"),
        "VAL-822": ("pantoprazole", "continue"),
        "VAL-823": ("apixaban", "hold"),
        "VAL-824": ("aspirin", "stop"),
    }
    for key, (drug, action) in expected.items():
        case_id = key.split("-dose")[0] if key.endswith("-dose") else key
        if key == "VAL-814-dose":
            case_id = "VAL-814"
        evaluator = _load(REVISED / f"{case_id}_evaluator.json")
        matches = [
            item["action"]
            for item in evaluator["reference_discharge_plan"]["medications"]
            if drug in str(item["medication"]).casefold()
        ]
        assert matches == [action], f"{case_id} {drug} {matches}"
    valg = _load(REVISED / "VAL-814_evaluator.json")
    dose = next(
        item["dose"]
        for item in valg["reference_discharge_plan"]["medications"]
        if "valganciclovir" in item["medication"]
    )
    assert dose == "450 MG"
    resident_803 = json.dumps(_load(REVISED / "VAL-803_resident.json")).casefold()
    assert "seven-day supply" not in resident_803
    assert "7-day supply" not in resident_803


def test_application_table_and_codebook_use_the_frozen_instrument() -> None:
    with (REPO / "docs" / "revision" / "revision_framework_application.csv").open(
        encoding="utf-8", newline=""
    ) as handle:
        rows = list(csv.DictReader(handle))
    by_case: dict[str, list[str]] = {}
    for row in rows:
        by_case.setdefault(row["case_id"], []).append(row["framework_domain"])
    assert set(by_case) == set(REMAINING)
    for case_id, domains in by_case.items():
        for domain in DOMAINS:
            assert domains.count(domain) == 1
        if case_id in {"VAL-802", "VAL-803"}:
            assert "reviewer_1_case_comment" in domains
        else:
            assert "reviewer_1_case_comment" not in domains
    document = Document(
        str(REPO / "docs" / "validation" / "CliniProof_Revised_Remaining20_Codebook.docx")
    )
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    cells = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    combined = text + "\n" + cells
    for case_id in REMAINING:
        assert case_id in combined
    for case_id in OVERLAP:
        assert f"CASE {case_id}" not in text
    assert "C1 — Clinical plausibility" in combined
    assert "C5 — Expected learner difficulty" in combined
    assert "Clean control" in combined
    assert "not clinically validated" in text.casefold()
    assert "reference_discharge_plan" not in combined
    log = (REPO / "docs" / "revision" / "remaining_20_revision_log.md").read_text(encoding="utf-8")
    assert "not clinically validated" in log.casefold()
    assert NO_COMMENT in log or "no case-specific clinician comments" in log.casefold()
