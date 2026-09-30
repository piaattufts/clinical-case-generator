"""Validation casebooks pair each chart with its canonical answer-key reference."""

from __future__ import annotations

import hashlib
from pathlib import Path

from app.services.validation_casebook import (
    TAG_C1_RESULT,
    TAG_C1_SCORE,
    TAG_C2_RESULT,
    TAG_C3_RESULT,
    TAG_C4_RESULT,
    TAG_C5,
    TAG_CASE_DATE,
    TAG_INITIALS,
    TAG_RECOMMENDATION,
    TAG_REVIEW_DATE,
    TAG_REVIEWER_CODE,
    export_casebooks,
    load_prepared_cases,
    stats_for,
)
from app.services.word_controls import CHECKED_CHAR, audit_form_controls
from app.services.word_export import load_active_batches
from app.services.word_facts import medication_tuple, rows_for_context
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "docs" / "resident_review_package" / "files"


def _blocks(path: Path) -> list[tuple[str, str, object]]:
    document = Document(str(path))
    found: list[tuple[str, str, object]] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            style = paragraph.style.name if paragraph.style is not None else ""
            found.append(("p", style, paragraph.text.strip()))
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            found.append(("t", "", rows))
    return found


def _chunks(blocks: list[tuple[str, str, object]]) -> list[tuple[str, str]]:
    chunks: list[tuple[str, str]] = []
    current = ""
    buffer: list[str] = []
    for kind, style, payload in blocks:
        text = str(payload)
        if kind == "p" and style == "Heading 1" and text.startswith("CASE VAL-"):
            if current:
                chunks.append((current, "\n".join(buffer)))
            current = text.removeprefix("CASE ")
            buffer = [text]
            continue
        if not current:
            continue
        if kind == "p":
            buffer.append(text)
        else:
            rows = payload if isinstance(payload, list) else []
            for row in rows:
                buffer.extend(str(cell) for cell in row)
    if current:
        chunks.append((current, "\n".join(buffer)))
    return chunks


def _source_hashes() -> dict[str, str]:
    names = (
        "resident_validation_cases.json",
        "investigator_answer_key.json",
        "batch_plan.json",
    )
    hashes: dict[str, str] = {}
    for folder in ("balanced", "seed_guided"):
        for name in names:
            path = ROOT / "data" / "case_sets" / folder / name
            hashes[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    return hashes


def test_casebooks_match_plan_and_answer_key_without_changing_sources(tmp_path: Path) -> None:
    before = _source_hashes()
    written = export_casebooks(ROOT, tmp_path)
    assert _source_hashes() == before
    balanced, seed = load_active_batches(ROOT)
    expected = {
        balanced: "CliniProof_Balanced_Validation_Casebook.docx",
        seed: "CliniProof_SeedGuided_Validation_Casebook.docx",
    }
    for batch, filename in expected.items():
        cases = load_prepared_cases(batch)
        stats = stats_for(batch, cases)
        assert stats.total == 24
        assert stats.controls == 4
        assert stats.error_bearing == 20
        assert stats.family_1 + stats.family_2 == 20
        path = written[filename]
        document = Document(str(path))
        assert document.paragraphs
        chunks = _chunks(_blocks(path))
        assert [case_id for case_id, _text in chunks] == [item.case_id for item in cases]
        by_id = {item.case_id: item for item in cases}
        for case_id, text in chunks:
            item = by_id[case_id]
            chart, after_c1 = text.split("C1 — Clinical plausibility", 1)
            _c1, reference_onward = after_c1.split("Validation Reference", 1)
            reference, after_reference = reference_onward.split(
                "C2 — Intended assessment problem", 1
            )
            assert "C2 — Intended assessment problem" in reference_onward
            assert "C3 — Detectability" in after_reference
            assert "C4 — Absence of unintended competing problems" in after_reference
            assert "C5 — Expected learner difficulty" in after_reference
            assert "Overall recommendation" in after_reference
            assert "Accept" in after_reference
            assert "Revise" in after_reference
            assert "Exclude" in after_reference
            assert "Presentation and demographics" in text
            for label in (
                "C1 reviewer comments",
                "C2 reviewer comments",
                "C3 reviewer comments",
                "C4 reviewer comments",
                "C5 reviewer comments",
                "Overall comments / suggested revisions",
            ):
                assert label in text
            assert "Validation Reference" not in chart
            for context in ("home", "inpatient", "discharge"):
                for row in rows_for_context(item.resident, context):
                    drug, dose, route, frequency = medication_tuple(row)
                    assert drug in chart
                    assert dose in chart
                    assert route in chart
                    assert frequency in chart
            if item.is_control:
                assert "Clean control" in reference
                assert "Error-bearing" not in reference
            else:
                assert item.error["error_description"] in reference
                assert item.error["correct_action"] in reference
                assert item.error["error_category"] in reference
                assert "Error-bearing" in reference
        whole = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "Saving your completed review" in whole
        assert "agreed study communication channel" in whole
        assert "do not need to edit GitHub" in whole
        assert "will not be shown to residents during the later assessment study" in whole
        assert "Please select one response per item unless otherwise indicated." in whole


def test_generated_casebooks_use_word_checkbox_controls(tmp_path: Path) -> None:
    written = export_casebooks(ROOT, tmp_path)
    for path in written.values():
        _assert_fillable(path)
        _assert_entries_survive_resave(path, tmp_path)


def test_readme_describes_fillable_casebooks() -> None:
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    package = (ROOT / "docs" / "resident_review_package" / "README.md").read_text(encoding="utf-8")
    sentence = (
        "The validation casebooks are fillable Microsoft Word documents. "
        "Click the checkboxes to select ratings and type comments directly into "
        "the provided fields."
    )
    for text in (readme, package):
        assert sentence in text
        assert "Please select one response per rating item." in text


def test_committed_casebooks_open() -> None:
    for name in (
        "CliniProof_Balanced_Validation_Casebook.docx",
        "CliniProof_SeedGuided_Validation_Casebook.docx",
    ):
        path = PACKAGE / name
        document = Document(str(path))
        headings = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.style is not None
            and paragraph.style.name == "Heading 1"
            and paragraph.text.startswith("CASE VAL-")
        ]
        assert len(headings) == 24
        assert len(set(headings)) == 24
        _assert_fillable(path)


def _assert_fillable(path: Path) -> None:
    audit = audit_form_controls(path)
    assert audit.xml_parts_well_formed
    assert audit.checkbox_controls == 24 * 47
    assert audit.malformed_checkboxes == 0
    assert audit.unchecked_checkboxes == audit.checkbox_controls
    assert audit.w14_checkbox_start_tags == audit.checkbox_controls
    assert audit.static_ballot_glyphs == 0
    assert audit.display_ballot_glyphs == audit.checkbox_controls
    assert audit.wingdings_fonts == 0
    assert audit.symbol_elements == 0
    assert audit.macros is False
    assert audit.document_protection is False
    assert audit.activex_parts == 0
    assert audit.duplicate_sdt_ids == 0
    assert audit.malformed_plain_text == 0
    assert audit.compatibility_mode == "16"
    assert audit.rich_text_controls == 24 * 6
    assert audit.bounding_box_controls == (
        audit.checkbox_controls + audit.plain_text_controls + audit.rich_text_controls
    )
    assert audit.plain_text_controls == 2 + (24 * 2)
    assert "COVER" in audit.plain_text_by_case
    assert audit.checkboxes_by_case.get("COVER", {}) == {}
    assert audit.plain_text_by_case["COVER"][TAG_REVIEWER_CODE] == 1
    assert audit.plain_text_by_case["COVER"][TAG_REVIEW_DATE] == 1
    case_ids = [case_id for case_id in audit.checkboxes_by_case if case_id != "COVER"]
    assert len(case_ids) == 24
    for case_id in case_ids:
        counts = audit.checkboxes_by_case[case_id]
        assert counts[TAG_C1_SCORE] == 32
        assert counts[TAG_C1_RESULT] == 2
        assert counts[TAG_C2_RESULT] == 2
        assert counts[TAG_C3_RESULT] == 2
        assert counts[TAG_C4_RESULT] == 2
        assert counts[TAG_C5] == 4
        assert counts[TAG_RECOMMENDATION] == 3
        assert sum(counts.values()) == 47
        fields = audit.plain_text_by_case[case_id]
        assert fields[TAG_INITIALS] == 1
        assert fields[TAG_CASE_DATE] == 1


def _assert_entries_survive_resave(path: Path, tmp_path: Path) -> None:
    document = Document(str(path))
    checkbox = next(document.element.iter(qn("w14:checkbox")))
    checked = checkbox.find(qn("w14:checked"))
    assert checked is not None
    checked.set(qn("w14:val"), "1")
    control = checkbox.getparent().getparent()
    display = control.find(".//" + qn("w:t"))
    assert display is not None
    display.text = CHECKED_CHAR
    commented = False
    for table in document.tables:
        if table.rows and "C1 reviewer comments" in table.cell(0, 0).text:
            table.cell(0, 0).add_paragraph("Persisted reviewer comment")
            commented = True
            break
    assert commented
    coded = False
    for control in document.element.iter(qn("w:sdt")):
        properties = control.find(qn("w:sdtPr"))
        if properties is None:
            continue
        tag = properties.find(qn("w:tag"))
        if tag is None or tag.get(qn("w:val")) != TAG_REVIEWER_CODE:
            continue
        text = control.find(".//" + qn("w:t"))
        assert text is not None
        text.text = "R01"
        coded = True
        break
    assert coded
    saved = tmp_path / f"resaved-{path.name}"
    document.save(str(saved))
    reopened = Document(str(saved))
    reopened_box = next(reopened.element.iter(qn("w14:checkbox")))
    reopened_checked = reopened_box.find(qn("w14:checked"))
    assert reopened_checked is not None
    assert reopened_checked.get(qn("w14:val")) == "1"
    reopened_control = reopened_box.getparent().getparent()
    reopened_text = reopened_control.find(".//" + qn("w:t"))
    assert reopened_text is not None
    assert reopened_text.text == CHECKED_CHAR
    assert any("Persisted reviewer comment" in table.cell(0, 0).text for table in reopened.tables)
    reviewer_code = ""
    for control in reopened.element.iter(qn("w:sdt")):
        properties = control.find(qn("w:sdtPr"))
        if properties is None:
            continue
        tag = properties.find(qn("w:tag"))
        if tag is None or tag.get(qn("w:val")) != TAG_REVIEWER_CODE:
            continue
        text = control.find(".//" + qn("w:t"))
        reviewer_code = "" if text is None else (text.text or "")
        break
    assert reviewer_code == "R01"
    again = audit_form_controls(saved)
    assert again.checkbox_controls == 24 * 47
    assert again.static_ballot_glyphs == 0
    assert again.macros is False
    assert again.document_protection is False
    assert again.xml_parts_well_formed
