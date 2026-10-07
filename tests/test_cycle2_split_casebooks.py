"""Split Cycle 2 codebooks keep clinical files unchanged and copy Round 1 marks."""

from __future__ import annotations

import hashlib
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from app.services.cycle2_split_casebooks import (
    CLEAN_IDS,
    REVISED_IDS,
    write_split_casebooks,
)
from app.services.round1_feedback import NOT_COMPLETED, parse_round1_casebook
from app.services.word_controls import (
    append_checkbox,
    append_plain_text,
    append_rich_text_field,
    finalize_word_form,
    prepare_form_document,
)
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
FROZEN = (
    ROOT / "data" / "case_sets" / "seed_guided" / "resident_validation_cases.json",
    ROOT / "data" / "case_sets" / "balanced" / "resident_validation_cases.json",
    ROOT / "exports" / "clean_balanced_seed_set" / "VAL-801_resident.json",
    ROOT / "exports" / "ko_revised_cases_v1" / "VAL-801_resident.json",
)


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _filled_round1(path: Path) -> None:
    document = Document()
    prepare_form_document(document)
    reviewer = document.add_paragraph("Reviewer code: ")
    append_plain_text(reviewer, document, tag="reviewer-code", alias="Reviewer code")
    dated = document.add_paragraph("Review date: ")
    append_plain_text(dated, document, tag="review-date", alias="Review date")
    _set_control_text(document, "reviewer-code", "KO")
    _set_control_text(document, "review-date", "10/5/2026")
    document.add_paragraph("CASE VAL-801")
    for _domain in range(8):
        paragraph = document.add_paragraph()
        for _score in range(4):
            append_checkbox(paragraph, document, tag="c1-score", alias="C1 domain rating")
    _check_nth(document, "c1-score", 1)
    _result(document, "c1-result", "C1 overall result")
    _comment(document, "c1-comments", "C1 reviewer comments", "cause of delirium was not explained")
    _result(document, "c2-result", "C2 result")
    _comment(
        document,
        "c2-comments",
        "C2 reviewer comments",
        "hospital course does not explain improvement",
    )
    _result(document, "c3-result", "C3 result")
    _comment(document, "c3-comments", "C3 reviewer comments", "")
    _result(document, "c4-result", "C4 result")
    _comment(document, "c4-comments", "C4 reviewer comments", "")
    paragraph = document.add_paragraph()
    for _choice in range(4):
        append_checkbox(paragraph, document, tag="c5-difficulty", alias="C5 difficulty")
    _comment(document, "c5-comments", "C5 reviewer comments", "")
    paragraph = document.add_paragraph()
    for _choice in range(3):
        append_checkbox(paragraph, document, tag="recommendation", alias="Overall recommendation")
    _check_nth(document, "recommendation", 1)
    _comment(
        document,
        "overall-comments",
        "Overall comments / suggested revisions",
        "ibuprofen discontinuation lacked justification",
    )
    initials = document.add_paragraph("Reviewer initials/code: ")
    append_plain_text(
        initials,
        document,
        tag="reviewer-initials",
        alias="Reviewer initials or code",
    )
    case_date = document.add_paragraph("Date: ")
    append_plain_text(case_date, document, tag="case-date", alias="Date")
    _set_control_text(document, "reviewer-initials", "KO")
    _set_control_text(document, "case-date", "10/5/2026")
    document.add_paragraph("CASE VAL-803")
    document.save(str(path))
    finalize_word_form(path)


def _result(document: Document, tag: str, alias: str) -> None:
    paragraph = document.add_paragraph()
    append_checkbox(paragraph, document, tag=tag, alias=alias)
    append_checkbox(paragraph, document, tag=tag, alias=alias)


def _comment(document: Document, tag: str, alias: str, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    append_rich_text_field(table.cell(0, 0), document, tag=tag, alias=alias, lines=2)
    if text:
        _set_control_text(document, tag, text)


def _controls(document: Document, tag: str) -> list[object]:
    matches = []
    for control in document.element.findall(f".//{W}sdt"):
        properties = control.find(f"{W}sdtPr")
        marker = None if properties is None else properties.find(f"{W}tag")
        if marker is not None and marker.get(qn("w:val")) == tag:
            matches.append(control)
    return matches


def _set_control_text(document: Document, tag: str, text: str) -> None:
    control = _controls(document, tag)[-1]
    node = control.find(f".//{W}t")  # type: ignore[union-attr]
    assert node is not None
    node.text = text


def _check_nth(document: Document, tag: str, index: int) -> None:
    control = _controls(document, tag)[index]
    checked = control.find(f".//{W14}checked")  # type: ignore[union-attr]
    assert checked is not None
    checked.set(qn("w14:val"), "1")


def _xml(path: Path) -> ET.Element:
    with ZipFile(path) as package:
        return ET.fromstring(package.read("word/document.xml"))


def _text(path: Path) -> str:
    root = _xml(path)
    body = root.find(f"{W}body")
    assert body is not None
    return "\n".join(
        "".join(node.text or "" for node in child.findall(f".//{W}t")) for child in list(body)
    )


def test_parser_copies_completed_marks_and_leaves_blanks(tmp_path: Path) -> None:
    source = tmp_path / "KO Casebook Validation.docx"
    _filled_round1(source)
    parsed = parse_round1_casebook(source)
    first = parsed["VAL-801"]
    assert first.reviewer == "KO"
    assert first.review_date == "10/5/2026"
    assert first.c1_scores[0].selected == ("2",)
    assert first.c1_comment == "cause of delirium was not explained"
    assert first.c2_comment == "hospital course does not explain improvement"
    assert first.c3_comment is None
    assert first.recommendation.selected == ("Revise",)
    assert first.overall_comment == "ibuprofen discontinuation lacked justification"
    third = parsed["VAL-803"]
    assert third.recommendation.completed is False
    assert third.c1_comment is None
    assert third.overall_comment is None


def test_split_codebooks_separate_reviews_and_copy_history(tmp_path: Path) -> None:
    before = {path: _sha(path) for path in FROZEN}
    source = tmp_path / "KO Casebook Validation.docx"
    _filled_round1(source)
    revised_dir = tmp_path / "revised"
    clean_dir = tmp_path / "clean"
    revised_path, clean_path = write_split_casebooks(
        revised_dir,
        clean_dir,
        round1_casebook=source,
    )
    assert {path: _sha(path) for path in FROZEN} == before
    revised_text = _text(revised_path)
    clean_text = _text(clean_path)
    revised_ids = [
        line.removeprefix("CASE ").strip()
        for line in revised_text.splitlines()
        if line.startswith("CASE VAL-")
    ]
    clean_ids = [
        line.removeprefix("CASE ").strip()
        for line in clean_text.splitlines()
        if line.startswith("CASE VAL-")
    ]
    assert revised_ids == list(REVISED_IDS)
    assert clean_ids == list(CLEAN_IDS)
    assert set(revised_ids).isdisjoint(clean_ids)
    assert sorted(revised_ids + clean_ids) == [f"VAL-{number}" for number in range(801, 825)]
    assert "PREVIOUS ROUND 1 CLINICIAN FEEDBACK" not in clean_text
    assert "Round 2 clean case" in clean_text
    assert "unreviewed" not in clean_text.casefold()
    first, second = revised_text.split("CASE VAL-802", 1)
    assert first.index("CURRENT ROUND 2 CASE") < first.index("C1 — Clinical plausibility")
    assert first.index("C1 — Clinical plausibility") < first.index("Overall recommendation")
    assert first.index("Overall recommendation") < first.index(
        "PREVIOUS ROUND 1 CLINICIAN FEEDBACK"
    )
    assert first.index("PREVIOUS ROUND 1 CLINICIAN FEEDBACK") < first.index(
        "ROUND 2 RESPONSE TO ROUND 1 FEEDBACK"
    )
    assert "cause of delirium was not explained" in first
    assert "ibuprofen discontinuation lacked justification" in first
    assert "Reviewer: KO" in first
    assert "Date: 10/5/2026" in first
    history = first.split("PREVIOUS ROUND 1 CLINICIAN FEEDBACK", 1)[1]
    assert "☒" in history
    assert "☒" not in first.split("PREVIOUS ROUND 1 CLINICIAN FEEDBACK", 1)[0]
    third = revised_text.split("CASE VAL-803", 1)[1].split("CASE VAL-805", 1)[0]
    assert NOT_COMPLETED in third
    assert "ROUND 2 RESPONSE TO ROUND 1 FEEDBACK" in third
    assert "CURRENT ROUND 2 CASE" in second
    _assert_blank_controls(revised_path)
    _assert_blank_controls(clean_path)
    _assert_borders_and_breaks(revised_path, historical=True)
    _assert_borders_and_breaks(clean_path, historical=False)
    _assert_form_tables(revised_path)
    _assert_form_tables(clean_path)


def _assert_blank_controls(path: Path) -> None:
    root = _xml(path)
    checked = [node.get(f"{W14}val") for node in root.findall(f".//{W14}checked")]
    assert checked
    assert set(checked) == {"0"}


def _assert_borders_and_breaks(path: Path, *, historical: bool) -> None:
    root = _xml(path)
    tables = root.findall(f".//{W}tbl")
    assert tables
    for table in tables:
        borders = table.find(f"{W}tblPr/{W}tblBorders")
        assert borders is not None
        for edge in ("top", "left", "bottom", "right"):
            element = borders.find(f"{W}{edge}")
            assert element is not None
            assert int(element.get(f"{W}sz") or "0") >= 6
    breaks = [
        "".join(node.text or "" for node in paragraph.findall(f".//{W}t"))
        for paragraph in root.findall(f".//{W}p")
        if paragraph.find(f".//{W}pageBreakBefore") is not None
    ]
    assert [line for line in breaks if line.startswith("CASE VAL-")]
    historical_breaks = [line for line in breaks if line.startswith("PREVIOUS ROUND 1")]
    assert bool(historical_breaks) is historical


def _assert_form_tables(path: Path) -> None:
    root = _xml(path)
    blobs = []
    for table in root.findall(f".//{W}tbl"):
        blobs.append("".join(node.text or "" for node in table.findall(f".//{W}t")))
    assert any(blob.startswith("Domain1234Presentation and demographics") for blob in blobs)
    assert any("Current reference action" in blob for blob in blobs)
    assert any("Alternative action" in blob for blob in blobs)
    assert any("C1 reviewer comments" in blob for blob in blobs)
    assert any("Overall comments / suggested revisions" in blob for blob in blobs)
