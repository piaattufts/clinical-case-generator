"""The four-case revision stays downstream of the clean base."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

from docx import Document

REPO = Path(__file__).resolve().parents[1]
CLEAN = REPO / "data" / "case_sets" / "seed_guided" / "CLEAN_BASE"
REVISED = REPO / "data" / "case_sets" / "seed_guided" / "REVISED" / "overlap_4"
CASE_IDS = ("VAL-801", "VAL-805", "VAL-809", "VAL-813")
LEAKS = (
    "reference_discharge_plan",
    "revision_provenance",
    "answer_key",
    "error_category",
    "synthetic_facts",
)


def _load(path: Path) -> dict[str, Any]:
    payload: Any = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise AssertionError(path)
    return payload


def test_clean_base_and_review_sources_are_unchanged() -> None:
    text = (REPO / "docs" / "source_integrity.md").read_text(encoding="utf-8")
    rows = re.findall(
        r"`(data/case_sets/seed_guided/CLEAN_BASE/[^`]+\.json)` \| `([0-9a-f]{64})`",
        text,
    )
    assert len(rows) == 48
    for relative, digest in rows:
        actual = hashlib.sha256((REPO / relative).read_bytes()).hexdigest()
        assert actual == digest, relative
    comparison = REPO / "docs" / "clinical_feedback" / "reviewer_comparison.csv"
    assert "SECOND_REVISION_CANDIDATE" in comparison.read_text(encoding="utf-8")
    template = REPO / "docs" / "validation" / "CliniProof_Clinical_Validation_Template.docx"
    assert template.is_file()
    assert template.stat().st_size > 0


def test_revised_charts_hide_the_reference_and_record_synthetic_facts() -> None:
    for case_id in CASE_IDS:
        resident = _load(REVISED / f"{case_id}_resident.json")
        evaluator = _load(REVISED / f"{case_id}_evaluator.json")
        blob = json.dumps(resident)
        for leak in LEAKS:
            assert leak not in blob
        assert all(row.get("context") != "discharge" for row in resident["CaseMedication"])
        stripped = dict(evaluator)
        reference = stripped.pop("reference_discharge_plan")
        provenance = stripped.pop("revision_provenance")
        assert stripped == resident
        assert provenance["clinically_validated"] is False
        assert provenance["status"] == "READY FOR NEXT CLINICIAN REVIEW"
        assert provenance["baseline"].endswith(f"{case_id}_resident.json")
        assert provenance["synthetic_facts"]
        assert reference["medications"]
        chart = blob.casefold()
        for item in reference["medications"]:
            name = str(item["medication"]).split()[0].casefold()
            assert name in chart, f"{case_id} chart does not mention {name}"


def test_revised_references_follow_the_visible_chart() -> None:
    expectations = {
        "VAL-801": ("restart", "lisinopril", "1.0 mg/dl", "not restarted"),
        "VAL-805": ("start", "enalapril", "dry weight", "missed"),
        "VAL-809": ("restart", "lisinopril", "four weeks", "no indication for valve surgery"),
        "VAL-813": ("hold", "mycophenolate", "not a home medicine", "3.2"),
    }
    for case_id, (action, drug, *needles) in expectations.items():
        resident = _load(REVISED / f"{case_id}_resident.json")
        evaluator = _load(REVISED / f"{case_id}_evaluator.json")
        chart = json.dumps(resident).casefold()
        for needle in needles:
            assert needle in chart, f"{case_id} missing {needle}"
        actions = {
            str(item["medication"]).casefold(): item["action"]
            for item in evaluator["reference_discharge_plan"]["medications"]
        }
        match = [value for key, value in actions.items() if drug in key]
        assert match == [action]


def test_codebook_uses_the_existing_instrument_for_these_four_cases() -> None:
    codebook = REPO / "docs" / "validation" / "CliniProof_Revised_Overlap4_Codebook.docx"
    document = Document(str(codebook))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    cells = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    combined = text + "\n" + cells
    for case_id in CASE_IDS:
        assert case_id in combined
    for heading in (
        "C1 — Clinical plausibility",
        "C2 — Intended assessment problem",
        "C3 — Detectability",
        "C4 — Absence of unintended competing problems",
        "C5 — Expected learner difficulty",
        "Accept",
        "Revise",
        "Exclude",
        "Clean control",
    ):
        assert heading in combined
    assert "not clinically validated" in text.casefold()
    assert "reference_discharge_plan" not in combined
    log = (REPO / "docs" / "revision" / "overlap_4_revision_log.md").read_text(encoding="utf-8")
    assert "Reviewer 1 concern" in log
    assert "Reviewer 2 concern" in log
    assert "not clinically validated" in log.casefold()
