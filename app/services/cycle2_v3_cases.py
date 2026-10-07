# ruff: noqa: E501
"""Canonical Set 1 for the six cases with completed Round 1 feedback.

Version 3 starts from the recovered clean charts. It keeps a version-1 or
version-2 detail only when that detail has a source and does not tell the
resident the discharge action. It does not edit the other 18 cases.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx.document import Document as WordDocument

from app.services.cycle2_split_casebooks import _bordered_table, _chart, _round2_form
from app.services.cycle2_v2_cases import _historical_section
from app.services.round1_six_case_review import completed_reviews, write_feedback_files
from app.services.word_controls import finalize_word_form, prepare_form_document
from app.services.word_export import NAVY, _add_body, _add_heading, _new_document, _set_run_font

ROOT = Path(__file__).resolve().parents[2]
CLEAN = ROOT / "exports" / "clean_balanced_seed_set"
CASE_IDS = ("VAL-801", "VAL-802", "VAL-803", "VAL-805", "VAL-809", "VAL-813")

DIRECT_LEAKS = (
    "no reason to stop",
    "no discharge medicine is stopped",
    "verified list at discharge",
    "verified collateral medication list at discharge",
    "intended at discharge",
    "intended to continue",
    "intended after discharge",
    "should continue",
    "should stop",
    "should be restarted",
    "complete the planned",
    "take discharge medications exactly as listed",
    "take exactly as listed",
    "are continued",
    "planned outpatient regimen",
    "planned antibiotic course",
    "four weeks of intravenous",
)

REVIEW_ITEMS: tuple[dict[str, str], ...] = (
    {
        "id": "VAL-801-C01",
        "case": "VAL-801",
        "comment": "C1 overall Fail. Presentation, fit, and hospital course are 2. Comment: the cause of delirium is not revealed, improvement is unexplained, and it is not clear why ibuprofen was stopped. The reviewer offered infection or gastrointestinal bleeding as examples.",
        "revision": "Named poor oral intake, dry mucous membranes, and a 20 mmHg orthostatic fall as the volume-depletion course. Did not add infection or bleeding. Removed every sentence that tells the resident what to do with ibuprofen.",
        "location": "presentation; admission note; hospital course",
        "old": "The note names delirium and says it improves with treatment. It does not name a precipitant. The history says ibuprofen was stopped.",
        "new": "Poor oral intake, dry mucous membranes, and a 20 mmHg fall from the recorded supine pressure of 136 mmHg are described. Confusion clears as intake resumes. Glucose is 163 mg/dL then 103 mg/dL. Creatinine is 1.1 mg/dL then 1.0 mg/dL. Ibuprofen remains on the medication list. The chart does not say whether to stop it.",
        "reasoning": "The diagnosis says delirium has a physiological cause, but the original course does not show one, so the medication implications of the stay are not interpretable. Volume depletion is a reversible precipitant and uses the recorded laboratories and the recorded supine pressure. The dry membranes, poor-intake history, and orthostatic fall are synthetic. They were not measured in the source chart. Infection and gastrointestinal bleeding were not added because either one would open a separate medication pathway. The resident chart shows the illness. It does not state the discharge action.",
        "evidence": "[E1] [E2] [E3] [E4]",
        "status": "REVISED",
    },
    {
        "id": "VAL-801-C02",
        "case": "VAL-801",
        "comment": "C2 Fail: it was not clear why ibuprofen should be stopped. C4 Fail: ibuprofen stopped on admission, in the home medication list, with no obvious indication.",
        "revision": "The ibuprofen hold was removed from the medication rows. The reference continues it. The resident chart does not say to continue it.",
        "location": "home medications; inpatient medications; reference discharge plan",
        "old": "Ibuprofen status held. Reference action stop. Discharge instruction: ibuprofen was stopped. Geriatrics: use the verified list at discharge.",
        "new": "Ibuprofen 400 mg every 8 hours as needed is active on the home and inpatient lists. Indication remains symptomatic analgesia. Reference action continue. Geriatrics records that a collateral history was verified and does not say to use that list as the discharge order.",
        "reasoning": "Creatinine is 1.1 mg/dL then 1.0 mg/dL, and the chart does not describe bleeding. Stopping ibuprofen is not required by those facts. Continuing it is the reference because the drug is on the verified list and no stop indication is stored. A resident can still stop an as-needed NSAID in a 75-year-old; that alternative is not forbidden by the chart. The chart itself no longer announces either choice.",
        "evidence": "[E1] [E4]",
        "status": "REVISED",
    },
    {
        "id": "VAL-801-C03",
        "case": "VAL-801",
        "comment": "C3 has both Pass and Fail selected, with a blank comment. C5 is Inappropriate / outlier: the case is unclear. Overall Revise: more clinical complexity is needed, for example an infection, and the medication list should be clearer.",
        "revision": "The conflict and the incomplete C3 comment are preserved in the historical record. No infection was added.",
        "location": "round 1 historical record",
        "old": "Pass and Fail both selected. Comment blank.",
        "new": "Pass and Fail both selected. Comment blank.",
        "reasoning": "Two selected C3 boxes and a blank comment do not identify a medication. Choosing one box, or adding the infection used as an example, would invent the review.",
        "evidence": "[E1]",
        "status": "SOURCE_SELECTION_AMBIGUITY",
    },
    {
        "id": "VAL-802-C01",
        "case": "VAL-802",
        "comment": "C1 overall Fail. Fit and hospital course are 1. The cause is not revealed, the course looks empty, creatinine changed slightly, and the baseline creatinine is unknown.",
        "revision": "Named poor oral intake. Stated the recorded creatinine pair and blood pressure. Did not add a baseline creatinine or a potassium pair.",
        "location": "presentation; admission note; hospital course; laboratories",
        "old": "No precipitant. Creatinine 1.3 mg/dL then 1.2 mg/dL is stored and not discussed. No earlier creatinine.",
        "new": "Poor oral intake for two days. The same creatinine pair. Blood pressure 138/69 mmHg then 124/68 mmHg. No creatinine from before the admission.",
        "reasoning": "The resident cannot connect the stay to a medication decision without a precipitant. Poor intake does not require a new laboratory. That history is synthetic. Version 1 added a clinic creatinine of 1.2 mg/dL six weeks earlier and potassium 4.2 then 4.1 mmol/L. Neither value is in the clean chart, so neither was restored. The README sentence about that baseline describes version 1, not a recovered measurement.",
        "evidence": "[E1] [E2] [E5]",
        "status": "REVISED",
    },
    {
        "id": "VAL-802-C02",
        "case": "VAL-802",
        "comment": "The statin was described as taken but not continued. C2 Fail: was the trainee supposed to continue the statin? C3 Fail: no reasoning for stopping the statin.",
        "revision": "Removed the sentence that a continued statin had been confirmed. Atorvastatin stays on the list. The chart does not say to continue it.",
        "location": "reconciliation note; medication list; reference discharge plan",
        "old": "The reconciliation note said a pharmacy history confirmed a continued statin, while other sentences mentioned stopped medicines.",
        "new": "The note says a pharmacy fill history was used to verify the home list. Atorvastatin 40 mg daily is on that list. Reference action remains continue.",
        "reasoning": "The clean medication rows continue atorvastatin. The prose made a stop look possible. The revision removes that contradiction and does not replace it with an instruction.",
        "evidence": "[E1] [E5]",
        "status": "REVISED",
    },
    {
        "id": "VAL-802-C03",
        "case": "VAL-802",
        "comment": "C4 Fail: creatinine elevated but lisinopril continued. If the creatinine is acute kidney injury, lisinopril should be stopped. C5 Inappropriate / outlier. Overall Exclude.",
        "revision": "Reference continues lisinopril. Hold is an acceptable alternative. The uncertainty is marked on the evaluator plan. The resident chart shows the numbers and does not state the action.",
        "location": "reference discharge plan; acceptable alternatives",
        "old": "Reference continues lisinopril. Version 2 already stored a hold alternative. Version 1 did not.",
        "new": "Reference continues lisinopril and records the decision as clinically ambiguous. Hold is acceptable.",
        "reasoning": "The only creatinine values are 1.3 mg/dL and 1.2 mg/dL. The comment says the change is slight and that a baseline is missing, and also says to stop lisinopril if this is acute kidney injury. Those statements do not establish the threshold. Continuing the recorded order is the reference because the change is 0.1 mg/dL and blood pressure fell from 138 to 124 mmHg. Holding it is acceptable because no pre-admission creatinine exists. The original Exclude recommendation is preserved in the historical record.",
        "evidence": "[E1] [E5] [E6]",
        "status": "REVISED",
    },
    {
        "id": "VAL-803-C01",
        "case": "VAL-803",
        "comment": "C1 overall Fail. Fit and hospital course are 1. The cause is not revealed. The medication list was called acceptable and no discharge change was requested, then the reviewer asked why lisinopril was supplied for 7 days and what the creatinine did. C2 through the overall recommendation were not completed.",
        "revision": "Added a synthetic sodium course and an inpatient hold of hydrochlorothiazide so the delirium has a visible precipitant. Removed the sentence that said every medicine is continued and none is stopped. Kept the 30-day lisinopril supply on the hidden reference.",
        "location": "presentation; admission note; hospital course; laboratories; inpatient medications; reference discharge plan",
        "old": "No sodium values. Hydrochlorothiazide is active. The version 2 course says the four medicines are continued and no discharge medicine is stopped. The reviewed readable case gave lisinopril 7 days.",
        "new": "Sodium is 128 mmol/L on admission and 135 mmol/L at discharge. Hydrochlorothiazide is a home medicine and is held on the inpatient list while the sodium is 128 mmol/L. Confusion clears as the sodium rises. Creatinine is 1.0 mg/dL then 1.2 mg/dL. The reference stops hydrochlorothiazide. A hold pending an outpatient sodium check is an acceptable alternative. Lisinopril duration is 30 days. The resident chart does not state either discharge action.",
        "reasoning": "Hydrochlorothiazide 25 mg daily is already on the clean medication list. Thiazides can cause hyponatremia, and hyponatremia is a recognized delirium precipitant. The sodium values are not in the clean laboratories. They are a synthetic course chosen so the resident can see a mechanism and a treatment response. Sodium 128 mmol/L is low enough to explain confusion and is a value at which holding the offending thiazide is ordinary. Sodium 135 mmol/L is a recovery into the normal range, paired with the recorded statement that cognition returned. The version 2 sentence that listed the discharge actions is removed because it is the answer. The 7-day supply is the injected discrepancy. The preserved reference uses 30 days. Later Round 1 items stay incomplete.",
        "evidence": "[E1] [E2] [E7] [E8]",
        "status": "REVISED",
    },
    {
        "id": "VAL-805-C01",
        "case": "VAL-805",
        "comment": "C1 hospital course is 2. Medication regimen is both 2 and 3. The discharge weight is not close to the dry weight. Inpatient furosemide would be intravenous and increased. More intake-and-output days and an echocardiogram or cardiology recommendation were requested. Overall Revise: the resident should decide about additional heart-failure medicines.",
        "revision": "Kept the source weights, the single intake-and-output day, the oral 40 mg inpatient order, and the source laboratories. Did not add an intravenous dose, extra fluid days, or an ejection fraction. The hidden reference intensifies furosemide to 40 mg oral twice daily. Continuing once daily is an acceptable alternative. The codebook question uses only facts that are in the chart.",
        "location": "hospital course; consultations; reference discharge plan; acceptable alternatives",
        "old": "Version 2 states oral 40 mg once daily and weights 81 kg, 78 kg, and dry weight 73 kg, and the reference continues that dose. Version 1 replaces the weights with 86 kg, 83 kg, and 80 kg, changes inpatient furosemide to 40 mg intravenous twice daily, and cites an ejection fraction of 30 percent that is not stored.",
        "new": "Admission weight 81 kg, discharge weight 78 kg, dry weight 73 kg. Hospital day 3 intake 1418 mL, output 2463 mL, net -1045 mL. Inpatient furosemide remains 40 mg oral once daily. Creatinine 1.7 then 0.9 mg/dL. Potassium 4.7 then 4.3 mmol/L. Natriuretic peptide 1120 then 369 pg/mL. Blood pressure 109/82 then 110/84 mmHg. No ejection fraction. Reference frequency is twice daily. Continue once daily is acceptable. An additional drug class is acceptable and not required.",
        "reasoning": "The recorded weights do not show a return to dry weight, so a reference that simply continues the same oral dose treats an unfinished diuresis as complete. That is clinically inconsistent. Intravenous 40 mg twice daily and the 86-to-80 kg series are not in the clean chart, and the project furosemide regimen is oral 40 mg once daily and excludes the injection product. Those version 1 numbers were not kept. Furosemide labeling allows a 20 to 80 mg dose to be repeated. Twice-daily oral 40 mg is an intensification inside that labeled range. It is one defensible discharge response to a weight that is still 5 kg above dry weight. Continuing the current dose is also defensible because natriuretic peptide, oxygen saturation, and creatinine improved, so it is an acceptable alternative rather than a hidden contradiction. No ejection fraction was added. The diagnosis already says systolic heart failure. Guideline-directed classes remain optional at a discharge pressure of 110/84 mmHg after an admission creatinine of 1.7 mg/dL. The resident sees the evidence and not the chosen dose.",
        "evidence": "[E1] [E9] [E10] [E11] [E12]",
        "status": "REVISED",
    },
    {
        "id": "VAL-809-C01",
        "case": "VAL-809",
        "comment": "C1 presentation and hospital course are 2. The case should include fever and a predisposition such as a mechanical valve or poor dentition. C2 and C4 Fail: lisinopril was continued despite acute kidney injury and should not be continued if acute kidney injury is present. Overall Revise. The home list is short.",
        "revision": "Wrote the recorded fatigue, temperature, cultures, vegetation, ceftriaxone start, and creatinine pair into the course. Removed the instruction to complete the antibiotic course. Did not add fever, a valve, dental work, or a new organism. Lisinopril remains active on the inpatient list. The reference continues it, and a hold is an acceptable alternative.",
        "location": "admission note; hospital course; consultations; reference discharge plan",
        "old": "Version 2 consult says to complete the planned parenteral course. Version 1 replaces the admission temperature of 36.80°C with 38.6°C, changes the organism to viridans group streptococcus, adds a dental extraction and a murmur, holds lisinopril, and says four weeks of ceftriaxone is the planned course.",
        "new": "Temperature remains 36.80°C. The organism remains gram-positive cocci. A later culture shows no growth. The echocardiogram shows a vegetation. Ceftriaxone 2 g IV daily was started. A PICC is present. Creatinine is 1.3 mg/dL then 0.8 mg/dL. The consult lists those findings and the follow-up. It does not give the discharge order.",
        "reasoning": "The recorded temperature is 36.80°C. Replacing it with a fever erases a source vital sign. A mechanical valve is not in the chart. A dental extraction and a viridans label are not in the culture or the archetype, and they would imply an oral source the microbiology does not establish. The existing bacteremia, vegetation, inpatient antibiotic, culture clearance, and PICC are enough for the resident to decide the outpatient antimicrobial plan. The creatinine fall is visible and the baseline is not. Continuing lisinopril matches the recorded order at a discharge creatinine of 0.8 mg/dL. Holding it is acceptable because the admission value was higher. Neither action is written into the resident chart. No extra home medicine was added.",
        "evidence": "[E1] [E6] [E13] [E14] [E15]",
        "status": "REVISED",
    },
    {
        "id": "VAL-813-C01",
        "case": "VAL-813",
        "comment": "C1 medication regimen is 1. A new CMV diagnosis with valganciclovir already present on admission does not make sense. The drug should start after diagnosis and should not be an admission medicine. The regimen is too simple. Potassium changes without a cause. C4: valganciclovir should not be an admission medication. Overall Exclude, with a blank comment.",
        "revision": "Removed home valganciclovir and started 900 mg twice daily after the admission viral-load result. Did not add mycophenolate. Removed the codebook question about mycophenolate. Stated the potassium pair without inventing a cause.",
        "location": "home medications; inpatient medications; admission note; hospital course; reference discharge plan",
        "old": "The clean chart lists valganciclovir as a home medicine. Version 1 adds mycophenolate 1000 mg twice daily, which is not on the source list, and the codebook asks the reviewer to judge it.",
        "new": "Home medicines are tacrolimus 1 mg every 12 hours, amlodipine 5 mg daily, and atorvastatin 40 mg daily. Valganciclovir 900 mg twice daily starts after the recorded viral-load detection. The reference action is start. A temporary tacrolimus reduction is an acceptable alternative. Mycophenolate is absent, and the adjudication question does not mention it.",
        "reasoning": "The admission diagnosis is CMV disease, and the procedure text says the admission viral-load review detected CMV viral burden and supported antiviral treatment. A pre-admission valganciclovir row contradicts that sequence. Mycophenolate is not in the clean chart. Adding it because labeling recommends it for kidney transplantation would invent a home medicine and would force the reviewer to score a drug the resident cannot see. The dose remains 900 mg twice daily because that is the project regimen when creatinine is 1.2 mg/dL then 1.0 mg/dL. The 450 mg tablet is the product strength. RxNorm was not used to choose the dose. The resident chart describes the inpatient start. It does not say to discharge on that dose.",
        "evidence": "[E1] [E16] [E17] [E18] [E20]",
        "status": "REVISED",
    },
)


EVIDENCE_NOTES: dict[str, str] = {
    "[E1]": "KO Casebook Validation.docx, reviewer KO, 10/5/2026. Supports the comment text, the double selections, and the items left blank. It does not by itself prove a medication change.",
    "[E2]": "Inouye SK, Westendorp RGJ, Saczynski JS. Delirium in elderly people. Lancet. 2014;383:911-922. doi:10.1016/S0140-6736(13)60688-1. Supports treating a precipitant of delirium, including metabolic precipitants. It does not establish which precipitant this synthetic patient had.",
    "[E3]": "Freeman R, Wieling W, Axelrod FB, et al. Consensus statement on the definition of orthostatic hypotension. Clin Auton Res. 2011;21:69-72. doi:10.1007/s10286-011-0119-5. Supports the 20 mmHg systolic threshold used for the synthetic orthostatic finding. It does not show that this patient had that finding.",
    "[E4]": "Clean pre-injection VAL-801. Glucose 163 mg/dL then 103 mg/dL. Creatinine 1.1 mg/dL then 1.0 mg/dL. Supine blood pressure 136/78 mmHg. Ibuprofen is on the verified list. No bleeding diagnosis is stored.",
    "[E5]": "Clean pre-injection VAL-802. Creatinine 1.3 mg/dL then 1.2 mg/dL. Blood pressure 138/69 mmHg then 124/68 mmHg. No earlier creatinine and no potassium series are stored. Atorvastatin is on the verified list.",
    "[E6]": "KDIGO Clinical Practice Guideline for Acute Kidney Injury. Kidney Int Suppl. 2012;2:1-138. Supports reviewing ACE-inhibitor exposure when kidney function may have worsened. It does not create a baseline creatinine.",
    "[E7]": "Clean pre-injection VAL-803. Hydrochlorothiazide 25 mg daily is on the list. Creatinine 1.0 mg/dL then 1.2 mg/dL. Potassium 4.4 mmol/L then 4.2 mmol/L. No sodium is stored. Evaluator lisinopril duration is 30 days. The reviewed readable case gave lisinopril 7 days.",
    "[E8]": "Liamis G, Milionis H, Elisaf M. A review of drug-induced hyponatremia. Am J Kidney Dis. 2008;52:144-153. doi:10.1053/j.ajkd.2008.03.004. Thiazides are a common cause of hyponatremia. Hydrochlorothiazide labeling, DailyMed setid 9f0beacd-4c41-432d-b7d6-e779ae4c1b99. The citations support plausibility of the synthetic sodium course. They do not show that those sodium values were recovered from a source chart.",
    "[E9]": "Clean pre-injection VAL-805. Furosemide 40 mg oral once daily at home and in the hospital. Weights 81 kg, 78 kg, dry weight 73 kg. Hospital day 3 intake 1418 mL, output 2463 mL, net -1045 mL. Creatinine 1.7 then 0.9 mg/dL. Potassium 4.7 then 4.3 mmol/L. Natriuretic peptide 1120 then 369 pg/mL. Blood pressure 109/82 then 110/84 mmHg. No ejection fraction.",
    "[E10]": "Heidenreich PA, et al. 2022 AHA/ACC/HFSA Guideline for the Management of Heart Failure. Circulation. 2022;145:e895-e1032. doi:10.1161/CIR.0000000000001063. Supports using congestion, kidney function, potassium, and blood pressure when judging diuretic and other heart-failure therapy. It was not used to invent a weight or an ejection fraction.",
    "[E11]": "Repository regimen FUROSEMIDE_40_DAILY in data/bootstrap/medication_regimens.json. Oral 40 mg once daily. The entry excludes the injection product. Furosemide tablet labeling, DailyMed setid 571a52ed-5258-46d5-a0d2-9a984cf73895: the usual initial dose is 20 to 80 mg, and the same dose may be repeated 6 to 8 hours later. Supports the charted oral dose and a twice-daily reference frequency. It does not supply an intravenous order.",
    "[E12]": "Mullens W, et al. The use of diuretics in heart failure with congestion. Eur J Heart Fail. 2019;21:137-155. doi:10.1002/ejhf.1369. Supports judging decongestion with weight and urine output. It was not used to replace the recorded weights.",
    "[E13]": "Clean pre-injection VAL-809. Temperature 36.80°C on admission and at discharge. Symptom: fatigue. Gram-positive cocci, later no growth. Vegetation with preserved ventricular function. Creatinine 1.3 then 0.8 mg/dL. Ceftriaxone 2 g IV daily. PICC present. No valve or dental history.",
    "[E14]": "Baddour LM, et al. Infective Endocarditis in Adults. Circulation. 2015;132:1435-1486. doi:10.1161/CIR.0000000000000296. Supports using bacteremia, echocardiography, and antimicrobial therapy in the representation. It was not used to add a valve, a dental procedure, a fever, or a species the culture does not name.",
    "[E15]": "Repository regimen CEFTRIAXONE_ENDOCARDITIS_OPAT. Ceftriaxone 2000 mg intravenous once daily. Supports the dose already charted.",
    "[E16]": "Clean pre-injection VAL-813. Procedure text: admission viral burden detected and supported antiviral treatment; later burden lower. Creatinine 1.2 then 1.0 mg/dL. Potassium 4.7 then 3.9 mmol/L. Home list includes tacrolimus, amlodipine, atorvastatin, and valganciclovir. No mycophenolate row.",
    "[E17]": "Repository regimen VALGANCICLOVIR_CMV_TREATMENT. 900 mg oral twice daily, given as 450 mg tablets, when creatinine is in the range these profiles keep. DailyMed setid 89a934f0-85a3-44c1-82e5-d09d1738e08d. The dose was not taken from RxNorm.",
    "[E18]": "Kotton CN, et al. The Third International Consensus Guidelines on the Management of Cytomegalovirus in Solid-organ Transplantation. Transplantation. 2018;102:900-931. doi:10.1097/TP.0000000000002191. Supports starting treatment after laboratory evidence of CMV and considering the intensity of immunosuppression. It was not used to invent a viral-load number or a mycophenolate row.",
    "[E20]": "CELLCEPT labeling, DailyMed setid 37241e87-4af4-4dc3-a1aa-ea6f20d8dc40, recommends 1 g orally twice daily for adult kidney transplantation. The source case has no mycophenolate row, so the labeled dose was not added.",
}


def write_v3_package(directory: Path) -> Path:
    """Write the canonical six-case package and return the codebook path."""
    directory.mkdir(parents=True, exist_ok=True)
    write_feedback_files(directory)
    (directory / "SET1_VERSION_RECONCILIATION.md").write_text(_reconciliation(), encoding="utf-8")
    (directory / "CLINICAL_REVISION_LOG.md").write_text(_revision_log(), encoding="utf-8")
    (directory / "REVISION_DIFF.md").write_text(_revision_diff(), encoding="utf-8")
    (directory / "REVISION_EVIDENCE_LEDGER.md").write_text(_ledger(), encoding="utf-8")
    for case_id in CASE_IDS:
        resident, evaluator = _build_case(case_id)
        _assert_resident_clean(case_id, resident)
        (directory / f"{case_id}_resident.json").write_text(
            json.dumps(resident, indent=2) + "\n",
            encoding="utf-8",
        )
        (directory / f"{case_id}_evaluator.json").write_text(
            json.dumps(evaluator, indent=2) + "\n",
            encoding="utf-8",
        )
    (directory / "ANSWER_LEAK_AUDIT.md").write_text(_leak_audit(directory), encoding="utf-8")
    (directory / "REFERENCE_EVIDENCE_AUDIT.md").write_text(_reference_audit(), encoding="utf-8")
    (directory / "KATIE_REAUDIT.md").write_text(_reaudit(), encoding="utf-8")
    (directory / "MANIFEST.md").write_text(_manifest(), encoding="utf-8")
    (directory / "README.md").write_text(_readme(), encoding="utf-8")
    path = _write_codebook(directory)
    return path


def _build_case(case_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
    resident = json.loads((CLEAN / f"{case_id}_resident.json").read_text(encoding="utf-8"))
    evaluator = json.loads((CLEAN / f"{case_id}_evaluator.json").read_text(encoding="utf-8"))
    _PATCH[case_id](resident)
    _PATCH[case_id](evaluator)
    _REFERENCE[case_id](evaluator)
    resident.pop("reference_discharge_plan", None)
    if "reference_discharge_plan" in resident:
        raise ValueError(case_id)
    return resident, evaluator


def _set_note(case: dict[str, Any], note_type: str, text: str) -> None:
    for note in case.get("CaseNote") or []:
        if note.get("note_type") == note_type:
            note["note_text"] = text
            return
    raise KeyError(note_type)


def _set_hpi(case: dict[str, Any], text: str) -> None:
    case["ClinicalCase"]["presentation"]["hpi"] = text


def _patch_801(case: dict[str, Any]) -> None:
    _set_hpi(
        case,
        "A 75-year-old Male is admitted with delirium after several days of confusion, "
        "fatigue, and poor oral intake. Mucous membranes are dry. The recorded supine "
        "blood pressure is 136/78 mmHg, and the orthostatic systolic pressure falls "
        "20 mmHg from that supine value. Home medicines include lisinopril, atorvastatin, "
        "metformin, and ibuprofen 400 mg every 8 hours as needed. He could not give a "
        "reliable medication history at admission. A collateral list was verified later. "
        "Glucose is 163 mg/dL on admission and 103 mg/dL at discharge. Creatinine is "
        "1.1 mg/dL then 1.0 mg/dL.",
    )
    _set_note(
        case,
        "admission",
        "Admission note for a 75-year-old Male with delirium. Poor oral intake, dry "
        "mucous membranes, and a 20 mmHg orthostatic fall from the supine pressure of "
        "136 mmHg are present. Home medicines include lisinopril, atorvastatin, metformin, "
        "and ibuprofen as needed. A collateral medication history was verified after admission.",
    )
    _set_note(
        case,
        "hospital_course",
        "Poor oral intake and volume depletion were treated with fluids, and oral intake "
        "resumed. Confusion cleared as intake improved. Glucose is 163 mg/dL then "
        "103 mg/dL. Creatinine is 1.1 mg/dL then 1.0 mg/dL. Ibuprofen 400 mg every "
        "8 hours as needed is on the medication list for symptomatic analgesia.",
    )
    for medication in case.get("CaseMedication") or []:
        if "ibuprofen" not in str(medication.get("drug")):
            continue
        medication["status"] = "home" if medication.get("context") == "home" else "active"
        medication["held_reason"] = None
        if medication.get("indication") in {None, ""}:
            medication["indication"] = "symptomatic analgesia"
    for consult in case.get("CaseConsult") or []:
        if consult.get("service") == "geriatrics":
            consult["assessment"] = "Delirium improved toward baseline."
            consult["recommendation"] = (
                "A collateral medication history was verified after admission."
            )
    for item in case.get("CaseInstruction") or []:
        text = str(item.get("instruction_text") or item.get("text") or "")
        if "ibuprofen" in text.lower() or "exactly as listed" in text.lower():
            item["instruction_text"] = "Medication instructions were reviewed with the patient."
            if "text" in item:
                item["text"] = item["instruction_text"]


def _patch_802(case: dict[str, Any]) -> None:
    _set_hpi(
        case,
        "A 81-year-old Female is admitted with delirium after two days of confusion and "
        "poor oral intake. No other precipitant is recorded. Home medicines are lisinopril "
        "10 mg daily and atorvastatin 40 mg daily. Creatinine is 1.3 mg/dL on admission "
        "and 1.2 mg/dL at discharge. No creatinine from before this admission is recorded. "
        "Blood pressure is 138/69 mmHg then 124/68 mmHg.",
    )
    _set_note(
        case,
        "admission",
        "Admission note for a 81-year-old Female with delirium. Confusion and poor oral "
        "intake for two days. Home medicines are lisinopril and atorvastatin. No creatinine "
        "from before this admission is recorded.",
    )
    _set_note(
        case,
        "hospital_course",
        "Poor oral intake was the only precipitant recorded. Confusion improved as intake "
        "resumed. Creatinine is 1.3 mg/dL then 1.2 mg/dL. Blood pressure is 138/69 mmHg "
        "then 124/68 mmHg. Atorvastatin 40 mg daily is on the verified home list. No "
        "creatinine from before this admission is recorded.",
    )
    for reconciliation in case.get("CaseMedicationReconciliation") or []:
        reconciliation["notes"] = (
            "A pharmacy fill history was used to verify the home list after admission. "
            "Atorvastatin 40 mg daily is on that list."
        )


def _patch_803(case: dict[str, Any]) -> None:
    _set_hpi(
        case,
        "A 71-year-old Male is admitted with one week of fatigue and confusion. Home "
        "medicines include hydrochlorothiazide 25 mg daily, lisinopril, atorvastatin, "
        "and metformin. Serum sodium is 128 mmol/L on admission. Hydrochlorothiazide "
        "is held during the admission. Sodium is 135 mmol/L at discharge. Confusion "
        "clears as the sodium rises. Creatinine is 1.0 mg/dL then 1.2 mg/dL. Potassium "
        "is 4.4 mmol/L then 4.2 mmol/L.",
    )
    _set_note(
        case,
        "admission",
        "Admission note for a 71-year-old Male with fatigue and confusion for one week. "
        "Hydrochlorothiazide 25 mg daily is a home medicine. Serum sodium is 128 mmol/L. "
        "Creatinine is 1.0 mg/dL. Potassium is 4.4 mmol/L.",
    )
    _set_note(
        case,
        "hospital_course",
        "Serum sodium is 128 mmol/L on admission and 135 mmol/L at discharge. "
        "Hydrochlorothiazide is held during the admission. Confusion clears as the "
        "sodium rises. Creatinine is 1.0 mg/dL then 1.2 mg/dL. Potassium is 4.4 mmol/L "
        "then 4.2 mmol/L. Glucose is 149 mg/dL then 129 mg/dL.",
    )
    for medication in case.get("CaseMedication") or []:
        if "hydrochlorothiazide" not in str(medication.get("drug")):
            continue
        if medication.get("context") == "inpatient":
            medication["status"] = "held"
            medication["held_reason"] = "Serum sodium was 128 mmol/L on admission."
        else:
            medication["status"] = "home"
            medication["held_reason"] = None
    labs = case.setdefault("CaseLab", [])
    if not any("Sodium" in str(lab.get("test_name")) for lab in labs):
        labs.extend(
            [
                {
                    "lab_id": "LAB-VAL803-NA1",
                    "case_id": "VAL-803",
                    "timepoint": "admission",
                    "test_name": "Sodium [Moles/volume] in Serum or Plasma",
                    "value": 128,
                    "value_text": None,
                    "unit": "mmol/L",
                    "status": "final",
                    "source_reference": "synthetic_v3",
                },
                {
                    "lab_id": "LAB-VAL803-NA2",
                    "case_id": "VAL-803",
                    "timepoint": "discharge",
                    "test_name": "Sodium [Moles/volume] in Serum or Plasma",
                    "value": 135,
                    "value_text": None,
                    "unit": "mmol/L",
                    "status": "final",
                    "source_reference": "synthetic_v3",
                },
            ]
        )


def _patch_805(case: dict[str, Any]) -> None:
    _set_hpi(
        case,
        "A 68-year-old Female is admitted with acute systolic heart failure. Dyspnea, "
        "edema, and orthopnea have been present for one week. Home medicines are oral "
        "furosemide 40 mg once daily, atorvastatin, and metoprolol succinate 25 mg once "
        "daily. The inpatient furosemide order is also 40 mg oral once daily. Weight is "
        "81 kg on admission and 78 kg at discharge. Dry weight is 73 kg.",
    )
    _set_note(
        case,
        "admission",
        "Admission note for a 68-year-old Female with acute systolic heart failure. "
        "Symptoms are dyspnea, edema, and orthopnea for one week. Home and inpatient "
        "furosemide are both recorded as 40 mg oral once daily. No echocardiogram is recorded.",
    )
    _set_note(
        case,
        "hospital_course",
        "Dyspnea, edema, and orthopnea were present for one week. The chest radiograph "
        "shows pulmonary edema without pneumonia. The recorded inpatient furosemide order "
        "is 40 mg oral once daily. Weight is 81 kg on admission and 78 kg at discharge. "
        "Dry weight is 73 kg. One intake-and-output day is stored, hospital day 3, with "
        "intake 1418 mL, output 2463 mL, and net -1045 mL. Creatinine is 1.7 mg/dL then "
        "0.9 mg/dL. Potassium is 4.7 mmol/L then 4.3 mmol/L. B-type natriuretic peptide "
        "is 1120 pg/mL then 369 pg/mL. Blood pressure is 109/82 mmHg then 110/84 mmHg. "
        "Oxygen saturation is 92 percent then 98 percent. No echocardiogram and no "
        "ejection fraction are recorded. Metoprolol succinate 25 mg daily was administered.",
    )
    for consult in case.get("CaseConsult") or []:
        if consult.get("service") == "cardiology":
            consult["assessment"] = "Pulmonary edema is present on the admission radiograph."
            consult["recommendation"] = "No echocardiogram is recorded in this chart."


def _patch_809(case: dict[str, Any]) -> None:
    text = (
        "A 82-year-old Male is admitted with infective endocarditis. The recorded symptom "
        "is fatigue for one week. Temperature is 36.80°C on admission and at discharge. "
        "Home medicines are lisinopril and atorvastatin. Ceftriaxone 2000 mg intravenously "
        "once daily was started during the admission. A PICC is in place. Transthoracic "
        "echocardiogram shows a mobile echodensity consistent with a vegetation and "
        "preserved ventricular function. Admission blood culture grew gram-positive cocci. "
        "A later culture showed no growth. Creatinine is 1.3 mg/dL on admission and "
        "0.8 mg/dL at discharge. No creatinine from before this admission is recorded."
    )
    _set_hpi(case, text)
    _set_note(case, "admission", text)
    _set_note(
        case,
        "hospital_course",
        "Ceftriaxone 2000 mg intravenously once daily was given during the admission. "
        "The echocardiogram shows a vegetation with preserved ventricular function. "
        "Blood culture grew gram-positive cocci, and a later culture showed no growth. "
        "A PICC is in place for parenteral therapy. Creatinine is 1.3 mg/dL then "
        "0.8 mg/dL. Lisinopril is on the inpatient medication list. No creatinine from "
        "before this admission is recorded. Temperature is 36.80°C on admission and at discharge.",
    )
    for consult in case.get("CaseConsult") or []:
        if consult.get("service") == "infectious disease":
            consult["assessment"] = (
                "Blood culture grew gram-positive cocci. A later culture showed no growth. "
                "Echocardiogram shows a vegetation."
            )
            consult["recommendation"] = (
                "Ceftriaxone was administered intravenously during the admission. "
                "Weekly blood count and creatinine monitoring are recorded. "
                "Infectious-disease follow-up is scheduled."
            )


def _patch_813(case: dict[str, Any]) -> None:
    case["CaseMedication"] = [
        medication
        for medication in case.get("CaseMedication") or []
        if not (
            medication.get("context") == "home" and "valganciclovir" in str(medication.get("drug"))
        )
    ]
    for medication in case.get("CaseMedication") or []:
        if "valganciclovir" in str(medication.get("drug")):
            medication["notes"] = (
                "Started after the admission viral-load result. Given as 450 MG tablets."
            )
    text = (
        "A 64-year-old Female with a kidney transplant is admitted with diarrhea for "
        "several days. Home medicines are tacrolimus 1 mg every 12 hours, amlodipine "
        "5 mg daily, and atorvastatin 40 mg daily. She was not taking valganciclovir "
        "before admission. The admission viral-load review detected CMV viral burden."
    )
    _set_hpi(case, text)
    _set_note(case, "admission", text)
    _set_note(
        case,
        "hospital_course",
        "The admission viral-load review detected CMV viral burden. Valganciclovir "
        "900 mg orally twice daily, given as 450 mg tablets, was started after that "
        "result. A later review showed a lower viral burden. Diarrhea improved. "
        "Tacrolimus 1 mg every 12 hours was administered during the hospitalization. "
        "Creatinine is 1.2 mg/dL then 1.0 mg/dL. Potassium is 4.7 mmol/L then "
        "3.9 mmol/L. No cause for the potassium change is recorded.",
    )
    for consult in case.get("CaseConsult") or []:
        if consult.get("service") == "infectious disease":
            consult["recommendation"] = (
                "The admission viral-load review detected CMV viral burden. "
                "Valganciclovir was started after that result."
            )
        elif consult.get("service") == "transplant":
            consult["recommendation"] = (
                "Tacrolimus 1 mg every 12 hours is recorded. "
                "No other immunosuppressant is on the medication list."
            )


_PATCH = {
    "VAL-801": _patch_801,
    "VAL-802": _patch_802,
    "VAL-803": _patch_803,
    "VAL-805": _patch_805,
    "VAL-809": _patch_809,
    "VAL-813": _patch_813,
}


def _set_alternative(
    evaluator: dict[str, Any], medication: str, action: str, rationale: str
) -> None:
    plan = evaluator["reference_discharge_plan"]
    plan["acceptable_alternatives"] = [
        {"medication": medication, "action": action, "rationale": rationale}
    ]
    plan["evidence_uncertainty"] = "CLINICALLY_AMBIGUOUS"


def _reference_801(evaluator: dict[str, Any]) -> None:
    for item in evaluator["reference_discharge_plan"]["medications"]:
        if "ibuprofen" in str(item.get("medication")):
            item["action"] = "continue"
            item["indication"] = "symptomatic analgesia"
            item["rationale"] = (
                "Ibuprofen is on the verified list. Creatinine is 1.1 mg/dL then "
                "1.0 mg/dL. No bleeding diagnosis is stored. The resident chart does "
                "not state this action."
            )
            item["evidence_class"] = "SUFFICIENT_VISIBLE_EVIDENCE"


def _reference_802(evaluator: dict[str, Any]) -> None:
    for item in evaluator["reference_discharge_plan"]["medications"]:
        if "lisinopril" in str(item.get("medication")):
            item["evidence_class"] = "CLINICALLY_AMBIGUOUS"
            item["rationale"] = (
                "Creatinine is 1.3 mg/dL then 1.2 mg/dL and blood pressure is "
                "138/69 mmHg then 124/68 mmHg. No pre-admission creatinine is stored. "
                "Continuation is the reference. Hold is acceptable."
            )
    _set_alternative(
        evaluator,
        "lisinopril 10 MG Oral Tablet",
        "hold",
        "Hold is acceptable because no creatinine from before this admission is recorded.",
    )


def _reference_803(evaluator: dict[str, Any]) -> None:
    for item in evaluator["reference_discharge_plan"]["medications"]:
        if "hydrochlorothiazide" in str(item.get("medication")):
            item["action"] = "stop"
            item["evidence_class"] = "CLINICALLY_AMBIGUOUS"
            item["rationale"] = (
                "Home hydrochlorothiazide is recorded. Synthetic sodium is 128 mmol/L "
                "then 135 mmol/L, and the inpatient row is held. Stopping it at discharge "
                "follows that course. A hold pending an outpatient sodium check is acceptable. "
                "The sodium values were not in the source chart."
            )
        if "lisinopril" in str(item.get("medication")):
            item["duration"] = "30 days"
            item["evidence_class"] = "SUFFICIENT_VISIBLE_EVIDENCE"
            item["rationale"] = (
                "The preserved chart uses 30 days. The reviewed case used 7 days. "
                "Creatinine 1.0 mg/dL then 1.2 mg/dL does not set the duration."
            )
    _set_alternative(
        evaluator,
        "hydrochlorothiazide 25 MG Oral Tablet",
        "hold",
        "Holding through an outpatient sodium check is acceptable. The sodium course is synthetic.",
    )


def _reference_805(evaluator: dict[str, Any]) -> None:
    for item in evaluator["reference_discharge_plan"]["medications"]:
        if "furosemide" in str(item.get("medication")):
            item["action"] = "change"
            item["dose"] = "40 MG"
            item["route"] = "oral"
            item["frequency"] = "twice daily"
            item["evidence_class"] = "CLINICALLY_AMBIGUOUS"
            item["rationale"] = (
                "Discharge weight is 78 kg and dry weight is 73 kg after 40 mg oral "
                "once daily. Furosemide labeling allows a 20 to 80 mg dose to be repeated. "
                "Twice daily is the reference. Continuing once daily is acceptable because "
                "natriuretic peptide, creatinine, and oxygen saturation improved. "
                "No intravenous dose is in the project regimen."
            )
    _set_alternative(
        evaluator,
        "furosemide 40 MG Oral Tablet",
        "continue",
        "Continuing 40 mg oral once daily is acceptable. Natriuretic peptide fell from 1120 to 369 pg/mL and oxygen saturation rose from 92 to 98 percent.",
    )
    plan = evaluator["reference_discharge_plan"]
    plan["acceptable_alternatives"].append(
        {
            "medication": "Additional systolic heart-failure therapy",
            "action": "start",
            "rationale": (
                "An ACE inhibitor, ARNI, SGLT2 inhibitor, or mineralocorticoid antagonist "
                "is acceptable and not required. No ejection fraction is recorded. "
                "Discharge blood pressure is 110/84 mmHg, creatinine is 0.9 mg/dL, "
                "and potassium is 4.3 mmol/L. Admission creatinine was 1.7 mg/dL."
            ),
        }
    )


def _reference_809(evaluator: dict[str, Any]) -> None:
    for item in evaluator["reference_discharge_plan"]["medications"]:
        if "lisinopril" in str(item.get("medication")):
            item["action"] = "continue"
            item["evidence_class"] = "CLINICALLY_AMBIGUOUS"
            item["rationale"] = (
                "Creatinine is 1.3 mg/dL then 0.8 mg/dL. No pre-admission creatinine "
                "is stored. The inpatient row was not held. Continuation is the reference. "
                "A hold is acceptable."
            )
        if "ceftriaxone" in str(item.get("medication")):
            item["evidence_class"] = "SUFFICIENT_VISIBLE_EVIDENCE"
            item["rationale"] = (
                "Vegetation, gram-positive cocci, later culture clearance, inpatient "
                "ceftriaxone 2 g IV daily, and a PICC are visible. The resident chart "
                "does not state the outpatient duration."
            )
    _set_alternative(
        evaluator,
        "lisinopril 10 MG Oral Tablet",
        "hold",
        "A hold is acceptable because the admission creatinine is 1.3 mg/dL and no baseline is stored.",
    )


def _reference_813(evaluator: dict[str, Any]) -> None:
    for item in evaluator["reference_discharge_plan"]["medications"]:
        if "valganciclovir" in str(item.get("medication")):
            item["action"] = "start"
            item["dose"] = "900 MG"
            item["frequency"] = "twice daily"
            item["evidence_class"] = "SUFFICIENT_VISIBLE_EVIDENCE"
            item["rationale"] = (
                "Start 900 mg orally twice daily after the admission viral-load result. "
                "Creatinine is 1.2 mg/dL then 1.0 mg/dL. The 450 mg tablet is the product strength."
            )
        if "tacrolimus" in str(item.get("medication")):
            item["evidence_class"] = "CLINICALLY_AMBIGUOUS"
            item["rationale"] = (
                "Tacrolimus 1 mg every 12 hours is the only recorded immunosuppressant. "
                "No trough is stored. Continuation is the reference. A temporary reduction "
                "is acceptable."
            )
    _set_alternative(
        evaluator,
        "BX Rating tacrolimus 1 MG Oral Capsule",
        "modify",
        "A temporary reduction is acceptable. No trough and no requested change are recorded.",
    )


_REFERENCE = {
    "VAL-801": _reference_801,
    "VAL-802": _reference_802,
    "VAL-803": _reference_803,
    "VAL-805": _reference_805,
    "VAL-809": _reference_809,
    "VAL-813": _reference_813,
}


def _assert_resident_clean(case_id: str, resident: dict[str, Any]) -> None:
    if "reference_discharge_plan" in resident:
        raise ValueError(case_id)
    blob = json.dumps(resident).lower()
    for phrase in DIRECT_LEAKS:
        if phrase in blob:
            raise ValueError(f"{case_id} direct leak: {phrase}")


def _cell(value: str) -> str:
    return value.replace("|", "/").replace("\n", " ")


def _revision_log() -> str:
    lines = [
        "# Clinical revision log",
        "",
        "Canonical Set 1, version 3. Each row is one Round 1 concern.",
        "",
        "| ID | Comment | Revision | Location | Old | New | Clinical reasoning | Evidence |",
        "| --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for item in REVIEW_ITEMS:
        cells = [
            item["id"],
            item["comment"],
            item["revision"],
            item["location"],
            item["old"],
            item["new"],
            item["reasoning"],
            item["evidence"],
        ]
        lines.append("| " + " | ".join(_cell(cell) for cell in cells) + " |")
    lines.append("")
    return "\n".join(lines)


def _revision_diff() -> str:
    lines = [
        "# Revision diff",
        "",
        "Executed changes from the clean pre-injection chart to version 3.",
        "",
        "| Case | Location | Old | New | Comment addressed | Clinical reasoning | Evidence |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for item in REVIEW_ITEMS:
        if item["status"] == "SOURCE_SELECTION_AMBIGUITY":
            continue
        cells = [
            item["case"],
            item["location"],
            item["old"],
            item["new"],
            item["comment"],
            item["reasoning"],
            item["evidence"],
        ]
        lines.append("| " + " | ".join(_cell(cell) for cell in cells) + " |")
    lines.append("")
    return "\n".join(lines)


def _ledger() -> str:
    rows = [
        (
            "VAL-801",
            "Poor intake, dry mucous membranes, and a 20 mmHg orthostatic fall.",
            "Not in the source chart. Supine pressure 136/78 mmHg is stored.",
            "Described in the note. Not added as a second vital-sign row.",
            "SYNTHETIC_ADDITION_SUPPORTED_BY_LITERATURE [E2] [E3]. Supine pressure is CASE_SOURCE [E4].",
        ),
        (
            "VAL-801",
            "Ibuprofen hold removed.",
            "Status held. Reference stop.",
            "Status active. Reference continue. No discharge instruction in the resident chart.",
            "CASE_SOURCE [E4].",
        ),
        (
            "VAL-802",
            "Poor oral intake added. Baseline creatinine not added.",
            "No precipitant. Creatinine 1.3 then 1.2 mg/dL only.",
            "Poor intake named. Same creatinine pair. Version 1 baseline 1.2 mg/dL was not restored.",
            "SYNTHETIC_ADDITION_SUPPORTED_BY_LITERATURE for the intake history [E2]. Creatinine is CASE_SOURCE [E5].",
        ),
        (
            "VAL-803",
            "Sodium 128 then 135 mmol/L and inpatient hydrochlorothiazide hold.",
            "No sodium. Hydrochlorothiazide active.",
            "Synthetic sodium course. Inpatient row held. Reference stop, with hold acceptable. Resident chart does not state the discharge action.",
            "SYNTHETIC_ADDITION_SUPPORTED_BY_LITERATURE [E8]. Hydrochlorothiazide presence is CASE_SOURCE [E7].",
        ),
        (
            "VAL-805",
            "Furosemide reference frequency twice daily. Source dose and weights kept.",
            "40 mg oral once daily. Weights 81 / 78 / dry 73 kg.",
            "Same charted order and weights. Reference 40 mg oral twice daily. Continue once daily is acceptable. No intravenous dose and no ejection fraction.",
            "CASE_SOURCE [E9]. CLINICAL_LITERATURE [E10] [E11] [E12]. Version 1 intravenous dose and 86-to-80 kg series were not used.",
        ),
        (
            "VAL-809",
            "Antibiotic consult no longer states the discharge plan. Source temperature and organism kept.",
            "Temperature 36.80°C. Gram-positive cocci. Consult said to complete the planned course.",
            "Same temperature and organism. Consult lists findings and follow-up. Lisinopril hold is an alternative.",
            "CASE_SOURCE [E13]. CLINICAL_LITERATURE [E14]. Version 1 fever, dental extraction, and viridans label were not used.",
        ),
        (
            "VAL-813",
            "Home valganciclovir removed. Mycophenolate not added.",
            "Valganciclovir on the home list. No mycophenolate.",
            "Valganciclovir starts after the viral-load result at 900 mg twice daily. Mycophenolate remains absent. The adjudication question does not mention it.",
            "CASE_SOURCE procedure text [E16]. REPOSITORY_SOURCE dose [E17]. Labeling [E20] was not used to add mycophenolate.",
        ),
    ]
    lines = [
        "# Revision evidence ledger",
        "",
        "| Case | Change | Old | New | Evidence class |",
        "| --- | --- | --- | --- | --- |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(_cell(cell) for cell in row) + " |")
    lines.extend(["", "## Evidence notes", ""])
    for key, text in EVIDENCE_NOTES.items():
        lines.append(f"{key} {text}")
        lines.append("")
    return "\n".join(lines)


def _reconciliation() -> str:
    return """# Set 1 version reconciliation

Version 1 is `exports/ko_revised_cases_v1/`. Its audit describes dehydration, a creatinine baseline, thiazide hyponatremia, intravenous furosemide, fever, and mycophenolate. Those details were written during an earlier revision. They are not in `exports/clean_balanced_seed_set/`.

Version 2 is `exports/ko_cycle2_revised_validation_v2/`. It kept source measurements and refused invented numbers. It also left answer-telling sentences in VAL-801, VAL-803, and VAL-809, and it left VAL-803 without a precipitant.

Version 3 starts from the clean chart. A version-1 detail is used only when the drug was already on the source list and the added course is labeled synthetic. Invented weights, an intravenous furosemide dose, an ejection fraction, a replacement temperature, a new organism, a dental extraction, a baseline creatinine, and mycophenolate are not used.

| Case | Field | V1 | V2 | Provenance of v1 | Provenance of v2 | Better supported | Why | Starting value for v3 |
| --- | --- | --- | --- | --- | --- | --- | --- | --- |
| VAL-801 | Precipitant | Dehydration narrative, and geriatrics says no bleeding was found | Volume depletion, and geriatrics says to use the verified list at discharge | Synthetic narrative | Synthetic narrative plus a source blood pressure | Neither sentence | Both versions reveal a medication conclusion | Clean chart plus a synthetic volume course, with the conclusion sentences removed |
| VAL-801 | Ibuprofen | Continue, no discharge instruction in the later v1 consult | Continue, but the course says there is no reason to stop it | Clean medication row | Clean medication row | The medication row | The stop is unsupported. The conclusion must stay in the reference | Active ibuprofen row. Reference continue. No resident instruction |
| VAL-802 | Baseline creatinine | 1.2 mg/dL six weeks earlier, plus potassium 4.2 then 4.1 | No baseline. Creatinine 1.3 then 1.2 only | Not in the clean chart | Clean laboratories | V2 | A baseline was not recovered | V2 laboratories. Poor intake added and labeled synthetic |
| VAL-803 | Sodium and HCTZ | Sodium 128 then 135. HCTZ held. Reference stop | No sodium. Course says all medicines are continued and none is stopped | HCTZ is source. Sodium is synthetic | Clean laboratories, plus an answer leak | V1 mechanism, without v1's home-row hold label and without v2's answer sentence | The missing precipitant is the Round 1 problem. The sodium values must stay labeled synthetic | Inpatient HCTZ held. Synthetic sodium. Reference stop. Resident chart does not state the action |
| VAL-803 | Lisinopril supply | 30 days | 30 days | Clean evaluator. The reviewed case had 7 days | Clean evaluator | Both | The 7-day supply is the injected discrepancy | 30 days on the hidden reference only |
| VAL-805 | Weights and diuretic | 86 to 80 kg, dry weight 80 kg, furosemide 40 mg IV twice daily, EF 30 percent | 81 kg, 78 kg, dry weight 73 kg, oral 40 mg daily, no EF | Not in the clean chart. No IV regimen in the repository | Clean chart | V2 measurements | Invented weights and an intravenous dose were not recovered | Source measurements. Reference oral 40 mg twice daily, with once daily acceptable |
| VAL-809 | Presentation | Fever 38.6°C, dental extraction, viridans, murmur, lisinopril held, four-week course stated | Temperature 36.80°C, gram-positive cocci, consult says to complete the course | Temperature, organism, and dental history are not the source values | Clean vital sign, culture, and imaging | V2 facts, without the consult sentence | Replacing 36.80°C or the organism erases source data | Source facts. Neutral consult. Lisinopril hold is an alternative |
| VAL-813 | Mycophenolate and valganciclovir | Mycophenolate added. Valganciclovir starts after the viral load | No mycophenolate. Valganciclovir starts after the viral load. Codebook still asks about mycophenolate | Mycophenolate is not in the clean list. The viral-load procedure is source | Same procedure text. Question came from the shared form written for v1 | V2 medication list | A reviewer cannot score a drug that is absent | No mycophenolate. Question rewritten. Valganciclovir 900 mg twice daily after the recorded result |

The uploaded KO Casebook Validation.docx was read for the six reviews. The checkbox states match the extracted Round 1 JSON: VAL-801 C3 is both Pass and Fail, VAL-805 medication regimen is both 2 and 3, and VAL-803 stops after C1.
"""


def _leak_audit(directory: Path) -> str:
    lines = [
        "# Answer-leak audit for version 3",
        "",
        "Resident JSON files were scanned for direct discharge instructions.",
        "A phrase that describes an inpatient event or a recorded order is clinical evidence.",
        "",
        "| Case | Phrase checked | Classification |",
        "| --- | --- | --- |",
    ]
    for case_id in CASE_IDS:
        blob = (directory / f"{case_id}_resident.json").read_text(encoding="utf-8").lower()
        hits = [phrase for phrase in DIRECT_LEAKS if phrase in blob]
        if hits:
            lines.append(f"| {case_id} | {', '.join(hits)} | DIRECT_ANSWER_LEAK |")
        else:
            lines.append(f"| {case_id} | none of the direct-leak phrases | CLINICAL_EVIDENCE |")
    lines.extend(
        [
            "",
            "Reviewed strong hints that were removed: geriatrics telling the reader to use the verified list at discharge; the sentence that there is no reason to stop ibuprofen; the sentence that no discharge medicine is stopped; and the instruction to complete the planned parenteral course.",
            "",
            "Remaining clinical wording that is not a discharge order: oral intake resumed; hydrochlorothiazide held during the admission while sodium was 128 mmol/L; inpatient furosemide recorded as 40 mg oral once daily; ceftriaxone administered during the admission; valganciclovir started after the viral-load result.",
            "",
            "DIRECT_ANSWER_LEAK count: 0.",
            "",
        ]
    )
    return "\n".join(lines)


def _reference_audit() -> str:
    return """# Reference evidence audit for version 3

No reference action is HIDDEN_REFERENCE_DEPENDENCY or CLINICALLY_INCONSISTENT.
Ambiguous actions are marked on the evaluator plan and have an acceptable alternative.

| Case | Medication | Reference action | Visible evidence | Class | Alternative |
| --- | --- | --- | --- | --- | --- |
| VAL-801 | Ibuprofen | continue | On the list for symptomatic analgesia. Creatinine 1.1 then 1.0 mg/dL. No bleeding diagnosis. | SUFFICIENT_VISIBLE_EVIDENCE | A resident may still stop an as-needed NSAID. The chart does not forbid it and does not encode a second required action. |
| VAL-801 | Lisinopril, metformin, atorvastatin | continue | Active diagnoses and inpatient administration. | SUFFICIENT_VISIBLE_EVIDENCE | None required. |
| VAL-802 | Lisinopril | continue | Creatinine 1.3 then 1.2 mg/dL. Blood pressure 138/69 then 124/68 mmHg. No baseline. | CLINICALLY_AMBIGUOUS | Hold. |
| VAL-802 | Atorvastatin | continue | On the verified list. | SUFFICIENT_VISIBLE_EVIDENCE | None. |
| VAL-803 | Hydrochlorothiazide | stop | Home drug. Synthetic sodium 128 then 135 mmol/L. Held while sodium was 128 mmol/L. Confusion cleared. | CLINICALLY_AMBIGUOUS | Hold pending an outpatient sodium check. |
| VAL-803 | Lisinopril | continue, 30 days | Clean reference duration. Creatinine 1.0 then 1.2 mg/dL. | SUFFICIENT_VISIBLE_EVIDENCE | None for the duration. The 7-day supply is not restored. |
| VAL-805 | Furosemide | change to 40 mg oral twice daily | Weight 81 then 78 kg, dry weight 73 kg, one negative fluid day, oral 40 mg once daily charted. | CLINICALLY_AMBIGUOUS | Continue 40 mg once daily. |
| VAL-805 | Metoprolol, atorvastatin | continue | Administered and matched to systolic heart failure or hyperlipidemia. | SUFFICIENT_VISIBLE_EVIDENCE | Additional drug classes are acceptable, not required. |
| VAL-809 | Ceftriaxone | start / continue 2 g IV daily | Vegetation, gram-positive cocci, culture clearance, inpatient start, PICC. | SUFFICIENT_VISIBLE_EVIDENCE | Duration is not invented. |
| VAL-809 | Lisinopril | continue | Creatinine 1.3 then 0.8 mg/dL. No baseline. Inpatient row was active. | CLINICALLY_AMBIGUOUS | Hold. |
| VAL-813 | Valganciclovir | start 900 mg twice daily | Not a home medicine. Started after viral-load detection. Creatinine 1.2 then 1.0 mg/dL. | SUFFICIENT_VISIBLE_EVIDENCE | None for the start. |
| VAL-813 | Tacrolimus | continue 1 mg every 12 hours | Only recorded immunosuppressant. No trough. | CLINICALLY_AMBIGUOUS | Temporary reduction. |
| VAL-813 | Mycophenolate | absent | Not on the source list and not added. | Not scored | The codebook does not ask about it. |
"""


def _reaudit() -> str:
    return """# Katie re-audit of version 3

These six cases are prepared for another clinician review. They are not clinically validated.

| Case | Round 1 concern addressed? | Clean chart? | Answer leakage? | Reference supported? | Multiple reasonable answers represented? | Katie status |
| --- | --- | --- | --- | --- | --- | --- |
| VAL-801 | Partial. A precipitant is visible. Infection was not added. | Yes | No direct leak | Yes for continuation of the recorded medicines | Stopping ibuprofen remains possible and is not forbidden | PASS_WITH_MINOR_CONCERN |
| VAL-802 | Partial. No baseline was invented. The Exclude recommendation is preserved historically. | Yes | No | Ambiguous lisinopril decision is marked | Hold is encoded | PASS |
| VAL-803 | Partial. The cause is now visible. The review after C1 stays incomplete. | Yes. Sodium is labeled synthetic. | No | Stop versus hold is marked | Hold is encoded | PASS_WITH_MINOR_CONCERN |
| VAL-805 | Partial. The weight gap and the oral inpatient dose are visible. Intravenous therapy and an ejection fraction were not invented. | Yes | No | Ambiguous diuretic intensity is marked | Continue once daily, and optional extra therapy, are encoded | PASS |
| VAL-809 | Partial. Source findings are visible. Fever and a dental source were not invented. | Yes | No | Antibiotic evidence is sufficient. Lisinopril is ambiguous and marked | Hold is encoded | PASS |
| VAL-813 | Partial. Chronology is coherent. Mycophenolate was not invented, and the question about it was removed. | Yes | No | Antiviral start is supported. Tacrolimus intensity is marked ambiguous | Temporary reduction is encoded | PASS |

DIRECT_ANSWER_LEAK = 0

HIDDEN_REFERENCE_DEPENDENCY = 0

CLINICALLY_INCONSISTENT = 0

CODEBOOK REFERENCES ABSENT FACT = 0

README/V3 MISMATCH is checked by tests/test_set1_v3.py.

OLD ERROR-INJECTION LANGUAGE = 0 in the six resident charts.
"""


def _manifest() -> str:
    lines = [
        "# Set 1 version 3 manifest",
        "",
        "This is the current Set 1 package. Version 1 and version 2 are superseded revision attempts.",
        "",
        "| Case | Resident file | Evaluator file |",
        "| --- | --- | --- |",
    ]
    for case_id in CASE_IDS:
        lines.append(f"| {case_id} | `{case_id}_resident.json` | `{case_id}_evaluator.json` |")
    lines.append("")
    return "\n".join(lines)


def _readme() -> str:
    return "\n".join(
        [
            "# Set 1 revised cases, version 3",
            "",
            "Open [CliniProof_Cycle2_Revised_Cases_Validation.docx](CliniProof_Cycle2_Revised_Cases_Validation.docx).",
            "",
            "This is the current Set 1 package for VAL-801, VAL-802, VAL-803, VAL-805, VAL-809, and VAL-813.",
            "Version 1 and version 2 are superseded revision attempts. They are not the current review.",
            "The eighteen-case Set 2 codebook is separate.",
            "",
            "Ready for clinician review does not mean clinically validated.",
            "",
        ]
    )


def _write_codebook(directory: Path) -> Path:
    document = _new_document("CliniProof Cycle 2  |  Set 1 revised cases, version 3")
    prepare_form_document(document)
    _add_heading(document, "CliniProof", 0)
    _add_body(document, "Set 1. Six revised cases prepared for clinician re-review.")
    _add_heading(document, "SET 1", 1)
    _add_heading(document, "REVISED CASES FOLLOWING ROUND 1 CLINICIAN FEEDBACK", 2)
    _add_body(document, "Cases: VAL-801, VAL-802, VAL-803, VAL-805, VAL-809, VAL-813")
    _add_body(
        document,
        "The chart is the information a resident would see. Complete C1 and C2 before you read "
        "the hidden reference. The reference is not shown to residents. Historical Round 1 "
        "feedback follows the blank form. Ready for clinician review does not mean clinically validated.",
    )
    reviews = completed_reviews()
    for case_id in CASE_IDS:
        resident, evaluator = _build_case(case_id)
        row = {"case_id": case_id, "resident": resident, "evaluator": evaluator}
        banner = document.add_paragraph()
        banner.paragraph_format.page_break_before = True
        banner.paragraph_format.keep_with_next = True
        _set_run_font(banner.add_run(f"CASE {case_id}"), size=16, bold=True, color=NAVY)
        _add_body(document, "Round 2 revised case")
        _add_heading(document, "RESIDENT-FACING CHART", 2)
        _add_body(document, "This section is the same information the resident would see.")
        _chart(document, resident)
        _round2_form(document, row)
        _historical_section(document, reviews[case_id])
        _revision_section(document, case_id)
    path = directory / "CliniProof_Cycle2_Revised_Cases_Validation.docx"
    document.save(str(path))
    finalize_word_form(path)
    return path


def _revision_section(document: WordDocument, case_id: str) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True
    _set_run_font(heading.add_run("CLINICAL REVISION RECORD"), size=14, bold=True, color=NAVY)
    rows = tuple(
        (
            item["comment"],
            item["revision"],
            item["location"],
            item["old"],
            item["new"],
            item["reasoning"],
            item["evidence"],
        )
        for item in REVIEW_ITEMS
        if item["case"] == case_id
    )
    _bordered_table(
        document,
        (
            "Round 1 comment",
            "Revision",
            "Location",
            "Old",
            "New",
            "Clinical reasoning",
            "Evidence",
        ),
        rows,
        (1.05, 0.95, 0.85, 0.95, 0.95, 1.15, 0.9),
        header=True,
        size=8,
    )
    _add_heading(document, "Evidence", 3)
    for key, text in EVIDENCE_NOTES.items():
        if any(key in item["evidence"] for item in REVIEW_ITEMS if item["case"] == case_id):
            _add_body(document, f"{key} {text}")
