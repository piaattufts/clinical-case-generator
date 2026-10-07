"""Deterministic Word export of the current resident-facing case sets.

The renderer reads resident_validation_cases.json and writes chart text.
It does not call a language model and does not add clinical facts.
"""

from __future__ import annotations

import json
import shutil
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.services.word_facts import (
    clinical,
    exact_text,
    field_text,
    is_missing,
    medication_note,
    note_text,
    presentation,
    rows_for_context,
)

FONT = "Calibri"
CONTENT_WIDTH = 6.8
HEADER_FILL = "D9E2F3"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)

EMPTY_MEDICATIONS = "No medications were specified for this list."
EMPTY_MONITORING = "No scheduled monitoring was specified."
EMPTY_FOLLOWUP = "No follow-up appointments were specified."
EMPTY_INSTRUCTIONS = "No discharge instructions were specified."
EMPTY_VITALS = "No vital signs were specified."
EMPTY_LABS = "No laboratory results were specified."
EMPTY_COURSE = "No hospital-course details were specified."
EMPTY_MEDREC = "No medication-reconciliation documentation was specified."

BALANCED_STATEMENT = (
    "These are synthetic clinical cases generated for medication-reconciliation "
    "and transition-of-care assessment. They are currently intended for "
    "clinician/resident review and validation."
)
SEED_STATEMENT = (
    "This case set was developed using six resident-provided clinical examples as design inputs. "
    "Those examples were abstracted into clinical archetypes and used to generate new synthetic "
    "encounters. The current cases are not copies of the original resident cases."
)

MED_HEADERS = ("Medication", "Dose", "Route", "Frequency", "Indication / relevant note")
MED_WIDTHS = (2.35, 0.85, 0.8, 1.05, 1.75)


class WordExportError(ValueError):
    """The active case files do not match the validation registry."""


@dataclass(frozen=True)
class ActiveBatch:
    batch_code: str
    generation_strategy: str
    directory: Path
    resident_json: Path
    first_id: str
    last_id: str
    case_count: int
    status_label: str
    approach_label: str
    title: str
    simple_filename: str
    coded_filename: str
    cover_statement: str
    footer_label: str


def load_active_batches(root: Path) -> tuple[ActiveBatch, ActiveBatch]:
    registry_path = root / "data" / "validation_registry.json"
    registry = json.loads(registry_path.read_text(encoding="utf-8"))
    codes = registry.get("active_validation_batches")
    batches = registry.get("batches")
    if not isinstance(codes, list) or not isinstance(batches, dict):
        raise WordExportError("validation registry is missing active batches")
    loaded = tuple(_batch_from_registry(root, str(code), batches) for code in codes)
    if len(loaded) != 2:
        raise WordExportError(f"expected two active batches, found {len(loaded)}")
    by_strategy = {item.generation_strategy: item for item in loaded}
    try:
        return (by_strategy["balanced_structured"], by_strategy["resident_seed_guided"])
    except KeyError as exc:
        raise WordExportError(f"active batches are not the two current strategies: {exc}") from exc


def _batch_from_registry(root: Path, code: str, batches: dict[str, Any]) -> ActiveBatch:
    raw = batches.get(code)
    if not isinstance(raw, dict) or raw.get("status") != "active":
        raise WordExportError(f"{code} is not an active batch")
    strategy = str(raw.get("generation_strategy") or "")
    directory = root / str(raw.get("directory") or "")
    readme = _readme_fields(directory / "README.md")
    if readme.get("Batch") != code:
        raise WordExportError(f"{directory} README batch does not match {code}")
    first_id = str(raw.get("public_id_first") or "")
    last_id = str(raw.get("public_id_last") or "")
    if first_id not in readme.get("Cases", "") or last_id not in readme.get("Cases", ""):
        raise WordExportError(f"{code} README case range does not match the registry")
    case_count = int(raw.get("case_count") or 0)
    if readme.get("Number of cases") != str(case_count):
        raise WordExportError(f"{code} README case count does not match the registry")
    if strategy == "balanced_structured":
        title = "CliniProof Balanced Structured Case Set"
        simple = "CliniProof_Balanced_Case_Set.docx"
        coded = f"CliniProof_Balanced_{code}.docx"
        statement = BALANCED_STATEMENT
        footer = "CliniProof balanced case set"
    elif strategy == "resident_seed_guided":
        title = "CliniProof Resident-Seed-Guided Case Set"
        simple = "CliniProof_Seed_Guided_Case_Set.docx"
        coded = f"CliniProof_Seed_Guided_{code}.docx"
        statement = SEED_STATEMENT
        footer = "CliniProof resident-seed-guided case set"
    else:
        raise WordExportError(f"unsupported generation strategy {strategy}")
    return ActiveBatch(
        batch_code=code,
        generation_strategy=strategy,
        directory=directory,
        resident_json=directory / "resident_validation_cases.json",
        first_id=first_id,
        last_id=last_id,
        case_count=case_count,
        status_label=readme.get("Status", ""),
        approach_label=readme.get("Generation approach", ""),
        title=title,
        simple_filename=simple,
        coded_filename=coded,
        cover_statement=statement,
        footer_label=footer,
    )


def _readme_fields(path: Path) -> dict[str, str]:
    fields: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if ":" not in line or line.startswith(("#", "-", "|", " ")):
            continue
        key, value = line.split(":", 1)
        key = key.strip()
        if key in fields:
            continue
        if key in {"Batch", "Cases", "Number of cases", "Status", "Generation approach"}:
            fields[key] = value.strip().strip("`")
    return fields


def load_resident_cases(batch: ActiveBatch) -> list[dict[str, Any]]:
    payload = json.loads(batch.resident_json.read_text(encoding="utf-8"))
    if payload.get("batch_code") != batch.batch_code:
        raise WordExportError(f"{batch.resident_json} batch_code does not match the registry")
    rows = payload.get("cases")
    if not isinstance(rows, list):
        raise WordExportError(f"{batch.resident_json} has no cases list")
    cases = [row for row in rows if isinstance(row, dict)]
    ids = [str(row.get("case_id_code") or "") for row in cases]
    if ids[0] != batch.first_id or ids[-1] != batch.last_id or len(ids) != batch.case_count:
        raise WordExportError(f"{batch.batch_code} resident file range does not match the registry")
    if len(ids) != len(set(ids)):
        raise WordExportError(f"{batch.batch_code} resident file has duplicate case ids")
    return cases


def export_all(root: Path, output_dir: Path, *, export_date: date) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    balanced, seed = load_active_batches(root)
    written: dict[str, Path] = {}
    for batch in (balanced, seed):
        document = build_case_document(batch, load_resident_cases(batch), export_date)
        simple = output_dir / batch.simple_filename
        coded = output_dir / batch.coded_filename
        document.save(str(simple))
        shutil.copyfile(simple, coded)
        written[batch.simple_filename] = simple
        written[batch.coded_filename] = coded
    codebook = output_dir / "CliniProof_Resident_Validation_Codebook.docx"
    build_codebook(balanced, seed, export_date).save(str(codebook))
    written[codebook.name] = codebook
    return written


def build_case_document(
    batch: ActiveBatch, cases: list[dict[str, Any]], export_date: date
) -> WordDocument:
    document = _new_document(f"{batch.footer_label}  |  {batch.batch_code}")
    _add_title(document, batch.title)
    _add_body(document, f"Batch code: {batch.batch_code}")
    _add_body(document, f"Case range: {batch.first_id}–{batch.last_id}")
    _add_body(document, f"Number of cases: {batch.case_count}")
    _add_body(document, f"Status: {batch.status_label}")
    _add_body(document, f"Generation approach: {batch.approach_label}")
    _add_body(document, f"Date of export: {export_date.strftime('%d %B %Y')}")
    _add_body(document, batch.cover_statement)
    _add_heading(document, "Cases", 2)
    for case in cases:
        _add_body(document, str(case.get("case_id_code") or ""))
    for case in cases:
        _render_case(document, case)
    return document


def _render_case(
    document: WordDocument,
    case: dict[str, Any],
    *,
    page_break: bool = True,
    heading_text: str | None = None,
) -> None:
    case_id = str(case.get("case_id_code") or "")
    heading = _add_heading(document, heading_text or case_id, 1)
    heading.paragraph_format.page_break_before = page_break
    info = clinical(case)
    presented = presentation(case)
    _add_heading(document, "Patient overview", 2)
    _add_pairs(
        document,
        [
            ("Age", exact_text(info.get("patient_age"))),
            ("Sex/gender", exact_text(info.get("patient_gender"))),
            ("Weight", _with_unit(exact_text(info.get("weight_kg")), "kg")),
            ("Clinical setting/specialty", exact_text(info.get("specialty"))),
            ("Admission diagnosis", exact_text(info.get("admission_dx"))),
            ("Disposition", exact_text(info.get("disposition_status"))),
            ("One-liner", exact_text(info.get("one_liner"))),
        ],
    )
    _add_heading(document, "Reason for hospitalization", 2)
    symptoms = presented.get("presenting_symptoms")
    symptom_text = exact_text(symptoms) if isinstance(symptoms, list) else exact_text(symptoms)
    _add_pairs(
        document,
        [
            (
                "Chief complaint",
                exact_text(presented.get("chief_complaint") or info.get("chief_complaint")),
            ),
            ("Symptoms", symptom_text),
            ("Symptom duration", exact_text(presented.get("symptom_duration"))),
            ("Symptom course", exact_text(presented.get("symptom_course"))),
        ],
    )
    hpi = exact_text(presented.get("hpi"))
    if hpi:
        _add_heading(document, "History of present illness", 3)
        _add_body(document, hpi)
    review = exact_text(presented.get("review_of_systems"))
    if review:
        _add_pairs(document, [("Review of systems", review)])
    admission = note_text(case, "admission")
    if admission:
        _add_heading(document, "Admission note", 3)
        _add_body(document, admission)

    _add_heading(document, "Relevant medical history", 2)
    history = info.get("medical_history")
    if isinstance(history, list) and not is_missing(history):
        _add_pairs(document, [("Past medical history", exact_text(history))])
    allergies = info.get("allergies")
    if isinstance(allergies, list) and not is_missing(allergies):
        _add_pairs(document, [("Allergies", exact_text(allergies))])
    _add_record_table(
        document,
        _dict_rows(case, "CaseDiagnosis"),
        (
            ("diagnosis", "Diagnosis"),
            ("diagnosis_type", "Type"),
            ("status", "Status"),
            ("context", "Context"),
        ),
    )
    problems = _dict_rows(case, "CaseProblemList")
    if problems:
        _add_heading(document, "Problem list", 3)
        _add_record_table(
            document,
            problems,
            (
                ("problem", "Problem"),
                ("problem_type", "Type"),
                ("priority", "Priority"),
                ("status", "Status"),
            ),
        )
    _add_social(document, info)

    _add_heading(document, "Hospital course", 2)
    course = note_text(case, "hospital_course")
    if course:
        _add_body(document, course)
    elif not _dict_rows(case, "CaseIntakeOutput") and is_missing(info.get("discharge_planning")):
        _add_body(document, EMPTY_COURSE)
    io_rows = _dict_rows(case, "CaseIntakeOutput")
    if io_rows:
        _add_heading(document, "Intake and output", 3)
        _add_record_table(
            document,
            io_rows,
            (
                ("timepoint", "Timepoint"),
                ("intake_ml", "Intake (mL)"),
                ("output_ml", "Output (mL)"),
                ("net_ml", "Net (mL)"),
                ("notes", "Notes"),
            ),
        )
    _add_discharge_planning(document, info.get("discharge_planning"))

    _add_status_block(document, case, "Admission clinical status", "admission")
    _add_status_block(document, case, "Discharge / most recent clinical status", "discharge")

    _add_medication_section(document, "Home medications", rows_for_context(case, "home"))
    _add_medication_section(
        document, "Medications during hospitalization", rows_for_context(case, "inpatient")
    )
    _add_medication_section(document, "Discharge medications", rows_for_context(case, "discharge"))
    _add_other_medication_contexts(document, case)

    _add_heading(document, "Medication reconciliation", 2)
    _add_medrec(document, _dict_rows(case, "CaseMedicationReconciliation"))

    _add_heading(document, "Monitoring", 2)
    monitoring = _dict_rows(case, "CaseMonitoring")
    if monitoring:
        _add_record_table(
            document,
            monitoring,
            (
                ("parameter", "Parameter"),
                ("frequency", "Frequency"),
                ("target", "Target"),
                ("trigger_for_action", "Trigger for action"),
                ("duration", "Duration"),
                ("responsible_service", "Responsible service"),
            ),
        )
    else:
        _add_body(document, EMPTY_MONITORING)

    _add_heading(document, "Follow-up", 2)
    followup = _dict_rows(case, "CaseFollowup")
    if followup:
        _add_record_table(
            document,
            followup,
            (("item", "Item"), ("timing", "Timing"), ("with_service", "With service")),
        )
    else:
        _add_body(document, EMPTY_FOLLOWUP)

    _add_heading(document, "Discharge instructions", 2)
    instructions = _dict_rows(case, "CaseInstruction")
    if instructions:
        for row in instructions:
            category = field_text(row, "category")
            text = exact_text(row.get("instruction_text"))
            if category:
                _add_pairs(document, [("Category", category)])
            if text:
                _add_body(document, text)
    else:
        _add_body(document, EMPTY_INSTRUCTIONS)

    _add_optional_section(
        document,
        "Imaging",
        _dict_rows(case, "CaseImaging"),
        (
            ("study_type", "Study"),
            ("timepoint", "Timepoint"),
            ("body_site", "Body site"),
            ("finding", "Finding"),
        ),
    )
    _add_optional_section(
        document,
        "Procedures",
        _dict_rows(case, "CaseProcedure"),
        (
            ("procedure_name", "Procedure"),
            ("procedure_type", "Type"),
            ("time", "Time"),
            ("findings", "Findings"),
            ("complications", "Complications"),
            ("laterality", "Laterality"),
        ),
    )
    _add_optional_section(
        document,
        "Consultations",
        _dict_rows(case, "CaseConsult"),
        (
            ("service", "Service"),
            ("timepoint", "Timepoint"),
            ("assessment", "Assessment"),
            ("recommendation", "Recommendation"),
        ),
    )
    _add_optional_section(
        document,
        "Microbiology",
        _dict_rows(case, "CaseMicrobiology"),
        (
            ("timepoint", "Timepoint"),
            ("specimen", "Specimen"),
            ("test", "Test"),
            ("organism", "Organism"),
            ("result", "Result"),
            ("status", "Status"),
            ("notes", "Notes"),
        ),
    )
    _add_optional_section(
        document,
        "Devices",
        _dict_rows(case, "CaseDevice"),
        (
            ("device_type", "Device"),
            ("site", "Site"),
            ("placement_timepoint", "Placement"),
            ("status", "Status"),
            ("tip_location_or_confirmation", "Confirmation"),
            ("care_instructions", "Care instructions"),
            ("removal_plan", "Removal plan"),
        ),
    )
    _add_optional_section(
        document,
        "Return precautions",
        _dict_rows(case, "CaseReturnPrecaution"),
        (
            ("symptom", "Symptom"),
            ("reason", "Reason"),
            ("action", "Action"),
            ("severity", "Severity"),
            ("patient_instruction", "Patient instruction"),
        ),
    )
    _add_optional_section(
        document,
        "Therapy restrictions",
        _dict_rows(case, "CaseTherapyRestriction"),
        (("item", "Item"), ("category", "Category"), ("status", "Status")),
    )


def _add_status_block(
    document: WordDocument, case: dict[str, Any], title: str, timepoint: str
) -> None:
    _add_heading(document, title, 2)
    vitals = [
        row for row in _dict_rows(case, "CaseVital") if str(row.get("timepoint")) == timepoint
    ]
    labs = [row for row in _dict_rows(case, "CaseLab") if str(row.get("timepoint")) == timepoint]
    weights = [
        row for row in _dict_rows(case, "CaseWeight") if str(row.get("timepoint")) == timepoint
    ]
    _add_heading(document, "Vital signs", 3)
    if vitals:
        _add_fixed_table(
            document, ("Measure", "Value", "Unit"), _vital_rows(vitals), (3.2, 1.8, 1.8)
        )
    else:
        _add_body(document, EMPTY_VITALS)
    _add_heading(document, "Laboratory findings", 3)
    if labs:
        _add_fixed_table(
            document,
            ("Laboratory test", "Result", "Unit"),
            _lab_rows(labs),
            (4.0, 1.3, 1.5),
        )
    else:
        _add_body(document, EMPTY_LABS)
    if weights:
        _add_heading(document, "Weight", 3)
        _add_record_table(
            document,
            weights,
            (("weight_kg", "Weight (kg)"), ("dry_weight_kg", "Dry weight (kg)")),
        )


def _vital_rows(rows: list[dict[str, Any]]) -> list[tuple[str, str, str]]:
    rendered: list[tuple[str, str, str]] = []
    for row in rows:
        temp = exact_text(row.get("temp_c"))
        systolic = exact_text(row.get("bp_systolic"))
        diastolic = exact_text(row.get("bp_diastolic"))
        heart = exact_text(row.get("heart_rate"))
        respiratory = exact_text(row.get("resp_rate"))
        spo2 = exact_text(row.get("spo2_percent"))
        oxygen = exact_text(row.get("oxygen_support"))
        if temp:
            rendered.append(("Temperature", temp, "°C"))
        if systolic and diastolic:
            rendered.append(("Blood pressure", f"{systolic}/{diastolic}", "mmHg"))
        elif systolic:
            rendered.append(("Systolic blood pressure", systolic, "mmHg"))
        elif diastolic:
            rendered.append(("Diastolic blood pressure", diastolic, "mmHg"))
        if heart:
            rendered.append(("Heart rate", heart, "beats/min"))
        if respiratory:
            rendered.append(("Respiratory rate", respiratory, "breaths/min"))
        if spo2:
            rendered.append(("SpO2", spo2, "%"))
        if oxygen:
            rendered.append(("Oxygen support", oxygen, ""))
    return rendered


def _lab_rows(rows: list[dict[str, Any]]) -> list[tuple[str, str, str]]:
    rendered: list[tuple[str, str, str]] = []
    for row in rows:
        result = exact_text(row.get("value"))
        if result is None:
            result = exact_text(row.get("value_text")) or ""
        rendered.append(
            (
                exact_text(row.get("test_name")) or "",
                result,
                exact_text(row.get("unit")) or "",
            )
        )
    return rendered


def _add_medication_section(document: WordDocument, title: str, rows: list[dict[str, Any]]) -> None:
    _add_heading(document, title, 2)
    if not rows:
        _add_body(document, EMPTY_MEDICATIONS)
        return
    table_rows = [
        (
            exact_text(row.get("drug")) or exact_text(row.get("reported_name")) or "",
            exact_text(row.get("dose")) or "",
            exact_text(row.get("route")) or "",
            exact_text(row.get("frequency")) or "",
            medication_note(row),
        )
        for row in rows
    ]
    _add_fixed_table(document, MED_HEADERS, table_rows, MED_WIDTHS)


def _add_other_medication_contexts(document: WordDocument, case: dict[str, Any]) -> None:
    known = {"home", "inpatient", "discharge"}
    contexts = {
        str(row.get("context") or "")
        for row in _dict_rows(case, "CaseMedication")
        if str(row.get("context") or "") not in known
    }
    for context in sorted(contexts):
        _add_medication_section(
            document,
            f"Other medications ({context.replace('_', ' ')})",
            rows_for_context(case, context),
        )


def _add_medrec(document: WordDocument, rows: list[dict[str, Any]]) -> None:
    if not rows:
        _add_body(document, EMPTY_MEDREC)
        return
    for row in rows:
        _add_pairs(
            document,
            [
                ("Best possible medication history source", field_text(row, "bpmh_source")),
                ("Interviewer", exact_text(row.get("bpmh_interviewer"))),
                ("Date", exact_text(row.get("bpmh_date"))),
                ("Reconciliation status", field_text(row, "medrec_status")),
                ("Patient able to participate", exact_text(row.get("patient_able_to_participate"))),
                (
                    "Unverified medications present",
                    exact_text(row.get("unverified_medications_present")),
                ),
                ("Pharmacist review", exact_text(row.get("pharmacist_review"))),
                (
                    "High-alert medications identified",
                    exact_text(row.get("high_alert_meds_identified")),
                ),
                ("Discrepancy types", exact_text(row.get("discrepancy_types"))),
            ],
        )
        notes = exact_text(row.get("notes"))
        if notes:
            _add_heading(document, "Reconciliation notes", 3)
            _add_body(document, notes)


def _add_discharge_planning(document: WordDocument, planning: object) -> None:
    if not isinstance(planning, dict):
        return
    pairs = [
        ("Disposition", exact_text(planning.get("disposition"))),
        ("Disposition detail", exact_text(planning.get("disposition_detail"))),
        ("Transportation needed", exact_text(planning.get("transportation_needed"))),
        ("Barriers to discharge", exact_text(planning.get("barriers_to_discharge"))),
        ("Discharge readiness", exact_text(planning.get("discharge_readiness"))),
        ("Anticipated discharge date", exact_text(planning.get("anticipated_discharge_date"))),
        ("Home health ordered", exact_text(planning.get("home_health_ordered"))),
        ("DME needed", exact_text(planning.get("dme_needed"))),
    ]
    if any(value for _, value in pairs):
        _add_heading(document, "Discharge planning", 3)
        _add_pairs(document, pairs)


def _add_social(document: WordDocument, info: dict[str, Any]) -> None:
    support = info.get("social_support")
    support_row = support if isinstance(support, dict) else {}
    living = exact_text(support_row.get("living_situation"))
    social_context = exact_text(info.get("social_context"))
    pairs: list[tuple[str, str | None]] = []
    if living:
        pairs.append(("Living situation", living))
    elif social_context:
        pairs.append(("Living situation", social_context))
    if social_context and living and social_context != living:
        pairs.append(("Social context", social_context))
    pairs.extend(
        [
            ("Caregiver support", exact_text(support_row.get("caregiver_support"))),
            ("Transportation", exact_text(support_row.get("transportation"))),
            ("Financial barriers", exact_text(support_row.get("financial_barriers"))),
            ("Health literacy", exact_text(support_row.get("health_literacy"))),
            ("Language preference", exact_text(support_row.get("language_preference"))),
            ("Substance use", exact_text(support_row.get("substance_use"))),
            ("Advance directive", exact_text(support_row.get("advance_directive"))),
        ]
    )
    if any(value for _, value in pairs):
        _add_heading(document, "Social context", 3)
        _add_pairs(document, pairs)


def _add_optional_section(
    document: WordDocument,
    title: str,
    rows: list[dict[str, Any]],
    columns: tuple[tuple[str, str], ...],
) -> None:
    if not rows:
        return
    _add_heading(document, title, 2)
    _add_record_table(document, rows, columns)


def _add_record_table(
    document: WordDocument,
    rows: list[dict[str, Any]],
    columns: tuple[tuple[str, str], ...],
) -> None:
    used = [
        (key, label)
        for key, label in columns
        if any(field_text(row, key) is not None for row in rows)
    ]
    if not used:
        return
    rendered = [tuple(field_text(row, key) or "" for key, _label in used) for row in rows]
    headers = tuple(label for _key, label in used)
    _add_fixed_table(document, headers, rendered, None)


def _dict_rows(case: dict[str, Any], key: str) -> list[dict[str, Any]]:
    rows = case.get(key)
    if not isinstance(rows, list):
        return []
    return [row for row in rows if isinstance(row, dict)]


def _with_unit(value: str | None, unit: str) -> str | None:
    if value is None:
        return None
    return f"{value} {unit}"


def _new_document(footer_text: str) -> WordDocument:
    document = Document()
    _configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.footer_distance = Inches(0.35)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_run_font(footer.add_run(f"{footer_text}    "), size=9, color=NAVY)
    _add_page_field(footer)
    return document


def _configure_styles(document: WordDocument) -> None:
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in (
        ("Title", 22, 0, 12),
        ("Heading 1", 16, 0, 8),
        ("Heading 2", 13, 10, 4),
        ("Heading 3", 12, 8, 2),
    ):
        style = document.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def _add_title(document: WordDocument, text: str) -> None:
    paragraph = document.add_heading(text, level=0)
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.color.rgb = NAVY


def _add_heading(document: WordDocument, text: str, level: int) -> Any:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.color.rgb = NAVY
    return paragraph


def _add_body(document: WordDocument, text: str) -> None:
    for block in text.split("\n\n"):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        _set_run_font(paragraph.add_run(block), size=11)


def _add_pairs(document: WordDocument, pairs: list[tuple[str, str | None]]) -> None:
    for label, value in pairs:
        if value is None:
            continue
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.space_before = Pt(0)
        label_run = paragraph.add_run(f"{label}: ")
        _set_run_font(label_run, size=11, bold=True)
        _set_run_font(paragraph.add_run(value), size=11)


def _add_fixed_table(
    document: WordDocument,
    headers: tuple[str, ...],
    rows: Sequence[tuple[str, ...]],
    widths: tuple[float, ...] | None,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    if widths is None:
        widths = tuple(CONTENT_WIDTH / len(headers) for _header in headers)
    _set_table_width(table, sum(widths))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _write_cell(cell, header, bold=True, fill=HEADER_FILL)
        cell.width = Inches(widths[index])
    _mark_header_row(table.rows[0])
    _keep_row_together(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            _write_cell(cells[index], value, bold=False, fill=None)
            cells[index].width = Inches(widths[index])
        _keep_row_together(table.rows[-1])
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _write_cell(
    cell: Any,
    text: str,
    *,
    bold: bool,
    fill: str | None,
    size: int = 10,
) -> None:
    cell.text = text
    if fill:
        _shade(cell, fill)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.space_before = Pt(1)
        for run in paragraph.runs:
            _set_run_font(run, size=size, bold=bold)


def _set_run_font(
    run: Any, *, size: int, bold: bool = False, color: RGBColor | None = None
) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs"):
        r_fonts.set(qn(attribute), FONT)


def _shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_table_width(table: Any, width_in: float) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(width_in * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _mark_header_row(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    tr_pr.append(header)


def _keep_row_together(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run2 = paragraph.add_run()
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    run2._r.append(instruction)
    run3 = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run3._r.append(end)
    for item in (run, run2, run3):
        _set_run_font(item, size=9, color=NAVY)


def build_codebook(balanced: ActiveBatch, seed: ActiveBatch, export_date: date) -> WordDocument:
    document = _new_document("CliniProof resident validation codebook")
    _add_title(document, "CliniProof Resident Validation Codebook")
    _add_body(document, f"Date of export: {export_date.strftime('%d %B %Y')}")
    _add_body(
        document,
        "This codebook explains how to read and validate the synthetic cases. "
        "It does not identify the assessment target of any individual case.",
    )
    _add_heading(document, "Purpose", 1)
    _add_body(
        document,
        "The cases are synthetic inpatient charts created for medication-reconciliation "
        "and transition-of-care assessment. Reviewers are asked to decide whether each "
        "chart is clinically coherent and usable for the planned study.",
    )
    _add_body(
        document,
        "Validation is one review of the complete case. The same reading produces "
        "ratings C1 through C5 and one recommendation. There is no separate "
        "plausibility-only pass.",
    )
    _add_heading(document, "The two case sets", 1)
    _add_heading(document, "Balanced structured set", 2)
    _add_body(
        document,
        f"{balanced.batch_code}, {balanced.first_id} through {balanced.last_id} "
        f"({balanced.case_count} cases). Status: {balanced.status_label}.",
    )
    _add_body(
        document,
        "This set was generated from predefined structured clinical profiles to provide "
        "controlled coverage across several inpatient scenarios. The profiles specify a "
        "presentation, a medication role, a hospital course, and follow-up before any "
        "assessment problem is considered. The distribution is a study design. It is not "
        "an estimate of disease prevalence.",
    )
    _add_heading(document, "Resident-seed-guided set", 2)
    _add_body(
        document,
        f"{seed.batch_code}, {seed.first_id} through {seed.last_id} "
        f"({seed.case_count} cases). Status: {seed.status_label}.",
    )
    _add_body(
        document,
        "This set was generated from clinical archetypes abstracted from six resident-provided "
        "example cases. The generated cases are new synthetic encounters and are not copies "
        "of the resident-provided source cases.",
    )
    _add_body(
        document,
        "The two sets ask different design questions. This codebook does not rank "
        "one approach above the other.",
    )
    _add_heading(document, "How these cases connect to the CliniProof dashboard", 1)
    _add_body(
        document,
        "The same underlying case structure is intended to populate the dashboard. "
        "A resident review dashboard is planned. It is not implemented in the current software.",
    )
    _add_body(
        document,
        "What is implemented now is the case generator and a reference-terminology search "
        "service for medications, laboratory tests, diagnoses, and symptoms. "
        "There is no resident-review application in this repository.",
    )
    _add_body(
        document,
        "Residents using the planned dashboard would review parts of the synthetic chart, "
        "including "
        "clinical presentation, medical history, laboratory results, the home medication list, "
        "the inpatient medication list, the discharge medication list, monitoring, follow-up, "
        "and discharge information. These Word documents are the format for validating that "
        "clinical content before any dashboard administration.",
    )
    _add_body(
        document,
        "The planned dashboard is the presentation and interaction layer. "
        "The case generator provides the underlying clinical content. "
        "The concealed scoring reference is stored separately and is not part of these charts.",
    )
    _add_heading(document, "What to look for when reading a case", 1)
    _add_body(
        document,
        "Medication reconciliation compares the home regimen, the medications used during "
        "the hospitalization, and the discharge regimen. Also read indication, dose, route, "
        "frequency, intentional holds or stops, newly started medications, monitoring, supply, "
        "restart decisions, and follow-up.",
    )
    _add_body(
        document,
        "A clinic appointment and a laboratory monitoring task are different items. "
        "Supply is different from the drug name. A medication limited to the hospital stay "
        "is different from a medication intended to continue. Do not assume that every chart "
        "contains a medication problem.",
    )
    _add_body(
        document,
        "Types of medication-list problems that this project can assess include a medication "
        "omitted from the discharge list, a medication continued after it was meant to stop, "
        "a dose mismatch, a route mismatch, a frequency mismatch, and an unexplained "
        "therapeutic substitution.",
    )
    _add_body(
        document,
        "Types of transition-of-care gaps include required monitoring that was not arranged, "
        "a held medication without a restart plan, a supply that does not last until follow-up, "
        "hospital-only therapy still listed at discharge, a temporary substitution that was not "
        "addressed, and a pending medication decision without follow-up.",
    )
    _add_body(
        document,
        "Those are types only. This codebook does not say which problem, if any, "
        "is present in a given case.",
    )

    rubric = _add_heading(document, "Validation rubric", 1)
    rubric.paragraph_format.page_break_before = True
    _add_body(
        document,
        "Use the definitions below. They follow the current clinical-validation method and the "
        "validation rubric shipped with the case sets. Do not substitute a different scale.",
    )
    _add_heading(document, "C1 — Clinical plausibility", 2)
    _add_body(
        document,
        "Could this reasonably represent a patient encountered in the stated inpatient "
        "clinical setting? "
        "Could this hospitalization occur as charted?",
    )
    _add_body(
        document,
        "Rate these domains from the chart: presentation and demographics; fit between "
        "presentation "
        "and diagnosis; vital signs; laboratory findings; medication regimen; hospital course; "
        "consistency across the chart; discharge plan and follow-up.",
    )
    _add_body(
        document,
        "A rating of 1 means implausible. A rating of 2 means questionable and requires revision. "
        "A rating of 3 means plausible with minor concern. A rating of 4 means fully plausible.",
    )
    _add_body(
        document,
        "Any domain rated 1 or 2 must include a written explanation that identifies the specific "
        "clinical concern. Also answer whether the hospitalization could reasonably represent a "
        "patient in the stated setting. The allowed overall answers are Yes or No.",
    )
    _add_body(
        document,
        "C1 passes when all clinically relevant domains are rated 3 or 4 and the overall "
        "answer is Yes. Clinical plausibility is not the same as optimal management or "
        "complete guideline concordance.",
    )
    _add_heading(document, "C2 — Intended assessment problem", 2)
    _add_body(
        document,
        "Does the case actually contain the medication-reconciliation or "
        "transition-of-care problem "
        "it was designed to assess? Is the predetermined discrepancy actually present, and does it "
        "match the answer key?",
    )
    _add_body(
        document,
        "The blinded case documents do not contain the intended target. For formal validation, "
        "that information is provided in the clinician validation packet for the case set, "
        "separately from the chart in this Word file.",
    )
    _add_body(
        document,
        "Record Pass or Fail. C2 is a hard requirement. A written explanation is required "
        "for a failure.",
    )
    _add_heading(document, "C3 — Detectability", 2)
    _add_body(
        document,
        "Could an internal medicine resident identify and resolve the intended problem using only "
        "the clinical information provided in the case? Can an internal-medicine resident see the "
        "problem from resident-visible information and say what should change, without the chart "
        "announcing the answer?",
    )
    _add_body(
        document,
        "Consider whether there is enough evidence, whether important information is missing, "
        "whether the case is ambiguous, and whether wording or formatting gives away the answer.",
    )
    _add_body(
        document,
        "Record Pass or Fail. C3 is a hard requirement. A written explanation is required "
        "for a failure.",
    )
    _add_heading(document, "C4 — Absence of unintended problems", 2)
    _add_body(
        document,
        "Apart from the intended assessment problem, does the case contain another clinically "
        "meaningful medication-reconciliation or transition-of-care problem? Is there a second "
        "medication, monitoring, temporal, or diagnostic problem that could be an "
        "alternative answer?",
    )
    _add_body(
        document,
        "Search the whole chart for a second problem a reasonable resident could treat as "
        "the target. "
        "Record Pass or Fail. If the result is Fail, name the additional problem, the "
        "medication or "
        "clinical issue involved, and why it is clinically meaningful.",
    )
    _add_heading(document, "C5 — Expected learner difficulty", 2)
    _add_body(
        document,
        "How difficult would this case likely be for an internal medicine resident? "
        "How hard should this chart be for the intended learner?",
    )
    _add_body(
        document,
        "Provide a provisional expert estimate using Easy, Moderate, Hard, or "
        "Inappropriate / outlier. "
        "C5 is advisory. It should not by itself cause a case to fail validation. "
        "Actual difficulty should ultimately be determined from resident performance.",
    )
    _add_heading(document, "Recommendation", 1)
    _add_body(document, "After C1 through C5, record one recommendation.")
    _add_body(
        document,
        "Accept. The case is suitable for use without clinically meaningful revision.",
    )
    _add_body(
        document,
        "Revise. The case requires one or more changes before it should be used. "
        "Do not silently edit a frozen batch. A revision is a new freeze.",
    )
    _add_body(
        document,
        "Exclude. The case should not be used, even if software checks passed, because "
        "its problems "
        "cannot be reasonably corrected without substantially reconstructing it.",
    )
    _add_heading(document, "Comments", 1)
    _add_body(
        document,
        "When a rating needs explanation, name the exact clinical concern, the case section, "
        "and the medication, laboratory, or problem involved. Where a revision is recommended, "
        "give a suggested correction. Reviewers do not need to rewrite the case.",
    )
    _add_body(
        document,
        "Comments are recommendations for the study team. They are stored separately "
        "from the frozen case.",
    )
    _add_heading(document, "What this codebook does not contain", 1)
    _add_body(
        document,
        "It does not list the intended problem for any case. It does not label a case as a control "
        "or as an error-bearing chart. The validation packet, not this codebook and not the case "
        "document, is where the intended target is shown to a validator.",
    )
    _add_heading(document, "Source of these definitions", 2)
    _add_body(
        document,
        "Clinical questions and decisions follow docs/clinical_validation.md. "
        "Domain list, 1–4 plausibility scale, Pass/Fail rules, and the difficulty labels follow "
        "the validation rubric distributed with each current case set.",
    )
    return document


def main() -> None:
    root = Path(__file__).resolve().parents[2]
    output_dir = root / "exports" / "word"
    export_date = date.today()
    export_all(root, output_dir, export_date=export_date)
    from app.services.word_export_qa import write_qa_report

    write_qa_report(root, output_dir)


if __name__ == "__main__":
    main()
