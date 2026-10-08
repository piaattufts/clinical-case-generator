"""Clinician validation casebooks: one Word file per active case set.

Clinical charts come from the resident-facing case JSON. Validation references
come from the investigator answer key and the batch plan. Rubric wording is
fixed methodology text. Nothing here calls a language model or edits case data.
"""

from __future__ import annotations

import json
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from docx.document import Document as WordDocument
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.services.error_taxonomy import FAMILY_FOR_CATEGORY, IMPLEMENTABLE_CATEGORIES
from app.services.readable_packets import (
    CLINICAL_CATEGORY_LABELS,
    _format_state,
    _location_text,
)
from app.services.word_controls import (
    append_checkbox,
    append_plain_text,
    append_rich_text_field,
    finalize_word_form,
    prepare_form_document,
)
from app.services.word_export import (
    HEADER_FILL,
    ActiveBatch,
    _add_body,
    _add_fixed_table,
    _add_heading,
    _keep_row_together,
    _mark_header_row,
    _new_document,
    _render_case,
    _set_run_font,
    _set_table_width,
    _write_cell,
    load_active_batches,
    load_resident_cases,
)
from app.services.word_facts import exact_text

BANNER_FILL = "F4E8C1"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)

C1_DOMAINS = (
    "Presentation and demographics",
    "Fit between presentation and diagnosis",
    "Vital signs",
    "Laboratory findings",
    "Medication regimen",
    "Hospital course",
    "Consistency across the chart",
    "Discharge plan and follow-up",
)

FAMILY_1 = "family_1"
FAMILY_2 = "family_2"
OPTION_GAP = "\u2003\u2003\u2003"
TAG_C1_SCORE = "c1-score"
TAG_C1_RESULT = "c1-result"
TAG_C2_RESULT = "c2-result"
TAG_C3_RESULT = "c3-result"
TAG_C4_RESULT = "c4-result"
TAG_C5 = "c5-difficulty"
TAG_RECOMMENDATION = "recommendation"
TAG_REVIEWER_CODE = "reviewer-code"
TAG_REVIEW_DATE = "review-date"
TAG_INITIALS = "reviewer-initials"
TAG_CASE_DATE = "case-date"


class ValidationCasebookError(ValueError):
    """Canonical plan, answer key, and casebook reference do not agree."""


@dataclass(frozen=True)
class CasebookStats:
    batch_code: str
    first_id: str
    last_id: str
    total: int
    controls: int
    error_bearing: int
    family_1: int
    family_2: int


@dataclass(frozen=True)
class PreparedCase:
    case_id: str
    resident: dict[str, Any]
    investigator: dict[str, Any]
    plan: dict[str, Any]

    @property
    def error(self) -> dict[str, Any]:
        raw = self.investigator.get("error")
        return raw if isinstance(raw, dict) else {}

    @property
    def is_control(self) -> bool:
        status = str(self.investigator.get("control_error_status") or "")
        category = str(self.error.get("error_category") or "none")
        return status == "NO INTENTIONAL ERROR" or category == "none"


def load_prepared_cases(batch: ActiveBatch) -> list[PreparedCase]:
    residents = {str(row.get("case_id_code")): row for row in load_resident_cases(batch)}
    investigator = json.loads(
        batch.directory.joinpath("investigator_answer_key.json").read_text(encoding="utf-8")
    )
    plan = json.loads(batch.directory.joinpath("batch_plan.json").read_text(encoding="utf-8"))
    if (
        investigator.get("batch_code") != batch.batch_code
        or plan.get("batch_code") != batch.batch_code
    ):
        raise ValidationCasebookError(f"{batch.batch_code} plan or answer key batch code mismatch")
    inv_rows = investigator.get("cases")
    plan_rows = plan.get("cases")
    if not isinstance(inv_rows, list) or not isinstance(plan_rows, list):
        raise ValidationCasebookError(f"{batch.batch_code} is missing case lists")
    plan_by = {
        str(row.get("validation_case_id")): row
        for row in plan_rows
        if isinstance(row, dict)
    }
    prepared: list[PreparedCase] = []
    for row in inv_rows:
        if not isinstance(row, dict):
            continue
        case_id = str(row.get("validation_case_id") or "")
        resident = residents.get(case_id)
        assignment = plan_by.get(case_id)
        if resident is None or assignment is None:
            raise ValidationCasebookError(
                f"{case_id} is missing from the resident file or batch plan"
            )
        item = PreparedCase(case_id, resident, row, assignment)
        _require_consistent(item)
        prepared.append(item)
    expected = [str(row.get("case_id_code")) for row in load_resident_cases(batch)]
    if [item.case_id for item in prepared] != expected:
        raise ValidationCasebookError(
            f"{batch.batch_code} case order does not match the resident file"
        )
    return prepared


def _require_consistent(item: PreparedCase) -> None:
    planned = item.plan.get("error_category")
    planned_family = item.plan.get("error_family")
    planned_text = "none" if planned in (None, "") else str(planned)
    key_category = str(item.error.get("error_category") or "none")
    key_family = str(item.error.get("error_family") or "none")
    if item.is_control:
        if item.plan.get("inject_error") or planned_text not in {"none", "None"}:
            raise ValidationCasebookError(f"{item.case_id} control plan is not clean")
        if key_category != "none" or key_family != "none":
            raise ValidationCasebookError(f"{item.case_id} control answer key is not none")
        if str(item.investigator.get("control_error_status")) != "NO INTENTIONAL ERROR":
            raise ValidationCasebookError(
                f"{item.case_id} control status is not NO INTENTIONAL ERROR"
            )
        return
    if not item.plan.get("inject_error"):
        raise ValidationCasebookError(f"{item.case_id} error-bearing plan has inject_error false")
    changes = item.error.get("intentional_changes")
    change_category = None
    if isinstance(changes, list) and changes and isinstance(changes[0], dict):
        change_category = changes[0].get("error_category")
    if planned_text != key_category or change_category != key_category:
        raise ValidationCasebookError(
            f"{item.case_id} category mismatch plan={planned_text} "
            f"key={key_category} change={change_category}"
        )
    if planned_family != key_family or FAMILY_FOR_CATEGORY.get(key_category) != key_family:
        raise ValidationCasebookError(f"{item.case_id} family mismatch")
    for field_name in (
        "error_description",
        "correct_action",
        "evidence_required",
        "evidence_location",
    ):
        if not item.error.get(field_name):
            raise ValidationCasebookError(f"{item.case_id} missing {field_name}")


def stats_for(batch: ActiveBatch, cases: list[PreparedCase]) -> CasebookStats:
    counts: Counter[str] = Counter()
    controls = 0
    for item in cases:
        if item.is_control:
            controls += 1
        else:
            counts[str(item.error.get("error_family"))] += 1
    return CasebookStats(
        batch_code=batch.batch_code,
        first_id=batch.first_id,
        last_id=batch.last_id,
        total=len(cases),
        controls=controls,
        error_bearing=len(cases) - controls,
        family_1=counts[FAMILY_1],
        family_2=counts[FAMILY_2],
    )


def build_casebook(
    batch: ActiveBatch,
    cases: list[PreparedCase],
    *,
    title: str,
    filename_pattern: str,
) -> WordDocument:
    document = _new_document(title)
    prepare_form_document(document)
    _write_front_matter(document, batch, stats_for(batch, cases), title, filename_pattern)
    for item in cases:
        _start_case_section(document, title, item.case_id)
        _render_case(
            document,
            item.resident,
            page_break=False,
            heading_text=f"CASE {item.case_id}",
        )
        _write_validation(document, item)
    return document


def export_casebooks(root: Path, output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    balanced, seed = load_active_batches(root)
    written: dict[str, Path] = {}
    specs = (
        (
            balanced,
            "CliniProof Balanced Case Validation",
            "CliniProof_Balanced_Validation_Casebook.docx",
            "CliniProof_Balanced_Validation_[ReviewerCode]_[YYYY-MM-DD].docx",
        ),
        (
            seed,
            "CliniProof Resident-Seed-Guided Case Validation",
            "CliniProof_SeedGuided_Validation_Casebook.docx",
            "CliniProof_SeedGuided_Validation_[ReviewerCode]_[YYYY-MM-DD].docx",
        ),
    )
    for batch, title, filename, pattern in specs:
        cases = load_prepared_cases(batch)
        document = build_casebook(batch, cases, title=title, filename_pattern=pattern)
        path = output_dir / filename
        document.save(str(path))
        finalize_word_form(path)
        written[filename] = path
    return written


def _write_front_matter(
    document: WordDocument,
    batch: ActiveBatch,
    stats: CasebookStats,
    title: str,
    filename_pattern: str,
) -> None:
    _add_heading(document, title, 0)
    _add_body(document, f"Case set: {batch.batch_code}")
    _add_body(document, f"Cases: {batch.first_id}–{batch.last_id}")
    _add_body(document, f"Number of cases: {stats.total}")
    _add_body(
        document,
        "Purpose: Clinical validation of synthetic medication-reconciliation "
        "and transition-of-care assessment cases.",
    )
    _add_text_field(document, "Reviewer code", tag=TAG_REVIEWER_CODE, alias="Reviewer code")
    _add_text_field(document, "Review date", tag=TAG_REVIEW_DATE, alias="Review date")
    _add_body(document, "Please select one response per item unless otherwise indicated.")
    _add_body(
        document,
        "These are clinician-validation documents. They contain case-specific "
        "validation references that will not be shown to residents during the "
        "later assessment study.",
    )
    _add_heading(document, "This case set", 1)
    _add_fixed_table(
        document,
        ("Count", "Number"),
        [
            ("Cases", str(stats.total)),
            ("Clean controls", str(stats.controls)),
            ("Error-bearing cases", str(stats.error_bearing)),
            ("Family 1", str(stats.family_1)),
            ("Family 2", str(stats.family_2)),
        ],
        (4.2, 2.6),
    )
    _add_body(
        document,
        "The counts above describe the set. They do not say which case is which. "
        "Read each clinical chart before you read that case's validation reference.",
    )
    _add_heading(document, "What you are validating", 1)
    _add_body(
        document,
        "Each case is a synthetic inpatient clinical chart intended for "
        "medication-reconciliation and transition-of-care assessment. "
        "Some cases are clean controls. Some cases are error-bearing cases.",
    )
    _add_heading(document, "Clean control", 2)
    _add_body(
        document,
        "A clean-control case contains no deliberately introduced "
        "medication-reconciliation or transition-of-care discrepancy. "
        "Still look actively for unintended clinical problems.",
    )
    _add_heading(document, "Error-bearing case", 2)
    _add_body(
        document,
        "An error-bearing case begins as a clean clinical case. After that clean "
        "case passes the implemented checks, exactly one pre-specified eligible "
        "assessment discrepancy is introduced.",
    )
    _add_body(
        document,
        "Clean clinical case, then automated checks. A clean control stops there, "
        "with no planted discrepancy. An error-bearing case then receives one "
        "pre-specified discrepancy, followed by post-injection checks.",
    )
    _add_body(
        document,
        "You are being asked whether that intended design actually works clinically.",
    )
    _add_heading(document, "Family 1 — medication-regimen discrepancies", 2)
    _add_body(document, "Problems involving the medication regimen itself. Implemented categories:")
    for label in _category_labels(FAMILY_1):
        _add_body(document, f"• {label}")
    _add_heading(document, "Family 2 — transition-of-care gaps", 2)
    _add_body(
        document,
        "Problems involving the plan surrounding medication transitions. "
        "Implemented categories:",
    )
    for label in _category_labels(FAMILY_2):
        _add_body(document, f"• {label}")
    _add_body(
        document,
        "A category that is defined but not implemented is not listed here and is "
        "not used in these case sets.",
    )
    _add_heading(document, "How to review each case", 1)
    _add_body(document, "Step 1. Read the clinical case first.")
    _add_body(
        document,
        "Step 2. Complete C1 — Clinical plausibility before using the case-specific "
        "validation reference.",
    )
    _add_body(document, "Step 3. Read the Validation Reference immediately following C1.")
    _add_body(document, "Step 4. Complete C2 through C5.")
    _add_body(document, "Step 5. Select Accept, Revise, or Exclude.")
    _add_body(document, "Step 6. Add comments where required or useful.")
    _add_body(
        document,
        "C1 should reflect your impression of the patient-facing case. "
        "Do not use the validation reference to decide C1.",
    )
    _add_heading(document, "Saving your completed review", 1)
    _add_body(
        document,
        "Please save a separate copy of this document before entering your ratings. "
        "Do not overwrite the original review file.",
    )
    _add_body(document, f"Filename: {filename_pattern}")
    _add_body(
        document,
        "Use the reviewer code assigned by the study team. Do not put a full name "
        "in the filename.",
    )
    _add_heading(document, "Returning your completed validation", 1)
    _add_body(
        document,
        "Return the completed Word document to the study investigator using the "
        "agreed study communication channel.",
    )
    _add_body(
        document,
        "Reviewers do not need to edit GitHub, open a pull request, or upload their "
        "completed validation to the public repository.",
    )
    _add_body(document, "Cases in this document:")
    _add_body(document, ", ".join(f"{batch_case}" for batch_case in _case_id_span(batch)))


def _case_id_span(batch: ActiveBatch) -> list[str]:
    start = int(batch.first_id.removeprefix("VAL-"))
    end = int(batch.last_id.removeprefix("VAL-"))
    return [f"VAL-{index:03d}" for index in range(start, end + 1)]


def _category_labels(family: str) -> list[str]:
    labels = [
        CLINICAL_CATEGORY_LABELS[category]
        for category in sorted(IMPLEMENTABLE_CATEGORIES)
        if FAMILY_FOR_CATEGORY.get(category) == family and category in CLINICAL_CATEGORY_LABELS
    ]
    return labels


def _write_validation(
    document: WordDocument,
    item: PreparedCase,
    *,
    reference_plan: dict[str, Any] | None = None,
    counterpart: str | None = None,
) -> None:
    _add_heading(document, "Clinical validation", 2)
    _add_heading(document, "C1 — Clinical plausibility", 3)
    _add_body(
        document,
        "Could this reasonably represent a patient encountered in the stated "
        "inpatient clinical setting?",
    )
    _add_body(
        document,
        "A rating of 1 means implausible. A rating of 2 means questionable and "
        "requires revision. A rating of 3 means plausible with minor concern. "
        "A rating of 4 means fully plausible. Any domain rated 1 or 2 needs a "
        "written explanation. C1 passes when every clinically relevant domain "
        "is rated 3 or 4.",
    )
    _add_c1_rating_table(document)
    _add_choice_line(
        document,
        ("Pass", "Fail"),
        tag=TAG_C1_RESULT,
        alias="C1 overall result",
        prefix="Overall C1:",
    )
    _comment_box(document, "C1 reviewer comments", tag="c1-comments")
    _add_heading(document, "Validation Reference", 2)
    _banner(
        document,
        "For clinician validation only. This section would not be shown to a "
        "resident completing the later assessment.",
    )
    if item.is_control:
        _control_reference(document)
        if reference_plan:
            _add_reference_discharge_plan(document, reference_plan)
        if counterpart:
            _add_body(document, f"Matched counterpart: {counterpart}")
        _add_heading(document, "C2 — Intended assessment problem", 3)
        _add_body(
            document,
            "Does this case appropriately contain no deliberately introduced "
            "medication-reconciliation or transition-of-care problem?",
        )
    else:
        _error_reference(document, item)
        _add_heading(document, "C2 — Intended assessment problem", 3)
        _add_body(
            document,
            "Does the clinical case actually contain the intended "
            "medication-reconciliation or transition-of-care problem described "
            "above, and does it match the intended category?",
        )
    _add_choice_line(document, ("Pass", "Fail"), tag=TAG_C2_RESULT, alias="C2 result")
    _comment_box(document, "C2 reviewer comments", tag="c2-comments")
    _add_heading(document, "C3 — Detectability", 3)
    if item.is_control:
        _add_body(
            document,
            "Does the case avoid misleading cues suggesting that an error must exist?",
        )
    else:
        _add_body(
            document,
            "Could an internal-medicine resident identify and resolve the intended "
            "problem using only the clinical information provided in the "
            "patient-facing case?",
        )
        _add_body(
            document,
            "Consider whether the necessary evidence is present, whether key "
            "information is missing, whether the target is ambiguous, whether "
            "wording gives away the answer, and whether the expected clinical "
            "action can reasonably be determined.",
        )
    _add_choice_line(document, ("Pass", "Fail"), tag=TAG_C3_RESULT, alias="C3 result")
    _comment_box(document, "C3 reviewer comments", tag="c3-comments")
    _add_heading(document, "C4 — Absence of unintended competing problems", 3)
    if item.is_control:
        _add_body(
            document,
            "Does the case contain any clinically meaningful medication-reconciliation "
            "or transition-of-care problem that should not be present?",
        )
    else:
        _add_body(
            document,
            "Apart from the intended assessment problem, does the case contain another "
            "clinically meaningful medication-reconciliation or transition-of-care "
            "problem that a reasonable resident could interpret as an alternative target?",
        )
    _add_choice_line(document, ("Pass", "Fail"), tag=TAG_C4_RESULT, alias="C4 result")
    _add_heading(document, "If C4 fails", 3)
    _add_fixed_table(
        document,
        ("Field", "Reviewer entry"),
        [
            ("Medication / clinical issue", ""),
            ("Where it appears in the case", ""),
            ("Why it is clinically meaningful", ""),
        ],
        (2.8, 4.0),
    )
    _comment_box(document, "C4 reviewer comments", tag="c4-comments")
    _add_heading(document, "C5 — Expected learner difficulty", 3)
    _add_body(
        document,
        "This is an expert estimate only. Actual difficulty will ultimately be "
        "determined from resident performance.",
    )
    _add_choice_line(
        document,
        ("Easy", "Moderate", "Hard", "Inappropriate / outlier"),
        tag=TAG_C5,
        alias="C5 difficulty",
    )
    _comment_box(document, "C5 reviewer comments", tag="c5-comments", lines=3)
    _add_heading(document, "Overall recommendation", 2)
    _add_body(
        document,
        "Accept: the case is suitable for use without clinically meaningful revision.",
    )
    _add_body(
        document,
        "Revise: the case requires one or more changes before it should be used.",
    )
    _add_body(
        document,
        "Exclude: the case should not be used, because its problems cannot be "
        "reasonably corrected without substantially reconstructing it.",
    )
    _add_choice_line(
        document,
        ("Accept", "Revise", "Exclude"),
        tag=TAG_RECOMMENDATION,
        alias="Overall recommendation",
    )
    _comment_box(
        document,
        "Overall comments / suggested revisions",
        tag="overall-comments",
        lines=6,
    )
    _add_text_field(
        document,
        "Reviewer initials/code",
        tag=TAG_INITIALS,
        alias="Reviewer initials or code",
    )
    _add_text_field(document, "Date", tag=TAG_CASE_DATE, alias="Date")


def _plan_text(value: object) -> str:
    rendered = exact_text(value)
    return "" if rendered is None else rendered


def _add_reference_discharge_plan(document: WordDocument, plan: dict[str, Any]) -> None:
    """Copy the evaluator reference. Do not add doses, labs, or other facts."""
    _add_heading(document, "Hidden reference discharge plan", 3)
    _add_body(document, "For clinician validation only — not shown to residents.")
    raw_rows = plan.get("medications")
    if not isinstance(raw_rows, list):
        raw_rows = plan.get("actions")
    medications = [
        row for row in raw_rows if isinstance(row, dict)
    ] if isinstance(raw_rows, list) else []
    table_rows = [
        (
            _plan_text(row.get("medication")),
            _plan_text(row.get("action")),
            _plan_text(row.get("dose")),
            _plan_text(row.get("route")),
            _plan_text(row.get("frequency")),
            _plan_text(row.get("indication")),
            _plan_text(row.get("rationale")),
        )
        for row in medications
    ]
    _add_fixed_table(
        document,
        ("Medication", "Action", "Dose", "Route", "Frequency", "Indication", "Rationale"),
        table_rows,
        (1.3, 0.8, 0.7, 0.7, 0.9, 1.2, 1.2),
    )
    action_follow = [
        (
            _plan_text(row.get("medication")),
            _plan_text(row.get("monitoring")),
            _plan_text(row.get("follow_up")),
        )
        for row in medications
        if _plan_text(row.get("monitoring")) or _plan_text(row.get("follow_up"))
    ]
    if action_follow:
        _add_fixed_table(
            document,
            ("Medication", "Monitoring", "Follow-up"),
            action_follow,
            (2.2, 2.3, 2.3),
        )
    monitoring = plan.get("monitoring_requirements")
    follow_up = plan.get("follow_up_requirements")
    if isinstance(monitoring, list) and monitoring:
        monitor_rows = [
            (
                _plan_text(row.get("parameter")),
                _plan_text(row.get("frequency")),
                _plan_text(row.get("target")),
                _plan_text(row.get("duration")),
            )
            for row in monitoring
            if isinstance(row, dict)
        ]
        _add_fixed_table(
            document,
            ("Monitoring parameter", "Frequency", "Target", "Duration"),
            monitor_rows,
            (2.2, 1.8, 1.6, 1.2),
        )
    if isinstance(follow_up, list) and follow_up:
        follow_rows = [
            (
                _plan_text(row.get("item")),
                _plan_text(row.get("timing")),
                _plan_text(row.get("with_service")),
            )
            for row in follow_up
            if isinstance(row, dict)
        ]
        _add_fixed_table(
            document,
            ("Follow-up", "Timing", "Service"),
            follow_rows,
            (3.4, 1.6, 1.8),
        )


def _control_reference(document: WordDocument) -> None:
    _add_fixed_table(
        document,
        ("Field", "Validation reference"),
        [
            ("Case type", "Clean control"),
            ("Error family", "None"),
            ("Error category", "None"),
            ("Planted assessment discrepancy", "None"),
            (
                "Expected condition",
                "No deliberately introduced medication-reconciliation or "
                "transition-of-care discrepancy",
            ),
        ],
        (2.6, 4.2),
    )
    _add_body(
        document,
        "The reviewer should actively inspect this case for any unintended "
        "clinically meaningful discrepancy.",
    )


def _error_reference(document: WordDocument, item: PreparedCase) -> None:
    error = item.error
    category = str(error.get("error_category") or "")
    family = str(error.get("error_family") or "")
    changes = error.get("intentional_changes")
    first: dict[str, Any] = {}
    if isinstance(changes, list) and changes and isinstance(changes[0], dict):
        first = changes[0]
    if family == FAMILY_1:
        family_name = "Family 1"
    elif family == FAMILY_2:
        family_name = "Family 2"
    else:
        family_name = family
    label = CLINICAL_CATEGORY_LABELS.get(category, category)
    clean_state = _format_state(first.get("clean_expected_state"))
    injected = _format_state(first.get("injected_state"))
    original = first.get("original_clean_state")
    if isinstance(original, dict) and original:
        clean_state = f"{clean_state}\nOriginal clean state: {_format_state(original)}"
    post = first.get("post_injection_state")
    if post not in (None, "", [], {}):
        injected = f"{injected}\nPost-injection state: {_format_state(post)}"
    drug = _drug_names(error.get("trigger_meds") or error.get("affected_medication"))
    evidence = (
        f"{exact_text(error.get('evidence_required')) or ''}\n"
        f"Evidence location: {_location_text(error.get('evidence_location'))}\n"
        f"Detectability location: {_location_text(error.get('detectability_location'))}"
    )
    changed = exact_text(first.get("changed_field") or error.get("changed_field")) or ""
    _add_fixed_table(
        document,
        ("Field", "Validation reference"),
        [
            ("Case type", "Error-bearing"),
            ("Error family", family_name),
            ("Error category", f"{label} ({category})"),
            ("Intended assessment target", exact_text(error.get("error_description")) or ""),
            ("Medication involved", drug.replace("- ", "").strip()),
            ("Clean expected state", f"Changed field: {changed}\n{clean_state}"),
            ("What appears in this case", injected),
            ("Evidence available in the chart", evidence),
            ("Expected clinical action", exact_text(error.get("correct_action")) or ""),
        ],
        (2.6, 4.2),
    )


def _drug_names(value: object) -> str:
    rows = value if isinstance(value, list) else []
    names = [
        str(row.get("drug")).strip()
        for row in rows
        if isinstance(row, dict) and row.get("drug")
    ]
    return "\n".join(names) if names else "Not specified in the investigator record"


def _add_c1_rating_table(document: WordDocument) -> None:
    headers = ("Domain", "1", "2", "3", "4")
    widths = (3.6, 0.8, 0.8, 0.8, 0.8)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    _set_table_width(table, sum(widths))
    _set_grid_widths(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _write_cell(cell, header, bold=True, fill=HEADER_FILL)
        cell.width = Inches(widths[index])
        if index:
            _align_cell(cell, center=True)
    _mark_header_row(table.rows[0])
    _keep_row_together(table.rows[0])
    for domain in C1_DOMAINS:
        row = table.add_row()
        cells = row.cells
        _write_cell(cells[0], domain, bold=False, fill=None)
        cells[0].width = Inches(widths[0])
        _align_cell(cells[0], center=False)
        for index in range(1, 5):
            cell = cells[index]
            cell.text = ""
            cell.width = Inches(widths[index])
            paragraph = cell.paragraphs[0]
            paragraph.clear()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            append_checkbox(paragraph, document, tag=TAG_C1_SCORE, alias="C1 domain rating")
            _align_cell(cell, center=True)
        _set_row_height(row, 420)
        _keep_row_together(row)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_choice_line(
    document: WordDocument,
    labels: tuple[str, ...],
    *,
    tag: str,
    alias: str,
    prefix: str = "",
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    if prefix:
        _set_run_font(paragraph.add_run(f"{prefix}  "), size=11)
    for index, label in enumerate(labels):
        append_checkbox(paragraph, document, tag=tag, alias=alias)
        _set_run_font(paragraph.add_run(f"\u00A0\u00A0{label}"), size=11)
        if index != len(labels) - 1:
            _set_run_font(paragraph.add_run(OPTION_GAP), size=11)


def _add_text_field(document: WordDocument, label: str, *, tag: str, alias: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    _set_run_font(paragraph.add_run(f"{label}: "), size=11)
    append_plain_text(paragraph, document, tag=tag, alias=alias)


def _set_grid_widths(table: Any, widths: tuple[float, ...]) -> None:
    columns = table._tbl.tblGrid.gridCol_lst
    if len(columns) != len(widths):
        raise ValidationCasebookError("C1 rating table column count does not match its widths")
    for column, width in zip(columns, widths, strict=True):
        column.w = Inches(width)


def _align_cell(cell: Any, *, center: bool) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:vAlign"))
    if existing is not None:
        tc_pr.remove(existing)
    align = OxmlElement("w:vAlign")
    align.set(qn("w:val"), "center")
    tc_pr.append(align)
    if center:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _comment_box(document: WordDocument, label: str, *, tag: str, lines: int = 4) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _set_run_font(paragraph.add_run(label), size=11, bold=True)
    append_rich_text_field(cell, document, tag=tag, alias=label, lines=lines)
    _set_row_height(table.rows[0], 360 * lines)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def _banner(document: WordDocument, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = text
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), BANNER_FILL)
    cell._tc.get_or_add_tcPr().append(shading)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            _set_run_font(run, size=11, bold=True, color=NAVY)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def _set_row_height(row: Any, twips: int) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(twips))
    height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(height)
    _keep_row_together(row)


def _start_case_section(
    document: WordDocument,
    title: str,
    case_id: str,
    *,
    header_text: str | None = None,
) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = True
    paragraph = section.header.paragraphs[0]
    paragraph.text = ""
    shown = header_text if header_text is not None else f"{title}    {case_id}"
    _set_run_font(paragraph.add_run(shown), size=9, color=NAVY)


def main() -> None:
    root = Path(__file__).resolve().parents[2]
    export_casebooks(root, root / "docs" / "resident_review_package" / "files")
    _ = date.today()


if __name__ == "__main__":
    main()
