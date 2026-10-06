"""Clinician-review packages for the recovered VAL-801–VAL-824 clean charts.

Set A revises six cases from earlier clinician feedback. Nine other charts are
corrected from the recovered clean source. Charts that pass the evidence audit
join the untouched ready cases in the fresh-review package. Charts that still
fail stay in a held audit and are not placed in a clinician-review document.
Historical frozen VAL files are not overwritten.
"""

from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document

from app.services.word_controls import (
    append_checkbox,
    append_rich_text_field,
    finalize_word_form,
    prepare_form_document,
)

SOURCE_DIR = Path(__file__).resolve().parents[2] / "exports" / "clean_balanced_seed_set"
FROZEN_RESIDENT = (
    Path(__file__).resolve().parents[2]
    / "data"
    / "case_sets"
    / "seed_guided"
    / "resident_validation_cases.json"
)

REVISED_IDS = (
    "VAL-801",
    "VAL-802",
    "VAL-803",
    "VAL-805",
    "VAL-809",
    "VAL-813",
)
REMAINING_IDS = (
    "VAL-804",
    "VAL-806",
    "VAL-807",
    "VAL-808",
    "VAL-810",
    "VAL-811",
    "VAL-812",
    "VAL-814",
    "VAL-815",
    "VAL-816",
    "VAL-817",
    "VAL-818",
    "VAL-819",
    "VAL-820",
    "VAL-821",
    "VAL-822",
    "VAL-823",
    "VAL-824",
)
FRESH_UNTOUCHED_IDS = (
    "VAL-804",
    "VAL-810",
    "VAL-811",
    "VAL-812",
    "VAL-818",
    "VAL-819",
    "VAL-822",
    "VAL-823",
    "VAL-824",
)
REPAIR_IDS = (
    "VAL-806",
    "VAL-807",
    "VAL-808",
    "VAL-814",
    "VAL-815",
    "VAL-816",
    "VAL-817",
    "VAL-820",
    "VAL-821",
)

Revision = tuple[dict[str, Any], dict[str, Any], list[dict[str, Any]], str]

SUFFICIENT = "SUFFICIENT_EVIDENCE"
WEAK = "WEAK_EVIDENCE"
HIDDEN = "HIDDEN_ANSWER_DEPENDENCY"
INCONSISTENT = "CLINICALLY_INCONSISTENT"
READY_REVIEW = "READY_FOR_CLINICIAN_REVIEW"
NOT_READY = "NOT_READY"
FRESH = "READY_FOR_FRESH_REVIEW"
NEEDS_FIX = "NEEDS_PRE_REVIEW_FIX"

_LEAKS = (
    "reference_discharge_plan",
    "exactly as listed",
    "clean_expected_state",
    "error_category",
    "error_family",
    "planted",
    "f1_",
    "f2_",
)


def write_ko_review_sets(
    revised_dir: Path,
    fresh_dir: Path,
    held_dir: Path,
) -> dict[str, Any]:
    """Write the revised, fresh-review, and held packages."""
    revised_dir.mkdir(parents=True, exist_ok=True)
    fresh_dir.mkdir(parents=True, exist_ok=True)
    held_dir.mkdir(parents=True, exist_ok=True)
    revised_rows = [_write_revised(case_id, revised_dir) for case_id in REVISED_IDS]
    untouched_rows = [_write_remaining(case_id, fresh_dir) for case_id in FRESH_UNTOUCHED_IDS]
    fresh_rows = list(untouched_rows)
    held_rows: list[dict[str, Any]] = []
    repair_status: dict[str, str] = {}
    for case_id in REPAIR_IDS:
        row = _build_repaired(case_id)
        repair_status[case_id] = row["status"]
        if row["status"] == READY_REVIEW:
            _dump(fresh_dir, case_id, row["resident"], row["evaluator"])
            fresh_rows.append(row)
        else:
            _dump(held_dir, case_id, row["resident"], row["evaluator"])
            held_rows.append(row)
    fresh_rows.sort(key=lambda item: item["case_id"])
    _purge_case_files(fresh_dir, {row["case_id"] for row in fresh_rows})
    _purge_case_files(held_dir, {row["case_id"] for row in held_rows})
    for stale in held_dir.glob("*.docx"):
        stale.unlink()
    _write_docx(
        revised_dir / "KO_REVISED_CASES_REVIEW.docx",
        "Revised cases for clinician review",
        revised_rows,
    )
    _write_docx(
        fresh_dir / "KO_CLEAN_CASES_REVIEW.docx",
        "Clean cases ready for clinician review",
        fresh_rows,
        notice=(
            "These are clean cases ready for clinician review. "
            "Ready for clinician review does not mean clinically validated."
        ),
    )
    (revised_dir / "AUDIT.md").write_text(_revised_audit(revised_rows), encoding="utf-8")
    (fresh_dir / "AUDIT.md").write_text(_fresh_audit(fresh_rows), encoding="utf-8")
    held_text = _held_audit(held_rows)
    (held_dir / "KO_HELD_CASES_AUDIT.md").write_text(held_text, encoding="utf-8")
    (held_dir / "AUDIT.md").write_text(held_text, encoding="utf-8")
    return {
        "revised": [row["case_id"] for row in revised_rows],
        "fresh": [row["case_id"] for row in fresh_rows],
        "held": [row["case_id"] for row in held_rows],
        "remaining": [row["case_id"] for row in fresh_rows],
        "revised_status": {row["case_id"]: row["status"] for row in revised_rows},
        "fresh_status": {row["case_id"]: row["status"] for row in fresh_rows},
        "held_status": {row["case_id"]: row["status"] for row in held_rows},
        "repair_status": repair_status,
        "remaining_status": {row["case_id"]: row["status"] for row in fresh_rows},
    }


def _purge_case_files(directory: Path, keep_ids: set[str]) -> None:
    for path in directory.glob("VAL-*_*.json"):
        case_id = path.name.split("_", 1)[0]
        if case_id not in keep_ids:
            path.unlink()


def _load_pair(case_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
    resident = json.loads((SOURCE_DIR / f"{case_id}_resident.json").read_text(encoding="utf-8"))
    evaluator = json.loads(
        (SOURCE_DIR / f"{case_id}_evaluator.json").read_text(encoding="utf-8")
    )
    return resident, evaluator


def _write_revised(case_id: str, directory: Path) -> dict[str, Any]:
    resident, _evaluator = _load_pair(case_id)
    revised, reference, trace, summary = _REVISIONS[case_id](deepcopy(resident))
    _strip_resident_answers(revised)
    classes = [item["evidence_class"] for item in trace]
    status = READY_REVIEW
    if any(item in {HIDDEN, INCONSISTENT} for item in classes):
        status = NOT_READY
    evaluator = deepcopy(revised)
    evaluator["reference_discharge_plan"] = reference
    evaluator["evidence_trace"] = trace
    evaluator["clinician_review_status"] = status
    evaluator["revision_summary"] = summary
    _dump(directory, case_id, revised, evaluator)
    return {
        "case_id": case_id,
        "resident": revised,
        "evaluator": evaluator,
        "summary": summary,
        "status": status,
        "classes": classes,
    }


def _write_remaining(case_id: str, directory: Path) -> dict[str, Any]:
    resident, evaluator = _load_pair(case_id)
    resident = deepcopy(resident)
    evaluator = deepcopy(evaluator)
    _strip_resident_answers(resident)
    finding = _precheck(resident, evaluator)
    status = _remaining_status(finding)
    evaluator.pop("reference_discharge_plan", None)
    # Keep the recovered reference; do not rebuild it.
    source_eval = json.loads(
        (SOURCE_DIR / f"{case_id}_evaluator.json").read_text(encoding="utf-8")
    )
    evaluator = deepcopy(resident)
    evaluator["reference_discharge_plan"] = source_eval["reference_discharge_plan"]
    evaluator["clinician_review_status"] = status
    evaluator["precheck_finding"] = finding
    _dump(directory, case_id, resident, evaluator)
    return {
        "case_id": case_id,
        "resident": resident,
        "evaluator": evaluator,
        "finding": finding,
        "status": status,
        "summary": finding,
        "classes": [],
    }


def _dump(
    directory: Path,
    case_id: str,
    resident: dict[str, Any],
    evaluator: dict[str, Any],
) -> None:
    (directory / f"{case_id}_resident.json").write_text(
        json.dumps(resident, indent=2) + "\n",
        encoding="utf-8",
    )
    (directory / f"{case_id}_evaluator.json").write_text(
        json.dumps(evaluator, indent=2) + "\n",
        encoding="utf-8",
    )


def _strip_resident_answers(chart: dict[str, Any]) -> None:
    chart.pop("reference_discharge_plan", None)
    chart.pop("evidence_trace", None)
    chart.pop("clinician_review_status", None)
    chart.pop("revision_summary", None)
    chart.pop("precheck_finding", None)
    chart["CaseMedication"] = [
        row for row in chart.get("CaseMedication") or [] if row.get("context") != "discharge"
    ]
    chart["CaseInstruction"] = [
        row
        for row in chart.get("CaseInstruction") or []
        if "exactly as listed" not in (row.get("instruction_text") or "").casefold()
        and "was stopped during this admission"
        not in (row.get("instruction_text") or "").casefold()
    ]


def _precheck(resident: dict[str, Any], evaluator: dict[str, Any]) -> str:
    """Flag objective defects. Do not rewrite the chart to make them pass."""
    findings: list[str] = []
    blob = json.dumps(resident).casefold()
    for needle in _LEAKS:
        if needle in blob:
            findings.append(f"answer leakage:{needle}")
    if "reference_discharge_plan" not in evaluator:
        findings.append("unsupported:evaluator is missing the hidden reference plan")
    if "in this profile" in blob or "the discharge list is the intended" in blob:
        findings.append("answer leakage:resident narrative states an intended discharge list")
    if "resume when holding" in blob:
        findings.append(
            "answer leakage:an instruction directs the resident to resume a held medicine"
        )
    diagnoses = " ".join(
        (item.get("diagnosis") or "") for item in resident.get("CaseDiagnosis") or []
    ).casefold()
    admission = (resident.get("ClinicalCase") or {}).get("admission_dx") or ""
    medications = resident.get("CaseMedication") or []
    _precheck_stops(findings, medications, diagnoses)
    _precheck_volume(findings, resident, diagnoses)
    _precheck_regimen_claims(findings, blob, medications)
    _precheck_monitoring(findings, resident, medications)
    if "cytomegalo" in admission.casefold() and "prophylaxis" not in blob:
        if any(
            "valganciclovir" in (row.get("drug") or "").casefold() and row.get("context") == "home"
            for row in medications
        ):
            findings.append(
                "contradictory timeline:valganciclovir is already a home medicine while the "
                "admission is for CMV disease, and the chart does not describe prophylaxis"
            )
    if not findings:
        return "none"
    return "; ".join(dict.fromkeys(findings))


def _precheck_stops(
    findings: list[str],
    medications: list[dict[str, Any]],
    diagnoses: str,
) -> None:
    for row in medications:
        reason = (row.get("held_reason") or "").casefold()
        if reason != "stopped during this admission.":
            continue
        if any(token in diagnoses for token in ("bleed", "hemorrhage")):
            continue
        findings.append(f"unexplained stop:{row.get('drug')} has no clinical reason")


def _precheck_volume(
    findings: list[str],
    resident: dict[str, Any],
    diagnoses: str,
) -> None:
    if "heart failure" not in diagnoses:
        return
    for weight in resident.get("CaseWeight") or []:
        if not weight.get("dry_weight_kg") or weight.get("weight_kg") is None:
            continue
        actual = float(weight["weight_kg"])
        dry = float(weight["dry_weight_kg"])
        if actual > dry + 2:
            findings.append(
                f"inconsistent:discharge weight {actual:.0f} kg remains above "
                f"dry weight {dry:.0f} kg"
            )


def _named(
    medications: list[dict[str, Any]],
    token: str,
    context: str,
) -> dict[str, Any] | None:
    return next(
        (
            row
            for row in medications
            if token in (row.get("drug") or "").casefold() and row.get("context") == context
        ),
        None,
    )


def _same_regimen(left: dict[str, Any] | None, right: dict[str, Any] | None) -> bool:
    if left is None or right is None:
        return False
    return all(left.get(key) == right.get(key) for key in ("dose", "route", "frequency"))


def _precheck_regimen_claims(
    findings: list[str],
    blob: str,
    medications: list[dict[str, Any]],
) -> None:
    if "potassium repletion" in blob and not any(
        "potassium" in (row.get("drug") or "").casefold() for row in medications
    ):
        findings.append(
            "unsupported:hospital course describes potassium repletion but no potassium "
            "medication is listed"
        )
    if "diuretic plan was adjusted" in blob:
        if _same_regimen(
            _named(medications, "furosemide", "home"),
            _named(medications, "furosemide", "inpatient"),
        ):
            findings.append(
                "contradictory timeline:narrative says the diuretic plan was adjusted but "
                "the home and inpatient loop-diuretic regimens match"
            )
    if "tacrolimus" in blob and "dose adjustment" in blob:
        if _same_regimen(
            _named(medications, "tacrolimus", "home"),
            _named(medications, "tacrolimus", "inpatient"),
        ):
            findings.append(
                "contradictory timeline:narrative describes a tacrolimus dose adjustment "
                "but the home and inpatient doses match"
            )
    if "enoxaparin" in blob and _named(medications, "enoxaparin", "inpatient") is None:
        findings.append(
            "contradictory timeline:narrative describes enoxaparin but it is not on "
            "the medication list"
        )
    if "prior baseline" in blob and "baseline of" not in blob:
        findings.append(
            "unsupported:a restart instruction refers to a creatinine baseline that is "
            "not in the chart"
        )
    enoxaparin = _named(medications, "enoxaparin", "inpatient")
    if (
        enoxaparin is not None
        and "fibrillation" in (enoxaparin.get("indication") or "").casefold()
        and "prophylaxis" in blob
    ):
        findings.append(
            "indication mismatch:enoxaparin is labeled for atrial fibrillation while the "
            "narrative describes venous-thromboembolism prophylaxis"
        )


def _precheck_monitoring(
    findings: list[str],
    resident: dict[str, Any],
    medications: list[dict[str, Any]],
) -> None:
    warfarin = any("warfarin" in (row.get("drug") or "").casefold() for row in medications)
    inr_lab = any(
        "inr" in (row.get("test_name") or "").casefold() for row in resident.get("CaseLab") or []
    )
    inr_task = any(
        "inr" in (row.get("parameter") or "").casefold()
        for row in resident.get("CaseMonitoring") or []
    )
    if warfarin and not inr_lab:
        findings.append("inconsistent monitoring:warfarin is present without an INR result")
    if inr_task and not warfarin:
        findings.append("inconsistent monitoring:INR monitoring is present without warfarin")


def _remaining_status(finding: str) -> str:
    if finding == "none":
        return FRESH
    parts = [item.strip() for item in finding.split(";")]
    serious = (
        "inconsistent:",
        "contradictory timeline:",
        "unexplained",
        "inconsistent monitoring:",
    )
    if any(part.startswith(serious) for part in parts):
        return INCONSISTENT
    return NEEDS_FIX


def _set_text(chart: dict[str, Any], *, hpi: str, admission: str, course: str) -> None:
    chart["ClinicalCase"]["presentation"]["hpi"] = hpi
    chart["ClinicalCase"]["chief_complaint"] = chart["ClinicalCase"].get("chief_complaint")
    for note in chart.get("CaseNote") or []:
        if note.get("note_type") == "admission":
            note["note_text"] = admission
        elif note.get("note_type") == "hospital_course":
            note["note_text"] = course


def _meds(chart: dict[str, Any], drug_prefix: str) -> list[dict[str, Any]]:
    return [
        row
        for row in chart.get("CaseMedication") or []
        if (row.get("drug") or "").startswith(drug_prefix)
    ]


def _decision(
    chart: dict[str, Any],
    *,
    medication: str,
    action: str,
    indication: str,
    rationale: str,
    diagnosis: str,
    event: str,
    labs: str,
    vitals: str,
    history: str,
    location: str,
    phrases: list[str],
    dose: str | None = None,
    route: str | None = None,
    frequency: str | None = None,
    duration: str | None = None,
    monitoring: str | None = None,
    weak: bool = False,
) -> tuple[dict[str, Any], dict[str, Any]]:
    source: dict[str, Any] = next(
        (
            row
            for row in chart.get("CaseMedication") or []
            if row.get("drug") == medication and row.get("context") in {"inpatient", "home"}
        ),
        {},
    )
    plan = {
        "medication": medication,
        "action": action,
        "dose": dose if dose is not None else source.get("dose"),
        "route": route if route is not None else source.get("route"),
        "frequency": frequency if frequency is not None else source.get("frequency"),
        "duration": duration,
        "indication": indication,
        "rationale": rationale,
        "monitoring": monitoring,
        "required": action in {"continue", "start", "change", "restart"},
    }
    visible = json.dumps(
        {key: value for key, value in chart.items() if key != "reference_discharge_plan"}
    ).casefold()
    missing = [phrase for phrase in phrases if phrase.casefold() not in visible]
    if missing:
        evidence_class = HIDDEN
    elif weak:
        evidence_class = WEAK
    else:
        evidence_class = SUFFICIENT
    trace = {
        "medication": medication,
        "action": action,
        "indication": indication,
        "supporting_diagnosis": diagnosis,
        "supporting_hospital_event": event,
        "supporting_labs": labs,
        "supporting_vitals": vitals,
        "supporting_medication_history": history,
        "rationale": rationale,
        "resident_visible_evidence_location": location,
        "evidence_class": evidence_class,
    }
    return plan, trace


def _plan_bundle(
    plans: list[dict[str, Any]],
    followups: list[dict[str, Any]],
    monitoring: list[dict[str, Any]] | None = None,
    alternatives: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    return {
        "medications": plans,
        "medications_to_stop": [
            item["medication"] for item in plans if item["action"] == "stop"
        ],
        "acceptable_alternatives": alternatives or [],
        "contraindications": [],
        "monitoring_requirements": monitoring or [],
        "follow_up_requirements": [
            {
                "item": item.get("item"),
                "timing": item.get("timing"),
                "with_service": item.get("with_service"),
            }
            for item in followups
        ],
    }


def _revise_801(chart: dict[str, Any]) -> Revision:
    hpi = (
        "A 75-year-old man with diabetes, hypertension, and hyperlipidemia is admitted "
        "with confusion and fatigue for several days. He had been eating and drinking poorly. "
        "Home medicines are lisinopril, atorvastatin, metformin, and ibuprofen as needed for "
        "knee pain. There is no gastrointestinal bleeding, melena, or recent kidney injury."
    )
    admission = (
        "Admission note. Confusion and fatigue for several days, with poor oral intake. "
        "Mucous membranes are dry, and the orthostatic systolic pressure falls 20 mmHg from "
        "the recorded supine pressure of 136 mmHg. Collateral history later confirmed the "
        "home medication list, including ibuprofen for knee pain. No bleeding symptoms."
    )
    course = (
        "Delirium was attributed to dehydration from poor oral intake. Intravenous fluids "
        "and nutrition were given. Glucose fell from 163 mg/dL to 103 mg/dL as eating resumed; "
        "this was not a hyperosmolar state. Creatinine stayed 1.1 mg/dL then 1.0 mg/dL. "
        "Confusion cleared as intake improved. Knee pain was still present. No bleeding was "
        "found, and kidney function stayed stable."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for row in _meds(chart, "ibuprofen"):
        row["status"] = "home" if row.get("context") == "home" else "active"
        row["held_reason"] = None
        row["indication"] = "knee pain; not a treatment for the admission diagnosis"
    consults = chart.get("CaseConsult") or []
    if consults:
        consults[0]["assessment"] = (
            "Delirium improved after rehydration and the glucose fell from 163 mg/dL to 103 mg/dL."
        )
        consults[0]["recommendation"] = (
            "The collateral medication list was verified. No bleeding or kidney injury was found."
        )
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    specs: list[tuple[str, str, str, str, str, str, str, str, str, list[str]]] = [
        (
            "atorvastatin 40 MG Oral Tablet",
            "continue",
            "Mixed hyperlipidemia",
            "Hyperlipidemia is active and there is no statin intolerance.",
            "atorvastatin was confirmed on the collateral home list",
            "none required for this continuation",
            "none",
            "Home atorvastatin for mixed hyperlipidemia",
            "diagnoses; home medications",
            ["mixed hyperlipidemia", "atorvastatin"],
        ),
        (
            "lisinopril 10 MG Oral Tablet",
            "continue",
            "Essential (primary) hypertension",
            "Blood pressure remained adequate and creatinine did not rise.",
            "Confusion cleared as intake and glucose improved.",
            "creatinine stayed 1.1 mg/dL then 1.0 mg/dL",
            "discharge systolic blood pressure 121 mmHg",
            "Home lisinopril",
            "diagnoses; labs; vitals",
            ["lisinopril", "1.0", "hypertension"],
        ),
        (
            "metformin hydrochloride 500 MG Oral Tablet",
            "continue",
            "Type 2 diabetes mellitus without complications",
            "Glucose improved and kidney function remained stable.",
            "Glucose improved with fluids and nutrition.",
            "glucose fell from 163 mg/dL to 103 mg/dL",
            "none",
            "Home metformin",
            "diagnoses; labs",
            ["metformin", "103", "diabetes"],
        ),
        (
            "ibuprofen 400 MG Oral Tablet",
            "continue",
            "knee pain; not a treatment for the admission diagnosis",
            "No bleeding and stable creatinine, so the as-needed analgesic remains.",
            "no gastrointestinal bleeding",
            "creatinine 1.0 mg/dL",
            "none",
            "Home ibuprofen for knee pain",
            "history of present illness; hospital course; labs",
            ["ibuprofen", "knee pain", "no gastrointestinal bleeding"],
        ),
    ]
    for spec in specs:
        plan, trace = _decision(
            chart,
            medication=spec[0],
            action=spec[1],
            indication=spec[2],
            rationale=spec[3],
            diagnosis=spec[2],
            event=spec[4],
            labs=spec[5],
            vitals=spec[6],
            history=spec[7],
            location=spec[8],
            phrases=spec[9],
            duration=None,
        )
        plans.append(plan)
        traces.append(trace)
    reference = _plan_bundle(plans, chart.get("CaseFollowup") or [])
    summary = (
        "Delirium is attributed to dehydration from poor oral intake, with dry mucous "
        "membranes and an orthostatic drop. Glucose fell from 163 to 103 mg/dL as intake "
        "resumed and is not treated as the cause. Ibuprofen is continued for knee pain "
        "because the chart shows no bleeding and creatinine stays 1.1 then 1.0 mg/dL."
    )
    return chart, reference, traces, summary


def _add_lab(
    chart: dict[str, Any],
    *,
    lab_id: str,
    timepoint: str,
    test_name: str,
    value: float,
    unit: str,
) -> None:
    chart.setdefault("CaseLab", []).append(
        {
            "lab_id": lab_id,
            "case_id": chart["case_id_code"],
            "timepoint": timepoint,
            "test_name": test_name,
            "value": value,
            "value_text": None,
            "unit": unit,
            "status": "final",
            "source_reference": None,
        }
    )


def _revise_802(chart: dict[str, Any]) -> Revision:
    _add_lab(
        chart,
        lab_id="LAB-VAL802-K1",
        timepoint="admission",
        test_name="Potassium [Moles/volume] in Serum or Plasma",
        value=4.2,
        unit="mmol/L",
    )
    _add_lab(
        chart,
        lab_id="LAB-VAL802-K2",
        timepoint="discharge",
        test_name="Potassium [Moles/volume] in Serum or Plasma",
        value=4.1,
        unit="mmol/L",
    )
    hpi = (
        "An 81-year-old woman with hypertension and hyperlipidemia is admitted with two days "
        "of confusion. She had been drinking poorly. A clinic creatinine six weeks earlier was "
        "1.2 mg/dL. Home medicines are lisinopril and atorvastatin. There is no focal weakness."
    )
    admission = (
        "Admission note. Confusion for two days with poor oral intake. Outpatient creatinine "
        "was 1.2 mg/dL six weeks earlier. Home lisinopril and atorvastatin were confirmed "
        "from a collateral list."
    )
    course = (
        "Delirium was attributed to dehydration from poor oral intake. Intravenous fluids "
        "were given. Creatinine was 1.3 mg/dL on admission and 1.2 mg/dL at discharge, matching "
        "the recent baseline of 1.2 mg/dL. Potassium was 4.2 mmol/L then 4.1 mmol/L. Systolic "
        "blood pressure was 138 mmHg on admission and 124 mmHg at discharge. Confusion cleared "
        "after rehydration. Kidney function and potassium were back at the recent baseline, "
        "and the discharge systolic pressure was 124 mmHg."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    plans = []
    traces = []
    for medication, indication, phrases, rationale, labs, vitals in (
        (
            "lisinopril 10 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["lisinopril", "1.2 mg/dl", "4.1", "124"],
            "Creatinine returned to the documented baseline and potassium "
            "and blood pressure were acceptable.",
            "creatinine 1.3 then 1.2 mg/dL; potassium 4.2 then 4.1 mmol/L",
            "systolic blood pressure 124 mmHg at discharge",
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
            "Hyperlipidemia remains active and no statin adverse effect is described.",
            "none required",
            "none",
        ),
    ):
        plan, trace = _decision(
            chart,
            medication=medication,
            action="continue",
            indication=indication,
            rationale=rationale,
            diagnosis=indication,
            event="Confusion cleared after rehydration.",
            labs=labs,
            vitals=vitals,
            history=f"Home {medication.split()[0]}",
            location="hospital course; labs; vitals; diagnoses",
            phrases=phrases,
        )
        plans.append(plan)
        traces.append(trace)
    summary = (
        "Delirium is attributed to poor intake and dehydration. Creatinine is interpreted "
        "against a stated baseline of 1.2 mg/dL, with potassium 4.1 mmol/L and discharge "
        "systolic pressure 124 mmHg. Lisinopril and atorvastatin stay on the reference "
        "plan. Atorvastatin was already on the recovered chart and is not re-omitted."
    )
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or []), traces, summary


def _revise_803(chart: dict[str, Any]) -> Revision:
    sodium = "Sodium [Moles/volume] in Serum or Plasma"
    _add_lab(
        chart,
        lab_id="LAB-VAL803-NA1",
        timepoint="admission",
        test_name=sodium,
        value=128,
        unit="mmol/L",
    )
    _add_lab(
        chart,
        lab_id="LAB-VAL803-NA2",
        timepoint="discharge",
        test_name=sodium,
        value=135,
        unit="mmol/L",
    )
    hpi = (
        "A 71-year-old man with hypertension, diabetes, and hyperlipidemia is admitted with "
        "one week of fatigue and confusion. Home medicines include hydrochlorothiazide, "
        "lisinopril, atorvastatin, and metformin. He has been drinking more water than usual "
        "and has no focal weakness."
    )
    admission = (
        "Admission note. Subacute confusion and fatigue. Hydrochlorothiazide is on the home "
        "list. Serum sodium on arrival is 128 mmol/L."
    )
    course = (
        "Delirium was attributed to thiazide-associated hyponatremia. Hydrochlorothiazide was "
        "held because the serum sodium was 128 mmol/L. Isotonic fluid was given. Sodium rose "
        "to 135 mmol/L and confusion cleared. Creatinine was 1.0 mg/dL then 1.2 mg/dL and "
        "potassium was 4.4 mmol/L then 4.2 mmol/L."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for row in _meds(chart, "hydrochlorothiazide"):
        row["status"] = "held"
        row["held_reason"] = (
            "Serum sodium 128 mmol/L on admission while taking hydrochlorothiazide."
        )
    plans = []
    traces = []
    stop_plan, stop_trace = _decision(
        chart,
        medication="hydrochlorothiazide 25 MG Oral Tablet",
        action="stop",
        indication="Essential (primary) hypertension",
        rationale="Sodium was 128 mmol/L on the thiazide and rose to 135 mmol/L after it was held.",
        diagnosis="Essential (primary) hypertension",
        event="Hydrochlorothiazide was held because the serum sodium was 128 mmol/L.",
        labs="sodium 128 mmol/L then 135 mmol/L",
        vitals="none",
        history="Home hydrochlorothiazide",
        location="hospital course; labs",
        phrases=["128 mmol/l", "135 mmol/l", "hydrochlorothiazide"],
    )
    plans.append(stop_plan)
    traces.append(stop_trace)
    for medication, indication, phrases in (
        (
            "lisinopril 10 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["lisinopril", "4.2", "1.2"],
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
        ),
        (
            "metformin hydrochloride 500 MG Oral Tablet",
            "Type 2 diabetes mellitus without complications",
            ["metformin", "diabetes"],
        ),
    ):
        plan, trace = _decision(
            chart,
            medication=medication,
            action="continue",
            indication=indication,
            rationale=(
                "This medicine was not implicated in the hyponatremia, and glucose, "
                "creatinine, and potassium remained acceptable."
            ),
            diagnosis=indication,
            event="Confusion cleared as the sodium rose to 135 mmol/L.",
            labs="creatinine 1.2 mg/dL; potassium 4.2 mmol/L",
            vitals="none",
            history=f"Home {medication.split()[0]}",
            location="diagnoses; hospital course; labs",
            phrases=phrases,
            duration="30 days",
        )
        plans.append(plan)
        traces.append(trace)
    summary = (
        "Delirium is attributed to thiazide-associated hyponatremia (sodium 128 mmol/L, "
        "later 135 mmol/L). Hydrochlorothiazide is stopped. Lisinopril supply stays the "
        "recovered 30-day clean value, not the old 7-day injected supply."
    )
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or []), traces, summary


def _revise_805(chart: dict[str, Any]) -> Revision:
    chart["CaseWeight"] = [
        {
            "weight_id": "WT-VAL805-001",
            "case_id": "VAL-805",
            "timepoint": "admission",
            "weight_kg": "86.000",
            "dry_weight_kg": None,
            "source_reference": None,
        },
        {
            "weight_id": "WT-VAL805-002",
            "case_id": "VAL-805",
            "timepoint": "hospital_day_2",
            "weight_kg": "83.000",
            "dry_weight_kg": None,
            "source_reference": None,
        },
        {
            "weight_id": "WT-VAL805-003",
            "case_id": "VAL-805",
            "timepoint": "discharge",
            "weight_kg": "80.000",
            "dry_weight_kg": "80.000",
            "source_reference": None,
        },
    ]
    chart["CaseIntakeOutput"] = [
        {
            "io_id": f"IO-VAL805-00{index}",
            "case_id": "VAL-805",
            "timepoint": f"hospital_day_{index}",
            "intake_ml": intake,
            "output_ml": output,
            "net_ml": intake - output,
            "notes": None,
            "source_reference": None,
        }
        for index, intake, output in ((1, 1100, 2800), (2, 1300, 2600), (3, 1500, 1900))
    ]
    chart.setdefault("CaseImaging", []).append(
        {
            "study_id": "IMG-VAL805-ECHO",
            "case_id": "VAL-805",
            "timepoint": "admission",
            "study_type": "Transthoracic echocardiogram",
            "body_site": "heart",
            "finding": "Left ventricular ejection fraction 30 percent. No valvular vegetation.",
            "source_reference": None,
        }
    )
    for row in _meds(chart, "furosemide"):
        if row.get("context") == "inpatient":
            row["route"] = "intravenous"
            row["frequency"] = "twice daily"
            row["notes"] = (
                "Intravenous furosemide 40 MG twice daily during volume overload. "
                "The home regimen was 40 MG orally once daily."
            )
    hpi = (
        "A 68-year-old woman with systolic heart failure and hyperlipidemia is admitted with "
        "one week of dyspnea, edema, and orthopnea. Home medicines are oral furosemide 40 MG "
        "once daily, metoprolol succinate 25 MG once daily, and atorvastatin. She is hypoxic "
        "and volume overloaded. Dry weight is 80 kg."
    )
    admission = (
        "Admission note. Acute decompensated heart failure. Weight 86 kg, above a dry weight "
        "of 80 kg. Oxygen saturation 92 percent. Home oral furosemide is not the inpatient dose."
    )
    course = (
        "Pulmonary edema was treated with intravenous furosemide 40 MG twice daily, which is "
        "distinct from the home oral 40 MG once daily. Intake and output were net negative on "
        "hospital days 1, 2, and 3. Weight fell from 86 kg to 83 kg and then to 80 kg, equal "
        "to the dry weight. Creatinine fell from 1.7 mg/dL to 0.9 mg/dL and potassium from "
        "4.7 mmol/L to 4.3 mmol/L. B-type natriuretic peptide fell from 1120 pg/mL to 369 pg/mL. "
        "Discharge systolic blood pressure was 110 mmHg and heart rate 69. Oxygen saturation "
        "was 92 percent on admission and 98 percent at discharge. Ejection fraction was 30 "
        "percent. The home list is oral furosemide, metoprolol succinate, and atorvastatin."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    consults = chart.get("CaseConsult") or []
    if consults:
        consults[0]["assessment"] = (
            "Volume overload responded to intravenous diuresis. Ejection fraction is 30 percent. "
            "Discharge weight equals the 80 kg dry weight."
        )
        consults[0]["recommendation"] = (
            "Blood pressure at discharge is 110 mmHg systolic. Kidney function improved."
        )
    plans = []
    traces = []
    for medication, indication, phrases, rationale, event, labs, vitals in (
        (
            "furosemide 40 MG Oral Tablet",
            "Acute systolic (congestive) heart failure",
            ["intravenous furosemide", "80 kg", "dry weight"],
            "Congestion resolved and weight returned to the 80 kg dry weight, "
            "so the home oral 40 MG daily dose is resumed.",
            "Intravenous furosemide 40 MG twice daily, distinct from home oral therapy.",
            "creatinine 1.7 to 0.9 mg/dL; natriuretic peptide 1120 to 369 pg/mL",
            "oxygen saturation 98 percent at discharge",
        ),
        (
            "24 HR metoprolol succinate 25 MG Extended Release Oral Tablet",
            "Acute systolic (congestive) heart failure",
            ["ejection fraction 30", "110 mmhg", "heart rate 69"],
            "Ejection fraction is 30 percent, discharge heart rate is 69, "
            "and systolic pressure is 110 mmHg.",
            "Beta blocker was continued through diuresis.",
            "potassium 4.3 mmol/L",
            "systolic blood pressure 110 mmHg; heart rate 69",
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
            "Hyperlipidemia is a separate active problem and no statin intolerance is described.",
            "none",
            "none",
            "none",
        ),
    ):
        plan, trace = _decision(
            chart,
            medication=medication,
            action="continue",
            indication=indication,
            rationale=rationale,
            diagnosis=indication,
            event=event,
            labs=labs,
            vitals=vitals,
            history=f"Home {medication.split()[0]}",
            location="hospital course; weights; labs; imaging",
            phrases=phrases,
            route="oral",
            frequency="once daily",
        )
        plans.append(plan)
        traces.append(trace)
    summary = (
        "The hospitalization is diuresis with intravenous furosemide, serial weights from "
        "86 kg to an 80 kg dry weight, and three days of negative intake and output. "
        "The reference plan resumes the home oral diuretic and beta blocker. No new "
        "guideline-directed class was added at a systolic pressure of 110 mmHg; a low-dose "
        "ACE inhibitor, ARB, or SGLT2 inhibitor is recorded only as an acceptable alternative."
    )
    alternatives = [
        {
            "medication": "ACE inhibitor, ARB, or SGLT2 inhibitor",
            "action": "start",
            "rationale": (
                "Ejection fraction is 30 percent and creatinine recovered to 0.9 mg/dL. "
                "The reference plan does not start one of these classes because none was "
                "a home medicine and the discharge systolic pressure is 110 mmHg. A reviewer "
                "may still accept a low-dose start."
            ),
        }
    ]
    return (
        chart,
        _plan_bundle(plans, chart.get("CaseFollowup") or [], alternatives=alternatives),
        traces,
        summary,
    )


def _revise_809(chart: dict[str, Any]) -> Revision:
    for vital in chart.get("CaseVital") or []:
        if vital.get("timepoint") == "admission":
            vital["temp_c"] = "38.60"
            vital["heart_rate"] = 104
            vital["resp_rate"] = 22
            vital["bp_systolic"] = 128
        elif vital.get("timepoint") == "discharge":
            vital["temp_c"] = "36.80"
            vital["heart_rate"] = 76
            vital["bp_systolic"] = 118
    _add_lab(
        chart,
        lab_id="LAB-VAL809-CR3",
        timepoint="hospital_day_3",
        test_name="Creatinine [Mass/volume] in Serum or Plasma",
        value=1.0,
        unit="mg/dL",
    )
    for row in chart.get("CaseMicrobiology") or []:
        if row.get("result") == "growth":
            row["organism"] = "Viridans group streptococcus"
            row["notes"] = (
                "Identified from the admission blood culture after a dental "
                "extraction three weeks earlier."
            )
    hpi = (
        "An 82-year-old man with hypertension and hyperlipidemia is admitted with one week of "
        "fever and fatigue. He had a dental extraction three weeks earlier. A new murmur is "
        "present. Home medicines are lisinopril and atorvastatin. Outpatient creatinine last "
        "month was 0.8 mg/dL."
    )
    admission = (
        "Admission note. Fever 38.6 C, heart rate 104, and a new murmur after a recent dental "
        "extraction. Blood cultures grew gram-positive cocci. Creatinine is 1.3 mg/dL, above "
        "the 0.8 mg/dL baseline."
    )
    course = (
        "Blood cultures grew Viridans group streptococcus. A transthoracic echocardiogram showed "
        "a vegetation. Ceftriaxone 2000 MG intravenously once daily was started. Later blood "
        "cultures showed no growth, and fever resolved. Acute kidney injury was attributed to "
        "the infection: creatinine 1.3 mg/dL on admission, 1.0 mg/dL on hospital day 3, and "
        "0.8 mg/dL at discharge. Lisinopril was held while creatinine was 1.3 mg/dL. At "
        "discharge the creatinine is 0.8 mg/dL and the systolic blood pressure is 118 mmHg."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for row in _meds(chart, "lisinopril"):
        if row.get("context") == "inpatient":
            row["status"] = "held"
            row["held_reason"] = "Held while creatinine was 1.3 mg/dL."
    consults = chart.get("CaseConsult") or []
    if consults:
        consults[0]["assessment"] = (
            "Endocarditis with a vegetation and Viridans group streptococcus bacteremia after "
            "a dental extraction. Fever has resolved and follow-up cultures show no growth."
        )
        consults[0]["recommendation"] = (
            "Four weeks of intravenous ceftriaxone is the planned course for viridans "
            "endocarditis with a vegetation. Kidney function has returned to the "
            "0.8 mg/dL baseline."
        )
    plans = []
    traces = []
    for medication, action, indication, phrases, rationale, event, labs, vitals in (
        (
            "ceftriaxone 2000 MG Injection",
            "start",
            "Acute and subacute infective endocarditis",
            ["ceftriaxone", "vegetation", "viridans"],
            "Started for endocarditis after a vegetation and streptococcal "
            "bacteremia were documented.",
            "Ceftriaxone was started during the admission.",
            "none",
            "fever resolved",
        ),
        (
            "lisinopril 10 MG Oral Tablet",
            "restart",
            "Essential (primary) hypertension",
            ["lisinopril was held", "0.8 mg/dl", "118"],
            "Held during the creatinine rise and restarted after creatinine returned to 0.8 mg/dL.",
            "Held while creatinine was 1.3 mg/dL.",
            "creatinine 1.3, then 1.0, then 0.8 mg/dL",
            "systolic blood pressure 118 mmHg",
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "continue",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
            "Hyperlipidemia is unchanged and no statin adverse effect is described.",
            "none",
            "none",
            "none",
        ),
    ):
        plan, trace = _decision(
            chart,
            medication=medication,
            action=action,
            indication=indication,
            rationale=rationale,
            diagnosis=indication,
            event=event,
            labs=labs,
            vitals=vitals,
            history="Home lisinopril and atorvastatin; ceftriaxone was not a home medicine",
            location="history of present illness; hospital course; labs; microbiology; imaging",
            phrases=phrases,
            duration="4 weeks" if medication.startswith("ceftriaxone") else None,
        )
        plans.append(plan)
        traces.append(trace)
    summary = (
        "Endocarditis is supported by fever, a recent dental extraction, viridans bacteremia, "
        "and a vegetation. Lisinopril was held at creatinine 1.3 mg/dL and restarted at 0.8 mg/dL."
    )
    monitoring = chart.get("CaseMonitoring") or []
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or [], monitoring), traces, summary


def _revise_813(chart: dict[str, Any]) -> Revision:
    tacrolimus = next(
        row for row in _meds(chart, "BX Rating tacrolimus") if row.get("context") == "home"
    )
    mycophenolate_home = deepcopy(tacrolimus)
    mycophenolate_home.update(
        {
            "medication_id": "MED-VAL813-MMF-H",
            "drug": "mycophenolate mofetil 500 MG Oral Tablet",
            "reported_name": "mycophenolate mofetil 500 MG Oral Tablet",
            "dose": "1000 MG",
            "route": "oral",
            "frequency": "twice daily",
            "indication": "Kidney transplant status",
            "status": "home",
            "context": "home",
            "notes": "Given as 500 MG tablets.",
            "held_reason": None,
        }
    )
    mycophenolate_inpatient = deepcopy(mycophenolate_home)
    mycophenolate_inpatient.update(
        {
            "medication_id": "MED-VAL813-MMF-I",
            "context": "inpatient",
            "status": "active",
        }
    )
    kept = []
    for row in chart.get("CaseMedication") or []:
        drug = (row.get("drug") or "").casefold()
        if "valganciclovir" in drug and row.get("context") == "home":
            continue
        kept.append(row)
    for row in kept:
        if "valganciclovir" in (row.get("drug") or "").casefold():
            row["notes"] = (
                "Started after the CMV viral load was detected during this admission. "
                "Given as 450 MG tablets. Not a pre-admission medicine."
            )
            row["indication"] = "Other cytomegaloviral diseases"
    chart["CaseMedication"] = [mycophenolate_home, mycophenolate_inpatient, *kept]
    hpi = (
        "A 64-year-old woman with a kidney transplant is admitted with several days of diarrhea. "
        "Maintenance immunosuppression is tacrolimus and mycophenolate. She also takes amlodipine "
        "for hypertension and atorvastatin for hyperlipidemia. She was not taking valganciclovir "
        "before this admission."
    )
    admission = (
        "Admission note. Diarrhea after kidney transplantation. Home immunosuppression is "
        "tacrolimus and mycophenolate. Valganciclovir is not a home medicine."
    )
    course = (
        "CMV viral load drawn after admission was detected, and valganciclovir was started "
        "only after that result, 900 MG twice daily given as 450 MG tablets. Diarrhea then "
        "improved, and a later viral-load review was lower than the admission result. "
        "Creatinine was 1.2 mg/dL then 1.0 mg/dL and potassium was 4.7 mmol/L then 3.9 mmol/L. "
        "Tacrolimus, mycophenolate, amlodipine, and atorvastatin were continued."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for consult in chart.get("CaseConsult") or []:
        if consult.get("service") == "infectious disease":
            consult["assessment"] = (
                "CMV viral load was detected after admission in a transplant "
                "recipient with diarrhea."
            )
            consult["recommendation"] = (
                "Antiviral therapy was begun after the viral-load result. Symptoms improved."
            )
        if consult.get("service") == "transplant":
            consult["assessment"] = (
                "Kidney transplant maintained on tacrolimus and mycophenolate."
            )
            consult["recommendation"] = (
                "Immunosuppression was kept active while the new antiviral was started."
            )
    plans = []
    traces = []
    start_plan, start_trace = _decision(
        chart,
        medication="valganciclovir 450 MG Oral Tablet",
        action="start",
        dose="900 MG",
        frequency="twice daily",
        indication="Other cytomegaloviral diseases",
        rationale="Started only after CMV viral load was detected during the admission.",
        diagnosis="Other cytomegaloviral diseases",
        event="Valganciclovir was started after the CMV viral load was detected.",
        labs="creatinine 1.0 mg/dL at discharge supports the labeled dose",
        vitals="none",
        history="Not a home medicine",
        location="history of present illness; hospital course; procedures",
        phrases=["not taking valganciclovir", "viral load", "started"],
    )
    plans.append(start_plan)
    traces.append(start_trace)
    for medication, indication, phrases in (
        (
            "BX Rating tacrolimus 1 MG Oral Capsule",
            "Kidney transplant status",
            ["tacrolimus", "transplant"],
        ),
        (
            "mycophenolate mofetil 500 MG Oral Tablet",
            "Kidney transplant status",
            ["mycophenolate", "transplant"],
        ),
        (
            "amlodipine 5 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["amlodipine", "hypertension"],
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
        ),
    ):
        plan, trace = _decision(
            chart,
            medication=medication,
            action="continue",
            indication=indication,
            rationale=(
                "This pre-admission medicine matches an active diagnosis "
                "and was not the new antiviral."
            ),
            diagnosis=indication,
            event="Immunosuppression and cardiovascular medicines were kept active.",
            labs="creatinine 1.2 then 1.0 mg/dL; potassium 3.9 mmol/L",
            vitals="discharge systolic blood pressure 134 mmHg",
            history=f"Home {medication.split()[0]}",
            location="diagnoses; history of present illness; labs",
            phrases=phrases,
        )
        plans.append(plan)
        traces.append(trace)
    summary = (
        "Chronology is corrected inside the same transplant/CMV scenario: diarrhea on "
        "tacrolimus and mycophenolate, CMV diagnosed after admission, and valganciclovir "
        "started only after the viral load. Not marked REPLACE_WITH_MATCHED_CASE."
    )
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or []), traces, summary


def _drop_direct_orders(chart: dict[str, Any]) -> None:
    chart["CaseInstruction"] = [
        row
        for row in chart.get("CaseInstruction") or []
        if "resume when holding" not in (row.get("instruction_text") or "").casefold()
        and "exactly as listed" not in (row.get("instruction_text") or "").casefold()
    ]
    for row in chart.get("CaseMedication") or []:
        goal = (row.get("target_or_goal") or "").casefold()
        if "resume " in goal or "restart" in goal:
            row["target_or_goal"] = None


def _repair_blockers(chart: dict[str, Any], traces: list[dict[str, Any]]) -> list[str]:
    findings: list[str] = []
    blob = json.dumps(chart).casefold()
    needles = (
        *_LEAKS,
        "resume when holding",
        "intended to restart",
        "in this profile",
        "the discharge list is the intended",
        "intended outpatient",
        "intended adjusted outpatient",
        "intended regimen",
        "correct regimen",
        "expected state",
        "scenario rule",
        "test case",
    )
    for needle in needles:
        if needle in blob:
            findings.append(f"answer leakage:{needle}")
    if any(item["evidence_class"] == HIDDEN for item in traces):
        findings.append(HIDDEN)
    if any(item["evidence_class"] == INCONSISTENT for item in traces):
        findings.append(INCONSISTENT)
    if not _cmv_chronology_ok(chart):
        findings.append("contradictory timeline:CMV and valganciclovir do not agree")
    return findings


def _cmv_chronology_ok(chart: dict[str, Any]) -> bool:
    admission = ((chart.get("ClinicalCase") or {}).get("admission_dx") or "").casefold()
    if "cytomegalo" not in admission:
        return True
    blob = json.dumps(chart).casefold()
    home = any(
        "valganciclovir" in (row.get("drug") or "").casefold() and row.get("context") == "home"
        for row in chart.get("CaseMedication") or []
    )
    if home:
        return "before this admission" in blob
    return "viral load" in blob and "not taking valganciclovir" in blob


def _build_repaired(case_id: str) -> dict[str, Any]:
    resident, _evaluator = _load_pair(case_id)
    revised, reference, trace, summary = _REPAIRS[case_id](deepcopy(resident))
    _strip_resident_answers(revised)
    _drop_direct_orders(revised)
    blockers = _repair_blockers(revised, trace)
    weak = [
        str(item["medication"])
        for item in trace
        if item["evidence_class"] == WEAK
    ]
    status = READY_REVIEW if not blockers else NOT_READY
    finding = "none" if not blockers else "; ".join(blockers)
    if weak:
        note = "WEAK_EVIDENCE: " + ", ".join(weak)
        finding = note if finding == "none" else f"{finding}; {note}"
    evaluator = deepcopy(revised)
    evaluator["reference_discharge_plan"] = reference
    evaluator["evidence_trace"] = trace
    evaluator["clinician_review_status"] = status
    evaluator["revision_summary"] = summary
    evaluator["precheck_finding"] = finding
    return {
        "case_id": case_id,
        "resident": revised,
        "evaluator": evaluator,
        "summary": summary,
        "status": status,
        "classes": [item["evidence_class"] for item in trace],
        "finding": evaluator["precheck_finding"],
    }


def _continue_home(
    chart: dict[str, Any],
    plans: list[dict[str, Any]],
    traces: list[dict[str, Any]],
    medication: str,
    indication: str,
    phrases: list[str],
    rationale: str,
    event: str,
    labs: str,
    vitals: str,
) -> None:
    plan, trace = _decision(
        chart,
        medication=medication,
        action="continue",
        indication=indication,
        rationale=rationale,
        diagnosis=indication,
        event=event,
        labs=labs,
        vitals=vitals,
        history=f"Home {medication.split()[0]}",
        location="diagnoses; medication list; hospital course; labs",
        phrases=phrases,
    )
    plans.append(plan)
    traces.append(trace)


def _repair_806(chart: dict[str, Any]) -> Revision:
    _drop_direct_orders(chart)
    for row in _meds(chart, "lisinopril"):
        if row.get("context") == "inpatient":
            row["status"] = "held"
            row["held_reason"] = (
                "Held for creatinine 2.8 mg/dL during decongestion. "
                "Discharge creatinine is 1.6 mg/dL."
            )
    hpi = (
        "An 82-year-old woman with heart failure, hypertension, and hyperlipidemia is "
        "admitted with several days of progressive dyspnea and edema. Home medicines are "
        "furosemide, atorvastatin, metoprolol succinate, and lisinopril. No prior creatinine "
        "from before this illness is documented."
    )
    admission = (
        "Admission note. Volume overload with creatinine 2.8 mg/dL. Lisinopril was held "
        "because the creatinine rose during decongestion."
    )
    course = (
        "Lisinopril was held while creatinine was 2.8 mg/dL. Creatinine then fell to "
        "1.6 mg/dL. Potassium was 3.8 mmol/L on admission and 4.5 mmol/L at discharge. "
        "Discharge systolic blood pressure was 116 mmHg. No prior creatinine is available "
        "for comparison."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for consult in chart.get("CaseConsult") or []:
        consult["assessment"] = "Creatinine fell from 2.8 mg/dL to 1.6 mg/dL during diuresis."
        consult["recommendation"] = "No creatinine from before this illness is in the record."
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    restarted, restarted_trace = _decision(
        chart,
        medication="lisinopril 10 MG Oral Tablet",
        action="restart",
        indication="Essential (primary) hypertension",
        rationale=(
            "No pre-illness creatinine is in the source. The visible course is a fall "
            "from 2.8 mg/dL to 1.6 mg/dL, potassium 4.5 mmol/L, and systolic pressure "
            "116 mmHg, with cardiology follow-up in 3 days. That is enough to end the "
            "hold, but the remaining creatinine elevation makes the decision weak."
        ),
        diagnosis="Essential (primary) hypertension",
        event="Lisinopril was held while creatinine was 2.8 mg/dL.",
        labs="creatinine 2.8 mg/dL then 1.6 mg/dL; potassium 4.5 mmol/L",
        vitals="discharge systolic blood pressure 116 mmHg",
        history="Home lisinopril",
        location="hospital course; labs; vitals; follow-up",
        phrases=["lisinopril was held", "1.6", "2.8", "4.5", "116"],
        weak=True,
    )
    plans.append(restarted)
    traces.append(restarted_trace)
    for medication, indication, phrases in (
        (
            "furosemide 40 MG Oral Tablet",
            "Acute systolic (congestive) heart failure",
            ["furosemide", "heart failure"],
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
        ),
        (
            "24 HR metoprolol succinate 25 MG Extended Release Oral Tablet",
            "Acute systolic (congestive) heart failure",
            ["metoprolol", "heart failure"],
        ),
    ):
        _continue_home(
            chart,
            plans,
            traces,
            medication,
            indication,
            phrases,
            "The home medicine matches an active diagnosis and was not the held ACE inhibitor.",
            "Diuresis proceeded while lisinopril was held.",
            "creatinine 1.6 mg/dL; potassium 4.5 mmol/L",
            "systolic blood pressure 116 mmHg",
        )
    summary = (
        "No prior creatinine exists in the source data, so none was invented. Direct "
        "restart wording was removed. The hidden reference restarts lisinopril from the "
        "visible fall in creatinine to 1.6 mg/dL, potassium 4.5 mmol/L, and systolic "
        "pressure 116 mmHg. That decision is WEAK_EVIDENCE because creatinine is still "
        "elevated and no baseline is known."
    )
    alternatives = [
        {
            "medication": "lisinopril 10 MG Oral Tablet",
            "action": "hold",
            "rationale": (
                "A reviewer may keep lisinopril held because creatinine is still 1.6 mg/dL "
                "and no pre-illness baseline is documented."
            ),
        }
    ]
    return (
        chart,
        _plan_bundle(plans, chart.get("CaseFollowup") or [], alternatives=alternatives),
        traces,
        summary,
    )


def _repair_807(chart: dict[str, Any]) -> Revision:
    hpi = (
        "A 60-year-old man with heart failure and hyperlipidemia is admitted with two days "
        "of worsening dyspnea and orthopnea. Home medicines are oral furosemide 40 MG once "
        "daily, atorvastatin, and metoprolol succinate."
    )
    admission = (
        "Admission note. Decompensated heart failure. The medication list does not include "
        "a potassium replacement product."
    )
    course = (
        "Diuresis used the recorded loop diuretic. Potassium was 3.5 mmol/L on admission and "
        "3.8 mmol/L at discharge. Creatinine stayed 1.2 mg/dL. Natriuretic peptide fell from "
        "836 pg/mL to 585 pg/mL. Weight fell from 99 kg to 93 kg. No potassium medication "
        "was administered."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for consult in chart.get("CaseConsult") or []:
        consult["recommendation"] = (
            "Potassium was 3.5 mmol/L then 3.8 mmol/L without a potassium product."
        )
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    for medication, indication, phrases, rationale in (
        (
            "furosemide 40 MG Oral Tablet",
            "Acute systolic (congestive) heart failure",
            ["furosemide", "3.8", "1.2"],
            "Creatinine stayed 1.2 mg/dL and potassium was 3.8 mmol/L after diuresis.",
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
            "Hyperlipidemia remains active and no statin adverse effect is described.",
        ),
        (
            "24 HR metoprolol succinate 25 MG Extended Release Oral Tablet",
            "Acute systolic (congestive) heart failure",
            ["metoprolol", "heart failure"],
            "The beta blocker was already home therapy for systolic heart failure.",
        ),
    ):
        _continue_home(
            chart,
            plans,
            traces,
            medication,
            indication,
            phrases,
            rationale,
            "Diuresis used the recorded loop diuretic.",
            "potassium 3.5 then 3.8 mmol/L; creatinine 1.2 mg/dL",
            "discharge oxygen saturation 98 percent",
        )
    summary = (
        "The source medication list has no potassium product, so the repletion claim was "
        "removed. Potassium remains the measured 3.5 then 3.8 mmol/L. Home heart-failure "
        "medicines stay on the reference plan."
    )
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or []), traces, summary


def _repair_808(chart: dict[str, Any]) -> Revision:
    for weight in chart.get("CaseWeight") or []:
        if weight.get("timepoint") == "discharge":
            weight["dry_weight_kg"] = weight.get("weight_kg")
    hpi = (
        "A 66-year-old woman with systolic heart failure and hypertension is admitted with "
        "one week of worsening edema and orthopnea. Home medicines are oral furosemide 40 MG "
        "once daily, spironolactone, lisinopril, and metoprolol succinate. Admission weight "
        "is 94 kg."
    )
    admission = (
        "Admission note. Edema and orthopnea. Oral furosemide 40 MG once daily is both the "
        "home dose and the inpatient dose."
    )
    course = (
        "The same oral furosemide 40 MG once daily was used in the hospital. Weight fell "
        "from 94 kg to 92 kg. The dry weight is 92 kg, matching the discharge weight. "
        "The recorded intake and output do not support an 8 kg loss. Creatinine stayed "
        "0.9 mg/dL and potassium rose from 3.4 mmol/L to 3.7 mmol/L. Discharge systolic "
        "blood pressure was 142 mmHg."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for consult in chart.get("CaseConsult") or []:
        consult["assessment"] = "Discharge weight is 92 kg, equal to the dry weight."
        consult["recommendation"] = "Creatinine stayed 0.9 mg/dL and potassium was 3.7 mmol/L."
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    loop, loop_trace = _decision(
        chart,
        medication="furosemide 40 MG Oral Tablet",
        action="continue",
        indication="Acute systolic (congestive) heart failure",
        rationale=(
            "Home and inpatient furosemide are both oral 40 MG once daily. Weight fell "
            "from 94 kg to the 92 kg dry weight, with creatinine 0.9 mg/dL and potassium "
            "3.7 mmol/L. No dose change is supported."
        ),
        diagnosis="Acute systolic (congestive) heart failure",
        event="The same oral furosemide 40 MG once daily was used in the hospital.",
        labs="creatinine 0.9 mg/dL; potassium 3.7 mmol/L",
        vitals="discharge systolic blood pressure 142 mmHg",
        history="Home oral furosemide 40 MG once daily",
        location="hospital course; weights; medication list; labs",
        phrases=["oral furosemide 40 mg once daily", "92 kg", "0.9"],
    )
    plans.append(loop)
    traces.append(loop_trace)
    for medication, indication, phrases in (
        (
            "spironolactone 25 MG Oral Tablet",
            "Acute systolic (congestive) heart failure",
            ["spironolactone", "3.7"],
        ),
        (
            "lisinopril 10 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["lisinopril", "0.9", "142"],
        ),
        (
            "24 HR metoprolol succinate 25 MG Extended Release Oral Tablet",
            "Acute systolic (congestive) heart failure",
            ["metoprolol", "heart failure"],
        ),
    ):
        _continue_home(
            chart,
            plans,
            traces,
            medication,
            indication,
            phrases,
            "This home medicine matches an active diagnosis, and creatinine and potassium "
            "stayed acceptable during diuresis.",
            "Weight fell from 94 kg to the 92 kg dry weight.",
            "creatinine 0.9 mg/dL; potassium 3.7 mmol/L",
            "systolic blood pressure 142 mmHg",
        )
    summary = (
        "Answer-revealing regimen language was removed. Home and inpatient furosemide "
        "stay oral 40 MG once daily because the source doses match. The unsupported "
        "86 kg dry weight was the incorrect field: discharge weight stays 92 kg, and "
        "dry weight is corrected to 92 kg. The reference continues that oral dose."
    )
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or []), traces, summary


def _repair_814(chart: dict[str, Any]) -> Revision:
    _drop_direct_orders(chart)
    kept = []
    for row in chart.get("CaseMedication") or []:
        drug = (row.get("drug") or "").casefold()
        if "valganciclovir" in drug and row.get("context") == "home":
            continue
        if "valganciclovir" in drug:
            row["dose"] = "450 MG"
            row["frequency"] = "twice daily"
            row["notes"] = (
                "Started after the CMV viral load was detected during this admission. "
                "Given as one 450 MG tablet twice daily while creatinine is 1.6 mg/dL. "
                "Not a home medicine."
            )
            row["indication"] = "Other cytomegaloviral diseases"
        if "mycophenolate" in drug and row.get("context") == "inpatient":
            row["status"] = "held"
            row["held_reason"] = "Held during active CMV infection with diarrhea."
        kept.append(row)
    chart["CaseMedication"] = kept
    for procedure in chart.get("CaseProcedure") or []:
        if procedure.get("time") == "admission":
            procedure["findings"] = "CMV viral load drawn after admission was detected."
        elif procedure.get("time") == "inpatient":
            procedure["findings"] = (
                "CMV viral burden was lower than the admission result."
            )
    hpi = (
        "A 71-year-old man with a kidney transplant is admitted with one week of progressive "
        "diarrhea and nausea. Maintenance immunosuppression is tacrolimus and mycophenolate. "
        "He also takes amlodipine. He was not taking valganciclovir before this admission."
    )
    admission = (
        "Admission note. Diarrhea and nausea after kidney transplantation. Home "
        "immunosuppression is tacrolimus and mycophenolate. Valganciclovir is not a home "
        "medicine. Creatinine is 2.5 mg/dL."
    )
    course = (
        "CMV viral load drawn after admission was detected. Valganciclovir was started only "
        "after that result. While creatinine was 2.5 mg/dL the dose used was 450 MG once "
        "daily. Creatinine then fell to 1.6 mg/dL, and the inpatient dose is 450 MG twice "
        "daily, given as 450 MG tablets. A later viral-load review was lower than the "
        "admission result. Mycophenolate was held during the active CMV infection. "
        "Diarrhea was improving as the viral load fell."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for consult in chart.get("CaseConsult") or []:
        if consult.get("service") == "infectious disease":
            consult["assessment"] = "CMV viral load drawn after admission was detected."
            consult["recommendation"] = (
                "A later viral-load review was lower than the admission result."
            )
        if consult.get("service") == "transplant":
            consult["assessment"] = "Mycophenolate was held during active CMV infection."
            consult["recommendation"] = "Creatinine fell from 2.5 mg/dL to 1.6 mg/dL."
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    started, started_trace = _decision(
        chart,
        medication="valganciclovir 450 MG Oral Tablet",
        action="start",
        dose="450 MG",
        frequency="twice daily",
        indication="Other cytomegaloviral diseases",
        rationale=(
            "Started after the in-hospital viral load. The twice-daily 450 MG dose matches "
            "creatinine 1.6 mg/dL; 900 MG twice daily is not used at that creatinine."
        ),
        diagnosis="Other cytomegaloviral diseases",
        event="Valganciclovir was started only after the viral load was detected.",
        labs="creatinine 2.5 mg/dL then 1.6 mg/dL",
        vitals="none",
        history="Not a home medicine",
        location="history of present illness; hospital course; procedures",
        phrases=["not taking valganciclovir", "viral load", "450 mg twice daily", "1.6"],
    )
    plans.append(started)
    traces.append(started_trace)
    restarted, restarted_trace = _decision(
        chart,
        medication="mycophenolate mofetil 500 MG Oral Tablet",
        action="restart",
        indication="Kidney transplant status",
        rationale=(
            "Mycophenolate was held for CMV diarrhea. The later viral load is lower than "
            "the admission result, diarrhea was improving, and creatinine fell to 1.6 mg/dL."
        ),
        diagnosis="Kidney transplant status",
        event="Mycophenolate was held during the active CMV infection.",
        labs="creatinine 1.6 mg/dL; later viral load lower than admission",
        vitals="none",
        history="Home mycophenolate 1000 MG twice daily",
        location="hospital course; procedures; labs",
        phrases=["mycophenolate was held", "lower than the admission", "diarrhea was improving"],
    )
    plans.append(restarted)
    traces.append(restarted_trace)
    for medication, indication, phrases in (
        (
            "BX Rating tacrolimus 1 MG Oral Capsule",
            "Kidney transplant status",
            ["tacrolimus", "transplant"],
        ),
        (
            "amlodipine 5 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["amlodipine", "hypertension"],
        ),
    ):
        _continue_home(
            chart,
            plans,
            traces,
            medication,
            indication,
            phrases,
            "This home medicine matches an active diagnosis and is not the new antiviral.",
            "Tacrolimus and amlodipine were kept active during CMV treatment.",
            "creatinine 1.6 mg/dL",
            "none",
        )
    alternatives = [
        {
            "medication": "mycophenolate mofetil 500 MG Oral Tablet",
            "action": "hold",
            "rationale": (
                "A reviewer may keep mycophenolate held until the transplant visit because "
                "the viral load is lower but not described as cleared."
            ),
        }
    ]
    summary = (
        "Valganciclovir is no longer a home medicine. It starts after the in-hospital viral "
        "load, at 450 MG twice daily once creatinine is 1.6 mg/dL. Direct mycophenolate "
        "restart instructions were removed. The hidden reference still restarts mycophenolate "
        "because the viral load is lower and diarrhea was improving."
    )
    return (
        chart,
        _plan_bundle(
            plans,
            chart.get("CaseFollowup") or [],
            alternatives=alternatives,
        ),
        traces,
        summary,
    )


def _repair_815(chart: dict[str, Any]) -> Revision:
    kept = []
    for row in chart.get("CaseMedication") or []:
        drug = (row.get("drug") or "").casefold()
        if "valganciclovir" in drug and row.get("context") == "home":
            continue
        if "valganciclovir" in drug:
            row["notes"] = (
                "Started after the CMV viral load was detected during this admission. "
                "Not a home medicine. Given as 450 MG tablets."
            )
        kept.append(row)
    chart["CaseMedication"] = kept
    chart.setdefault("CaseProcedure", []).append(
        {
            "procedure_id": "PROC-VAL815-CMV",
            "case_id": "VAL-815",
            "procedure_name": "CMV viral-load review",
            "procedure_type": "diagnostic",
            "date": None,
            "time": "admission",
            "performed_by": None,
            "anesthesia_type": None,
            "findings": "CMV viral load drawn after admission was detected.",
            "complications": None,
            "duration_minutes": None,
            "laterality": None,
            "source_reference": None,
        }
    )
    hpi = (
        "A 68-year-old woman with a kidney transplant is admitted with two days of nausea. "
        "Home medicines are tacrolimus 1 MG every 12 hours, amlodipine, lisinopril, and "
        "atorvastatin. She was not taking valganciclovir before this admission. The "
        "tacrolimus dose is the same at home and in the hospital. No dose change is recorded."
    )
    admission = (
        "Admission note. Nausea after kidney transplantation. Tacrolimus is 1 MG every "
        "12 hours. Valganciclovir is not a home medicine."
    )
    course = (
        "CMV viral load drawn after admission was detected. Valganciclovir 900 MG twice "
        "daily was started after that result. Creatinine was 0.8 mg/dL then 0.9 mg/dL and "
        "potassium was 4.6 mmol/L then 4.5 mmol/L. Tacrolimus remained 1 MG every 12 hours. "
        "There was no tacrolimus dose change."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for consult in chart.get("CaseConsult") or []:
        consult["recommendation"] = (
            "Tacrolimus remained 1 MG every 12 hours. No dose change is recorded."
        )
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    started, started_trace = _decision(
        chart,
        medication="valganciclovir 450 MG Oral Tablet",
        action="start",
        dose="900 MG",
        frequency="twice daily",
        indication="Other cytomegaloviral diseases",
        rationale=(
            "Started after the in-hospital viral load. Creatinine 0.9 mg/dL supports "
            "the 900 MG twice-daily dose."
        ),
        diagnosis="Other cytomegaloviral diseases",
        event="Valganciclovir was started after the viral load was detected.",
        labs="creatinine 0.8 then 0.9 mg/dL",
        vitals="none",
        history="Not a home medicine",
        location="history of present illness; hospital course; procedures",
        phrases=["not taking valganciclovir", "viral load", "900 mg"],
    )
    plans.append(started)
    traces.append(started_trace)
    for medication, indication, phrases in (
        (
            "BX Rating tacrolimus 1 MG Oral Capsule",
            "Kidney transplant status",
            ["tacrolimus", "1 mg every 12 hours", "no dose change"],
        ),
        (
            "amlodipine 5 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["amlodipine", "hypertension"],
        ),
        (
            "lisinopril 10 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["lisinopril", "0.9", "4.5"],
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
        ),
    ):
        _continue_home(
            chart,
            plans,
            traces,
            medication,
            indication,
            phrases,
            "The recorded dose is unchanged and the labs do not show a reason to alter it.",
            "No tacrolimus dose change is recorded.",
            "creatinine 0.9 mg/dL; potassium 4.5 mmol/L",
            "none",
        )
    summary = (
        "No tacrolimus dose change exists in the medication data, so the adjustment claim "
        "was removed. CMV follows the VAL-813 chronology: valganciclovir is not a home "
        "medicine and starts after the in-hospital viral load, at 900 MG twice daily."
    )
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or []), traces, summary


def _repair_816(chart: dict[str, Any]) -> Revision:
    hpi = (
        "A 57-year-old man with a kidney transplant is admitted with one day of diarrhea and "
        "fatigue that are improving. This is not a new CMV diagnosis. CMV disease was "
        "diagnosed before this admission, and he was already taking valganciclovir. Home "
        "medicines also include tacrolimus and atorvastatin."
    )
    admission = (
        "Admission note. Reassessment of improving diarrhea and fatigue while already "
        "taking valganciclovir for CMV disease before this admission."
    )
    course = (
        "Symptoms improved during established treatment. Valganciclovir was already "
        "treatment for CMV disease before this admission and remained on the inpatient "
        "list at 900 MG twice daily. Creatinine was 1.0 mg/dL then 0.9 mg/dL. Tacrolimus "
        "remained 1 MG every 12 hours. Infectious-disease follow-up is scheduled to review "
        "the antiviral course."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for row in chart.get("CaseInstruction") or []:
        row["instruction_text"] = (
            "Infectious-disease follow-up is scheduled to review the antiviral course."
        )
    for follow in chart.get("CaseFollowup") or []:
        follow["item"] = "Infectious-disease follow-up to review antiviral therapy"
    for consult in chart.get("CaseConsult") or []:
        if "intended" in (consult.get("recommendation") or "").casefold():
            consult["recommendation"] = (
                "Antiviral therapy was already underway before this admission."
            )
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    for medication, indication, phrases in (
        (
            "valganciclovir 450 MG Oral Tablet",
            "Other cytomegaloviral diseases",
            ["valganciclovir", "before this admission", "900 mg"],
        ),
        (
            "BX Rating tacrolimus 1 MG Oral Capsule",
            "Kidney transplant status",
            ["tacrolimus", "transplant"],
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
        ),
    ):
        _continue_home(
            chart,
            plans,
            traces,
            medication,
            indication,
            phrases,
            "The recorded dose is unchanged and matches an active diagnosis.",
            "Symptoms improved during established CMV treatment.",
            "creatinine 1.0 then 0.9 mg/dL",
            "none",
        )
    summary = (
        "CMV is framed as established disease already treated with valganciclovir before "
        "this admission. The direct pending-decision instruction was replaced with follow-up "
        "context. The reference continues the recorded doses."
    )
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or []), traces, summary


def _repair_817(chart: dict[str, Any]) -> Revision:
    hpi = (
        "An 87-year-old man is admitted after operative repair of a femoral-neck fracture. "
        "Home medicines are lisinopril, warfarin for atrial fibrillation, and metformin. "
        "He had been fatigued for several days."
    )
    admission = (
        "Admission note. Postoperative femoral-neck fracture. Warfarin is a home medicine "
        "for atrial fibrillation. Hemoglobin is 9.0 g/dL and INR is 2.6."
    )
    course = (
        "Warfarin was held around the operation and is active again on the inpatient list. "
        "INR is 2.6 at discharge and hemoglobin rose from 9.0 g/dL to 10.1 g/dL. "
        "Glucose was 127 mg/dL then 113 mg/dL."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    for medication, indication, phrases, rationale, labs in (
        (
            "warfarin sodium 5 MG Oral Tablet",
            "Paroxysmal atrial fibrillation",
            ["warfarin", "2.6", "atrial fibrillation"],
            "Warfarin is active on the inpatient list and the discharge INR is 2.6.",
            "INR 2.6; hemoglobin 9.0 then 10.1 g/dL",
        ),
        (
            "lisinopril 10 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["lisinopril", "hypertension"],
            "Hypertension remains active and lisinopril was not the perioperative anticoagulant.",
            "none",
        ),
        (
            "metformin hydrochloride 500 MG Oral Tablet",
            "Type 2 diabetes mellitus without complications",
            ["metformin", "diabetes", "113"],
            "Glucose improved from 127 mg/dL to 113 mg/dL.",
            "glucose 127 then 113 mg/dL",
        ),
    ):
        _continue_home(
            chart,
            plans,
            traces,
            medication,
            indication,
            phrases,
            rationale,
            "Warfarin is active again on the inpatient list after the operation.",
            labs,
            "discharge systolic blood pressure 136 mmHg",
        )
    summary = (
        "The phrase 'in this profile' and the enoxaparin sentence were removed. Enoxaparin "
        "is not part of this medication record, and the INR is already 2.6 on warfarin, so "
        "no prophylactic enoxaparin row was added. The reference continues warfarin, "
        "lisinopril, and metformin."
    )
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or []), traces, summary


def _repair_820(chart: dict[str, Any]) -> Revision:
    prophylaxis = "Inpatient venous-thromboembolism prophylaxis"
    for row in _meds(chart, "0.4 ML enoxaparin"):
        row["indication"] = prophylaxis
        row["notes"] = (
            "Started during the hospitalization for venous-thromboembolism prophylaxis "
            "after hip-fracture surgery. Not a home medicine and not treatment for "
            "atrial fibrillation."
        )
    hpi = (
        "An 82-year-old woman is admitted after hip-fracture surgery. Home medicines are "
        "lisinopril, warfarin for atrial fibrillation, and metformin. Enoxaparin was started "
        "during the hospitalization for venous-thromboembolism prophylaxis."
    )
    admission = (
        "Admission note. Postoperative hip fracture. Enoxaparin 40 MG subcutaneous once "
        "daily is inpatient venous-thromboembolism prophylaxis. Warfarin is the home "
        "medicine for atrial fibrillation."
    )
    course = (
        "Enoxaparin was used for venous-thromboembolism prophylaxis. It was not a home "
        "medicine. INR was 3.2 on admission and 2.0 at discharge. Hemoglobin rose from "
        "8.3 g/dL to 11.1 g/dL. Warfarin remained the atrial-fibrillation medicine."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    for medication, indication, phrases, rationale in (
        (
            "warfarin sodium 5 MG Oral Tablet",
            "Paroxysmal atrial fibrillation",
            ["warfarin", "atrial fibrillation", "2.0"],
            "Warfarin treats atrial fibrillation and the discharge INR is 2.0.",
        ),
        (
            "lisinopril 10 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["lisinopril", "hypertension"],
            "Hypertension remains active.",
        ),
        (
            "metformin hydrochloride 500 MG Oral Tablet",
            "Type 2 diabetes mellitus without complications",
            ["metformin", "diabetes"],
            "Diabetes remains active and glucose was 124 mg/dL then 119 mg/dL.",
        ),
    ):
        _continue_home(
            chart,
            plans,
            traces,
            medication,
            indication,
            phrases,
            rationale,
            "Warfarin remained the atrial-fibrillation medicine.",
            "INR 3.2 then 2.0; hemoglobin 8.3 then 11.1 g/dL",
            "none",
        )
    stopped, stopped_trace = _decision(
        chart,
        medication="0.4 ML enoxaparin sodium 100 MG/ML Prefilled Syringe",
        action="stop",
        indication=prophylaxis,
        rationale=(
            "Enoxaparin was started in the hospital for venous-thromboembolism prophylaxis, "
            "not for atrial fibrillation. It is not a home medicine. Warfarin covers atrial "
            "fibrillation and the discharge INR is 2.0."
        ),
        diagnosis=(
            "Fracture of unspecified part of neck of right femur, "
            "initial encounter for closed fracture"
        ),
        event="Enoxaparin was started during the hospitalization for prophylaxis.",
        labs="INR 2.0; hemoglobin 11.1 g/dL",
        vitals="none",
        history="Not a home medicine",
        location="history of present illness; hospital course; medication list; labs",
        phrases=[
            "venous-thromboembolism prophylaxis",
            "not a home medicine",
            "enoxaparin",
            "2.0",
        ],
    )
    plans.append(stopped)
    traces.append(stopped_trace)
    summary = (
        "Enoxaparin indication is inpatient venous-thromboembolism prophylaxis in the "
        "narrative, the medication record, and the hidden stop decision. It is not labeled "
        "as atrial-fibrillation therapy."
    )
    return chart, _plan_bundle(plans, chart.get("CaseFollowup") or []), traces, summary


def _repair_821(chart: dict[str, Any]) -> Revision:
    _drop_direct_orders(chart)
    for row in _meds(chart, "apixaban"):
        if row.get("context") == "inpatient":
            row["status"] = "held"
            row["held_reason"] = "Held for gastrointestinal bleeding."
    hpi = (
        "A 78-year-old woman with atrial fibrillation is admitted with two days of fatigue "
        "from gastrointestinal bleeding. Home medicines are apixaban 5 MG twice daily, "
        "lisinopril, and atorvastatin. Apixaban was held for the bleeding."
    )
    admission = (
        "Admission note. Gastrointestinal bleeding. Hemoglobin is 8.7 g/dL. Apixaban, used "
        "for atrial fibrillation, was held for gastrointestinal bleeding."
    )
    course = (
        "Gastrointestinal bleeding settled. Hemoglobin rose from 8.7 g/dL to 10.9 g/dL. "
        "Apixaban remained held for gastrointestinal bleeding. Creatinine was 0.9 mg/dL "
        "then 1.2 mg/dL. Discharge systolic blood pressure was 140 mmHg. Gastroenterology "
        "follow-up is scheduled in 7 days."
    )
    _set_text(chart, hpi=hpi, admission=admission, course=course)
    for consult in chart.get("CaseConsult") or []:
        consult["assessment"] = "Gastrointestinal bleeding settled."
        consult["recommendation"] = "Hemoglobin rose from 8.7 g/dL to 10.9 g/dL."
    plans: list[dict[str, Any]] = []
    traces: list[dict[str, Any]] = []
    for medication, indication, phrases in (
        (
            "lisinopril 10 MG Oral Tablet",
            "Essential (primary) hypertension",
            ["lisinopril", "hypertension"],
        ),
        (
            "atorvastatin 40 MG Oral Tablet",
            "Mixed hyperlipidemia",
            ["atorvastatin", "hyperlipidemia"],
        ),
    ):
        _continue_home(
            chart,
            plans,
            traces,
            medication,
            indication,
            phrases,
            "This home medicine was not the held anticoagulant.",
            "Gastrointestinal bleeding settled.",
            "creatinine 1.2 mg/dL",
            "systolic blood pressure 140 mmHg",
        )
    restarted, restarted_trace = _decision(
        chart,
        medication="apixaban 5 MG Oral Tablet",
        action="restart",
        indication="Paroxysmal atrial fibrillation",
        rationale=(
            "Apixaban was held for gastrointestinal bleeding. Bleeding then settled and "
            "hemoglobin rose from 8.7 g/dL to 10.9 g/dL. Atrial fibrillation remains the "
            "indication, creatinine is 1.2 mg/dL, and weight is 83 kg."
        ),
        diagnosis="Paroxysmal atrial fibrillation",
        event="Apixaban was held for gastrointestinal bleeding.",
        labs="hemoglobin 8.7 g/dL then 10.9 g/dL; creatinine 1.2 mg/dL",
        vitals="discharge systolic blood pressure 140 mmHg",
        history="Home apixaban 5 MG twice daily for atrial fibrillation",
        location="history of present illness; hospital course; labs; diagnoses",
        phrases=[
            "apixaban",
            "gastrointestinal bleeding settled",
            "10.9",
            "atrial fibrillation",
            "held for gastrointestinal bleeding",
        ],
    )
    plans.append(restarted)
    traces.append(restarted_trace)
    alternatives = [
        {
            "medication": "apixaban 5 MG Oral Tablet",
            "action": "hold",
            "rationale": (
                "A reviewer may keep apixaban held until the gastroenterology visit in "
                "7 days because the admission hemoglobin was 8.7 g/dL."
            ),
        }
    ]
    summary = (
        "Direct instructions to resume apixaban were removed. The chart shows the bleed, "
        "hemoglobin rise from 8.7 to 10.9 g/dL, the atrial-fibrillation indication, and "
        "gastroenterology follow-up in 7 days. The hidden reference restarts apixaban."
    )
    return (
        chart,
        _plan_bundle(
            plans,
            chart.get("CaseFollowup") or [],
            alternatives=alternatives,
        ),
        traces,
        summary,
    )


_REPAIRS = {
    "VAL-806": _repair_806,
    "VAL-807": _repair_807,
    "VAL-808": _repair_808,
    "VAL-814": _repair_814,
    "VAL-815": _repair_815,
    "VAL-816": _repair_816,
    "VAL-817": _repair_817,
    "VAL-820": _repair_820,
    "VAL-821": _repair_821,
}


_REVISIONS = {
    "VAL-801": _revise_801,
    "VAL-802": _revise_802,
    "VAL-803": _revise_803,
    "VAL-805": _revise_805,
    "VAL-809": _revise_809,
    "VAL-813": _revise_813,
}


def _write_docx(
    path: Path,
    title: str,
    rows: list[dict[str, Any]],
    notice: str | None = None,
) -> None:
    document = Document()
    prepare_form_document(document)
    document.add_heading(title, 0)
    if notice:
        document.add_paragraph(notice)
    document.add_paragraph(
        "Read the resident-facing chart and complete the ratings before the clinician "
        "validation reference. That reference is not shown to residents."
    )
    document.add_paragraph(
        "These charts are clean clinical cases. The resident determines the discharge "
        "medication regimen from the chart."
    )
    for row in rows:
        _add_case(document, row)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    finalize_word_form(path)


def _add_case(document: Any, row: dict[str, Any]) -> None:
    case_id = row["case_id"]
    chart = row["resident"]
    clinical = chart["ClinicalCase"]
    document.add_heading(f"CASE {case_id}", 1)
    document.add_paragraph(clinical.get("one_liner") or "")
    document.add_heading("Resident-facing chart", 2)
    presentation = clinical.get("presentation") or {}
    document.add_paragraph(presentation.get("hpi") or "")
    document.add_paragraph("Diagnoses: " + "; ".join(
        item.get("diagnosis") or "" for item in chart.get("CaseDiagnosis") or []
    ))
    for note in chart.get("CaseNote") or []:
        document.add_paragraph(f"{note.get('note_type')}: {note.get('note_text')}")
    document.add_paragraph("Medications (home and inpatient only):")
    for med in chart.get("CaseMedication") or []:
        document.add_paragraph(
            f"{med.get('context')}: {med.get('drug')} {med.get('dose')} {med.get('route')} "
            f"{med.get('frequency')}. Indication: {med.get('indication') or 'none'}. "
            f"Held reason: {med.get('held_reason') or 'none'}. "
            f"Notes: {med.get('notes') or 'none'}.",
            style="List Bullet",
        )
    document.add_paragraph("Vital signs:")
    for vital in chart.get("CaseVital") or []:
        document.add_paragraph(
            f"{vital.get('timepoint')}: blood pressure {vital.get('bp_systolic')}/"
            f"{vital.get('bp_diastolic')}, heart rate {vital.get('heart_rate')}, "
            f"respiratory rate {vital.get('resp_rate')}, temperature {vital.get('temp_c')}, "
            f"oxygen saturation {vital.get('spo2_percent')}.",
            style="List Bullet",
        )
    document.add_paragraph("Laboratories:")
    for lab in chart.get("CaseLab") or []:
        document.add_paragraph(
            f"{lab.get('timepoint')}: {lab.get('test_name')} {lab.get('value')} {lab.get('unit')}.",
            style="List Bullet",
        )
    if chart.get("CaseWeight"):
        document.add_paragraph("Weights:")
        for weight in chart["CaseWeight"]:
            document.add_paragraph(
                f"{weight.get('timepoint')}: {weight.get('weight_kg')} kg"
                + (
                    f", dry weight {weight.get('dry_weight_kg')} kg"
                    if weight.get("dry_weight_kg")
                    else ""
                ),
                style="List Bullet",
            )
    if chart.get("CaseIntakeOutput"):
        document.add_paragraph("Intake and output:")
        for item in chart["CaseIntakeOutput"]:
            document.add_paragraph(
                f"{item.get('timepoint')}: intake {item.get('intake_ml')} mL, "
                f"output {item.get('output_ml')} mL, net {item.get('net_ml')} mL.",
                style="List Bullet",
            )
    for imaging in chart.get("CaseImaging") or []:
        document.add_paragraph(
            f"Imaging, {imaging.get('study_type')}: {imaging.get('finding')}"
        )
    for consult in chart.get("CaseConsult") or []:
        document.add_paragraph(
            f"Consult, {consult.get('service')}: {consult.get('assessment')} "
            f"{consult.get('recommendation')}"
        )
    for micro in chart.get("CaseMicrobiology") or []:
        document.add_paragraph(
            f"Microbiology: {micro.get('test')} {micro.get('result')} "
            f"{micro.get('organism') or ''} "
            f"{micro.get('notes') or ''}"
        )
    for procedure in chart.get("CaseProcedure") or []:
        document.add_paragraph(
            f"Procedure, {procedure.get('time') or procedure.get('timepoint')}: "
            f"{procedure.get('procedure_name')}. {procedure.get('findings') or ''}"
        )
    for item in chart.get("CaseMonitoring") or []:
        document.add_paragraph(
            f"Monitoring: {item.get('parameter')} ({item.get('frequency')}), "
            f"{item.get('responsible_service') or ''}."
        )
    for item in chart.get("CaseInstruction") or []:
        document.add_paragraph(f"Instruction: {item.get('instruction_text')}")
    for follow in chart.get("CaseFollowup") or []:
        document.add_paragraph(
            f"Follow-up: {follow.get('item')} in {follow.get('timing')} "
            f"with {follow.get('with_service')}."
        )
    document.add_heading("Clinician ratings", 2)
    document.add_paragraph("Complete these ratings before reading the validation reference.")
    _c1(document, case_id)
    _yes_no(document, case_id, "C2", "Decision sufficiency", ("Pass", "Fail"))
    _comment(document, case_id, "c2-comment", "C2 comment if Fail")
    _yes_no(document, case_id, "C3", "Reference-plan validity", ("Pass", "Fail"))
    document.add_paragraph("Disputed medication, if any:")
    _comment(
        document,
        case_id,
        "c3-medication",
        "Disputed medication, action, suggested action, and reason",
    )
    _yes_no(
        document,
        case_id,
        "C4",
        "Alternative acceptable answers",
        ("No", "Yes"),
    )
    document.add_paragraph("Could another discharge regimen also reasonably be correct?")
    _comment(document, case_id, "c4-alternatives", "Acceptable alternatives if Yes")
    document.add_heading("C5 — Missing or misleading information", 3)
    document.add_paragraph(
        "Note missing information, contradictions, misleading statements, unrealistic "
        "medication behavior, inappropriate monitoring, or implausible hospital-course details."
    )
    _comment(document, case_id, "c5-comment", "C5 comments")
    _yes_no(
        document,
        case_id,
        "C6",
        "Resident-level appropriateness",
        (
            "Too easy",
            "Appropriate",
            "Challenging but appropriate",
            "Too difficult or inappropriate",
        ),
    )
    _case_specific_questions(document, case_id)
    _yes_no(document, case_id, "overall", "Overall recommendation", ("Accept", "Revise", "Exclude"))
    _comment(document, case_id, "overall-comment", "Revision comments")
    document.add_heading("CLINICIAN VALIDATION REFERENCE", 2)
    document.add_paragraph("Not shown to residents in the assessment study.")
    document.add_paragraph("Complete the ratings above before using this reference.")
    reference = row["evaluator"]["reference_discharge_plan"]
    for item in reference.get("medications") or []:
        document.add_paragraph(
            f"{item.get('action')}: {item.get('medication')} {item.get('dose')} "
            f"{item.get('route')} {item.get('frequency')}. "
            f"Duration: {item.get('duration') or 'not specified'}. "
            f"Indication: {item.get('indication')}. Rationale: {item.get('rationale')}. "
            f"Monitoring: {item.get('monitoring') or 'see chart'}.",
            style="List Bullet",
        )
    for item in reference.get("acceptable_alternatives") or []:
        document.add_paragraph(
            f"Acceptable alternative: {item.get('action')} {item.get('medication')}. "
            f"{item.get('rationale')}",
            style="List Bullet",
        )
    for item in reference.get("follow_up_requirements") or []:
        document.add_paragraph(
            f"Reference follow-up: {item.get('item')} in {item.get('timing')} "
            f"with {item.get('with_service')}.",
            style="List Bullet",
        )
    for item in row["evaluator"].get("evidence_trace") or []:
        document.add_paragraph(
            f"Evidence for {item.get('medication')}: {item.get('evidence_class')}. "
            f"{item.get('supporting_hospital_event')} {item.get('supporting_labs')} "
            f"Visible in {item.get('resident_visible_evidence_location')}.",
            style="List Bullet",
        )


def _case_specific_questions(document: Any, case_id: str) -> None:
    if case_id == "VAL-805":
        document.add_heading("Case-specific question — HFrEF therapy", 3)
        document.add_paragraph(
            "Given the discharge clinical status, which additional heart-failure therapies, "
            "if any, should be considered required versus acceptable alternatives? "
            "Ejection fraction is 30 percent, systolic blood pressure is 110 mmHg, "
            "creatinine is 0.9 mg/dL, and potassium is 4.3 mmol/L."
        )
        _comment(document, case_id, "hfref-therapies", "Required versus acceptable HFrEF therapies")
    if case_id == "VAL-813":
        document.add_heading("Case-specific question — mycophenolate", 3)
        document.add_paragraph(
            "Given the CMV disease and transplant context, is continuation of mycophenolate "
            "at discharge clinically appropriate, or should the regimen be modified or "
            "temporarily held?"
        )
        _yes_no(
            document,
            case_id,
            "MMF",
            "Mycophenolate at discharge",
            ("Appropriate", "Not appropriate"),
        )
        _comment(document, case_id, "mmf-comment", "Mycophenolate comment")


def _c1(document: Any, case_id: str) -> None:
    document.add_heading("C1 — Clinical plausibility", 3)
    document.add_paragraph(
        "1 = implausible; 2 = substantial revision required; 3 = plausible with minor concern; "
        "4 = fully plausible."
    )
    domains = (
        "presentation",
        "diagnosis",
        "hospital course",
        "laboratory findings",
        "vital signs",
        "medication history",
        "inpatient treatment",
        "discharge context",
        "overall consistency",
    )
    for domain in domains:
        paragraph = document.add_paragraph(f"{domain}: ")
        for score in ("1", "2", "3", "4"):
            paragraph.add_run(f" {score} ")
            append_checkbox(
                paragraph,
                document,
                tag=f"{case_id}-c1-{domain}-{score}",
                alias=f"{case_id} C1 {domain} {score}",
            )


def _yes_no(document: Any, case_id: str, code: str, prompt: str, choices: tuple[str, ...]) -> None:
    document.add_heading(f"{code} — {prompt}", 3)
    paragraph = document.add_paragraph()
    for choice in choices:
        paragraph.add_run(choice + " ")
        append_checkbox(
            paragraph,
            document,
            tag=f"{case_id}-{code}-{choice}",
            alias=f"{case_id} {code} {choice}",
        )


def _comment(document: Any, case_id: str, tag: str, label: str) -> None:
    table = document.add_table(rows=1, cols=1)
    append_rich_text_field(
        table.cell(0, 0),
        document,
        tag=f"{case_id}-{tag}",
        alias=label,
        lines=3,
    )


def _revised_audit(rows: list[dict[str, Any]]) -> str:
    lines = [
        "# Revised seed-guided cases",
        "",
        "Built from the recovered clean charts. Historical frozen files were not edited.",
        "",
        "| Case | Original issue | Revision made | Reference plan rebuilt | "
        "Evidence sufficient | Status |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    issues = {
        "VAL-801": "Delirium unexplained; ibuprofen stopped without a reason",
        "VAL-802": "Delirium unexplained; creatinine and lisinopril hard to interpret",
        "VAL-803": "Delirium unexplained; little discharge decision-making",
        "VAL-805": "Heart-failure course too thin for a diuretic and GDMT decision",
        "VAL-809": "Endocarditis and AKI not convincing; lisinopril continued through AKI",
        "VAL-813": "Valganciclovir already a home medicine before CMV was diagnosed",
    }
    for row in rows:
        sufficient = (
            "yes"
            if row["classes"] and all(item == SUFFICIENT for item in row["classes"])
            else "no"
        )
        lines.append(
            "| {case} | {issue} | {revision} | yes | {sufficient} | {status} |".format(
                case=row["case_id"],
                issue=issues[row["case_id"]],
                revision=row["summary"].replace("|", "/"),
                sufficient=sufficient,
                status=row["status"],
            )
        )
    lines.append("")
    return "\n".join(lines)


def _fresh_audit(rows: list[dict[str, Any]]) -> str:
    lines = [
        "# Clean cases for fresh clinician review",
        "",
        "Untouched rows are the recovered charts that already passed the precheck. "
        "Repaired rows are corrected copies that passed the evidence audit.",
        "",
        "| Case | Clean source | Correction | Evidence | Status |",
        "| --- | --- | --- | --- | --- |",
    ]
    for row in rows:
        if row["case_id"] in REPAIR_IDS:
            if any(item == WEAK for item in row["classes"]):
                evidence = "WEAK_EVIDENCE flagged for inspection"
            elif row["classes"] and all(item == SUFFICIENT for item in row["classes"]):
                evidence = "sufficient"
            else:
                evidence = "not sufficient"
            correction = row["summary"].replace("|", "/")
        else:
            evidence = "recovered reference retained"
            correction = "none"
        lines.append(
            f"| {row['case_id']} | yes | {correction} | {evidence} | {row['status']} |"
        )
    lines.append("")
    return "\n".join(lines)


def _held_audit(rows: list[dict[str, Any]]) -> str:
    lines = [
        "# Held cases after repair",
        "",
        "Cases in this table are not included in a clinician-review document.",
        "Ready for clinician review does not mean clinically validated.",
        "",
        "| Case | Defect | Repair attempted | Remaining issue | Status | "
        "Recommended next action |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    if not rows:
        lines.append("| — | — | — | No case remains held. | — | — |")
        lines.append("")
        return "\n".join(lines)
    for row in rows:
        lines.append(
            "| {case} | {defect} | {repair} | {issue} | {status} | {action} |".format(
                case=row["case_id"],
                defect=row.get("defect", "pre-review defect").replace("|", "/"),
                repair=row["summary"].replace("|", "/"),
                issue=row["finding"].replace("|", "/"),
                status=row["status"],
                action="Keep out of clinician review until the remaining issue is resolved.",
            )
        )
    lines.append("")
    return "\n".join(lines)


def _remaining_audit(rows: list[dict[str, Any]]) -> str:
    lines = [
        "# Remaining recovered clean cases",
        "",
        "These charts were not rewritten. Each file starts from the recovered clean export.",
        "",
        "| Case | Clean source recovered | Error mutation absent | "
        "Answer leakage absent | Precheck finding | Status |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for row in rows:
        resident = row["resident"]
        blob = json.dumps(resident).casefold()
        mutation_absent = (
            "no"
            if "exactly as listed" in blob
            or any(
                item.get("context") == "discharge" for item in resident.get("CaseMedication") or []
            )
            else "yes"
        )
        leakage_absent = "no" if "answer leakage" in row["finding"] else "yes"
        lines.append(
            f"| {row['case_id']} | yes | {mutation_absent} | {leakage_absent} | "
            f"{row['finding'].replace('|', '/')} | {row['status']} |"
        )
    lines.append("")
    return "\n".join(lines)
