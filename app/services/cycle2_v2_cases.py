# ruff: noqa: E501
"""Evidence-gated Round 2 copies for the six cases with Round 1 feedback.

Copies the recovered clean charts and applies only changes that have a reviewer
item and a source. It does not edit frozen files, the recovered export, or the
other 18 cases.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx.document import Document as WordDocument
from docx.shared import RGBColor

from app.services.cycle2_split_casebooks import (
    _banner,
    _bordered_table,
    _chart,
    _round2_form,
)
from app.services.round1_six_case_review import (
    AMBIGUITY,
    completed_reviews,
    write_feedback_files,
)
from app.services.word_controls import finalize_word_form, prepare_form_document
from app.services.word_export import NAVY, _add_body, _add_heading, _new_document, _set_run_font

ROOT = Path(__file__).resolve().parents[2]
CLEAN = ROOT / "exports" / "clean_balanced_seed_set"
CASE_IDS = ("VAL-801", "VAL-802", "VAL-803", "VAL-805", "VAL-809", "VAL-813")

REVIEW_ITEMS: tuple[dict[str, str], ...] = (
    {
        "id": "VAL-801-R1",
        "case": "VAL-801",
        "feedback": "The cause is not revealed during the hospital course.",
        "interpretation": "The reviewer could not tell why the patient was delirious.",
        "response": (
            "We agree that the chart did not give a clinically interpretable precipitant. "
            "We did not add infection, bleeding, or new examination findings, because none "
            "of those facts are in the preserved chart. The revised course states that the "
            "precipitant is not documented and reports the glucose and creatinine values "
            "that are already recorded."
        ),
        "change": "Hospital course states that no precipitant is documented and cites the existing laboratories.",
        "evidence": "CASE_SOURCE laboratories; LITERATURE supports looking for a precipitant but does not name one for this patient.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_resident.json CaseLab; [E1]",
        "decision": "PARTIALLY_ACCEPTED",
    },
    {
        "id": "VAL-801-R2",
        "case": "VAL-801",
        "feedback": "We just learn that the delirium improves with treatment. There are no further details on why he improved.",
        "interpretation": "Improvement is asserted without a mechanism.",
        "response": (
            "We agree the original course did not explain the improvement. The preserved "
            "chart already says the symptoms are improving and records glucose falling from "
            "163 mg/dL to 103 mg/dL with creatinine remaining 1.1 mg/dL then 1.0 mg/dL. "
            "The revision uses those recorded facts and does not invent a treatment response."
        ),
        "change": "Narrative limited to the recorded symptom course and the recorded laboratories.",
        "evidence": "CASE_SOURCE presentation.symptom_course and CaseLab.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_resident.json",
        "decision": "PARTIALLY_ACCEPTED",
    },
    {
        "id": "VAL-801-R3",
        "case": "VAL-801",
        "feedback": "It is also not clear why the ibuprofen was stopped.",
        "interpretation": "The stop has no visible indication.",
        "response": (
            "We agree. The preserved chart stops ibuprofen without a bleeding, renal, or "
            "other indication, while creatinine stays 1.1 mg/dL then 1.0 mg/dL. We removed "
            "the stop from the medication rows and changed the hidden reference from stop "
            "to continue. We did not add a gastrointestinal bleed to justify the old stop."
        ),
        "change": "Ibuprofen held reason removed; reference action changed from stop to continue.",
        "evidence": "CASE_SOURCE creatinine and the absence of a stop indication in the chart.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_evaluator.json reference action stop",
        "decision": "ACCEPTED",
    },
    {
        "id": "VAL-801-R4",
        "case": "VAL-801",
        "feedback": "I would add that he is discovered to have an infection or GI bleed or something that would explain this better.",
        "interpretation": "The reviewer offered examples of a precipitant, not a required diagnosis.",
        "response": (
            "We agree that a precipitant would make the case easier to follow. We did not "
            "add infection or gastrointestinal bleeding. Either addition would create a new "
            "scenario and a new medication problem that the preserved chart does not contain."
        ),
        "change": "No infection and no gastrointestinal bleed were added.",
        "evidence": "No source field documents either diagnosis.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_resident.json CaseDiagnosis",
        "decision": "NOT_ACCEPTED",
    },
    {
        "id": "VAL-801-R5",
        "case": "VAL-801",
        "feedback": "Medication list should be clearer.",
        "interpretation": "The stop was mixed into the history and the medication rows.",
        "response": (
            "We agree. Ibuprofen remains on the home and inpatient lists without a hold "
            "label, together with lisinopril, atorvastatin, and metformin. The resident "
            "chart still does not print a discharge answer list."
        ),
        "change": "Ibuprofen status no longer says held or stopped.",
        "evidence": "CASE_SOURCE medication rows, with the unsupported hold removed.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_resident.json CaseMedication ibuprofen",
        "decision": "ACCEPTED",
    },
    {
        "id": "VAL-801-R6",
        "case": "VAL-801",
        "feedback": "C3 Pass and Fail are both selected, and the C3 comment is blank.",
        "interpretation": "The source selection is ambiguous and cannot be cleaned.",
        "response": (
            "We reproduced both selections and did not choose which one was intended. "
            "No clinical fact was changed because of this item."
        ),
        "change": "None.",
        "evidence": "Historical checkbox states.",
        "locator": "KO Casebook Validation.docx VAL-801 C3",
        "decision": "REQUIRES_CLINICIAN_ADJUDICATION",
    },
    {
        "id": "VAL-802-R1",
        "case": "VAL-802",
        "feedback": "Delirium cause is not revealed (delirium due to known physiologic condition doesn’t make sense).",
        "interpretation": "The diagnosis label does not identify a precipitant.",
        "response": (
            "We agree the label is not an explanation. The preserved chart does not contain "
            "a precipitant, and we did not invent one. The revision says that directly."
        ),
        "change": "Narrative states that no precipitant is recorded.",
        "evidence": "CASE_SOURCE has no precipitant field; [E1] supports seeking one but not naming one.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_resident.json",
        "decision": "PARTIALLY_ACCEPTED",
    },
    {
        "id": "VAL-802-R2",
        "case": "VAL-802",
        "feedback": "It doesn’t appear that anything happened in the hospital course.",
        "interpretation": "The course sentence is about medication history rather than the illness.",
        "response": (
            "We agree. The revision describes the recorded course: confusion improving, "
            "creatinine 1.3 mg/dL then 1.2 mg/dL, and blood pressure 138/69 mmHg then "
            "124/68 mmHg. No new event was added."
        ),
        "change": "Hospital course rewritten from existing vital signs and creatinine only.",
        "evidence": "CASE_SOURCE CaseVital and CaseLab.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_resident.json",
        "decision": "ACCEPTED",
    },
    {
        "id": "VAL-802-R3",
        "case": "VAL-802",
        "feedback": "Creatinine was the only lab value and changed slightly, but it was not clear what the baseline creatinine was.",
        "interpretation": "A baseline is missing, and the change is small.",
        "response": (
            "We agree. No earlier creatinine exists in the preserved chart, so none was "
            "added. The revision states that 1.3 mg/dL and 1.2 mg/dL are the only values."
        ),
        "change": "Narrative states that no baseline creatinine is recorded. No new laboratory was added.",
        "evidence": "CASE_SOURCE CaseLab contains only those two creatinine rows.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_resident.json CaseLab",
        "decision": "PARTIALLY_ACCEPTED",
    },
    {
        "id": "VAL-802-R4",
        "case": "VAL-802",
        "feedback": "Lots of mention of the statin and it said that she was taking the statin, but this was not continued at discharge for no obvious reason. Was the trainee supposed to continue the statin?",
        "interpretation": "The historical chart omitted atorvastatin at discharge and also hinted at the answer.",
        "response": (
            "The omission the reviewer saw was the historical planted discrepancy. The "
            "preserved clean chart already lists atorvastatin at home and in the hospital, "
            "and the hidden reference continues it. We removed sentences that said a "
            "continued statin had been confirmed, because those sentences revealed the "
            "discharge decision. We did not remove atorvastatin from the evidence lists."
        ),
        "change": "Answer-revealing statin sentences removed. Reference remains continue atorvastatin.",
        "evidence": "CASE_SOURCE clean medication rows and reference action continue.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_evaluator.json",
        "decision": "ACCEPTED",
    },
    {
        "id": "VAL-802-R5",
        "case": "VAL-802",
        "feedback": "If creatinine was elevated to point of AKI the lisinopril should be stopped.",
        "interpretation": "The reviewer would hold lisinopril if the creatinine represents AKI.",
        "response": (
            "The chart shows creatinine 1.3 mg/dL then 1.2 mg/dL and no baseline. That "
            "does not establish AKI. We left the hidden reference as continue lisinopril "
            "and recorded a hold as an acceptable alternative rather than inventing a baseline."
        ),
        "change": "Acceptable alternative added: hold lisinopril. Reference action remains continue.",
        "evidence": "CASE_SOURCE creatinine; [E3] supports reviewing an ACE inhibitor during AKI but does not prove AKI here.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_resident.json CaseLab",
        "decision": "REQUIRES_CLINICIAN_ADJUDICATION",
    },
    {
        "id": "VAL-803-R1",
        "case": "VAL-803",
        "feedback": "Delirium diagnosis is unclear and the cause is not revealed during the hospital course.",
        "interpretation": "Only C1 was completed. The precipitant is absent from the chart.",
        "response": (
            "We agree the completed C1 comment identifies a missing precipitant. C2 through "
            "the overall recommendation were not completed, and we did not infer them. "
            "We did not add a sodium series or a thiazide-toxicity story, because sodium "
            "is not in the preserved laboratories. The revision states that no precipitant "
            "is documented. The existing sentence that cognition returned to baseline is unchanged."
        ),
        "change": "One sentence added: the chart does not name a precipitant. No new laboratory.",
        "evidence": "CASE_SOURCE CaseLab has glucose, creatinine, and potassium only.",
        "locator": "exports/clean_balanced_seed_set/VAL-803_resident.json CaseLab",
        "decision": "PARTIALLY_ACCEPTED",
    },
    {
        "id": "VAL-803-R2",
        "case": "VAL-803",
        "feedback": "Her creatinine changes so this might be considered clinically relevant or why the patient receives only 7 days of lisinopril.",
        "interpretation": "The reviewer is asking about the 7-day supply on the historical injected chart.",
        "response": (
            "The 7-day supply is the historical injected discrepancy. The preserved clean "
            "reference already records lisinopril for 30 days, with creatinine 1.0 mg/dL "
            "then 1.2 mg/dL. We did not shorten that supply and did not treat the creatinine "
            "change as a reason to stop lisinopril."
        ),
        "change": "No medication change. The v2 case keeps the clean 30-day lisinopril duration.",
        "evidence": "CASE_SOURCE evaluator duration 30 days, not the historical 7-day injected value.",
        "locator": "exports/clean_balanced_seed_set/VAL-803_evaluator.json lisinopril duration",
        "decision": "ACCEPTED",
    },
    {
        "id": "VAL-805-R1",
        "case": "VAL-805",
        "feedback": "The weight should be closer to dry weight at the time of discharge. The discharge weight not being close to the patient’s dry weight is confusing.",
        "interpretation": "Discharge weight 78 kg and dry weight 73 kg were hard to interpret.",
        "response": (
            "We agree the gap matters. Those weights are the preserved measurements, so we "
            "did not replace them with a new series. The revision states 81 kg on admission, "
            "78 kg at discharge, and dry weight 73 kg, so the 5 kg gap remains visible."
        ),
        "change": "Narrative now quotes the recorded weights. Weights themselves are unchanged.",
        "evidence": "CASE_SOURCE CaseWeight.",
        "locator": "exports/clean_balanced_seed_set/VAL-805_resident.json CaseWeight",
        "decision": "ACCEPTED",
    },
    {
        "id": "VAL-805-R2",
        "case": "VAL-805",
        "feedback": "The lasix would be given IV and increased.",
        "interpretation": "The reviewer wants the inpatient loop diuretic to differ from the home oral dose.",
        "response": (
            "We agree that acute decompensated heart failure is often treated with intravenous "
            "loop diuretic. The preserved chart records oral furosemide 40 mg once daily at "
            "home and in the hospital, and the project case does not contain a different "
            "inpatient dose. We did not invent an intravenous dose or a dose increase. "
            "The reference continues the recorded oral dose. A different inpatient dose "
            "remains for clinician adjudication."
        ),
        "change": "No dose or route change. Narrative states that the recorded inpatient dose matches the home dose.",
        "evidence": "[E2] supports intravenous loop diuretic in principle. CASE_SOURCE has only oral 40 mg once daily.",
        "locator": "exports/clean_balanced_seed_set/VAL-805_resident.json CaseMedication furosemide",
        "decision": "REQUIRES_CLINICIAN_ADJUDICATION",
    },
    {
        "id": "VAL-805-R3",
        "case": "VAL-805",
        "feedback": "It would be more realistic to have several days of input and output data.",
        "interpretation": "The original prose claimed serial intake and output, but only one day is stored.",
        "response": (
            "We agree the course overstated the record. Only hospital day 3 is stored, with "
            "net -1045 mL. We removed the claim of serial days and did not invent additional days."
        ),
        "change": "Narrative corrected to the single recorded intake-and-output day.",
        "evidence": "CASE_SOURCE CaseIntakeOutput.",
        "locator": "exports/clean_balanced_seed_set/VAL-805_resident.json CaseIntakeOutput",
        "decision": "PARTIALLY_ACCEPTED",
    },
    {
        "id": "VAL-805-R4",
        "case": "VAL-805",
        "feedback": "It would be helpful to have cardiology notes or echocardiogram data, and the trainee should consider additional GDMT such as ACE/ARB, SGLT2i, or MRA.",
        "interpretation": "The reviewer wants both more cardiac data and a real medication decision.",
        "response": (
            "The preserved chart has a cardiology consultation and a chest radiograph showing "
            "pulmonary edema. It does not record an ejection fraction. We surfaced those "
            "existing items and did not add an ejection fraction or start a new drug class. "
            "Additional HFrEF therapy is an acceptable alternative, not a required reference "
            "medication, because the chart does not contain the measurement that would force one choice."
        ),
        "change": "Narrative cites the existing consult and radiograph. Acceptable alternative added for additional GDMT. No new drug and no ejection fraction.",
        "evidence": "CASE_SOURCE CaseConsult and CaseImaging; [E2] describes GDMT classes but does not assign this patient's ejection fraction.",
        "locator": "exports/clean_balanced_seed_set/VAL-805_resident.json CaseConsult",
        "decision": "PARTIALLY_ACCEPTED",
    },
    {
        "id": "VAL-805-R5",
        "case": "VAL-805",
        "feedback": "Medication regimen was marked both 2 and 3.",
        "interpretation": "The historical row has two selected scores.",
        "response": "We reproduced both scores and did not collapse them to one number.",
        "change": "None.",
        "evidence": "Historical checkbox states.",
        "locator": "KO Casebook Validation.docx VAL-805 C1 medication regimen",
        "decision": "REQUIRES_CLINICIAN_ADJUDICATION",
    },
    {
        "id": "VAL-809-R1",
        "case": "VAL-809",
        "feedback": "Should present with fevers, have a history of mechanical valve or poor dentition that would predispose to endocarditis.",
        "interpretation": "The reviewer wants a predisposition and a febrile presentation. Valve and dentition are examples.",
        "response": (
            "We agree a predisposition is not stated. We did not add fever, a prosthetic "
            "valve, or a dental procedure. Temperature is recorded as 36.80°C, and no valve "
            "or dental history is in the chart. Changing those facts would fabricate history. "
            "The revision instead makes the existing vegetation, blood-culture result, and "
            "ceftriaxone course visible."
        ),
        "change": "No new temperature, valve, or dental history. Existing endocarditis evidence is stated in the narrative.",
        "evidence": "CASE_SOURCE CaseVital temperature, CaseImaging vegetation, CaseMicrobiology; [E4].",
        "locator": "exports/clean_balanced_seed_set/VAL-809_resident.json",
        "decision": "PARTIALLY_ACCEPTED",
    },
    {
        "id": "VAL-809-R2",
        "case": "VAL-809",
        "feedback": "AKI was unexplained and lisinopril was continued despite AKI. Lisinopril should be held if the patient has an AKI.",
        "interpretation": "The reviewer reads creatinine 1.3 mg/dL as AKI.",
        "response": (
            "Creatinine is 1.3 mg/dL on admission and 0.8 mg/dL at discharge. No separate "
            "baseline is stored, so we did not label 1.3 mg/dL as proven AKI and did not "
            "change the reference from continue to hold. A hold while creatinine is 1.3 mg/dL "
            "is an acceptable alternative."
        ),
        "change": "Acceptable alternative added: hold lisinopril. Reference action remains continue.",
        "evidence": "CASE_SOURCE CaseLab; [E3].",
        "locator": "exports/clean_balanced_seed_set/VAL-809_resident.json CaseLab",
        "decision": "REQUIRES_CLINICIAN_ADJUDICATION",
    },
    {
        "id": "VAL-809-R3",
        "case": "VAL-809",
        "feedback": "Patient of this age would be very likely to have more than 2 home medications.",
        "interpretation": "The reviewer wants a longer home list.",
        "response": (
            "We agree the list is short. We did not add medicines. The preserved home list "
            "contains only lisinopril and atorvastatin, and a longer list would be a new fact."
        ),
        "change": "No home medication added.",
        "evidence": "CASE_SOURCE CaseMedication home context.",
        "locator": "exports/clean_balanced_seed_set/VAL-809_resident.json CaseMedication",
        "decision": "NOT_ACCEPTED",
    },
    {
        "id": "VAL-813-R1",
        "case": "VAL-813",
        "feedback": "Patient is admitted for new diagnosis of CMV colitis, but was already on treatment for this (valganciclovir) prior to admission, which does not make sense.",
        "interpretation": "A new diagnosis and pre-admission treatment cannot both be the story.",
        "response": (
            "We agree the combination is incoherent if the case is read as a first diagnosis. "
            "Valganciclovir 900 mg twice daily is a home medicine in the preserved chart. We "
            "did not delete it. The revision says the patient was already taking it before "
            "admission, so this stay is not the first dose for a diagnosis made during the stay. "
            "We also removed the sentence that called outpatient continuation the intended answer."
        ),
        "change": "Narrative chronology aligned with the home medication list. Valganciclovir not removed.",
        "evidence": "CASE_SOURCE home valganciclovir row; [E5] describes treatment after CMV disease and does not require deleting a recorded home dose.",
        "locator": "exports/clean_balanced_seed_set/VAL-813_resident.json CaseMedication valganciclovir",
        "decision": "ACCEPTED",
    },
    {
        "id": "VAL-813-R2",
        "case": "VAL-813",
        "feedback": "The medications are much too simplified for a post transplant patient.",
        "interpretation": "A typical regimen would include more immunosuppressants.",
        "response": (
            "We agree a transplant regimen often includes more than tacrolimus. Mycophenolate "
            "and prednisone are not in this chart. We did not add them. Whether the case is "
            "too narrow for the study remains for the clinician."
        ),
        "change": "No immunosuppressant added.",
        "evidence": "CASE_SOURCE home list is tacrolimus, valganciclovir, amlodipine, and atorvastatin.",
        "locator": "exports/clean_balanced_seed_set/VAL-813_resident.json CaseMedication",
        "decision": "REQUIRES_CLINICIAN_ADJUDICATION",
    },
    {
        "id": "VAL-813-R3",
        "case": "VAL-813",
        "feedback": "The labs are incomplete and the patient has a change in potassium without obvious cause or indication.",
        "interpretation": "Potassium falls from 4.7 mmol/L to 3.9 mmol/L without an explanation.",
        "response": (
            "We agree no cause is written. Both values are already in the chart. We stated "
            "that no cause is documented and did not add a new laboratory or an invented cause."
        ),
        "change": "Narrative quotes both potassium values and states that no cause is recorded.",
        "evidence": "CASE_SOURCE CaseLab potassium rows.",
        "locator": "exports/clean_balanced_seed_set/VAL-813_resident.json CaseLab",
        "decision": "PARTIALLY_ACCEPTED",
    },
    {
        "id": "VAL-813-R4",
        "case": "VAL-813",
        "feedback": "Valganciclovir should not be an admission medication. Since the patient was on valganciclovir on admission, it would make you think that you need to change to an alternative regimen.",
        "interpretation": "The reviewer would remove the home drug or switch therapy.",
        "response": (
            "We do not have a source basis for deleting the home drug or for choosing a "
            "different antiviral. The recorded dose remains 900 mg twice daily, which is the "
            "charted dose at creatinine 1.2 mg/dL then 1.0 mg/dL. Switching or stopping it "
            "is an acceptable alternative for the clinician, not a required reference change."
        ),
        "change": "Acceptable alternative added: stop or switch valganciclovir. Reference action remains continue.",
        "evidence": "CASE_SOURCE dose and creatinine; [E5].",
        "locator": "exports/clean_balanced_seed_set/VAL-813_resident.json",
        "decision": "REQUIRES_CLINICIAN_ADJUDICATION",
    },
)

LEDGER: tuple[dict[str, str], ...] = (
    {
        "id": "C801-1",
        "case": "VAL-801",
        "field": "home_medication",
        "old": "ibuprofen status held; held reason Stopped during this admission.",
        "new": "ibuprofen remains on the home and inpatient lists without a hold.",
        "why": "No stop indication is present, and creatinine is stable.",
        "item": "VAL-801-R3",
        "class": "CASE_SOURCE",
        "source": "Clean VAL-801 medication row and creatinine 1.1 then 1.0 mg/dL.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_resident.json",
        "inference": "No. The stop is removed because the chart does not support it.",
    },
    {
        "id": "C801-2",
        "case": "VAL-801",
        "field": "reference_action",
        "old": "stop ibuprofen",
        "new": "continue ibuprofen",
        "why": "The hidden stop was not supported by a documented indication.",
        "item": "VAL-801-R3",
        "class": "CASE_SOURCE",
        "source": "Clean evaluator reference action stop, with no matching diagnosis.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_evaluator.json",
        "inference": "No.",
    },
    {
        "id": "C801-3",
        "case": "VAL-801",
        "field": "hospital_course",
        "old": "Course discusses only a collateral medication list and stopped medicines.",
        "new": "Course states that no precipitant is documented and quotes the recorded glucose and creatinine.",
        "why": "The reviewer could not see a cause or an explanation of improvement.",
        "item": "VAL-801-R1",
        "class": "CASE_SOURCE",
        "source": "Existing CaseLab values. No new precipitant.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_resident.json CaseLab",
        "inference": "The statement that a precipitant is absent describes the chart. It is not a new diagnosis.",
    },
    {
        "id": "C801-4",
        "case": "VAL-801",
        "field": "hospital_course",
        "old": "Improvement is asserted without the recorded laboratory course.",
        "new": "The course quotes glucose 163 then 103 mg/dL and creatinine 1.1 then 1.0 mg/dL.",
        "why": "Those values are the only recorded explanation of the hospital stay.",
        "item": "VAL-801-R2",
        "class": "CASE_SOURCE",
        "source": "Existing CaseLab.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_resident.json CaseLab",
        "inference": "No.",
    },
    {
        "id": "C801-5",
        "case": "VAL-801",
        "field": "home_medication",
        "old": "Ibuprofen was labeled stopped inside the history.",
        "new": "The history no longer says ibuprofen was stopped.",
        "why": "The medication list should match the unsupported-stop removal.",
        "item": "VAL-801-R5",
        "class": "CASE_SOURCE",
        "source": "Same ibuprofen row as C801-1.",
        "locator": "exports/clean_balanced_seed_set/VAL-801_resident.json CaseMedication",
        "inference": "No.",
    },
    {
        "id": "C802-0",
        "case": "VAL-802",
        "field": "diagnostic_context",
        "old": "The diagnosis label is used without saying that no precipitant is recorded.",
        "new": "The narrative states that the chart does not name a precipitant.",
        "why": "The reviewer rejected the diagnosis label as an explanation.",
        "item": "VAL-802-R1",
        "class": "CASE_SOURCE",
        "source": "No precipitant is stored on the clean chart.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_resident.json",
        "inference": "No precipitant was invented.",
    },
    {
        "id": "C802-1",
        "case": "VAL-802",
        "field": "hospital_course",
        "old": "Course says a collateral list was verified and mentions medicines that had been stopped.",
        "new": "Course quotes confusion improving, creatinine 1.3 then 1.2 mg/dL, and blood pressure 138/69 then 124/68 mmHg, and states that no baseline creatinine is recorded.",
        "why": "The reviewer found the course empty and the baseline missing.",
        "item": "VAL-802-R2",
        "class": "CASE_SOURCE",
        "source": "Existing vital signs and creatinine. No baseline was added.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_resident.json",
        "inference": "No new measurement.",
    },
    {
        "id": "C802-1b",
        "case": "VAL-802",
        "field": "lab_result",
        "old": "Creatinine 1.3 then 1.2 mg/dL without a statement that baseline is absent.",
        "new": "The same two values, with an explicit statement that no baseline is recorded.",
        "why": "The reviewer asked for the baseline and none exists to add.",
        "item": "VAL-802-R2",
        "class": "CASE_SOURCE",
        "source": "CaseLab has two creatinine rows and no earlier value.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_resident.json CaseLab",
        "inference": "No new laboratory.",
    },
    {
        "id": "C802-1c",
        "case": "VAL-802",
        "field": "lab_result",
        "old": "Baseline creatinine was not discussed.",
        "new": "Narrative states that no baseline creatinine is recorded.",
        "why": "Reviewer item VAL-802-R3.",
        "item": "VAL-802-R3",
        "class": "CASE_SOURCE",
        "source": "CaseLab.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_resident.json CaseLab",
        "inference": "No.",
    },
    {
        "id": "C802-2",
        "case": "VAL-802",
        "field": "diagnostic_context",
        "old": "Narrative says a verified statin was confirmed and a continued statin was confirmed.",
        "new": "Those answer-revealing sentences are removed. Atorvastatin remains on the home and inpatient lists.",
        "why": "The reviewer could not tell whether the statin sentence was the answer.",
        "item": "VAL-802-R4",
        "class": "CASE_SOURCE",
        "source": "Clean medication rows already contain atorvastatin. The continue action is the clean reference.",
        "locator": "exports/clean_balanced_seed_set/VAL-802_evaluator.json",
        "inference": "No.",
    },
    {
        "id": "C802-3",
        "case": "VAL-802",
        "field": "reference_action",
        "old": "No acceptable alternative.",
        "new": "Hold lisinopril is an acceptable alternative. The reference still continues it.",
        "why": "AKI is not established without a baseline.",
        "item": "VAL-802-R5",
        "class": "LITERATURE",
        "source": "[E3] KDIGO AKI guidance on reviewing ACE inhibitors, applied only as an alternative.",
        "locator": "https://kdigo.org/guidelines/acute-kidney-injury/",
        "inference": "Yes. The hold is not required.",
    },
    {
        "id": "C803-1",
        "case": "VAL-803",
        "field": "hospital_course",
        "old": "The patient returned to cognitive baseline.",
        "new": "Same sentence, plus a statement that the chart does not name a precipitant.",
        "why": "The only completed comment says the cause was not revealed.",
        "item": "VAL-803-R1",
        "class": "CASE_SOURCE",
        "source": "CaseLab has no sodium and no named precipitant.",
        "locator": "exports/clean_balanced_seed_set/VAL-803_resident.json CaseLab",
        "inference": "The added sentence reports an absence. No laboratory was created.",
    },
    {
        "id": "C803-2",
        "case": "VAL-803",
        "field": "reference_action",
        "old": "Historical injected chart showed lisinopril supply of 7 days.",
        "new": "Clean reference duration of 30 days is kept. Creatinine is not used to shorten it.",
        "why": "The 7-day supply was the planted discrepancy, not the preserved chart.",
        "item": "VAL-803-R2",
        "class": "CASE_SOURCE",
        "source": "Evaluator lisinopril duration is 30 days.",
        "locator": "exports/clean_balanced_seed_set/VAL-803_evaluator.json",
        "inference": "No.",
    },
    {
        "id": "C805-1",
        "case": "VAL-805",
        "field": "hospital_course",
        "old": "Prose says daily weights and serial intake and output were used.",
        "new": "Prose quotes admission weight 81 kg, discharge weight 78 kg, dry weight 73 kg, and the single net -1045 mL day.",
        "why": "The reviewer could not reconcile discharge weight with dry weight, and the prose overstated the record.",
        "item": "VAL-805-R1",
        "class": "CASE_SOURCE",
        "source": "CaseWeight and CaseIntakeOutput.",
        "locator": "exports/clean_balanced_seed_set/VAL-805_resident.json",
        "inference": "No weight was changed.",
    },
    {
        "id": "C805-1b",
        "case": "VAL-805",
        "field": "hospital_course",
        "old": "Prose claimed serial intake and output.",
        "new": "Prose cites the single hospital-day-3 net of -1045 mL.",
        "why": "Additional intake-and-output days were not in the source.",
        "item": "VAL-805-R3",
        "class": "CASE_SOURCE",
        "source": "CaseIntakeOutput has one row.",
        "locator": "exports/clean_balanced_seed_set/VAL-805_resident.json CaseIntakeOutput",
        "inference": "No new fluid balance was created.",
    },
    {
        "id": "C805-2",
        "case": "VAL-805",
        "field": "consultation",
        "old": "Cardiology consult and chest radiograph were present but not described in the course.",
        "new": "Course names the existing radiograph and the existing cardiology recommendation. No ejection fraction was added.",
        "why": "The reviewer asked for cardiology or echocardiogram information that the chart can actually support.",
        "item": "VAL-805-R4",
        "class": "CASE_SOURCE",
        "source": "CaseConsult cardiology and CaseImaging chest radiograph.",
        "locator": "exports/clean_balanced_seed_set/VAL-805_resident.json",
        "inference": "No.",
    },
    {
        "id": "C805-3",
        "case": "VAL-805",
        "field": "reference_action",
        "old": "Continue oral furosemide, metoprolol succinate, and atorvastatin, with no alternative.",
        "new": "Same required medicines. Additional ACE inhibitor, ARB, SGLT2 inhibitor, or MRA is an acceptable alternative only.",
        "why": "The reviewer asked for a GDMT decision. The chart has no ejection fraction, so one class was not made mandatory.",
        "item": "VAL-805-R4",
        "class": "LITERATURE",
        "source": "[E2] 2022 AHA/ACC/HFSA heart-failure guideline, GDMT classes.",
        "locator": "https://doi.org/10.1161/CIR.0000000000001063",
        "inference": "Yes. The alternative is not a required drug and no new measurement was created.",
    },
    {
        "id": "C809-1",
        "case": "VAL-809",
        "field": "diagnostic_context",
        "old": "Narrative says the patient is discharged on parenteral therapy and does not mention the vegetation or cultures.",
        "new": "Narrative quotes the vegetation, gram-positive cocci then no growth, ceftriaxone 2000 mg IV daily, temperature 36.80°C, and creatinine 1.3 then 0.8 mg/dL. It does not add fever or a valve.",
        "why": "The reviewer asked for more endocarditis detail. Only recorded detail was added to the prose.",
        "item": "VAL-809-R1",
        "class": "CASE_SOURCE",
        "source": "CaseImaging, CaseMicrobiology, CaseMedication ceftriaxone, CaseVital, CaseLab.",
        "locator": "exports/clean_balanced_seed_set/VAL-809_resident.json",
        "inference": "No.",
    },
    {
        "id": "C809-2",
        "case": "VAL-809",
        "field": "reference_action",
        "old": "Continue lisinopril.",
        "new": "Continue lisinopril, with hold as an acceptable alternative while creatinine is 1.3 mg/dL.",
        "why": "The reviewer would hold lisinopril for AKI, but no baseline proves AKI.",
        "item": "VAL-809-R2",
        "class": "LITERATURE",
        "source": "[E3] as an alternative only.",
        "locator": "exports/clean_balanced_seed_set/VAL-809_resident.json CaseLab",
        "inference": "Yes. The hold is not required.",
    },
    {
        "id": "C813-1",
        "case": "VAL-813",
        "field": "diagnostic_context",
        "old": "Narrative can be read as a new CMV diagnosis while valganciclovir is already a home medicine, and it says outpatient continuation is the intended answer.",
        "new": "Narrative says valganciclovir was already a home medicine before admission. The intended-answer sentence was removed. Potassium 4.7 then 3.9 mmol/L is quoted and left unexplained.",
        "why": "The reviewer identified contradictory timing and an unexplained potassium change.",
        "item": "VAL-813-R1",
        "class": "CASE_SOURCE",
        "source": "Home valganciclovir row and potassium rows.",
        "locator": "exports/clean_balanced_seed_set/VAL-813_resident.json",
        "inference": "No drug was added or removed.",
    },
    {
        "id": "C813-1b",
        "case": "VAL-813",
        "field": "lab_result",
        "old": "Potassium 4.7 then 3.9 mmol/L without a statement that no cause is recorded.",
        "new": "The same values, with an explicit statement that no cause is recorded.",
        "why": "The reviewer noted the potassium change and no cause is in the chart.",
        "item": "VAL-813-R3",
        "class": "CASE_SOURCE",
        "source": "CaseLab potassium rows.",
        "locator": "exports/clean_balanced_seed_set/VAL-813_resident.json CaseLab",
        "inference": "No cause was invented.",
    },
    {
        "id": "C813-2",
        "case": "VAL-813",
        "field": "reference_action",
        "old": "Continue valganciclovir 900 mg twice daily.",
        "new": "Same required action. Stopping or switching the antiviral is an acceptable alternative.",
        "why": "The reviewer would not keep a drug that looks like pre-diagnosis treatment. The home row is source data, so deletion was not forced.",
        "item": "VAL-813-R4",
        "class": "LITERATURE",
        "source": "[E5] AST CMV guideline, treatment after documented disease.",
        "locator": "https://doi.org/10.1111/ctr.13512",
        "inference": "Yes. The alternative is not required.",
    },
)

EVIDENCE_NOTES = {
    "E1": "Inouye SK, Westendorp RGJ, Saczynski JS. Delirium in elderly people. Lancet. 2014;383:911-922. doi:10.1016/S0140-6736(13)60688-1. Delirium assessment looks for precipitants. It does not assign this patient's precipitant.",
    "E2": "Heidenreich PA, et al. 2022 AHA/ACC/HFSA Guideline for the Management of Heart Failure. Circulation. 2022;145:e895-e1032. doi:10.1161/CIR.0000000000001063. Intravenous loop diuretic is usual for congestion, and HFrEF therapy includes several drug classes. The guideline does not set this patient's dose or ejection fraction.",
    "E3": "KDIGO Clinical Practice Guideline for Acute Kidney Injury. Kidney International Supplements. 2012;2:1-138. https://kdigo.org/guidelines/acute-kidney-injury/. ACE inhibitors are reviewed during AKI. A single creatinine without a baseline does not by itself prove AKI.",
    "E4": "Baddour LM, et al. Infective Endocarditis in Adults: Diagnosis, Antimicrobial Therapy, and Management of Complications. A Scientific Statement From the American Heart Association. Circulation. 2015;132:1435-1486. doi:10.1161/CIR.0000000000000296. Predispositions include prosthetic valves and some dental settings. This chart's recorded findings are a vegetation and gram-positive cocci.",
    "E5": "Razonable RR, Humar A. Cytomegalovirus in solid organ transplant recipients—Guidelines of the American Society of Transplantation Infectious Diseases Community of Practice. Clin Transplant. 2019;33:e13512. doi:10.1111/ctr.13512. Treatment follows documented CMV disease. It does not require deleting a valganciclovir row that the chart already records as a home medicine.",
}


def write_revision_package(feedback_dir: Path, case_dir: Path) -> Path:
    """Write the extract, the matrix, the ledger, the six cases, and the codebook."""
    feedback_dir.mkdir(parents=True, exist_ok=True)
    write_feedback_files(feedback_dir)
    (feedback_dir / "ROUND1_REVIEWER_RESPONSE_MATRIX.md").write_text(
        _matrix(),
        encoding="utf-8",
    )
    (feedback_dir / "REVISION_EVIDENCE_LEDGER.md").write_text(_ledger(), encoding="utf-8")
    _evidence_gate()
    case_dir.mkdir(parents=True, exist_ok=True)
    for case_id in CASE_IDS:
        resident, evaluator = _build_case(case_id)
        if "reference_discharge_plan" in resident:
            raise ValueError(f"{case_id} resident chart contains the reference")
        (case_dir / f"{case_id}_resident.json").write_text(
            json.dumps(resident, indent=2) + "\n",
            encoding="utf-8",
        )
        (case_dir / f"{case_id}_evaluator.json").write_text(
            json.dumps(evaluator, indent=2) + "\n",
            encoding="utf-8",
        )
    document_path = _write_codebook(case_dir)
    (case_dir / "README.md").write_text(_readme(), encoding="utf-8")
    (case_dir / "MANIFEST.md").write_text(_manifest(), encoding="utf-8")
    return document_path


def _evidence_gate() -> None:
    """Accepted changes must name a modification, a source, and a locator."""
    for item in REVIEW_ITEMS:
        if item["decision"] not in {
            "ACCEPTED",
            "PARTIALLY_ACCEPTED",
            "NOT_ACCEPTED",
            "REQUIRES_CLINICIAN_ADJUDICATION",
        }:
            raise ValueError(item["id"])
        if item["decision"] in {"ACCEPTED", "PARTIALLY_ACCEPTED"}:
            if not item["change"] or not item["evidence"] or not item["locator"]:
                raise ValueError(item["id"])
            if item["change"] != "None." and not any(row["item"] == item["id"] for row in LEDGER):
                raise ValueError(f"{item['id']} has no ledger row")
    for row in LEDGER:
        if not row["locator"] or not row["source"]:
            raise ValueError(row["id"])


def _build_case(case_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
    resident = json.loads((CLEAN / f"{case_id}_resident.json").read_text(encoding="utf-8"))
    evaluator = json.loads((CLEAN / f"{case_id}_evaluator.json").read_text(encoding="utf-8"))
    _NARRATIVE[case_id](resident)
    _NARRATIVE[case_id](evaluator)
    _REFERENCE[case_id](evaluator)
    resident.pop("reference_discharge_plan", None)
    return resident, evaluator


def _set_note(case: dict[str, Any], note_type: str, text: str) -> None:
    for note in case.get("CaseNote") or []:
        if note.get("note_type") == note_type:
            note["note_text"] = text
            return
    raise KeyError(note_type)


def _set_hpi(case: dict[str, Any], text: str) -> None:
    case["ClinicalCase"]["presentation"]["hpi"] = text


def _patch_801(case: dict[str, Any]) -> None:
    _set_hpi(
        case,
        "A 75-year-old Male is admitted with delirium. Confusion and fatigue have been "
        "present for several days and are described as improving after treatment. Recorded "
        "home medicines are lisinopril, atorvastatin, metformin, and ibuprofen. The chart "
        "does not document a precipitant for the delirium. Admission glucose is 163 mg/dL "
        "and discharge glucose is 103 mg/dL. Creatinine is 1.1 mg/dL on admission and "
        "1.0 mg/dL at discharge. He could not give a reliable medication history at "
        "admission. A collateral home-medication list was verified later.",
    )
    _set_note(
        case,
        "admission",
        "Admission note for a 75-year-old Male with delirium. Symptoms are confusion and "
        "fatigue for several days, described as improving after treatment. Home medicines "
        "are lisinopril, atorvastatin, metformin, and ibuprofen. No precipitant is "
        "documented. A collateral medication list was verified after admission.",
    )
    _set_note(
        case,
        "hospital_course",
        "A collateral medication list was verified after admission. Confusion is described "
        "as improving. The chart does not name a precipitant. Glucose is 163 mg/dL then "
        "103 mg/dL. Creatinine is 1.1 mg/dL then 1.0 mg/dL. No indication for stopping "
        "ibuprofen is recorded.",
    )
    for medication in case.get("CaseMedication") or []:
        if "ibuprofen" in str(medication.get("drug")):
            if medication.get("context") == "home":
                medication["status"] = "home"
            elif medication.get("context") == "inpatient":
                medication["status"] = "active"
            medication["held_reason"] = None


def _patch_802(case: dict[str, Any]) -> None:
    _set_hpi(
        case,
        "A 81-year-old Female is admitted with delirium. Confusion has been present for "
        "two days and is described as improving after treatment. Recorded home medicines "
        "are lisinopril 10 MG Oral Tablet and atorvastatin 40 MG Oral Tablet. The chart "
        "does not name a precipitant. The only creatinine values are 1.3 mg/dL on "
        "admission and 1.2 mg/dL at discharge. No earlier baseline is recorded. Systolic "
        "blood pressure is 138 mmHg on admission and 124 mmHg at discharge.",
    )
    _set_note(
        case,
        "admission",
        "Admission note for a 81-year-old Female with delirium. Confusion for two days, "
        "described as improving after treatment. Home medicines are lisinopril and "
        "atorvastatin. No precipitant and no earlier creatinine are recorded.",
    )
    _set_note(
        case,
        "hospital_course",
        "No new diagnosis is recorded. Creatinine is 1.3 mg/dL then 1.2 mg/dL. Blood "
        "pressure is 138/69 mmHg then 124/68 mmHg. Atorvastatin remains on the home and "
        "inpatient medication lists. No baseline creatinine is recorded.",
    )


def _patch_803(case: dict[str, Any]) -> None:
    for note in case.get("CaseNote") or []:
        if note.get("note_type") == "hospital_course":
            text = str(note.get("note_text") or "")
            addition = " The chart does not name a precipitant for the delirium."
            if addition.strip() not in text:
                note["note_text"] = text.rstrip() + addition


def _patch_805(case: dict[str, Any]) -> None:
    course = (
        "Congestion was treated with inpatient diuresis. Recorded weight is 81 kg on "
        "admission and 78 kg at discharge, with dry weight 73 kg. One intake-and-output "
        "day is recorded, hospital day 3, net -1045 mL. Inpatient furosemide is oral "
        "40 mg once daily, the same dose and route recorded at home. Chest radiograph "
        "shows pulmonary edema without pneumonia. A cardiology consultation is recorded. "
        "No ejection fraction is recorded."
    )
    _set_note(case, "hospital_course", course)
    _set_hpi(
        case,
        "A 68-year-old Female is admitted with acute systolic heart failure. Dyspnea, "
        "edema, and orthopnea have been present for one week and are described as "
        "improving after treatment. Home medicines are oral furosemide 40 mg once daily, "
        "atorvastatin, and metoprolol succinate 25 mg once daily. The inpatient furosemide "
        "row is the same oral dose. Discharge weight is 78 kg and the documented dry "
        "weight is 73 kg.",
    )
    _set_note(
        case,
        "admission",
        "Admission note for a 68-year-old Female with acute systolic heart failure. "
        "Symptoms are dyspnea, edema, and orthopnea for one week. Home and inpatient "
        "furosemide are both oral 40 mg once daily in the recorded medication list.",
    )


def _patch_809(case: dict[str, Any]) -> None:
    text = (
        "A 82-year-old Male is admitted with infective endocarditis. The recorded symptom "
        "is fatigue for one week, described as improving after treatment. Temperature is "
        "36.80°C on admission and at discharge. Home medicines are lisinopril and "
        "atorvastatin. Ceftriaxone 2000 mg intravenously once daily was started during "
        "the admission. Transthoracic echocardiogram shows a mobile echodensity consistent "
        "with a vegetation and preserved ventricular function. Admission blood culture grew "
        "gram-positive cocci. A later culture showed no growth. Creatinine is 1.3 mg/dL "
        "on admission and 0.8 mg/dL at discharge. The chart does not record fever, a "
        "prosthetic valve, or a dental procedure."
    )
    _set_hpi(case, text)
    _set_note(case, "admission", text)
    _set_note(
        case,
        "hospital_course",
        "Ceftriaxone was given intravenously during the admission. The echocardiogram "
        "and blood-culture results above are the recorded microbiologic evidence. "
        "Creatinine fell from 1.3 mg/dL to 0.8 mg/dL. Lisinopril remained on the "
        "inpatient list. No baseline creatinine before this admission is recorded.",
    )


def _patch_813(case: dict[str, Any]) -> None:
    text = (
        "A 64-year-old Female with a kidney transplant is admitted with diarrhea for "
        "several days, described as improving. Before this admission she was taking "
        "valganciclovir 900 mg orally twice daily, tacrolimus 1 mg every 12 hours, "
        "amlodipine, and atorvastatin. The chart does not show valganciclovir being "
        "started for a first diagnosis during this stay. Potassium is 4.7 mmol/L on "
        "admission and 3.9 mmol/L at discharge, and no cause for that change is recorded. "
        "Creatinine is 1.2 mg/dL then 1.0 mg/dL."
    )
    _set_hpi(case, text)
    _set_note(case, "admission", text)
    _set_note(
        case,
        "hospital_course",
        "Diarrhea is described as improving. Valganciclovir was already a home medicine. "
        "Potassium changed from 4.7 mmol/L to 3.9 mmol/L without a recorded cause. "
        "Creatinine is 1.2 mg/dL then 1.0 mg/dL.",
    )


_NARRATIVE = {
    "VAL-801": _patch_801,
    "VAL-802": _patch_802,
    "VAL-803": _patch_803,
    "VAL-805": _patch_805,
    "VAL-809": _patch_809,
    "VAL-813": _patch_813,
}


def _set_alternative(evaluator: dict[str, Any], medication: str, action: str, rationale: str) -> None:
    plan = evaluator["reference_discharge_plan"]
    plan["acceptable_alternatives"] = [
        {"medication": medication, "action": action, "rationale": rationale}
    ]


def _reference_801(evaluator: dict[str, Any]) -> None:
    for item in evaluator["reference_discharge_plan"]["medications"]:
        if "ibuprofen" in str(item.get("medication")):
            item["action"] = "continue"
            item["rationale"] = (
                "No stop indication is documented. Creatinine is 1.1 mg/dL then "
                "1.0 mg/dL. The reference continues ibuprofen rather than adding a bleed "
                "or an infection."
            )


def _reference_802(evaluator: dict[str, Any]) -> None:
    _set_alternative(
        evaluator,
        "lisinopril 10 MG Oral Tablet",
        "hold",
        "Hold is acceptable if creatinine 1.3 mg/dL is judged to be acute kidney injury. "
        "No baseline is recorded, so the reference does not require a hold.",
    )


def _reference_803(evaluator: dict[str, Any]) -> None:
    return None


def _reference_805(evaluator: dict[str, Any]) -> None:
    _set_alternative(
        evaluator,
        "Additional HFrEF therapy (ACE inhibitor, ARB, SGLT2 inhibitor, or MRA)",
        "start",
        "An additional class is acceptable and is not required. No ejection fraction "
        "is recorded, and the reference does not add a drug that is absent from the chart.",
    )


def _reference_809(evaluator: dict[str, Any]) -> None:
    _set_alternative(
        evaluator,
        "lisinopril 10 MG Oral Tablet",
        "hold",
        "Hold while creatinine is 1.3 mg/dL is acceptable if that value is judged to be "
        "acute kidney injury. No prior baseline is recorded, so the reference continues lisinopril.",
    )


def _reference_813(evaluator: dict[str, Any]) -> None:
    _set_alternative(
        evaluator,
        "valganciclovir 450 MG Oral Tablet",
        "stop",
        "Stopping or switching the antiviral is acceptable if the clinician rejects "
        "pre-admission therapy for this presentation. The recorded home row is otherwise continued.",
    )


_REFERENCE = {
    "VAL-801": _reference_801,
    "VAL-802": _reference_802,
    "VAL-803": _reference_803,
    "VAL-805": _reference_805,
    "VAL-809": _reference_809,
    "VAL-813": _reference_813,
}


def _matrix() -> str:
    lines = [
        "# Round 1 reviewer-response matrix",
        "",
        "Each row is one concern from the completed review. A suggestion is not treated "
        "as a required new fact. Decisions are ACCEPTED, PARTIALLY_ACCEPTED, NOT_ACCEPTED, "
        "or REQUIRES_CLINICIAN_ADJUDICATION.",
        "",
        "| Review item | Round 1 reviewer feedback | Interpretation of concern | Author response | Proposed case change | Evidence/source | Exact source locator | Decision |",
        "| --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for item in REVIEW_ITEMS:
        cells = [
            item["id"],
            item["feedback"],
            item["interpretation"],
            item["response"],
            item["change"],
            item["evidence"],
            item["locator"],
            item["decision"],
        ]
        lines.append("| " + " | ".join(cell.replace("|", "/").replace("\n", " ") for cell in cells) + " |")
    lines.append("")
    return "\n".join(lines)


def _ledger() -> str:
    lines = [
        "# Revision evidence ledger",
        "",
        "Every executed v2 change has one row. Changes marked NOT_ACCEPTED in the matrix "
        "were not executed and do not appear here. Literature supports a principle. It is "
        "not used as a patient-specific measurement.",
        "",
        "| Change ID | Case | Field changed | Old value/text | New value/text | Why changed | Reviewer item | Evidence class | Source | Exact locator | Clinical inference required? |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for row in LEDGER:
        cells = [
            row["id"],
            row["case"],
            row["field"],
            row["old"],
            row["new"],
            row["why"],
            row["item"],
            row["class"],
            row["source"],
            row["locator"],
            row["inference"],
        ]
        lines.append("| " + " | ".join(cell.replace("|", "/").replace("\n", " ") for cell in cells) + " |")
    lines.extend(["", "## Evidence notes", ""])
    for key, text in EVIDENCE_NOTES.items():
        lines.append(f"[{key}] {text}")
        lines.append("")
    return "\n".join(lines)


def _write_codebook(directory: Path) -> Path:
    document = _new_document("CliniProof Cycle 2  |  Set 1 revised cases, version 2")
    prepare_form_document(document)
    _add_heading(document, "CliniProof", 0)
    _add_body(document, "Clinical Revision Round 2, version 2")
    _add_heading(document, "SET 1", 1)
    _add_heading(document, "REVISED CASES FOLLOWING ROUND 1 CLINICIAN FEEDBACK", 2)
    _add_body(document, "Cases: VAL-801, VAL-802, VAL-803, VAL-805, VAL-809, VAL-813")
    _add_body(
        document,
        "Complete the blank Round 2 form before you read the historical Round 1 record "
        "or the author response. Version 2 replaces the earlier revised-case codebook "
        "for these six cases. The other eighteen cases are not in this file. Ready for "
        "clinician review does not mean clinically validated.",
    )
    reviews = completed_reviews()
    for case_id in CASE_IDS:
        resident, evaluator = _build_case(case_id)
        row = {"case_id": case_id, "resident": resident, "evaluator": evaluator}
        banner = document.add_paragraph()
        banner.paragraph_format.page_break_before = True
        banner.paragraph_format.keep_with_next = True
        _set_run_font(banner.add_run(f"CASE {case_id}"), size=16, bold=True, color=NAVY)
        _add_body(document, "Round 2 revised case, version 2")
        _add_heading(document, "CURRENT ROUND 2 CASE", 2)
        _chart(document, resident)
        _round2_form(document, row)
        _historical_section(document, reviews[case_id])
        _response_section(document, case_id)
    path = directory / "CliniProof_Cycle2_Revised_Cases_Validation.docx"
    document.save(str(path))
    finalize_word_form(path)
    return path


def _historical_section(document: WordDocument, review: dict[str, Any]) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True
    _set_run_font(
        heading.add_run("ROUND 1 CLINICIAN FEEDBACK — HISTORICAL RECORD"),
        size=14,
        bold=True,
        color=NAVY,
    )
    _banner(
        document,
        "ROUND 1 CLINICIAN FEEDBACK — READ ONLY",
        fill="595959",
        font_color=RGBColor(0xFF, 0xFF, 0xFF),
    )
    _add_body(
        document,
        "This section reproduces the completed Round 1 form. Boxes are not editable. "
        "Blank items stay blank. If two boxes were selected, both are shown.",
    )
    rows = []
    for domain in review["c1"]["domains"]:
        marks = tuple("☒" if domain["scores"][str(score)] else "☐" for score in (1, 2, 3, 4))
        label = domain["domain"]
        if AMBIGUITY in domain["flags"]:
            label += " — source selected more than one score"
        rows.append((label, *marks))
    _bordered_table(
        document,
        ("Domain", "1", "2", "3", "4"),
        tuple(rows),
        (3.6, 0.8, 0.8, 0.8, 0.8),
        header=True,
        fill="F2F2F2",
    )
    _add_body(document, "C1 overall: " + _selection_line(review["c1"]["overall"], ("pass", "fail")))
    _quote(document, "C1 comments", review["c1"]["comment"])
    _add_body(document, "C2 question: " + review["c2"]["question"])
    _add_body(document, "C2: " + _selection_line(review["c2"]["selection"], ("pass", "fail")))
    _quote(document, "C2 comments", review["c2"]["comment"])
    _add_body(document, "C3 question: " + review["c3"]["question"])
    _add_body(document, "C3: " + _selection_line(review["c3"]["selection"], ("pass", "fail")))
    _quote(document, "C3 comments", review["c3"]["comment"])
    _add_body(document, "C4 question: " + review["c4"]["question"])
    _add_body(document, "C4: " + _selection_line(review["c4"]["selection"], ("pass", "fail")))
    _add_body(
        document,
        "C4 issue: "
        + (review["c4"]["medication_or_clinical_issue"] or "Not completed in Round 1"),
    )
    _add_body(
        document,
        "C4 location: " + (review["c4"]["where_it_appears"] or "Not completed in Round 1"),
    )
    _add_body(
        document,
        "C4 clinical meaning: "
        + (review["c4"]["why_clinically_meaningful"] or "Not completed in Round 1"),
    )
    _quote(document, "C4 comments", review["c4"]["comment"])
    _add_body(
        document,
        "C5: "
        + _selection_line(
            review["c5"]["selection"],
            ("easy", "moderate", "hard", "inappropriate_or_outlier"),
        ),
    )
    _quote(document, "C5 comments", review["c5"]["comment"])
    _add_body(
        document,
        "Overall recommendation: "
        + _selection_line(
            review["overall_recommendation"]["selection"],
            ("accept", "revise", "exclude"),
        ),
    )
    _quote(document, "Overall comments", review["overall_recommendation"]["comment"])
    reviewer = review["overall_recommendation"]["reviewer_code"] or "Not completed in Round 1"
    dated = review["overall_recommendation"]["date"] or "Not completed in Round 1"
    _add_body(document, f"Reviewer: {reviewer}. Date: {dated}.")


_SELECTION_LABELS = {
    "pass": "Pass",
    "fail": "Fail",
    "easy": "Easy",
    "moderate": "Moderate",
    "hard": "Hard",
    "inappropriate_or_outlier": "Inappropriate / outlier",
    "accept": "Accept",
    "revise": "Revise",
    "exclude": "Exclude",
}


def _selection_line(selection: dict[str, Any], names: tuple[str, ...]) -> str:
    if selection.get("pass") is None and "accept" not in selection and "easy" not in selection:
        return "Not completed in Round 1"
    chosen = [_SELECTION_LABELS[name] for name in names if selection.get(name)]
    if not chosen:
        return "Not completed in Round 1"
    text = ", ".join(chosen)
    if AMBIGUITY in selection.get("flags", []):
        text += ". Source document contains multiple selected responses; reproduced exactly."
    return text


def _quote(document: WordDocument, label: str, comment: str | None) -> None:
    _bordered_table(
        document,
        (label,),
        ((comment or "Not completed in Round 1",),),
        (6.8,),
        header=True,
        fill="F2F2F2",
    )


def _response_section(document: WordDocument, case_id: str) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True
    _set_run_font(
        heading.add_run("ROUND 2 RESPONSE TO ROUND 1 REVIEW"),
        size=14,
        bold=True,
        color=NAVY,
    )
    rows = tuple(
        (item["id"], item["feedback"], item["response"], item["change"], item["evidence"])
        for item in REVIEW_ITEMS
        if item["case"] == case_id
    )
    _bordered_table(
        document,
        ("#", "Reviewer comment / concern", "Author response", "Revision made", "Evidence / justification"),
        rows,
        (0.9, 1.4, 1.6, 1.4, 1.5),
        header=True,
    )
    _add_heading(document, "Evidence for changes", 3)
    for key, text in EVIDENCE_NOTES.items():
        if any(key in item["evidence"] or key in item["response"] for item in REVIEW_ITEMS if item["case"] == case_id) or any(
            key in row["source"] for row in LEDGER if row["case"] == case_id
        ):
            _add_body(document, f"[{key}] {text}")


def _readme() -> str:
    return "\n".join(
        [
            "# Set 1 revised cases, version 2",
            "",
            "Open [CliniProof_Cycle2_Revised_Cases_Validation.docx]"
            "(CliniProof_Cycle2_Revised_Cases_Validation.docx).",
            "",
            "This version contains VAL-801, VAL-802, VAL-803, VAL-805, VAL-809, and "
            "VAL-813 only. It does not replace the eighteen-case fresh-review codebook. "
            "The earlier six-case codebook remains in "
            "`exports/ko_cycle2_revised_validation/` and was not overwritten.",
            "",
            "Round 1 comments are reproduced after the blank Round 2 form. The author "
            "response and evidence list follow that historical record.",
            "",
            "Ready for clinician review does not mean clinically validated.",
            "",
        ]
    )


def _manifest() -> str:
    lines = [
        "# Set 1 version 2 manifest",
        "",
        "| Case | Resident file | Evaluator file |",
        "| --- | --- | --- |",
    ]
    for case_id in CASE_IDS:
        lines.append(
            f"| {case_id} | `{case_id}_resident.json` | `{case_id}_evaluator.json` |"
        )
    lines.append("")
    return "\n".join(lines)
