"""Two Cycle 2 clinician codebooks: six revised cases, then eighteen clean cases.

Reads the existing resident and evaluator JSON. It does not regenerate cases
and does not edit clinical content. Round 1 ratings are copied from a completed
casebook when that file is available. They are not inferred.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.services.round1_feedback import (
    C1_DOMAINS,
    NOT_COMPLETED,
    ChoiceGroup,
    Round1CaseFeedback,
    parse_round1_casebook,
)
from app.services.word_controls import (
    append_checkbox,
    append_plain_text,
    append_rich_text_field,
    finalize_word_form,
    prepare_form_document,
)
from app.services.word_export import (
    CONTENT_WIDTH,
    HEADER_FILL,
    NAVY,
    _add_body,
    _add_heading,
    _keep_row_together,
    _mark_header_row,
    _new_document,
    _set_run_font,
    _set_table_width,
    _shade,
    _write_cell,
)
from app.services.word_facts import (
    clinical,
    exact_text,
    field_text,
    medication_note,
    note_text,
    presentation,
    rows_for_context,
)

ROOT = Path(__file__).resolve().parents[2]
REVISED_DIR = ROOT / "exports" / "ko_revised_cases_v1"
CLEAN_DIR = ROOT / "exports" / "ko_clean_cases_for_review_v1"
REVISED_IDS = ("VAL-801", "VAL-802", "VAL-803", "VAL-805", "VAL-809", "VAL-813")
CLEAN_IDS = (
    "VAL-804",
    "VAL-806",
    "VAL-807",
    "VAL-808",
    "VAL-810",
    "VAL-811",
    "VAL-812",
    "VAL-814",
    "VAL-815",
    "VAL-816",
    "VAL-817",
    "VAL-818",
    "VAL-819",
    "VAL-820",
    "VAL-821",
    "VAL-822",
    "VAL-823",
    "VAL-824",
)
ROUND2_C1_DOMAINS = (
    "Presentation and demographics",
    "Fit between presentation and diagnosis",
    "Vital signs",
    "Laboratory findings",
    "Medication regimen",
    "Hospital course",
    "Chart consistency",
    "Discharge context and follow-up",
)
SPECIAL_QUESTIONS = {
    "VAL-805": (
        "Which additional HFrEF therapies should be required versus acceptable "
        "alternatives, given ejection fraction 30 percent, discharge systolic "
        "pressure 110 mmHg, creatinine 0.9 mg/dL, and potassium 4.3 mmol/L?"
    ),
    "VAL-806": (
        "Restart lisinopril versus continued hold, given creatinine 2.8 mg/dL "
        "falling to 1.6 mg/dL without a known baseline?"
    ),
    "VAL-813": (
        "Continue versus modify or temporarily hold mycophenolate, given the "
        "CMV disease and transplant context?"
    ),
    "VAL-814": "Restart versus continued hold of mycophenolate?",
    "VAL-820": (
        "Is enoxaparin VTE prophylaxis appropriate during concurrent warfarin "
        "therapy with the documented INR values?"
    ),
    "VAL-821": (
        "Restart apixaban at discharge versus continued hold pending gastroenterology follow-up?"
    ),
}
TIMEPOINT_ORDER = (
    "admission",
    "hospital_day_1",
    "hospital_day_2",
    "hospital_day_3",
    "inpatient",
    "discharge",
)
GRAY = "F2F2F2"
GRAY_HEADER = "595959"
SOURCE_NAME = "KO Casebook Validation.docx"

# Concern / change pairs taken from the Cycle 2 revision record. They are not
# a substitute for the verbatim Round 1 comments copied from the casebook.
RESPONSE_ROWS: dict[str, tuple[tuple[str, str], ...]] = {
    "VAL-801": (
        (
            "The cause of delirium was not explained.",
            "Dehydration from poor intake is now the stated cause.",
        ),
        (
            "The hospital course did not explain the improvement.",
            "Dry mucous membranes and a 20 mmHg orthostatic change support "
            "dehydration. Confusion clears after fluids and restored intake.",
        ),
        (
            "Ibuprofen discontinuation lacked justification.",
            "The unsupported ibuprofen discontinuation was removed. Creatinine "
            "stays 1.1 then 1.0 mg/dL, and the chart does not describe bleeding.",
        ),
    ),
    "VAL-802": (
        (
            "The cause of delirium was not explained.",
            "Poor intake and dehydration are now the stated cause.",
        ),
        (
            "Creatinine and lisinopril were hard to interpret without a baseline.",
            "A clinic creatinine of 1.2 mg/dL six weeks earlier is the baseline. "
            "Admission creatinine is 1.3 mg/dL and discharge creatinine is 1.2 mg/dL. "
            "Potassium is 4.2 then 4.1 mmol/L.",
        ),
        (
            "Atorvastatin was missing from the recovered chart's intended regimen "
            "after the earlier injection.",
            "Atorvastatin remains on the chart and on the reference plan.",
        ),
    ),
    "VAL-803": (
        (
            "The Round 1 review of this case was incomplete.",
            "Delirium is attributed to thiazide-associated hyponatremia. Sodium "
            "moves from 128 mmol/L to 135 mmol/L after hydrochlorothiazide is held, "
            "and confusion clears. The clean 30-day lisinopril supply is restored.",
        ),
    ),
    "VAL-805": (
        (
            "Discharge weight was not close enough to dry weight.",
            "Serial weights are 86 kg, then 83 kg, then 80 kg, and 80 kg is the "
            "documented dry weight.",
        ),
        (
            "Inpatient diuresis was not realistically represented, and serial "
            "intake and output were insufficient.",
            "Intake and output are net negative on hospital days 1 through 3. "
            "Inpatient furosemide is intravenous 40 mg twice daily, distinct from "
            "the home oral dose of 40 mg daily.",
        ),
        (
            "Cardiology and echocardiography information was missing.",
            "Ejection fraction is 30 percent. Creatinine, potassium, and natriuretic "
            "peptide trajectories are on the chart, with discharge systolic pressure "
            "110 mmHg.",
        ),
        (
            "The trainee needed a meaningful decision about additional heart-failure therapy.",
            "The reference continues oral furosemide, metoprolol succinate, and "
            "atorvastatin. The form asks whether additional HFrEF therapy is required "
            "or only acceptable. It was not added automatically.",
        ),
    ),
    "VAL-809": (
        (
            "There was insufficient clinical evidence for endocarditis, including "
            "fever and risk context.",
            "The chart includes fever of 38.6°C, heart rate 104, a dental extraction "
            "three weeks earlier, a new murmur, viridans group streptococcus, and a "
            "vegetation. Later cultures show no growth and the fever resolves.",
        ),
        (
            "The acute kidney injury was unexplained, and lisinopril handling was not appropriate.",
            "Baseline creatinine is 0.8 mg/dL. Creatinine is 1.3, then 1.0, then "
            "0.8 mg/dL. Lisinopril is held during the elevation. The reference "
            "restarts it at discharge, and the chart does not announce that restart.",
        ),
        (
            "The case was not complex enough for the intended decision.",
            "Ceftriaxone 2 g intravenously daily for four weeks is the reference "
            "start, with the endocarditis findings above available to the resident.",
        ),
    ),
    "VAL-813": (
        (
            "Valganciclovir was present at home before a new cytomegalovirus diagnosis.",
            "Valganciclovir is not a home medication.",
        ),
        (
            "The cytomegalovirus chronology was misleading.",
            "Viral load is detected after admission, and valganciclovir 900 mg twice "
            "daily begins after that detection.",
        ),
        (
            "The transplant regimen was too simplified, and the laboratories were incomplete.",
            "The chart keeps tacrolimus and mycophenolate, with creatinine 1.2 then "
            "1.0 mg/dL. The form asks the clinician to judge whether mycophenolate "
            "should continue. That reference action was not changed in advance.",
        ),
    ),
}


def locate_round1_casebook(root: Path | None = None) -> Path | None:
    """Find the completed Round 1 file by its file name. Do not guess another file."""
    start = root or ROOT
    matches = [
        path
        for path in start.rglob(SOURCE_NAME)
        if ".venv" not in path.parts and "pytest" not in str(path)
    ]
    return matches[0] if matches else None


def write_split_casebooks(
    revised_directory: Path,
    clean_directory: Path,
    *,
    round1_casebook: Path | None = None,
) -> tuple[Path, Path]:
    """Write both codebooks. A missing Round 1 file does not invent ratings."""
    feedback: dict[str, Round1CaseFeedback] = {}
    source_found = False
    if round1_casebook is not None and round1_casebook.is_file():
        feedback = parse_round1_casebook(round1_casebook)
        source_found = True
    revised_path = _write_package(
        revised_directory,
        REVISED_IDS,
        kind="revised",
        feedback=feedback,
        source_found=source_found,
    )
    clean_path = _write_package(
        clean_directory,
        CLEAN_IDS,
        kind="clean",
        feedback={},
        source_found=False,
    )
    return revised_path, clean_path


def _write_package(
    directory: Path,
    case_ids: tuple[str, ...],
    *,
    kind: str,
    feedback: dict[str, Round1CaseFeedback],
    source_found: bool,
) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    rows = [_load(case_id) for case_id in case_ids]
    document = _new_document(
        "CliniProof Cycle 2  |  Set 1 revised cases"
        if kind == "revised"
        else "CliniProof Cycle 2  |  Set 2 clean cases"
    )
    prepare_form_document(document)
    if kind == "revised":
        _revised_cover(document)
        filename = "CliniProof_Cycle2_Revised_Cases_Validation.docx"
    else:
        _clean_cover(document)
        filename = "CliniProof_Cycle2_Clean_Cases_Validation.docx"
    for row in rows:
        _add_case(
            document,
            row,
            kind=kind,
            feedback=feedback.get(row["case_id"]) if kind == "revised" else None,
            source_found=source_found and kind == "revised",
        )
    path = directory / filename
    document.save(str(path))
    finalize_word_form(path)
    (directory / "README.md").write_text(_readme(kind), encoding="utf-8")
    (directory / "MANIFEST.md").write_text(_manifest(rows, kind), encoding="utf-8")
    return path


def _load(case_id: str) -> dict[str, Any]:
    folder = REVISED_DIR if case_id in REVISED_IDS else CLEAN_DIR
    resident = json.loads((folder / f"{case_id}_resident.json").read_text(encoding="utf-8"))
    evaluator = json.loads((folder / f"{case_id}_evaluator.json").read_text(encoding="utf-8"))
    if not isinstance(resident, dict) or not isinstance(evaluator, dict):
        raise ValueError(f"{case_id} export is not an object")
    if "reference_discharge_plan" in resident:
        raise ValueError(f"{case_id} resident file contains the hidden reference")
    return {
        "case_id": case_id,
        "folder": folder.name,
        "resident": resident,
        "evaluator": evaluator,
    }


def _revised_cover(document: WordDocument) -> None:
    _add_heading(document, "CliniProof", 0)
    _add_body(document, "Clinical Revision Round 2")
    _add_heading(document, "SET 1", 1)
    _add_heading(document, "REVISED CASES FOLLOWING ROUND 1 CLINICIAN FEEDBACK", 2)
    _add_body(document, "Cases: VAL-801, VAL-802, VAL-803, VAL-805, VAL-809, VAL-813")
    _cover_identity(document)
    _add_body(
        document,
        "These six cases were revised because the first clinician-validation round "
        "identified substantive problems with the underlying clinical scenario. "
        "For each case, please review the current revised case and complete the "
        "blank Round 2 validation form before reading the reproduced Round 1 "
        "feedback that follows. The previous feedback is included so that the "
        "reviewer can then assess whether the revision adequately addressed the "
        "earlier concerns.",
    )
    _add_body(
        document,
        "Set 2 is a separate codebook. It contains the other eighteen cases and "
        "should be reviewed after this set. Ready for clinician review does not "
        "mean clinically validated. Every Round 2 checkbox below starts blank.",
    )


def _clean_cover(document: WordDocument) -> None:
    _add_heading(document, "CliniProof", 0)
    _add_body(document, "Clinical Revision Round 2")
    _add_heading(document, "SET 2", 1)
    _add_heading(document, "CLEAN CASES FOR FRESH CLINICIAN REVIEW", 2)
    _add_body(document, "18 cases")
    _add_body(
        document,
        "VAL-804, VAL-806, VAL-807, VAL-808, VAL-810, VAL-811, VAL-812, VAL-814, "
        "VAL-815, VAL-816, VAL-817, VAL-818, VAL-819, VAL-820, VAL-821, VAL-822, "
        "VAL-823, and VAL-824.",
    )
    _cover_identity(document)
    _add_body(
        document,
        "These 18 cases are clean cases prepared for review under the corrected "
        "Round 2 protocol. They are not presented as previously clinician-validated "
        "cases. Some underwent narrow pre-review consistency corrections, but they "
        "have not received the same substantive prior clinician review as the six "
        "cases in Set 1.",
    )
    _add_body(
        document,
        "Please review Set 1, the six revised cases, before this codebook. This "
        "document does not include Round 1 feedback. Ready for clinician review "
        "does not mean clinically validated. Every checkbox below starts blank.",
    )


def _cover_identity(document: WordDocument) -> None:
    paragraph = document.add_paragraph()
    _set_run_font(paragraph.add_run("Reviewer code: "), size=11, bold=True)
    append_plain_text(paragraph, document, tag="cover-reviewer-code", alias="Reviewer code")
    paragraph = document.add_paragraph()
    _set_run_font(paragraph.add_run("Review date: "), size=11, bold=True)
    append_plain_text(paragraph, document, tag="cover-review-date", alias="Review date")


def _add_case(
    document: WordDocument,
    row: dict[str, Any],
    *,
    kind: str,
    feedback: Round1CaseFeedback | None,
    source_found: bool,
) -> None:
    case_id = str(row["case_id"])
    banner = document.add_paragraph()
    banner.paragraph_format.page_break_before = True
    banner.paragraph_format.keep_with_next = True
    banner.paragraph_format.space_after = Pt(2)
    _set_run_font(banner.add_run(f"CASE {case_id}"), size=16, bold=True, color=NAVY)
    subtitle = "Round 2 revised case" if kind == "revised" else "Round 2 clean case"
    _add_body(document, subtitle)
    if kind == "revised":
        _add_heading(document, "CURRENT ROUND 2 CASE", 2)
    _chart(document, row["resident"])
    _round2_form(document, row)
    if kind == "revised":
        _historical_feedback(document, case_id, feedback, source_found=source_found)
        _response_table(document, case_id)


def _chart(document: WordDocument, chart: dict[str, Any]) -> None:
    info = clinical(chart)
    presented = presentation(chart)
    _add_heading(document, "Patient overview", 2)
    _pair_table(
        document,
        [
            ("Age", exact_text(info.get("patient_age"))),
            ("Sex", exact_text(info.get("patient_gender"))),
            ("Admission diagnosis", exact_text(info.get("admission_dx"))),
            ("Disposition", exact_text(info.get("disposition_status"))),
            ("One-liner", exact_text(info.get("one_liner"))),
        ],
    )
    hpi = exact_text(presented.get("hpi"))
    if hpi:
        _add_heading(document, "Presentation", 2)
        _add_body(document, hpi)
    admission = note_text(chart, "admission")
    if admission:
        _add_heading(document, "Admission note", 3)
        _add_body(document, admission)
    course = note_text(chart, "hospital_course")
    if course:
        _add_heading(document, "Hospital course", 2)
        _add_body(document, course)
    diagnoses = [
        row
        for row in chart.get("CaseDiagnosis", [])
        if isinstance(row, dict) and exact_text(row.get("diagnosis"))
    ]
    if diagnoses:
        _add_heading(document, "Diagnoses", 2)
        _bordered_table(
            document,
            ("Diagnosis", "Type", "Status"),
            tuple(
                (
                    exact_text(row.get("diagnosis")) or "",
                    field_text(row, "diagnosis_type") or "",
                    field_text(row, "status") or "",
                )
                for row in diagnoses
            ),
            (3.4, 1.7, 1.7),
            header=True,
        )
    _vital_table(document, chart)
    _lab_table(document, chart)
    _weight_table(document, chart)
    _io_table(document, chart)
    _medication_table(document, "Home medications", rows_for_context(chart, "home"))
    _medication_table(
        document,
        "Inpatient medications",
        rows_for_context(chart, "inpatient") + rows_for_context(chart, "inpatient_history"),
    )
    _optional_records(
        document,
        "Imaging",
        chart.get("CaseImaging"),
        (("study_type", "Study"), ("timepoint", "Time"), ("finding", "Finding")),
    )
    _optional_records(
        document,
        "Procedures",
        chart.get("CaseProcedure"),
        (("procedure_name", "Procedure"), ("findings", "Findings")),
    )
    _optional_records(
        document,
        "Consultations",
        chart.get("CaseConsult"),
        (
            ("service", "Service"),
            ("assessment", "Assessment"),
            ("recommendation", "Recommendation"),
        ),
    )
    _optional_records(
        document,
        "Follow-up",
        chart.get("CaseFollowup"),
        (("item", "Item"), ("timing", "Timing"), ("with_service", "With service")),
    )


def _pair_table(document: WordDocument, pairs: list[tuple[str, str | None]]) -> None:
    rows = tuple((label, value) for label, value in pairs if value)
    if rows:
        _bordered_table(document, ("Field", "Value"), rows, (2.2, 4.6), header=True)


def _vital_table(document: WordDocument, chart: dict[str, Any]) -> None:
    rows = [row for row in _rows(chart, "CaseVital")]
    if not rows:
        return
    _add_heading(document, "Vital signs", 2)
    rendered = tuple(
        (
            field_text(row, "timepoint") or "",
            _blood_pressure(row),
            exact_text(row.get("heart_rate")) or "",
            _with_unit(exact_text(row.get("temp_c")), "°C"),
            _with_unit(exact_text(row.get("spo2_percent")), "%"),
        )
        for row in _ordered(rows)
    )
    _bordered_table(
        document,
        ("Time", "BP", "HR", "Temp", "SpO2"),
        rendered,
        (1.6, 1.4, 1.0, 1.4, 1.4),
        header=True,
    )


def _lab_table(document: WordDocument, chart: dict[str, Any]) -> None:
    rows = _rows(chart, "CaseLab")
    if not rows:
        return
    _add_heading(document, "Laboratory findings", 2)
    timepoints = _timepoint_labels(rows)
    by_test: dict[str, dict[str, tuple[str, str]]] = {}
    order: list[str] = []
    for row in rows:
        name = exact_text(row.get("test_name")) or ""
        if name not in by_test:
            by_test[name] = {}
            order.append(name)
        when = field_text(row, "timepoint") or ""
        value = exact_text(row.get("value")) or exact_text(row.get("value_text")) or ""
        unit = exact_text(row.get("unit")) or ""
        by_test[name][when] = (value, unit)
    headers = ("Test", *timepoints, "Unit")
    width_left = 2.6
    width_unit = 0.8
    remaining = CONTENT_WIDTH - width_left - width_unit
    each = remaining / max(len(timepoints), 1)
    widths = (width_left, *tuple(each for _ in timepoints), width_unit)
    rendered = []
    for name in order:
        values = by_test[name]
        unit = next((pair[1] for pair in values.values() if pair[1]), "")
        rendered.append((name, *tuple(values.get(when, ("", ""))[0] for when in timepoints), unit))
    _bordered_table(document, headers, tuple(rendered), widths, header=True)


def _weight_table(document: WordDocument, chart: dict[str, Any]) -> None:
    rows = _rows(chart, "CaseWeight")
    if not rows:
        return
    _add_heading(document, "Weight", 3)
    rendered = tuple(
        (
            field_text(row, "timepoint") or "",
            exact_text(row.get("weight_kg")) or "",
            exact_text(row.get("dry_weight_kg")) or "",
        )
        for row in _ordered(rows)
    )
    _bordered_table(
        document,
        ("Time", "Weight (kg)", "Dry weight (kg)"),
        rendered,
        (2.2, 2.3, 2.3),
        header=True,
    )


def _io_table(document: WordDocument, chart: dict[str, Any]) -> None:
    rows = _rows(chart, "CaseIntakeOutput")
    if not rows:
        return
    _add_heading(document, "Intake and output", 3)
    rendered = tuple(
        (
            field_text(row, "timepoint") or "",
            exact_text(row.get("intake_ml")) or "",
            exact_text(row.get("output_ml")) or "",
            exact_text(row.get("net_ml")) or "",
        )
        for row in _ordered(rows)
    )
    _bordered_table(
        document,
        ("Time", "Intake (mL)", "Output (mL)", "Net (mL)"),
        rendered,
        (1.7, 1.7, 1.7, 1.7),
        header=True,
    )


def _medication_table(document: WordDocument, title: str, rows: list[dict[str, Any]]) -> None:
    _add_heading(document, title, 2)
    if not rows:
        _add_body(document, "No medications were specified for this list.")
        return
    rendered = tuple(
        (
            exact_text(row.get("drug")) or exact_text(row.get("reported_name")) or "",
            exact_text(row.get("dose")) or "",
            exact_text(row.get("route")) or "",
            exact_text(row.get("frequency")) or "",
            medication_note(row),
        )
        for row in rows
    )
    _bordered_table(
        document,
        ("Medication", "Dose", "Route", "Frequency", "Indication"),
        rendered,
        (2.1, 0.9, 0.8, 1.1, 1.9),
        header=True,
    )


def _optional_records(
    document: WordDocument,
    title: str,
    raw: object,
    columns: tuple[tuple[str, str], ...],
) -> None:
    rows = [row for row in raw if isinstance(row, dict)] if isinstance(raw, list) else []
    if not rows:
        return
    _add_heading(document, title, 2)
    headers = tuple(label for _key, label in columns)
    rendered = tuple(
        tuple(field_text(row, key) or exact_text(row.get(key)) or "" for key, _label in columns)
        for row in rows
    )
    each = CONTENT_WIDTH / len(headers)
    _bordered_table(document, headers, rendered, tuple(each for _ in headers), header=True)


def _round2_form(document: WordDocument, row: dict[str, Any]) -> None:
    case_id = str(row["case_id"])
    _add_heading(document, "C1 — Clinical plausibility", 2)
    _add_body(
        document,
        "1 = implausible. 2 = substantial revision required. "
        "3 = plausible with minor concern. 4 = fully plausible.",
    )
    _c1_table(document, case_id)
    _choice_row(document, case_id, "c1-overall", "Overall C1", ("Pass", "Fail"))
    _comment_box(document, case_id, "c1-comment", "C1 reviewer comments", 1.0)
    _add_heading(document, "C2 — Sufficiency for resident decision-making", 2)
    _add_body(
        document,
        "Does the resident-facing chart contain enough clinical information for an "
        "internal-medicine resident to independently determine an appropriate "
        "discharge medication regimen?",
    )
    _choice_row(document, case_id, "c2", "C2", ("Pass", "Fail"))
    _comment_box(document, case_id, "c2-comment", "C2 reviewer comments", 1.0)
    _add_heading(document, "CLINICIAN VALIDATION REFERENCE", 2)
    _add_body(document, "Not shown to residents in the assessment study.")
    _add_body(document, "Complete C1 and C2 before using this reference.")
    _reference_table(document, row["evaluator"])
    question = SPECIAL_QUESTIONS.get(case_id)
    if question:
        _add_heading(document, "Case-specific adjudication", 3)
        _add_body(document, question)
        _comment_box(document, case_id, "adjudication", "Case-specific adjudication", 1.0)
    _add_heading(document, "C3 — Reference-plan validity", 2)
    _add_body(document, "Is the hidden reference discharge medication plan clinically defensible?")
    _choice_row(document, case_id, "c3", "C3", ("Pass", "Fail"))
    _add_body(document, "Disputed medication")
    _blank_entry_table(
        document,
        ("Medication", "Current reference action", "Suggested action", "Reason"),
        (1.7, 1.7, 1.7, 1.7),
        rows=3,
        row_inches=0.45,
    )
    _add_heading(document, "C4 — Acceptable alternatives", 2)
    _add_body(document, "Could another discharge medication decision also reasonably be correct?")
    _choice_row(document, case_id, "c4", "C4", ("No", "Yes"))
    _blank_entry_table(
        document,
        ("Medication", "Alternative action", "Clinical rationale"),
        (2.2, 2.0, 2.6),
        rows=2,
        row_inches=0.55,
    )
    _add_heading(document, "C5 — Missing, misleading, or competing issues", 2)
    _add_body(
        document,
        "Does the chart contain another clinically meaningful issue that could "
        "change the expected discharge plan?",
    )
    _choice_row(document, case_id, "c5-issue", "C5 competing issue", ("No", "Yes"))
    _add_body(document, "Does any wording reveal the intended answer?")
    _choice_row(document, case_id, "c5-leak", "C5 answer reveal", ("No", "Yes"))
    _comment_box(document, case_id, "c5-comment", "C5 reviewer comments", 1.0)
    _add_heading(document, "C6 — Expected learner difficulty", 2)
    _choice_row(
        document,
        case_id,
        "c6",
        "C6",
        ("Easy", "Moderate", "Hard", "Inappropriate or outlier"),
    )
    _add_heading(document, "Overall recommendation", 2)
    _choice_row(document, case_id, "overall", "Overall", ("Accept", "Revise", "Exclude"))
    _comment_box(
        document,
        case_id,
        "overall-comment",
        "Overall comments / suggested revisions",
        1.5,
    )
    paragraph = document.add_paragraph()
    _set_run_font(paragraph.add_run("Reviewer code: "), size=11)
    append_plain_text(paragraph, document, tag=f"{case_id}-reviewer", alias="Reviewer code")
    paragraph = document.add_paragraph()
    _set_run_font(paragraph.add_run("Date: "), size=11)
    append_plain_text(paragraph, document, tag=f"{case_id}-date", alias="Review date")


def _c1_table(document: WordDocument, case_id: str) -> None:
    headers = ("Domain", "1", "2", "3", "4")
    widths = (3.6, 0.8, 0.8, 0.8, 0.8)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    _set_table_width(table, sum(widths))
    _set_borders(table)
    _header_row(table.rows[0], headers, widths)
    for domain in ROUND2_C1_DOMAINS:
        row = table.add_row()
        _write_cell(row.cells[0], domain, bold=False, fill=None)
        row.cells[0].width = Inches(widths[0])
        for index, score in enumerate(("1", "2", "3", "4"), start=1):
            cell = row.cells[index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            append_checkbox(
                paragraph,
                document,
                tag=f"{case_id}-c1-{domain}-{score}",
                alias=f"{case_id} C1 {domain} {score}",
            )
            cell.width = Inches(widths[index])
            _valign(cell, "center")
        _keep_row_together(row)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _choice_row(
    document: WordDocument,
    case_id: str,
    code: str,
    label: str,
    choices: tuple[str, ...],
) -> None:
    if len(choices) > 3:
        _wide_choice_rows(document, case_id, code, label, choices)
        return
    table = document.add_table(rows=1, cols=len(choices) + 1)
    table.style = "Table Grid"
    table.autofit = False
    label_width = 2.2
    choice_width = (CONTENT_WIDTH - label_width) / len(choices)
    widths = (label_width, *tuple(choice_width for _choice in choices))
    _set_table_width(table, sum(widths))
    _set_borders(table)
    _write_cell(table.rows[0].cells[0], label, bold=True, fill=HEADER_FILL)
    table.rows[0].cells[0].width = Inches(label_width)
    for index, choice in enumerate(choices, start=1):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        append_checkbox(
            paragraph,
            document,
            tag=f"{case_id}-{code}-{choice}",
            alias=f"{case_id} {code} {choice}",
        )
        _set_run_font(paragraph.add_run(f"  {choice}"), size=11)
        cell.width = Inches(choice_width)
        _valign(cell, "center")
    _keep_row_together(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _wide_choice_rows(
    document: WordDocument,
    case_id: str,
    code: str,
    label: str,
    choices: tuple[str, ...],
) -> None:
    """Put long labels in half-width cells so each checkbox stays beside its words."""
    rows = 1 + (len(choices) + 1) // 2
    table = document.add_table(rows=rows, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    _set_table_width(table, CONTENT_WIDTH)
    _set_borders(table)
    _write_cell(table.cell(0, 0), label, bold=True, fill=HEADER_FILL)
    table.cell(0, 0).merge(table.cell(0, 1))
    width = CONTENT_WIDTH / 2
    for index, choice in enumerate(choices):
        cell = table.rows[1 + index // 2].cells[index % 2]
        cell.text = ""
        cell.width = Inches(width)
        _place_choice(cell, document, case_id, code, choice)
        _valign(cell, "center")
    for row in table.rows:
        _keep_row_together(row)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _place_choice(
    cell: Any,
    document: WordDocument,
    case_id: str,
    code: str,
    choice: str,
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    append_checkbox(
        paragraph,
        document,
        tag=f"{case_id}-{code}-{choice}",
        alias=f"{case_id} {code} {choice}",
    )
    _set_run_font(paragraph.add_run(f"  {choice}"), size=11)


def _comment_box(
    document: WordDocument,
    case_id: str,
    tag: str,
    label: str,
    height_inches: float,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.autofit = False
    _set_table_width(table, CONTENT_WIDTH)
    _set_borders(table)
    _set_row_height(table.rows[0], int(height_inches * 1440))
    cell = table.cell(0, 0)
    cell.text = ""
    _valign(cell, "top")
    _cell_margins(cell)
    paragraph = cell.paragraphs[0]
    _set_run_font(paragraph.add_run(label), size=11, bold=True)
    append_rich_text_field(cell, document, tag=f"{case_id}-{tag}", alias=label, lines=3)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def _blank_entry_table(
    document: WordDocument,
    headers: tuple[str, ...],
    widths: tuple[float, ...],
    *,
    rows: int,
    row_inches: float,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    _set_table_width(table, sum(widths))
    _set_borders(table)
    _header_row(table.rows[0], headers, widths)
    for _ in range(rows):
        row = table.add_row()
        _set_row_height(row, int(row_inches * 1440))
        for index, width in enumerate(widths):
            cell = row.cells[index]
            cell.text = ""
            cell.width = Inches(width)
            _valign(cell, "top")
            _cell_margins(cell)
        _keep_row_together(row)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _reference_table(document: WordDocument, evaluator: dict[str, Any]) -> None:
    reference = evaluator.get("reference_discharge_plan")
    plan = reference if isinstance(reference, dict) else {}
    raw_medications = plan.get("medications")
    medications = raw_medications if isinstance(raw_medications, list) else []
    rendered = tuple(
        (
            exact_text(item.get("action")) or "",
            exact_text(item.get("medication")) or "",
            exact_text(item.get("dose")) or "",
            exact_text(item.get("route")) or "",
            exact_text(item.get("frequency")) or "",
            exact_text(item.get("indication")) or "",
        )
        for item in medications
        if isinstance(item, dict)
    )
    if rendered:
        _bordered_table(
            document,
            ("Action", "Medication", "Dose", "Route", "Frequency", "Indication"),
            rendered,
            (0.9, 1.8, 0.8, 0.7, 1.0, 1.6),
            header=True,
        )
    raw_alternatives = plan.get("acceptable_alternatives")
    alternatives = raw_alternatives if isinstance(raw_alternatives, list) else []
    if alternatives:
        _add_body(document, "Acceptable alternatives already recorded on the reference plan:")
        alt_rows = tuple(
            (
                exact_text(item.get("medication")) or "",
                exact_text(item.get("action")) or "",
                exact_text(item.get("rationale")) or "",
            )
            for item in alternatives
            if isinstance(item, dict)
        )
        _bordered_table(
            document,
            ("Medication", "Action", "Rationale"),
            alt_rows,
            (2.2, 1.4, 3.2),
            header=True,
        )


def _historical_feedback(
    document: WordDocument,
    case_id: str,
    feedback: Round1CaseFeedback | None,
    *,
    source_found: bool,
) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True
    _set_run_font(
        heading.add_run("PREVIOUS ROUND 1 CLINICIAN FEEDBACK"),
        size=14,
        bold=True,
        color=NAVY,
    )
    _banner(
        document,
        "ROUND 1 CLINICIAN FEEDBACK — READ ONLY",
        fill=GRAY_HEADER,
        font_color=RGBColor(0xFF, 0xFF, 0xFF),
    )
    _add_body(
        document,
        "This section reproduces feedback recorded during the first "
        "clinician-validation round. It is shown after the current Round 2 form "
        "so that the revised case can first be judged independently. Historical "
        "boxes are not editable.",
    )
    if not source_found:
        _add_body(
            document,
            f"The completed Round 1 file {SOURCE_NAME} was not in the repository "
            "when this codebook was built. Ratings and comments were not "
            "reconstructed, and no Round 1 selection was invented.",
        )
        return
    if feedback is None:
        _add_body(
            document,
            f"{case_id} was not present in the completed Round 1 casebook. {NOT_COMPLETED}.",
        )
        return
    _add_body(document, f"Reviewer: {feedback.reviewer or NOT_COMPLETED}")
    _add_body(document, f"Date: {feedback.review_date or NOT_COMPLETED}")
    _add_body(document, "Round 1 overall recommendation:")
    _static_choices(document, feedback.recommendation)
    _historical_c1(document, feedback)
    _historical_result(document, "C1 overall", feedback.c1_result, feedback.c1_comment)
    _historical_result(
        document,
        "C2 — Intended assessment problem",
        feedback.c2_result,
        feedback.c2_comment,
    )
    _historical_result(document, "C3 — Detectability", feedback.c3_result, feedback.c3_comment)
    _historical_result(
        document,
        "C4 — Absence of unintended competing problems",
        feedback.c4_result,
        feedback.c4_comment,
    )
    _historical_result(
        document,
        "C5 — Expected learner difficulty",
        feedback.c5_difficulty,
        feedback.c5_comment,
    )
    _historical_comment(document, "Overall comments", feedback.overall_comment)


def _historical_c1(document: WordDocument, feedback: Round1CaseFeedback) -> None:
    _add_heading(document, "Round 1 C1 domain ratings", 3)
    headers = ("Domain", "1", "2", "3", "4")
    widths = (3.6, 0.8, 0.8, 0.8, 0.8)
    rendered = []
    incomplete: list[str] = []
    for domain, group in zip(C1_DOMAINS, feedback.c1_scores, strict=True):
        marks = tuple("☒" if score in group.selected else "☐" for score in ("1", "2", "3", "4"))
        rendered.append((domain, *marks))
        if not group.completed:
            incomplete.append(domain)
    _bordered_table(document, headers, tuple(rendered), widths, header=True, fill=GRAY)
    if incomplete:
        _add_body(document, f"{NOT_COMPLETED}: {', '.join(incomplete)}.")


def _historical_result(
    document: WordDocument,
    title: str,
    group: ChoiceGroup,
    comment: str | None,
) -> None:
    _add_heading(document, title, 3)
    _static_choices(document, group)
    _historical_comment(document, "Comment", comment)


def _static_choices(document: WordDocument, group: ChoiceGroup) -> None:
    if not group.completed:
        _add_body(document, NOT_COMPLETED)
        return
    paragraph = document.add_paragraph()
    for index, label in enumerate(group.labels):
        mark = "☒" if label in group.selected else "☐"
        _set_run_font(paragraph.add_run(f"{mark}  {label}"), size=11)
        if index != len(group.labels) - 1:
            _set_run_font(paragraph.add_run("      "), size=11)


def _historical_comment(document: WordDocument, label: str, comment: str | None) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.autofit = False
    _set_table_width(table, CONTENT_WIDTH)
    _set_borders(table)
    _set_row_height(table.rows[0], 900)
    cell = table.cell(0, 0)
    _shade(cell, GRAY)
    _valign(cell, "top")
    _cell_margins(cell)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _set_run_font(paragraph.add_run(label), size=10, bold=True)
    body = cell.add_paragraph()
    _set_run_font(body.add_run(comment or NOT_COMPLETED), size=11)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def _response_table(document: WordDocument, case_id: str) -> None:
    _add_heading(document, "ROUND 2 RESPONSE TO ROUND 1 FEEDBACK", 2)
    _add_body(
        document,
        "This table records how the revised chart responds to the earlier review. "
        "It is not a new rating, and it does not replace the blank form above.",
    )
    rows = RESPONSE_ROWS[case_id]
    _bordered_table(
        document,
        ("Round 1 concern", "Round 2 change"),
        rows,
        (3.2, 3.6),
        header=True,
    )


def _banner(
    document: WordDocument,
    text: str,
    *,
    fill: str,
    font_color: RGBColor,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.autofit = False
    _set_table_width(table, CONTENT_WIDTH)
    _set_borders(table, color="595959")
    cell = table.cell(0, 0)
    _shade(cell, fill)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _set_run_font(paragraph.add_run(text), size=12, bold=True, color=font_color)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def _bordered_table(
    document: WordDocument,
    headers: tuple[str, ...],
    rows: tuple[tuple[str, ...], ...],
    widths: tuple[float, ...],
    *,
    header: bool,
    fill: str | None = None,
    size: int = 10,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    _set_table_width(table, sum(widths))
    _set_borders(table)
    _header_row(table.rows[0], headers, widths, size=size)
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            _write_cell(row.cells[index], value, bold=False, fill=fill, size=size)
            row.cells[index].width = Inches(widths[index])
            _valign(row.cells[index], "top")
        _keep_row_together(row)
    if header:
        _mark_header_row(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _header_row(
    row: Any,
    headers: tuple[str, ...],
    widths: tuple[float, ...],
    *,
    size: int = 10,
) -> None:
    for index, header in enumerate(headers):
        cell = row.cells[index]
        _write_cell(cell, header, bold=True, fill=HEADER_FILL, size=size)
        cell.width = Inches(widths[index])
    _keep_row_together(row)


def _set_borders(table: Any, *, color: str = "666666", size: str = "8") -> None:
    properties = table._tbl.tblPr
    existing = properties.find(qn("w:tblBorders"))
    if existing is not None:
        properties.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    properties.append(borders)


def _set_row_height(row: Any, twips: int) -> None:
    properties = row._tr.get_or_add_trPr()
    height = properties.find(qn("w:trHeight"))
    if height is None:
        height = OxmlElement("w:trHeight")
        properties.append(height)
    height.set(qn("w:val"), str(twips))
    height.set(qn("w:hRule"), "atLeast")


def _valign(cell: Any, value: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    existing = properties.find(qn("w:vAlign"))
    if existing is not None:
        properties.remove(existing)
    align = OxmlElement("w:vAlign")
    align.set(qn("w:val"), value)
    properties.append(align)


def _cell_margins(cell: Any) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "80")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def _rows(chart: dict[str, Any], key: str) -> list[dict[str, Any]]:
    raw = chart.get(key)
    if not isinstance(raw, list):
        return []
    return [row for row in raw if isinstance(row, dict)]


def _ordered(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rank = {name: index for index, name in enumerate(TIMEPOINT_ORDER)}

    def key(row: dict[str, Any]) -> tuple[int, str]:
        label = str(row.get("timepoint") or "")
        return (rank.get(label, len(rank)), label)

    return sorted(rows, key=key)


def _timepoint_labels(rows: list[dict[str, Any]]) -> tuple[str, ...]:
    labels: list[str] = []
    seen: set[str] = set()
    for row in _ordered(rows):
        label = field_text(row, "timepoint") or ""
        if label and label not in seen:
            seen.add(label)
            labels.append(label)
    return tuple(labels)


def _blood_pressure(row: dict[str, Any]) -> str:
    systolic = exact_text(row.get("bp_systolic"))
    diastolic = exact_text(row.get("bp_diastolic"))
    if systolic and diastolic:
        return f"{systolic}/{diastolic}"
    return systolic or diastolic or ""


def _with_unit(value: str | None, unit: str) -> str:
    if not value:
        return ""
    return f"{value} {unit}"


def _readme(kind: str) -> str:
    if kind == "revised":
        return "\n".join(
            [
                "# Set 1 — Revised cases",
                "",
                "Open [CliniProof_Cycle2_Revised_Cases_Validation.docx]"
                "(CliniProof_Cycle2_Revised_Cases_Validation.docx) first.",
                "",
                "These six cases were changed because Round 1 clinician feedback "
                "identified problems with the underlying clinical case. Complete the "
                "blank Round 2 form before reading the reproduced Round 1 feedback "
                "that follows each case.",
                "",
                "Ready for clinician review does not mean clinically validated.",
                "",
            ]
        )
    return "\n".join(
        [
            "# Set 2 — Clean cases for fresh review",
            "",
            "Open [CliniProof_Cycle2_Clean_Cases_Validation.docx]"
            "(CliniProof_Cycle2_Clean_Cases_Validation.docx) after Set 1.",
            "",
            "These 18 cases are clean cases prepared for review under the corrected "
            "Round 2 protocol. They are not presented as previously clinician-validated "
            "cases. This codebook does not include Round 1 feedback.",
            "",
            "Ready for clinician review does not mean clinically validated.",
            "",
        ]
    )


def _manifest(rows: list[dict[str, Any]], kind: str) -> str:
    title = "Set 1 revised cases" if kind == "revised" else "Set 2 clean cases"
    lines = [
        f"# {title}",
        "",
        "| Case | Resident file | Evaluator file |",
        "| --- | --- | --- |",
    ]
    for row in rows:
        case_id = row["case_id"]
        folder = row["folder"]
        lines.append(
            f"| {case_id} | `exports/{folder}/{case_id}_resident.json` | "
            f"`exports/{folder}/{case_id}_evaluator.json` |"
        )
    lines.append("")
    return "\n".join(lines)
