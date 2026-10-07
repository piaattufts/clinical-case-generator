# ruff: noqa: E501
"""Clinician codebooks for Round 2 version 4.

The current blank form is written first. Historical Katie and Alex comments
follow that form. They are read-only.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx.document import Document as WordDocument
from docx.shared import RGBColor

from app.services.cycle2_split_casebooks import (
    _banner,
    _bordered_table,
    _chart,
    _historical_comment,
    _round2_form,
)
from app.services.round1_feedback import NOT_COMPLETED
from app.services.word_controls import (
    audit_form_controls,
    finalize_word_form,
    prepare_form_document,
)
from app.services.word_export import NAVY, _add_body, _add_heading, _new_document, _set_run_font

ROOT = Path(__file__).resolve().parents[2]
KATIE_JSON = ROOT / "exports" / "ko_cycle2_revised_validation_v3" / "round1_feedback"
ALEX_CASES = {"VAL-801", "VAL-805", "VAL-809", "VAL-813"}
SET1 = {"VAL-801", "VAL-802", "VAL-803", "VAL-805", "VAL-809", "VAL-813"}


def write_codebooks(
    directory: Path,
    *,
    cases: dict[str, tuple[dict[str, Any], dict[str, Any]]],
    set1: list[str],
    set2: list[str],
    alex: dict[str, dict[str, Any]],
    revision_rows: list[dict[str, str]],
) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    katie = _katie_feedback()
    written: list[Path] = []
    if set1:
        path = directory / "CliniProof_Round2_v4_Set1_Clinician_Validation.docx"
        _write_book(
            path,
            case_ids=set1,
            cases=cases,
            title="Round 2 version 4, Set 1",
            introduction=(
                f"This codebook contains {len(set1)} revised cases: {', '.join(set1)}. "
                "Each case is prepared for clinician re-review. It is not clinically validated. "
                "Complete the blank form before reading the historical feedback."
            ),
            katie=katie,
            alex=alex,
            revision_rows=revision_rows,
        )
        written.append(path)
    if set2:
        path = directory / "CliniProof_Round2_v4_Set2_Clinician_Validation.docx"
        held_note = (
            f"This codebook contains {len(set2)} cases that passed the internal task-validity and clinical-sufficiency gates. "
            "Cases that failed those gates are listed in HELD_CASES.md and are not in this file."
        )
        _write_book(
            path,
            case_ids=set2,
            cases=cases,
            title=f"Round 2 version 4, Set 2 ({len(set2)} cases)",
            introduction=held_note + " These cases are prepared for clinician review. They are not clinically validated.",
            katie={},
            alex={},
            revision_rows=revision_rows,
        )
        written.append(path)
    combined_ids = [*set1, *set2]
    if combined_ids:
        path = directory / "CliniProof_Round2_v4_Combined_Investigator_Copy.docx"
        _write_book(
            path,
            case_ids=combined_ids,
            cases=cases,
            title="INVESTIGATOR CONVENIENCE COPY. NOT THE PRIMARY CLINICIAN WORKFLOW.",
            introduction=(
                "INVESTIGATOR CONVENIENCE COPY. NOT THE PRIMARY CLINICIAN WORKFLOW. "
                "Clinicians should use the Set 1 codebook and the Set 2 codebook."
            ),
            katie=katie,
            alex=alex,
            revision_rows=revision_rows,
        )
        written.append(path)
    return written


def _write_book(
    path: Path,
    *,
    case_ids: list[str],
    cases: dict[str, tuple[dict[str, Any], dict[str, Any]]],
    title: str,
    introduction: str,
    katie: dict[str, dict[str, Any]],
    alex: dict[str, dict[str, Any]],
    revision_rows: list[dict[str, str]],
) -> None:
    document = _new_document("CliniProof Round 2 version 4  |  Prepared for clinician review. Not clinically validated.")
    prepare_form_document(document)
    _add_heading(document, "CliniProof", 0)
    _add_heading(document, title, 1)
    _add_body(document, introduction)
    _add_body(
        document,
        "C1 asks whether the hospitalization is plausible and sufficiently detailed for meaningful clinical reasoning. "
        "C2 asks whether an internal-medicine resident could independently determine a defensible discharge medication regimen from the resident-facing chart.",
    )
    for case_id in case_ids:
        resident, evaluator = cases[case_id]
        row = {
            "case_id": case_id,
            "resident": resident,
            "evaluator": evaluator,
            "special_question": _special_question(case_id),
        }
        _add_v4_case(
            document,
            row,
            katie=katie.get(case_id),
            alex=alex.get(case_id),
            revision_rows=[item for item in revision_rows if item["case"] == case_id],
            include_katie=case_id in SET1,
        )
    document.save(str(path))
    finalize_word_form(path)


def _add_v4_case(
    document: WordDocument,
    row: dict[str, Any],
    *,
    katie: dict[str, Any] | None,
    alex: dict[str, Any] | None,
    revision_rows: list[dict[str, str]],
    include_katie: bool,
) -> None:
    case_id = str(row["case_id"])
    banner = document.add_paragraph()
    banner.paragraph_format.page_break_before = True
    banner.paragraph_format.keep_with_next = True
    _set_run_font(banner.add_run(f"CASE {case_id}"), size=16, bold=True, color=NAVY)
    _add_body(document, "Round 2 combined revision, version 4. Prepared for clinician review. Not clinically validated.")
    _add_heading(document, "CURRENT RESIDENT-FACING CASE", 2)
    _chart(document, row["resident"])
    _add_body(
        document,
        "C1 is clinical plausibility and information sufficiency. Judge whether the chart is detailed enough for meaningful decision-making. "
        "C2 asks whether the resident-facing chart contains enough clinically relevant information for an internal-medicine resident to independently determine a defensible discharge medication regimen.",
    )
    _round2_form(document, row)
    if include_katie:
        _katie_section(document, katie)
    if alex is not None:
        _alex_section(document, alex)
    _revision_section(document, case_id, revision_rows, alex_case=case_id in ALEX_CASES, katie_case=include_katie)


def _katie_section(document: WordDocument, feedback: dict[str, Any] | None) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.page_break_before = True
    _set_run_font(heading.add_run("HISTORICAL ROUND 1 FEEDBACK — KO"), size=14, bold=True, color=NAVY)
    _banner(
        document,
        "ROUND 1 CLINICIAN FEEDBACK — KO — READ ONLY",
        fill="595959",
        font_color=RGBColor(0xFF, 0xFF, 0xFF),
    )
    _add_body(
        document,
        "This section is the extracted Round 1 review from the KO casebook. It is shown only after the blank version 4 form. The boxes are not editable.",
    )
    if feedback is None:
        _add_body(document, "No KO review was stored for this case.")
        return
    overall = feedback["overall_recommendation"]
    _add_body(document, f"Reviewer code: {overall.get('reviewer_code') or NOT_COMPLETED}")
    _add_body(document, f"Date: {overall.get('date') or NOT_COMPLETED}")
    _add_body(document, "Source: KO Casebook Validation extract in the version 3 package. This is not an Alex review.")
    _add_heading(document, "Round 1 C1 domain ratings", 3)
    rendered = []
    for domain in feedback["c1"]["domains"]:
        scores = domain["scores"]
        rendered.append(
            (
                domain["domain"],
                "☒" if scores.get("1") else "☐",
                "☒" if scores.get("2") else "☐",
                "☒" if scores.get("3") else "☐",
                "☒" if scores.get("4") else "☐",
            )
        )
    _bordered_table(document, ("Domain", "1", "2", "3", "4"), tuple(rendered), (3.6, 0.8, 0.8, 0.8, 0.8), header=True)
    _add_body(document, "C1 overall: " + _pass_fail(feedback["c1"]["overall"]))
    _historical_comment(document, "C1 comment", feedback["c1"].get("comment"))
    for key, label in (
        ("c2", "C2"),
        ("c3", "C3"),
        ("c4", "C4"),
        ("c5", "C5"),
    ):
        block = feedback.get(key) or {}
        selection = block.get("selection") or {}
        if "pass" in selection:
            shown = _pass_fail(selection)
        else:
            shown = ", ".join(name for name, flag in selection.items() if flag and name != "flags") or "blank"
        _add_body(document, f"{label}: {shown}")
        _historical_comment(document, f"{label} comment", block.get("comment"))
    recommendation = overall.get("selection") or {}
    shown = ", ".join(name for name, flag in recommendation.items() if flag and name != "flags") or "blank"
    _add_body(document, f"Overall recommendation: {shown}")
    _historical_comment(document, "Overall comment", overall.get("comment"))


def _alex_section(document: WordDocument, feedback: dict[str, Any]) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.page_break_before = True
    _set_run_font(heading.add_run("HISTORICAL C1 FEEDBACK — ALEX / TEAM KAPOw"), size=14, bold=True, color=NAVY)
    _banner(
        document,
        "ROUND 1 CLINICAL PLAUSIBILITY FEEDBACK — ALEX / TEAM KAPOw — READ ONLY",
        fill="595959",
        font_color=RGBColor(0xFF, 0xFF, 0xFF),
    )
    _add_body(
        document,
        "Alex completed C1 for this clean-control case. C2 through C5 and Accept, Revise, or Exclude were blank in the source and are shown as blank. They were not inferred.",
    )
    _add_body(document, feedback["attribution_note"])
    rendered = []
    for domain in feedback["c1_domains"]:
        selected = domain["selected"]
        rendered.append(
            (
                domain["domain"],
                "☒" if "1" in selected else "☐",
                "☒" if "2" in selected else "☐",
                "☒" if "3" in selected else "☐",
                "☒" if "4" in selected else "☐",
            )
        )
    _bordered_table(document, ("Domain", "1", "2", "3", "4"), tuple(rendered), (3.6, 0.8, 0.8, 0.8, 0.8), header=True)
    _add_body(document, "C1 overall: " + (", ".join(feedback["c1_overall"]["selected"]) or "blank"))
    _historical_comment(document, "C1 comment", feedback.get("c1_comment"))
    for label, key, comment_key in (
        ("C2", "c2", "c2_comment"),
        ("C3", "c3", "c3_comment"),
        ("C4", "c4", "c4_comment"),
        ("C5", "c5", "c5_comment"),
    ):
        selected = ", ".join(feedback[key]["selected"]) or "blank"
        _add_body(document, f"{label}: {selected}")
        _historical_comment(document, f"{label} comment", feedback.get(comment_key))
    selected = ", ".join(feedback["recommendation"]["selected"]) or "blank"
    _add_body(document, f"Accept / Revise / Exclude: {selected}")
    _historical_comment(document, "Overall comment", feedback.get("overall_comment"))


def _revision_section(
    document: WordDocument,
    case_id: str,
    rows: list[dict[str, str]],
    *,
    alex_case: bool,
    katie_case: bool,
) -> None:
    _add_heading(document, "CLINICAL REVISION RECORD", 2)
    if alex_case:
        _add_body(document, "Rows that address Alex quote his C1 comment. Other rows are investigator rubric applications.")
    elif katie_case:
        _add_body(document, "Katie / KO feedback is historical task-alignment review. Clinical-sufficiency rubric applied; no Alex case-specific feedback.")
    else:
        _add_body(document, "INVESTIGATOR APPLICATION OF CLINICAL-SUFFICIENCY RUBRIC. Clinical-sufficiency rubric applied; no Alex case-specific feedback.")
    if not rows:
        _add_body(document, f"{case_id} had no text change recorded in the version 4 log.")
        return
    rendered = tuple(
        (
            row["concern"][:220],
            row["revision"][:180],
            row["location"][:80],
            row["old"][:160],
            row["new"][:180],
            row["reasoning"][:180],
            row["evidence"][:120],
        )
        for row in rows
    )
    _bordered_table(
        document,
        (
            "Reviewer/rubric concern",
            "Revision",
            "Location",
            "Old",
            "New",
            "Clinical reasoning",
            "Evidence/provenance",
        ),
        rendered,
        (1.1, 1.0, 0.8, 0.9, 1.0, 1.0, 1.0),
        header=True,
        size=8,
    )


def _pass_fail(selection: dict[str, Any]) -> str:
    chosen = [name for name in ("pass", "fail") if selection.get(name)]
    return ", ".join(chosen) or "blank"


def _katie_feedback() -> dict[str, dict[str, Any]]:
    found: dict[str, dict[str, Any]] = {}
    for path in KATIE_JSON.glob("VAL-*_round1_feedback.json"):
        payload = json.loads(path.read_text(encoding="utf-8"))
        found[str(payload["case_id"])] = payload
    return found


def _special_question(case_id: str) -> str | None:
    questions = {
        "VAL-805": (
            "Given oral furosemide 40 mg once daily, discharge weight 78 kg, and dry weight 73 kg, "
            "is intensification of the loop diuretic acceptable without a recorded intravenous dose or ejection fraction?"
        ),
        "VAL-806": "Restart lisinopril versus continued hold, given creatinine 2.8 mg/dL falling to 1.6 mg/dL and no pre-illness baseline?",
        "VAL-813": "Is temporary tacrolimus reduction acceptable when no trough is recorded, and is any immunosuppressant other than tacrolimus on the chart?",
        "VAL-814": "Restart versus continued hold of mycophenolate after the viral load fell and diarrhea was improving?",
        "VAL-818": "Continue warfarin versus hold, given a postoperative hemoglobin nadir of 7.9 g/dL and a discharge value of 10.3 g/dL?",
        "VAL-820": "Stop prophylactic enoxaparin versus a brief further overlap while warfarin is given and the discharge INR is 2.0?",
        "VAL-821": "Restart apixaban versus continued hold when hemoglobin rose and no endoscopy is recorded?",
        "VAL-822": "Stop hospital-started pantoprazole versus a short continuation after gastrointestinal bleeding?",
        "VAL-823": "Hold apixaban versus restart when hemoglobin rose and the lesion was not examined by endoscopy?",
        "VAL-824": "Stop aspirin versus continue it when no myocardial infarction and no coronary stent are recorded?",
    }
    return questions.get(case_id)


def visual_qa_markdown(directory: Path) -> str:
    """Inspect every case in each clinician codebook. PDF rendering is recorded separately."""
    from docx import Document

    lines = [
        "# Visual QA",
        "",
        "Every included case was inspected in the generated docx. Checkboxes are Word content controls and start unchecked.",
        "",
    ]
    for path in sorted(directory.glob("*.docx")):
        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        audit = audit_form_controls(path)
        lines.append(f"## {path.name}")
        lines.append("")
        lines.append(f"Unchecked checkboxes: {audit.unchecked_checkboxes} of {audit.checkbox_controls}.")
        lines.append(f"Malformed checkboxes: {audit.malformed_checkboxes}.")
        case_ids = [line.removeprefix("CASE ").strip() for line in text.splitlines() if line.startswith("CASE VAL-")]
        lines.append(f"Cases: {', '.join(case_ids)}.")
        for case_id in case_ids:
            section = text.split(f"CASE {case_id}", 1)[1]
            next_case = section.find("\nCASE VAL-")
            body = section if next_case < 0 else section[:next_case]
            order = [
                body.find("CURRENT RESIDENT-FACING CASE"),
                body.find("C1 — Clinical plausibility"),
                body.find("C2 — Sufficiency for resident decision-making"),
                body.find("CLINICIAN VALIDATION REFERENCE"),
                body.find("C3 — Reference-plan validity"),
            ]
            historical = body.find("HISTORICAL")
            revision = body.find("CLINICAL REVISION RECORD")
            problems: list[str] = []
            if order != sorted(order) or any(index < 0 for index in order):
                problems.append(f"section order {order}")
            if historical >= 0 and historical < order[3]:
                problems.append("historical feedback appears before the reference")
            if revision >= 0 and historical >= 0 and revision < historical:
                problems.append("revision record appears before historical feedback")
            if "Not shown to residents" not in body:
                problems.append("reference separation sentence missing")
            lines.append(f"- {case_id}: {'pass' if not problems else '; '.join(problems)}")
        lines.append("")
    lines.append("PDF rendering is appended after LibreOffice conversion.")
    lines.append("")
    return "\n".join(lines)
