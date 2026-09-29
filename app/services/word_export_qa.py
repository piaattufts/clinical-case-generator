"""Independent checks that a resident Word export matches its source JSON."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any

from docx import Document
from docx.document import Document as WordDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.word_export import (
    EMPTY_FOLLOWUP,
    EMPTY_INSTRUCTIONS,
    EMPTY_LABS,
    EMPTY_MEDICATIONS,
    EMPTY_MONITORING,
    EMPTY_VITALS,
    ActiveBatch,
)
from app.services.word_facts import (
    NUMBER_RE,
    exact_text,
    medication_tuple,
    note_text,
    rows_for_context,
    source_numbers,
    source_strings,
)

CASE_ID_RE = re.compile(r"^VAL-\d{3}$")
WORD_RE = re.compile(r"[A-Za-z][A-Za-z/]{3,}")
ARCHIVED_BATCHES = (
    "CLINIPROOF_TAXONOMY_V1",
    "CLINIPROOF_BALANCED_V2",
    "CLINIPROOF_BALANCED_V3",
    "CLINIPROOF_SEEDCASES_V1",
    "CLINIPROOF_SEEDCASES_V2",
)
ANSWER_KEY_TERMS = (
    "error_category",
    "error_family",
    "f1_",
    "f2_",
    "clean_expected",
    "injected_state",
    "correct_action",
    "is_clean_control",
    "answer key",
    "answer_key",
    "CaseAnswerKey",
    "investigator",
)
MED_SECTIONS = {
    "Home medications": "home",
    "Medications during hospitalization": "inpatient",
    "Discharge medications": "discharge",
}
STATUS_SECTIONS = {
    "Admission clinical status": "admission",
    "Discharge / most recent clinical status": "discharge",
}
LEXICON = frozenset(
    {
        "patient",
        "overview",
        "age",
        "sex",
        "gender",
        "sex/gender",
        "weight",
        "clinical",
        "setting",
        "specialty",
        "setting/specialty",
        "admission",
        "diagnosis",
        "disposition",
        "one-liner",
        "reason",
        "hospitalization",
        "chief",
        "complaint",
        "symptoms",
        "symptom",
        "duration",
        "course",
        "history",
        "present",
        "illness",
        "note",
        "review",
        "systems",
        "relevant",
        "medical",
        "past",
        "allergies",
        "problem",
        "list",
        "type",
        "status",
        "context",
        "priority",
        "hospital",
        "intake",
        "output",
        "planning",
        "readiness",
        "detail",
        "transportation",
        "needed",
        "barriers",
        "anticipated",
        "date",
        "home",
        "health",
        "ordered",
        "vital",
        "signs",
        "measure",
        "value",
        "unit",
        "blood",
        "pressure",
        "heart",
        "rate",
        "respiratory",
        "temperature",
        "laboratory",
        "findings",
        "test",
        "result",
        "medications",
        "during",
        "dose",
        "route",
        "frequency",
        "indication",
        "supply",
        "held",
        "monitoring",
        "target",
        "reconciliation",
        "best",
        "possible",
        "medication",
        "source",
        "pharmacist",
        "participate",
        "able",
        "follow-up",
        "follow",
        "instructions",
        "category",
        "imaging",
        "study",
        "body",
        "site",
        "finding",
        "timepoint",
        "consultations",
        "service",
        "assessment",
        "recommendation",
        "procedures",
        "procedure",
        "complications",
        "microbiology",
        "specimen",
        "organism",
        "devices",
        "device",
        "plan",
        "placement",
        "confirmation",
        "care",
        "removal",
        "precautions",
        "return",
        "severity",
        "action",
        "instruction",
        "therapy",
        "restrictions",
        "item",
        "social",
        "living",
        "situation",
        "caregiver",
        "support",
        "financial",
        "literacy",
        "language",
        "preference",
        "substance",
        "advance",
        "directive",
        "dry",
        "notes",
        "with",
        "parameter",
        "trigger",
        "responsible",
        "unverified",
        "interviewer",
        "high-alert",
        "high",
        "alert",
        "identified",
        "discrepancy",
        "types",
        "systolic",
        "diastolic",
        "oxygen",
        "scheduled",
        "specified",
        "discharge",
        "most",
        "recent",
        "beats",
        "beats/min",
        "breaths",
        "breaths/min",
        "spo2",
        "timing",
        "time",
        "liner",
        "recorded",
        "other",
        "none",
        "were",
        "this",
        "details",
        "documentation",
        "mmhg",
        "kg",
        "net",
    }
)
HEADINGS = frozenset(
    {
        "Patient overview",
        "Reason for hospitalization",
        "History of present illness",
        "Admission note",
        "Review of systems",
        "Relevant medical history",
        "Problem list",
        "Social context",
        "Hospital course",
        "Intake and output",
        "Discharge planning",
        "Admission clinical status",
        "Discharge / most recent clinical status",
        "Vital signs",
        "Laboratory findings",
        "Weight",
        "Home medications",
        "Medications during hospitalization",
        "Discharge medications",
        "Medication reconciliation",
        "Reconciliation notes",
        "Monitoring",
        "Follow-up",
        "Discharge instructions",
        "Imaging",
        "Procedures",
        "Consultations",
        "Microbiology",
        "Devices",
        "Return precautions",
        "Therapy restrictions",
    }
)
STATIC_SENTENCES = frozenset(
    {
        EMPTY_MEDICATIONS,
        EMPTY_MONITORING,
        EMPTY_FOLLOWUP,
        EMPTY_INSTRUCTIONS,
        EMPTY_VITALS,
        EMPTY_LABS,
        "No hospital-course details were specified.",
        "No medication-reconciliation documentation was specified.",
    }
)


@dataclass
class CaseAudit:
    case_id: str
    medication_mismatches: list[str] = field(default_factory=list)
    numeric_mismatches: list[str] = field(default_factory=list)
    unsupported: list[str] = field(default_factory=list)
    monitoring_ok: bool = True
    followup_ok: bool = True

    @property
    def ok(self) -> bool:
        return (
            not self.medication_mismatches
            and not self.numeric_mismatches
            and not self.unsupported
            and self.monitoring_ok
            and self.followup_ok
        )


@dataclass
class SetAudit:
    batch_code: str
    expected_ids: list[str]
    exported_ids: list[str]
    missing: list[str]
    duplicates: list[str]
    cases: list[CaseAudit]
    answer_key_hits: list[str]
    archived_hits: list[str]
    raw_json_hits: list[str]

    @property
    def passed(self) -> bool:
        return (
            not self.missing
            and not self.duplicates
            and not self.answer_key_hits
            and not self.archived_hits
            and not self.raw_json_hits
            and all(item.ok for item in self.cases)
            and len(self.exported_ids) == len(self.expected_ids)
        )


def audit_document(path: Any, batch: ActiveBatch, cases: list[dict[str, Any]]) -> SetAudit:
    document = Document(str(path))
    blocks = _blocks(document)
    exported = [block.text for block in blocks if block.kind == "heading" and block.level == 1]
    duplicates = sorted({case_id for case_id in exported if exported.count(case_id) > 1})
    expected = [str(case.get("case_id_code") or "") for case in cases]
    missing = [case_id for case_id in expected if case_id not in exported]
    extra = [case_id for case_id in exported if case_id not in expected]
    by_id = {str(case.get("case_id_code") or ""): case for case in cases}
    audits = [
        _audit_case(
            case_id,
            [block for block in blocks if block.case_id == case_id],
            by_id[case_id],
        )
        for case_id in exported
        if case_id in by_id
    ]
    for case_id in extra:
        audits.append(
            CaseAudit(case_id=case_id, unsupported=[f"unexpected case heading {case_id}"])
        )
    full_text = "\n".join(block.text for block in blocks)
    return SetAudit(
        batch_code=batch.batch_code,
        expected_ids=expected,
        exported_ids=exported,
        missing=missing,
        duplicates=duplicates,
        cases=audits,
        answer_key_hits=[term for term in ANSWER_KEY_TERMS if term in full_text],
        archived_hits=[code for code in ARCHIVED_BATCHES if code in full_text],
        raw_json_hits=_raw_json_hits(full_text),
    )


@dataclass
class _Block:
    kind: str
    level: int
    text: str
    rows: list[list[str]]
    case_id: str | None
    h2: str | None
    h3: str | None


def _blocks(document: WordDocument) -> list[_Block]:
    case_id: str | None = None
    h2: str | None = None
    h3: str | None = None
    blocks: list[_Block] = []
    for item in _iter_blocks(document):
        if isinstance(item, Paragraph):
            style = item.style.name if item.style is not None else ""
            text = item.text.strip()
            level = 0
            if style == "Heading 1" and CASE_ID_RE.fullmatch(text):
                case_id = text
                h2 = None
                h3 = None
                level = 1
                kind = "heading"
            elif style == "Heading 2":
                h2 = text
                h3 = None
                level = 2
                kind = "heading"
            elif style == "Heading 3":
                h3 = text
                level = 3
                kind = "heading"
            else:
                kind = "paragraph"
            blocks.append(_Block(kind, level, text, [], case_id, h2, h3))
        elif isinstance(item, Table):
            rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
            blocks.append(_Block("table", 0, "", rows, case_id, h2, h3))
    return blocks


def _iter_blocks(document: WordDocument) -> Any:
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _audit_case(case_id: str, blocks: list[_Block], case: dict[str, Any]) -> CaseAudit:
    audit = CaseAudit(case_id=case_id)
    plain = _plain(blocks)
    audit.medication_mismatches.extend(_medication_mismatches(blocks, case))
    audit.unsupported.extend(_lab_mismatches(blocks, case))
    audit.unsupported.extend(_missing_narratives(plain, case))
    audit.monitoring_ok = _section_values_match(
        blocks, "Monitoring", _dict_rows(case, "CaseMonitoring"), EMPTY_MONITORING
    )
    audit.followup_ok = _section_values_match(
        blocks, "Follow-up", _dict_rows(case, "CaseFollowup"), EMPTY_FOLLOWUP
    )
    if not audit.monitoring_ok:
        audit.unsupported.append("monitoring section does not match the source")
    if not audit.followup_ok:
        audit.unsupported.append("follow-up section does not match the source")
    rendered_numbers = set(NUMBER_RE.findall(plain))
    audit.numeric_mismatches.extend(sorted(rendered_numbers - source_numbers(case)))
    audit.unsupported.extend(_unknown_words(plain, case))
    return audit


def _plain(blocks: list[_Block]) -> str:
    parts: list[str] = []
    for block in blocks:
        if block.text:
            parts.append(block.text)
        for row in block.rows:
            parts.extend(cell for cell in row if cell)
    return "\n".join(parts)


def _medication_mismatches(blocks: list[_Block], case: dict[str, Any]) -> list[str]:
    mismatches: list[str] = []
    for title, context in MED_SECTIONS.items():
        expected = {medication_tuple(row) for row in rows_for_context(case, context)}
        rendered = _medication_rows(blocks, title)
        if rendered is None:
            mismatches.append(f"{title}: section missing")
            continue
        missing = sorted(expected - rendered)
        extra = sorted(rendered - expected)
        for item in missing:
            mismatches.append(f"{title} missing {item}")
        for item in extra:
            mismatches.append(f"{title} extra {item}")
    return mismatches


def _medication_rows(blocks: list[_Block], title: str) -> set[tuple[str, str, str, str]] | None:
    seen = False
    rows: set[tuple[str, str, str, str]] = set()
    for block in blocks:
        if block.h2 != title or block.h3 not in (None,):
            continue
        seen = True
        if block.kind == "table" and block.rows:
            for row in block.rows[1:]:
                padded = (row + ["", "", "", ""])[:4]
                rows.add((padded[0], padded[1], padded[2], padded[3]))
        if block.kind == "paragraph" and block.text == EMPTY_MEDICATIONS:
            return set()
    if not seen:
        return None
    return rows


def _lab_mismatches(blocks: list[_Block], case: dict[str, Any]) -> list[str]:
    mismatches: list[str] = []
    for title, timepoint in STATUS_SECTIONS.items():
        expected = {
            (
                exact_text(row.get("test_name")) or "",
                exact_text(row.get("value")) or exact_text(row.get("value_text")) or "",
                exact_text(row.get("unit")) or "",
            )
            for row in _dict_rows(case, "CaseLab")
            if str(row.get("timepoint") or "") == timepoint
        }
        rendered = _table_tuples(blocks, title, "Laboratory findings", 3)
        if expected != rendered:
            mismatches.append(
                f"{timepoint} labs differ: missing {sorted(expected - rendered)} "
                f"extra {sorted(rendered - expected)}"
            )
    return mismatches


def _table_tuples(
    blocks: list[_Block], h2: str, h3: str, width: int
) -> set[tuple[str, ...]]:
    found: set[tuple[str, ...]] = set()
    for block in blocks:
        if block.kind != "table" or block.h2 != h2 or block.h3 != h3:
            continue
        for row in block.rows[1:]:
            padded = tuple((row + [""] * width)[:width])
            found.add(padded)
    return found


def _section_values_match(
    blocks: list[_Block],
    heading: str,
    rows: list[dict[str, Any]],
    empty_sentence: str,
) -> bool:
    text = "\n".join(
        block.text
        for block in blocks
        if block.h2 == heading and block.kind == "paragraph" and block.text
    )
    table_text = "\n".join(
        cell
        for block in blocks
        if block.h2 == heading and block.kind == "table"
        for row in block.rows
        for cell in row
    )
    combined = f"{text}\n{table_text}"
    if not rows:
        return empty_sentence in text and table_text.strip() == ""
    if empty_sentence in combined:
        return False
    for row in rows:
        for key, value in row.items():
            if key.endswith("_id") or key in {"case_id", "source_reference"}:
                continue
            rendered = exact_text(value)
            if rendered and rendered not in combined and rendered.replace("_", " ") not in combined:
                return False
    return True


def _missing_narratives(plain: str, case: dict[str, Any]) -> list[str]:
    missing: list[str] = []
    for snippet in _required_snippets(case):
        if snippet not in plain:
            missing.append(f"source text absent from chart: {snippet[:80]}")
    return missing


def _required_snippets(case: dict[str, Any]) -> list[str]:
    snippets: list[str] = []
    info = case.get("ClinicalCase")
    if isinstance(info, dict):
        for key in ("one_liner", "admission_dx"):
            text = exact_text(info.get(key))
            if text:
                snippets.append(text)
        presented = info.get("presentation")
        if isinstance(presented, dict):
            for key in ("hpi", "chief_complaint"):
                text = exact_text(presented.get(key))
                if text:
                    snippets.append(text)
            symptoms = presented.get("presenting_symptoms")
            if isinstance(symptoms, list):
                snippets.extend(text for item in symptoms if (text := exact_text(item)))
    for note_type in ("admission", "hospital_course"):
        text = note_text(case, note_type)
        if text:
            snippets.append(text)
    for key in (
        "CaseInstruction",
        "CaseMedicationReconciliation",
        "CaseImaging",
        "CaseConsult",
        "CaseProcedure",
        "CaseMicrobiology",
        "CaseDevice",
        "CaseReturnPrecaution",
        "CaseDiagnosis",
        "CaseMedication",
        "CaseLab",
    ):
        for row in _dict_rows(case, key):
            for field_name in (
                "instruction_text",
                "notes",
                "finding",
                "assessment",
                "recommendation",
                "findings",
                "care_instructions",
                "removal_plan",
                "reason",
                "action",
                "patient_instruction",
                "diagnosis",
                "drug",
                "test_name",
                "held_reason",
                "indication",
            ):
                text = exact_text(row.get(field_name))
                if text and text != "synthetic_model_generated":
                    snippets.append(text)
    return snippets


def _unknown_words(plain: str, case: dict[str, Any]) -> list[str]:
    text = plain
    strings = sorted(source_strings(case), key=len, reverse=True)
    for snippet in strings:
        if len(snippet) >= 8:
            text = text.replace(snippet, " ")
        elif len(snippet) >= 2:
            text = re.sub(rf"(?<![A-Za-z0-9]){re.escape(snippet)}(?![A-Za-z0-9])", " ", text)
    for sentence in STATIC_SENTENCES:
        text = text.replace(sentence, " ")
    for heading in sorted(HEADINGS, key=len, reverse=True):
        text = text.replace(heading, " ")
    text = NUMBER_RE.sub(" ", text)
    unknown = sorted({word for word in WORD_RE.findall(text) if word.casefold() not in LEXICON})
    return [f"unsupported word: {word}" for word in unknown]


def _dict_rows(case: dict[str, Any], key: str) -> list[dict[str, Any]]:
    rows = case.get(key)
    if not isinstance(rows, list):
        return []
    return [row for row in rows if isinstance(row, dict)]


def _raw_json_hits(text: str) -> list[str]:
    hits: list[str] = []
    for token in ("ClinicalCase", "case_id_code", '"CaseMedication"', "error_category"):
        if token in text:
            hits.append(token)
    return hits


def write_qa_report(root: Any, output_dir: Any) -> Any:
    from app.services.word_export import load_active_batches, load_resident_cases

    balanced_batch, seed_batch = load_active_batches(root)
    balanced = audit_document(
        output_dir / "CliniProof_Balanced_Case_Set.docx",
        balanced_batch,
        load_resident_cases(balanced_batch),
    )
    seed = audit_document(
        output_dir / "CliniProof_Seed_Guided_Case_Set.docx",
        seed_batch,
        load_resident_cases(seed_batch),
    )
    notes, codebook_passed = _codebook_check(
        output_dir / "CliniProof_Resident_Validation_Codebook.docx"
    )
    report = format_qa_report(
        balanced, seed, codebook_notes=notes, codebook_passed=codebook_passed
    )
    path = output_dir / "DOCX_EXPORT_QA.md"
    path.write_text(report, encoding="utf-8")
    return path


def _codebook_check(path: Any) -> tuple[list[str], bool]:
    document = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    required = (
        "A rating of 1 means implausible.",
        "Easy, Moderate, Hard, or Inappropriate / outlier",
        "Could this hospitalization occur as charted?",
        "Accept. The case is suitable for use without clinically meaningful revision.",
        "There is no resident-review application in this repository.",
        "A resident review dashboard is planned.",
    )
    notes = [f"Phrase present: {phrase}" for phrase in required if phrase in text]
    missing = [phrase for phrase in required if phrase not in text]
    ids = set(re.findall(r"VAL-(\d{3})", text))
    leakage = [
        token
        for token in ("f1_", "f2_", "error_category", "error_family", "correct_action")
        if token in text
    ]
    passed = not missing and ids <= {"701", "724", "801", "824"} and not leakage
    if missing:
        notes.append("Missing rubric phrases: " + "; ".join(missing))
    notes.append(
        "Answer-key leakage check: "
        + ("none" if not leakage and ids <= {"701", "724", "801", "824"} else "fail")
    )
    notes.append(
        "Dashboard wording: the resident review dashboard is described as planned; "
        "the implemented software is the generator and reference search."
    )
    return notes, passed


def format_qa_report(
    balanced: SetAudit,
    seed: SetAudit,
    *,
    codebook_notes: list[str],
    codebook_passed: bool,
) -> str:
    return "\n".join(
        [
            "# DOCX export QA",
            "",
            _set_section("Balanced set", balanced),
            _set_section("Seed-guided set", seed),
            "## Codebook",
            "",
            "- Source documentation: `docs/clinical_validation.md` and the validation rubric",
            "  distributed with each current case set (`readable/validation_rubric.md`).",
            "- Scoring/rubric fidelity: " + ("pass" if codebook_passed else "fail"),
            *[f"- {note}" for note in codebook_notes],
            f"- Result: {'PASS' if codebook_passed else 'FAIL'}",
            "",
        ]
    )


def _set_section(title: str, audit: SetAudit) -> str:
    meds = sum(len(case.medication_mismatches) for case in audit.cases)
    numbers = sum(len(case.numeric_mismatches) for case in audit.cases)
    unsupported = sum(len(case.unsupported) for case in audit.cases)
    lines = [
        f"## {title}",
        "",
        f"- Batch: `{audit.batch_code}`",
        f"- Expected case count: {len(audit.expected_ids)}",
        f"- Exported case count: {len(audit.exported_ids)}",
        f"- Missing cases: {audit.missing or 'none'}",
        f"- Duplicate cases: {audit.duplicates or 'none'}",
        f"- Medication mismatches: {meds}",
        f"- Numeric mismatches: {numbers}",
        f"- Unsupported/generated facts: {unsupported}",
        f"- Answer-key leakage: {audit.answer_key_hits or 'none'}",
        f"- Archived-batch leakage: {audit.archived_hits or 'none'}",
        f"- Result: {'PASS' if audit.passed else 'FAIL'}",
        "",
    ]
    if not audit.passed:
        for case in audit.cases:
            if case.ok:
                continue
            lines.append(f"### {case.case_id}")
            lines.extend(f"- {item}" for item in case.medication_mismatches)
            lines.extend(f"- numeric {item}" for item in case.numeric_mismatches)
            lines.extend(f"- {item}" for item in case.unsupported[:30])
            lines.append("")
    return "\n".join(lines)
